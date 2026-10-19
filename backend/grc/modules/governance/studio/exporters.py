"""The studio's document as a file: Word, PDF, Markdown or HTML.

Word and PDF carry the organisation's name in the header and its logo in the header or at the foot of the page (left, centre
or right), the document code, version and classification and "Page X of Y" in the footer, an optional watermark across every
page ("DRAFT"), tables with a repeating header row, numbered lists that restart where each list starts, and every placeholder
still to be filled highlighted (so an unfinished document is visibly unfinished). The look is a professional document's: black
text and headings on white, light grey table headers and rules; only links and the placeholders still to fill have colour.
Styles (fonts, sizes, page layout) belong to the export, not to the Markdown: that is how the studio keeps formatting structured.

A Word file made from an uploaded Word document is not built here: it is the original with the edits written into it
(``docx_io``).
"""
from __future__ import annotations

import html as _html
import io
import os
import re
from typing import Any, Dict, List, Optional, Sequence, Tuple

from . import blocks as B
from .checks import _PLACEHOLDERS

INK = "000000"                  # the title, the headings and the entity name
GREY = "404040"                 # the footer and quotations
RULE = "BFBFBF"                 # the lines under the header, over the footer and for a rule
HEADER_FILL = "F2F2F2"          # a table's header row
PLACEHOLDER_FILL = "FFF2A8"
_MUST = [rx for _kind, required, rx in _PLACEHOLDERS if required]
_RTL = re.compile(r"[\u0590-\u08FF\uFB1D-\uFDFF\uFE70-\uFEFF]")
LOGO_POSITIONS = ("header-left", "header-center", "header-right", "footer-left", "footer-center", "footer-right")
WATERMARK_MAX = 40


class ExportProblem(ValueError):
    """The document cannot be exported in that format; the message is for the person."""


def logo_spot(position: Optional[str]) -> Tuple[str, str]:
    """Where a logo goes: ("header" | "footer", "left" | "center" | "right"). An unknown position is the header's left."""
    where, side = (position if position in LOGO_POSITIONS else "header-left").split("-")
    return where, side


def _split_placeholders(text: str) -> List[tuple]:
    """``[(text, is_placeholder)]``: the text cut at the placeholders that must be filled in."""
    spans: List[tuple] = []
    for rx in _MUST:
        spans += [(m.start(), m.end()) for m in rx.finditer(text)]
    spans.sort()
    out, pos = [], 0
    for s, e in spans:
        if s < pos:
            continue
        if s > pos:
            out.append((text[pos:s], False))
        out.append((text[s:e], True))
        pos = e
    if pos < len(text):
        out.append((text[pos:], False))
    return out or [(text, False)]


def _title_of(blocks: Sequence[B.Block], fallback: str) -> str:
    return next((b.text for b in blocks if b.kind == "heading" and b.level == 1), fallback)


# ── Markdown and HTML ────────────────────────────────────────────────────────
def to_markdown(markdown: str) -> bytes:
    return (markdown if markdown.endswith("\n") else markdown + "\n").encode("utf-8")


def _seg_html(segs: Sequence[B.Seg]) -> str:
    out = []
    for s in segs:
        for piece, is_placeholder in _split_placeholders(s.text):
            text = _html.escape(piece).replace("\n", "<br>")
            if is_placeholder:
                text = f"<mark>{text}</mark>"
            for flag, tag in ((s.code, "code"), (s.strike, "s"), (s.underline, "u"), (s.italic, "em"), (s.bold, "strong")):
                if flag:
                    text = f"<{tag}>{text}</{tag}>"
            if s.href and re.match(r"^(https?:|mailto:)", s.href, re.I):
                text = f'<a href="{_html.escape(s.href, quote=True)}">{text}</a>'
            out.append(text)
    return "".join(out)


def to_html(markdown: str, *, title: str = "", entity_name: str = "", logo: Optional[tuple] = None, logo_position: str = "header-left",
            watermark: str = "") -> bytes:
    """A self-contained HTML page (the logo is embedded). Everything is escaped: only the tags written here appear."""
    import base64

    blocks = B.parse(markdown)
    body: List[str] = []
    stack: List[str] = []

    def close_to(depth: int) -> None:
        while len(stack) > depth:
            body.append(f"</{stack.pop()}>")

    for b in blocks:
        if b.kind != "list_item":
            close_to(0)
        if b.kind == "heading":
            body.append(f"<h{b.level}>{_seg_html(b.segments)}</h{b.level}>")
        elif b.kind == "paragraph":
            body.append(f"<p>{_seg_html(b.segments)}</p>")
        elif b.kind == "list_item":
            tag = "ol" if b.ordered else "ul"
            close_to(b.level + 1)
            while len(stack) < b.level + 1:
                body.append(f"<{tag}>")
                stack.append(tag)
            body.append(f"<li>{_seg_html(b.segments)}</li>")
        elif b.kind == "table":
            rows = "".join("<tr>" + "".join(f"<{'th' if r == 0 else 'td'}>{_seg_html(c.segments) or _html.escape(c.text)}</{'th' if r == 0 else 'td'}>" for c in row) + "</tr>"
                           for r, row in enumerate(b.rows))
            body.append(f"<table>{rows}</table>")
        elif b.kind == "quote":
            body.append(f"<blockquote>{_html.escape(b.text)}</blockquote>")
        elif b.kind == "code":
            body.append(f"<pre>{_html.escape(b.text)}</pre>")
        elif b.kind == "hr":
            body.append("<hr>")
        else:
            body.append(f"<p>{_html.escape(b.text)}</p>")
    close_to(0)
    where, side = logo_spot(logo_position)
    img = f'<img src="data:{logo[1]};base64,{base64.b64encode(logo[0]).decode()}" alt="">' if logo else ""
    letterhead = footer = ""
    justify = {"center": "center", "right": "flex-end"}.get(side, "flex-start") if where == "header" else "flex-start"
    if (img and where == "header") or entity_name:
        letterhead = f'<header style="justify-content:{justify}">{img if where == "header" else ""}<span>{_html.escape(entity_name)}</span></header>'
    if img and where == "footer":
        footer = f'<footer style="text-align:{side}">{img}</footer>'
    css = ("body{font:15px/1.55 Calibri,Arial,sans-serif;max-width:820px;margin:32px auto;padding:0 20px;color:#000;background:#fff}header{display:flex;align-items:center;gap:14px;"
           "border-bottom:2px solid #bfbfbf;padding-bottom:10px;margin-bottom:24px}header img{height:48px}footer{border-top:1px solid #bfbfbf;margin-top:32px;padding-top:12px}"
           "footer img{height:40px}h1,h2,h3,h4{color:#000}table{border-collapse:collapse;width:100%;margin:12px 0}"
           "td,th{border:1px solid #000;padding:6px 8px;vertical-align:top}th{background:#f2f2f2;text-align:left}blockquote{border-left:4px solid #808080;margin:12px 0;padding:2px 14px;color:#404040;font-style:italic}"
           "pre{background:#f2f2f2;padding:10px;overflow:auto}code{background:#f2f2f2;padding:1px 4px}mark{background:#fff2a8}")
    mark = ""
    if watermark.strip():
        mark = f' data-watermark="{_html.escape(watermark.strip()[:WATERMARK_MAX], quote=True)}"'
        css += ("body[data-watermark]::before{content:attr(data-watermark);position:fixed;top:50%;left:50%;transform:translate(-50%,-50%) rotate(-35deg);"
                "font:700 110px Calibri,Arial,sans-serif;color:rgba(0,0,0,.06);white-space:nowrap;pointer-events:none;z-index:0}body>*{position:relative;z-index:1}")
    page = (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{_html.escape(title or _title_of(blocks, "Document"))}</title>'
            f"<style>{css}</style></head><body{mark}>{letterhead}{''.join(body)}{footer}</body></html>")
    return page.encode("utf-8")


# ── Word ─────────────────────────────────────────────────────────────────────
def _docx_add_runs(paragraph, segs: Sequence[B.Seg], highlight: bool = True, base_bold: bool = False) -> None:
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor

    for s in segs:
        pieces = _split_placeholders(s.text) if highlight else [(s.text, False)]
        for text, is_placeholder in pieces:
            if not text:
                continue
            if s.href and re.match(r"^(https?:|mailto:)", s.href, re.I):
                run = _hyperlink(paragraph, s.href, text)
            else:
                run = paragraph.add_run()
            parts = text.split("\n")
            for i, part in enumerate(parts):
                if i:
                    run.add_break()
                run.add_text(part)
            run.bold = True if (s.bold or base_bold) else None
            run.italic = True if s.italic else None
            run.underline = True if s.underline or s.href else None
            run.font.strike = True if s.strike else None
            if s.code:
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
            if s.href:
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)         # Word's own link blue, kept: a link must look like one
            if is_placeholder:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _hyperlink(paragraph, url: str, text: str):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.run import Run

    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    link.append(r)
    paragraph._p.append(link)
    return Run(r, paragraph)


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _border(paragraph, side: str, size: int = 8, color: str = RULE, space: int = 4) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    bdr = p_pr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        before = [c for c in p_pr if c.tag in {qn(f"w:{n}") for n in ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers")}]
        (before[-1].addnext(bdr) if before else p_pr.insert(0, bdr))              # where Word's schema puts it, not at the end
    el = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", str(space)), ("color", color)):
        el.set(qn(f"w:{k}"), v)
    bdr.append(el)


def _field(paragraph, instruction: str) -> None:
    """A Word field (PAGE, NUMPAGES) that Word fills in when it lays the page out."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        run = paragraph.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif text == instruction:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instruction} "
            run._r.append(it)
        else:
            run.add_text(text)
        run.font.size = Pt(8.5)


def _restart_numbering(doc, style_name: str, start: int = 1) -> int:
    """A fresh numbering instance that starts at ``start``, so each numbered list of the document counts from its own first number."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    base = numbering.num_having_numId(doc.styles[style_name].element.pPr.numPr.numId.val)
    num = numbering.add_num(base.abstractNumId.val)
    num.add_lvlOverride(ilvl=0).add_startOverride(max(1, int(start or 1)))
    return num.numId


def _set_numbering(paragraph, num_id: int, level: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    old = p_pr.find(qn("w:numPr"))
    if old is not None:
        p_pr.remove(old)
    num_pr = OxmlElement("w:numPr")
    for tag, val in (("ilvl", level), ("numId", num_id)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), str(val))
        num_pr.append(el)
    p_pr.insert(1 if p_pr.find(qn("w:pStyle")) is not None else 0, num_pr)


def black_headings(doc) -> None:
    """The title and headings of a python-docx document in black. Word's own template, which python-docx starts from, colours
    them blue with theme colours (on each style and on the character style linked to it) and rules a blue line under the Title."""
    from docx.oxml.ns import qn
    from docx.shared import RGBColor

    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        for style in [doc.styles[n] for n in (name, f"{name} Char") if n in doc.styles]:
            style.font.color.rgb = RGBColor.from_string(INK)            # replaces the theme colour, which Word would show instead
        for rule in doc.styles[name].element.findall(f"{qn('w:pPr')}/{qn('w:pBdr')}"):
            rule.getparent().remove(rule)


def _style_doc(doc) -> None:
    from docx.shared import Cm, Pt

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.1)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.15
    for name, size, before in (("Title", 24, 0), ("Heading 1", 16, 16), ("Heading 2", 13, 12), ("Heading 3", 11.5, 10), ("Heading 4", 10.5, 8)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Calibri", Pt(size), True
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(6)
        st.paragraph_format.keep_with_next = True
    black_headings(doc)


def _header_footer(doc, *, entity_name: str, logo_path: Optional[str], logo_position: str, footer_text: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt, RGBColor

    sec = doc.sections[0]
    head = sec.header.paragraphs[0]
    where, side = logo_spot(logo_position)
    align = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(side, WD_ALIGN_PARAGRAPH.LEFT)
    has_logo = bool(logo_path and os.path.exists(logo_path))
    head.alignment = align if where == "header" else WD_ALIGN_PARAGRAPH.LEFT
    if has_logo and where == "header":
        head.add_run().add_picture(logo_path, height=Cm(1.3))
        if entity_name:
            head.add_run("   ")
    if entity_name:
        run = head.add_run(entity_name)
        run.font.size, run.font.bold, run.font.color.rgb = Pt(9.5), True, RGBColor.from_string(INK)
    _border(head, "bottom", 6, RULE, 6)
    foot = sec.footer.paragraphs[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    foot.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    run = foot.add_run(footer_text)
    run.font.size, run.font.color.rgb = Pt(8.5), RGBColor.from_string(GREY)
    foot.add_run("\tPage ").font.size = Pt(8.5)
    _field(foot, "PAGE")
    foot.add_run(" of ").font.size = Pt(8.5)
    _field(foot, "NUMPAGES")
    _border(foot, "top", 6, RULE, 6)
    if has_logo and where == "footer":                  # at the foot of every page, under the document code and the page numbers
        mark = sec.footer.add_paragraph()
        mark.alignment = align
        mark.paragraph_format.space_before = Pt(4)
        mark.add_run().add_picture(logo_path, height=Cm(0.9))


_VML = 'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"'
# Word's own text watermark: a WordArt shape (type 136) in the header, centred on the page, rotated, faint, behind the text
_WATERMARK_SHAPETYPE = (
    '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">'
    '<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/>'
    '<v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/>'
    '<v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>'
    '<v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
    '<v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles><o:lock v:ext="edit" text="t" shapetype="t"/>'
    '</v:shapetype>')


def _watermark_run(text: str, n: int):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    width = min(470, 90 * len(text) + 40)                # points; A4 between the margins is ~470pt
    height = min(150, width / max(1, len(text)) * 1.5)
    return parse_xml(
        f'<w:r {nsdecls("w")} {_VML}><w:rPr><w:noProof/></w:rPr><w:pict>{_WATERMARK_SHAPETYPE}'
        f'<v:shape id="CompliverseWatermark{n}" o:spid="_x0000_s{2049 + n}" type="#_x0000_t136" o:allowincell="f" fillcolor="#c0c0c0" stroked="f" '
        f'style="position:absolute;margin-left:0;margin-top:0;width:{width:.0f}pt;height:{height:.0f}pt;rotation:315;z-index:-251654144;'
        'mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin">'
        f'<v:fill opacity=".5"/><v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="{_html.escape(text, quote=True)}"/>'
        '</v:shape></w:pict></w:r>')


def add_watermark(doc, text: str) -> None:
    """Put ``text`` across every page of a python-docx document, the way Word does: in each header of the document (the first
    page's and the even pages' too when it has them). A header that only repeats the section before it is left alone."""
    text = (text or "").strip()[:WATERMARK_MAX]
    if not text:
        return
    n = 0
    for i, sec in enumerate(doc.sections):
        heads = [sec.header]
        if sec.different_first_page_header_footer:
            heads.append(sec.first_page_header)
        if doc.settings.odd_and_even_pages_header_footer:
            heads.append(sec.even_page_header)
        for head in heads:
            if i > 0 and head.is_linked_to_previous:         # it shows the previous section's header, which has the watermark
                continue
            p = head.paragraphs[0] if head.paragraphs else head.add_paragraph()
            p._p.append(_watermark_run(text, n))
            n += 1


def watermark_docx(data: bytes, text: str) -> bytes:
    """A Word file's bytes with a watermark added (a file edited in place keeps everything else of its own)."""
    from docx import Document

    if not (text or "").strip():
        return data
    doc = Document(io.BytesIO(data))
    add_watermark(doc, text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def to_docx(markdown: str, *, title: str = "", entity_name: str = "", logo_path: Optional[str] = None, logo_position: str = "header-left",
            doc_code: str = "", version: str = "", classification: str = "", highlight_placeholders: bool = True, watermark: str = "") -> bytes:
    """A Word document built from the Markdown."""
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    blocks = B.parse(markdown)
    doc = Document()
    _style_doc(doc)
    _header_footer(doc, entity_name=entity_name, logo_path=logo_path, logo_position=logo_position,
                   footer_text=" · ".join(x for x in (doc_code, f"v{version}" if version else "", classification.title() if classification else "") if x) or (title or ""))
    add_watermark(doc, watermark)
    heads = [b for b in blocks if b.kind == "heading"]
    one_title = bool(heads) and heads[0].level == 1 and sum(1 for h in heads if h.level == 1) == 1
    cp = doc.core_properties
    cp.title, cp.author, cp.subject = title or _title_of(blocks, ""), entity_name or "", doc_code or ""
    list_state: Dict[str, Any] = {"num": None, "ordered": None}
    for b in blocks:
        if b.kind != "list_item":
            list_state = {"num": None, "ordered": None}
        if b.kind == "heading":
            if one_title and b.level == 1:
                p = doc.add_paragraph(style="Title")
            else:
                level = min(4, max(1, b.level - 1 if one_title else b.level))
                p = doc.add_paragraph(style=f"Heading {level}")
            _docx_add_runs(p, b.segments, highlight_placeholders)
        elif b.kind == "paragraph":
            _docx_add_runs(doc.add_paragraph(), b.segments, highlight_placeholders)
        elif b.kind == "list_item":
            depth = min(b.level, 2)
            if b.ordered:
                style = "List Number" if depth == 0 else f"List Number {depth + 1}"
                p = doc.add_paragraph(style=style)
                if depth == 0:
                    if list_state["ordered"] is not True or list_state["num"] is None:
                        list_state = {"num": _restart_numbering(doc, "List Number", b.number or 1), "ordered": True}
                    _set_numbering(p, list_state["num"], 0)
            else:
                p = doc.add_paragraph(style="List Bullet" if depth == 0 else f"List Bullet {depth + 1}")
                if depth == 0:
                    list_state = {"num": None, "ordered": False}
            p.paragraph_format.space_after = Pt(3)
            _docx_add_runs(p, b.segments, highlight_placeholders)
        elif b.kind == "table":
            _docx_table(doc, b, highlight_placeholders)
        elif b.kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            _border(p, "left", 18, "808080", 8)
            run = p.add_run(b.text)
            run.italic, run.font.color.rgb = True, RGBColor.from_string(GREY)
        elif b.kind == "code":
            for line in (b.text or " ").split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(line or " ")
                run.font.name, run.font.size = "Consolas", Pt(9)
        elif b.kind == "hr":
            _border(doc.add_paragraph(), "bottom", 6, RULE, 1)
        else:
            doc.add_paragraph(b.text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _docx_table(doc, b: B.Block, highlight: bool) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    rows = b.rows
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            _docx_add_runs(p, row[c].segments if c < len(row) else [], highlight, base_bold=(r == 0))
            if r == 0:
                _shade(cell, HEADER_FILL)
        if r == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            head = OxmlElement("w:tblHeader")
            head.set(qn("w:val"), "true")
            tr_pr.append(head)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ── PDF ──────────────────────────────────────────────────────────────────────
_FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf", "C:/Windows/Fonts/arialbi.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf"),
]


def _pdf_font() -> tuple:
    """``(regular, bold, italic, bold-italic)`` font names: a TrueType font with wide Unicode coverage when the machine has one,
    else the built-in Helvetica (Latin text only)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for paths in _FONT_CANDIDATES:
        if all(os.path.exists(p) for p in paths):
            names = ("StudioSans", "StudioSans-Bold", "StudioSans-Italic", "StudioSans-BoldItalic")
            try:
                for n, p in zip(names, paths):
                    if n not in pdfmetrics.getRegisteredFontNames():
                        pdfmetrics.registerFont(TTFont(n, p))
                pdfmetrics.registerFontFamily(names[0], normal=names[0], bold=names[1], italic=names[2], boldItalic=names[3])
                return names
            except Exception:       # noqa: BLE001 - a broken font file must not stop the export
                continue
    return ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique")


def _pdf_inline(segs: Sequence[B.Seg], highlight: bool, mono: str = "Courier") -> str:
    out = []
    for s in segs:
        for text, is_placeholder in (_split_placeholders(s.text) if highlight else [(s.text, False)]):
            t = _html.escape(text).replace("\n", "<br/>")
            if s.code:
                t = f'<font face="{mono}">{t}</font>'
            if s.bold:
                t = f"<b>{t}</b>"
            if s.italic:
                t = f"<i>{t}</i>"
            if s.underline:
                t = f"<u>{t}</u>"
            if s.strike:
                t = f"<strike>{t}</strike>"
            if s.href and re.match(r"^(https?:|mailto:)", s.href, re.I):
                t = f'<a href="{_html.escape(s.href, quote=True)}" color="#0563C1">{t}</a>'
            if is_placeholder:
                t = f'<font backColor="#{PLACEHOLDER_FILL}">{t}</font>'
            out.append(t)
    return "".join(out)


def to_pdf(markdown: str, *, title: str = "", entity_name: str = "", logo_path: Optional[str] = None, logo_position: str = "header-left",
           doc_code: str = "", version: str = "", classification: str = "", highlight_placeholders: bool = True, watermark: str = "") -> bytes:
    """A PDF built from the Markdown (A4). Right-to-left scripts are not supported here: use Word for those."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    if _RTL.search(markdown):
        raise ExportProblem("PDF export cannot show right-to-left text (Arabic, Hebrew, Urdu) correctly yet. Export to Word instead.")
    regular, bold, italic, bold_italic = _pdf_font()
    ink, grey, rule, fill = (colors.HexColor("#" + c) for c in (INK, GREY, RULE, HEADER_FILL))
    blocks = B.parse(markdown)
    heads = [b for b in blocks if b.kind == "heading"]
    one_title = bool(heads) and heads[0].level == 1 and sum(1 for h in heads if h.level == 1) == 1
    body = ParagraphStyle("body", fontName=regular, fontSize=10, leading=14.5, spaceAfter=6)
    styles = {"title": ParagraphStyle("title", parent=body, fontName=bold, fontSize=22, leading=26, textColor=ink, spaceAfter=12)}
    for n, (size, before) in enumerate(((16, 14), (13, 11), (11.5, 9), (10.5, 7)), start=1):
        styles[f"h{n}"] = ParagraphStyle(f"h{n}", parent=body, fontName=bold, fontSize=size, leading=size + 4, textColor=ink, spaceBefore=before, spaceAfter=5, keepWithNext=1)
    quote = ParagraphStyle("quote", parent=body, leftIndent=18, textColor=grey, fontName=italic, borderPadding=(0, 0, 0, 8))
    story: List[Any] = []
    pending: List[B.Block] = []

    def nested(items: List[B.Block], i: int, depth: int) -> tuple:
        """One list at ``depth`` starting at ``items[i]``, with the deeper items nested under their parent: ``(flowable, next index)``."""
        ordered, first = items[i].ordered, items[i].number or 1
        rows: List[List[Any]] = []
        while i < len(items) and items[i].level >= depth:
            it = items[i]
            if it.level > depth:
                sub, i = nested(items, i, it.level)
                rows[-1].append(sub) if rows else rows.append([sub])
                continue
            if it.ordered != ordered:
                break
            rows.append([Paragraph(_pdf_inline(it.segments, highlight_placeholders), body)])
            i += 1
        flow = ListFlowable([ListItem(r) for r in rows], bulletType="1" if ordered else "bullet", bulletFontName=regular, bulletFontSize=9, leftIndent=18,
                            **({"start": first, "bulletFormat": "%s."} if ordered else {}))
        return flow, i

    def flush_list() -> None:
        i = 0
        while i < len(pending):
            flow, i = nested(pending, i, pending[i].level)
            story.append(flow)
        pending.clear()

    for b in blocks:
        if b.kind != "list_item":
            flush_list()
        if b.kind == "heading":
            if one_title and b.level == 1:
                story.append(Paragraph(_pdf_inline(b.segments, highlight_placeholders), styles["title"]))
            else:
                level = min(4, max(1, b.level - 1 if one_title else b.level))
                story.append(Paragraph(_pdf_inline(b.segments, highlight_placeholders), styles[f"h{level}"]))
        elif b.kind == "paragraph":
            story.append(Paragraph(_pdf_inline(b.segments, highlight_placeholders), body))
        elif b.kind == "list_item":
            pending.append(b)
        elif b.kind == "table":
            width = max(len(r) for r in b.rows)
            data = [[Paragraph(_pdf_inline(row[c].segments if c < len(row) else [], highlight_placeholders), ParagraphStyle("c", parent=body, fontName=bold if r == 0 else regular, fontSize=9, leading=12, spaceAfter=0))
                     for c in range(width)] for r, row in enumerate(b.rows)]
            avail = A4[0] - 4.4 * cm
            lens = [max(8, min(60, max(len(row[c].text) if c < len(row) else 0 for row in b.rows))) for c in range(width)]
            widths = [avail * n / sum(lens) for n in lens]
            table = Table(data, colWidths=widths, repeatRows=1)
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, ink), ("BACKGROUND", (0, 0), (-1, 0), fill),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
            story += [table, Spacer(1, 8)]
        elif b.kind == "quote":
            story.append(Paragraph(_html.escape(b.text), quote))
        elif b.kind == "code":
            story.append(Preformatted(b.text, ParagraphStyle("code", parent=body, fontName="Courier", fontSize=8.5, leading=11, backColor=fill)))
        elif b.kind == "hr":
            story.append(HRFlowable(width="100%", color=rule, spaceBefore=6, spaceAfter=6))
        else:
            story.append(Paragraph(_html.escape(b.text), body))
    flush_list()
    footer = " · ".join(x for x in (doc_code, f"v{version}" if version else "", classification.title() if classification else "") if x) or (title or "")
    logo = None
    if logo_path and os.path.exists(logo_path):
        from reportlab.lib.utils import ImageReader

        logo = ImageReader(logo_path)

    class Canvas(rl_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages: List[dict] = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._chrome(total)
                super().showPage()
            super().save()

        def _chrome(self, total: int) -> None:
            w, h = A4
            self.setStrokeColor(rule)
            self.setLineWidth(0.5)
            x = 2.2 * cm
            if logo is not None and where == "header":
                iw, ih = logo.getSize()
                lh = 1.3 * cm
                lw = lh * iw / ih
                x = {"center": (w - lw) / 2, "right": w - 2.2 * cm - lw}.get(side, 2.2 * cm)
                self.drawImage(logo, x, h - 1.2 * cm - lh, width=lw, height=lh, mask="auto")
                x = x + lw + 8 if side == "left" else 2.2 * cm
            if entity_name:
                self.setFont(bold, 9.5)
                self.setFillColor(ink)
                self.drawString(x, h - 1.2 * cm - 0.9 * cm, entity_name)
            self.line(2.2 * cm, h - 2.55 * cm, w - 2.2 * cm, h - 2.55 * cm)
            low = logo is not None and where == "footer"         # the logo at the foot of the page: the footer line and text move up over it
            line_y, text_y = (1.95 * cm, 1.5 * cm) if low else (1.7 * cm, 1.2 * cm)
            self.line(2.2 * cm, line_y, w - 2.2 * cm, line_y)
            self.setFont(regular, 8.5)
            self.setFillColor(grey)
            self.drawString(2.2 * cm, text_y, footer)
            self.drawRightString(w - 2.2 * cm, text_y, f"Page {self._pageNumber} of {total}")
            if low:
                iw, ih = logo.getSize()
                lh = 0.8 * cm
                lw = lh * iw / ih
                lx = {"center": (w - lw) / 2, "right": w - 2.2 * cm - lw}.get(side, 2.2 * cm)
                self.drawImage(logo, lx, 0.35 * cm, width=lw, height=lh, mask="auto")

    mark = (watermark or "").strip()[:WATERMARK_MAX]

    def under(canvas, _doc) -> None:
        """The watermark, drawn first on every page so the text lies over it."""
        if not mark:
            return
        from reportlab.pdfbase.pdfmetrics import stringWidth

        w, h = A4
        size = max(28.0, min(120.0, 100 * 700 / max(1.0, stringWidth(mark, bold, 100))))
        canvas.saveState()
        canvas.setFillColor(ink)
        canvas.setFillAlpha(0.08)
        canvas.setFont(bold, size)
        canvas.translate(w / 2, h / 2)
        canvas.rotate(45)
        canvas.drawCentredString(0, -size / 3, mark)
        canvas.restoreState()

    where, side = logo_spot(logo_position)
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=3 * cm, bottomMargin=2.3 * cm,
                      title=title or _title_of(blocks, ""), author=entity_name or "").build(story or [Spacer(1, 1)], onFirstPage=under, onLaterPages=under,
                                                                                           canvasmaker=Canvas)
    return out.getvalue()


FORMATS = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    "pdf": ("application/pdf", ".pdf"),
    "md": ("text/markdown; charset=utf-8", ".md"),
    "html": ("text/html; charset=utf-8", ".html"),
}


def export(markdown: str, fmt: str, *, title: str = "", entity_name: str = "", logo_path: Optional[str] = None, logo_position: str = "header-left",
           doc_code: str = "", version: str = "", classification: str = "", highlight_placeholders: bool = True, watermark: str = "") -> bytes:
    if fmt == "md":
        return to_markdown(markdown)
    if fmt == "html":
        logo = None
        if logo_path and os.path.exists(logo_path):
            ext = os.path.splitext(logo_path)[1].lower().lstrip(".")
            with open(logo_path, "rb") as f:
                logo = (f.read(), "image/" + ("jpeg" if ext in ("jpg", "jpeg") else ext or "png"))
        return to_html(markdown, title=title, entity_name=entity_name, logo=logo, logo_position=logo_position, watermark=watermark)
    kwargs = dict(title=title, entity_name=entity_name, logo_path=logo_path, logo_position=logo_position, doc_code=doc_code, version=version,
                  classification=classification, highlight_placeholders=highlight_placeholders, watermark=watermark)
    if fmt == "docx":
        return to_docx(markdown, **kwargs)
    if fmt == "pdf":
        return to_pdf(markdown, **kwargs)
    raise ExportProblem("Choose Word, PDF, Markdown or HTML.")
