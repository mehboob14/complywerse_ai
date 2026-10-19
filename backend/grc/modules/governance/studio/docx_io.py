"""Word files in the studio.

``import_docx`` reads a .docx as Markdown (headings, paragraphs, list items, tables) and checks that the Markdown reads back as
the same blocks, so each block can be tied to its paragraph or table in the file. Tracked changes in the file are accepted
for the working text; the text before them is kept as the baseline, so the Changes view shows what the reviewers did.

``export_in_place`` writes the edited Markdown back into a copy of the file: only the words that changed are touched, inside
their own runs (so fonts, sizes, colours and styles around them stay), new paragraphs take the look of their neighbours,
deleted ones go, table cells and rows are edited where they are. Everything the studio does not show (headers, footers,
images, text boxes, comments, page setup, fields) stays exactly as it was. The result is read again and must give the same
blocks as the edited Markdown; if it does not, ``PatchProblem`` is raised and the caller builds a clean Word copy instead.
"""
from __future__ import annotations

import copy
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from typing import Any, Dict, Iterator, List, Optional, Sequence, Tuple

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.text.run import Run
from lxml import etree

from . import blocks as B
from .importers import Imported, ImportProblem

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_HEADING = re.compile(r"^heading\s*([1-9])$", re.I)
_ESC = re.compile(r"([\\`*_\[\]<>|~&])")
_TOKEN = re.compile(r"\s+|\w+|[^\w\s]", re.UNICODE)
EMPHASIS = ("b", "bCs", "i", "iCs", "u", "strike", "dstrike", "highlight", "shd", "vertAlign", "caps", "smallCaps", "rStyle")
Fmt = Tuple[bool, bool, bool, bool]                    # bold, italic, underline, strike as the Markdown shows them
NewFmt = Tuple[bool, bool, bool, bool, Optional[str]]  # the same, with the link of the text
NO_FMT: Fmt = (False, False, False, False)
NO_NEW: NewFmt = (False, False, False, False, None)
_PPR_BEFORE_NUMPR = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl")


class PatchProblem(ValueError):
    """The edits cannot be written into the file safely; the caller falls back to a clean Word copy."""


# ── xml helpers ──────────────────────────────────────────────────────────────
def _local(el) -> str:
    tag = getattr(el, "tag", None)
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _on(el) -> bool:
    v = el.get(qn("w:val"))
    return v is None or v.lower() not in ("0", "false", "off", "none")


def _unwrap(el) -> None:
    parent = el.getparent()
    for ch in list(el):
        el.addprevious(ch)
    parent.remove(el)


def _insert_after_any(parent, el, names: Sequence[str]) -> None:
    """Put ``el`` after the last of the ``names`` children of ``parent`` (the order Word's schema wants), else first."""
    last = None
    for ch in parent:
        if _local(ch) in names:
            last = ch
    (last.addnext(el) if last is not None else parent.insert(0, el))


# ── styles and numbering ─────────────────────────────────────────────────────
class Styles:
    def __init__(self, root):
        self.by_id: Dict[str, Any] = {}
        self.name: Dict[str, str] = {}
        self.default_p = "Normal"
        for st in root.findall(qn("w:style")):
            sid = st.get(qn("w:styleId"))
            self.by_id[sid] = st
            n = st.find(qn("w:name"))
            self.name[sid] = (n.get(qn("w:val")) if n is not None else sid) or sid
            if st.get(qn("w:type")) == "paragraph" and st.get(qn("w:default")) in ("1", "true"):
                self.default_p = sid

    def chain(self, sid: Optional[str]) -> Iterator[Any]:
        seen = set()
        while sid and sid in self.by_id and sid not in seen:
            seen.add(sid)
            st = self.by_id[sid]
            yield st
            b = st.find(qn("w:basedOn"))
            sid = b.get(qn("w:val")) if b is not None else None

    def lname(self, sid: Optional[str]) -> str:
        return (self.name.get(sid or "", sid or "") or "").lower()

    def flag(self, sid: Optional[str], name: str) -> Optional[bool]:
        for st in self.chain(sid):
            rpr = st.find(qn("w:rPr"))
            e = rpr.find(qn(f"w:{name}")) if rpr is not None else None
            if e is not None:
                return _on(e)
        return None

    def outline(self, sid: Optional[str]) -> Optional[int]:
        for st in self.chain(sid):
            e = st.find(qn("w:pPr") + "/" + qn("w:outlineLvl"))
            if e is not None:
                try:
                    return int(e.get(qn("w:val")))
                except (TypeError, ValueError):
                    return None
        return None

    def numpr(self, sid: Optional[str]) -> Tuple[Optional[str], Optional[int]]:
        num, lvl = None, None
        for st in self.chain(sid):
            np_ = st.find(qn("w:pPr") + "/" + qn("w:numPr"))
            if np_ is None:
                continue
            n, l = np_.find(qn("w:numId")), np_.find(qn("w:ilvl"))
            if num is None and n is not None:
                num = n.get(qn("w:val"))
            if lvl is None and l is not None:
                lvl = int(l.get(qn("w:val")) or 0)
            if num is not None:
                break
        return num, lvl

    def id_for(self, *names: str) -> Optional[str]:
        want = {n.lower() for n in names}
        return next((sid for sid, n in self.name.items() if n.lower() in want), None)


class Numbering:
    def __init__(self, root):
        self.root = root

    def _find(self, num_id: str, ilvl: int):
        if self.root is None:
            return None, None
        num = next((n for n in self.root.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id)), None)
        if num is None:
            return None, None
        over = next((o for o in num.findall(qn("w:lvlOverride")) if o.get(qn("w:ilvl")) == str(ilvl)), None)
        if over is not None and over.find(qn("w:lvl")) is not None:
            return over.find(qn("w:lvl")), over
        abs_id = num.find(qn("w:abstractNumId"))
        if abs_id is not None:
            absn = next((a for a in self.root.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abs_id.get(qn("w:val"))), None)
            if absn is not None:
                return next((l for l in absn.findall(qn("w:lvl")) if l.get(qn("w:ilvl")) == str(ilvl)), None), over
        return None, over

    def ordered(self, num_id: str, ilvl: int) -> Optional[bool]:
        """True for numbered, False for bulleted, None when the level shows nothing (or is unknown)."""
        lvl, _ = self._find(num_id, ilvl)
        fmt = lvl.find(qn("w:numFmt")) if lvl is not None else None
        v = fmt.get(qn("w:val")) if fmt is not None else None
        if v is None or v == "none":
            return None
        return v != "bullet"

    def start(self, num_id: str, ilvl: int) -> int:
        lvl, over = self._find(num_id, ilvl)
        if over is not None and over.find(qn("w:startOverride")) is not None:
            return int(over.find(qn("w:startOverride")).get(qn("w:val")) or 1)
        s = lvl.find(qn("w:start")) if lvl is not None else None
        return int(s.get(qn("w:val")) or 1) if s is not None else 1

    def abstract_of(self, num_id: str) -> Optional[str]:
        if self.root is None:
            return None
        num = next((n for n in self.root.findall(qn("w:num")) if n.get(qn("w:numId")) == str(num_id)), None)
        a = num.find(qn("w:abstractNumId")) if num is not None else None
        return a.get(qn("w:val")) if a is not None else None


class Ctx:
    def __init__(self, doc, has_title: bool = False):
        self.doc = doc
        self.has_title = has_title
        self.styles = Styles(doc.styles.element)
        try:
            self.numbering = Numbering(doc.part.part_related_by(RT.NUMBERING).element)
        except KeyError:
            self.numbering = Numbering(None)
        self.warnings: List[str] = []

    def warn(self, text: str) -> None:
        if text not in self.warnings:
            self.warnings.append(text)

    def href(self, link) -> Optional[str]:
        rid = link.get(qn("r:id")) if link is not None else None
        rel = self.doc.part.rels.get(rid) if rid else None
        return rel.target_ref if rel is not None and rel.is_external else None

    def _numbering_root(self):
        return self.doc.part.numbering_part.numbering_definitions._numbering

    def new_num(self, abstract_id: str, start: int = 1) -> str:
        """A fresh numbering instance of an abstract list that counts from ``start``."""
        root = self._numbering_root()
        num = root.add_num(int(abstract_id))
        num.add_lvlOverride(ilvl=0).add_startOverride(max(1, start))
        self.numbering = Numbering(root)
        return str(num.numId)

    def ensure_list(self, ordered: bool) -> str:
        """A numbering instance for a list the file has no example of."""
        root = self._numbering_root()
        ids = [int(a.get(qn("w:abstractNumId"))) for a in root.findall(qn("w:abstractNum"))]
        abs_id = (max(ids) + 1) if ids else 1
        levels = "".join(
            f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{"decimal" if ordered else "bullet"}"/>'
            f'<w:lvlText w:val="{("%" + str(i + 1) + ".") if ordered else ("•" if i % 2 == 0 else "o")}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{720 + 360 * i}" w:hanging="360"/></w:pPr></w:lvl>' for i in range(9))
        absn = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abs_id}"><w:multiLevelType w:val="hybridMultilevel"/>{levels}</w:abstractNum>')
        first_num = root.find(qn("w:num"))
        (first_num.addprevious(absn) if first_num is not None else root.append(absn))
        num = root.add_num(abs_id)
        self.numbering = Numbering(root)
        return str(num.numId)

    def heading_style(self, md_level: int) -> Optional[str]:
        """The style id for a heading of Markdown level ``md_level``, adding the style when the file does not define it."""
        if self.has_title and md_level == 1:
            return self.styles.id_for("title")
        n = md_level - 1 if self.has_title else md_level
        n = max(1, min(n, 9))
        sid = self.styles.id_for(f"heading {n}")
        if sid:
            return sid
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt

        try:
            style = self.doc.styles.add_style(f"Heading {n}", WD_STYLE_TYPE.PARAGRAPH)
        except Exception:       # noqa: BLE001 - a style with that name exists under another id
            return None
        try:
            style.base_style = self.doc.styles[self.styles.name[self.styles.default_p]]
        except KeyError:
            pass
        style.font.bold, style.font.size = True, Pt({1: 16, 2: 14, 3: 12}.get(n, 11))
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if n < 3 else 8)
        ppr = style.element.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(n - 1))
        ppr.append(outline)
        self.styles = Styles(self.doc.styles.element)
        return style.style_id


# ── tracked changes ──────────────────────────────────────────────────────────
_TRACKED = ("w:ins", "w:del", "w:moveFrom", "w:moveTo", "w:rPrChange", "w:pPrChange")
_RANGE_MARKS = ("w:moveFromRangeStart", "w:moveFromRangeEnd", "w:moveToRangeStart", "w:moveToRangeEnd")
_TABLE_CHANGES = ("w:sectPrChange", "w:tblPrChange", "w:trPrChange", "w:tcPrChange", "w:tblGridChange")


def count_tracked(body) -> int:
    return sum(len(list(body.iter(qn(t)))) for t in _TRACKED)


def _join_into_next(p) -> None:
    """A paragraph whose mark is gone: what it holds joins the paragraph after it (an empty one just goes)."""
    has_content = any(True for _ in p.iter(qn("w:t"))) or any(True for _ in p.iter(qn("w:drawing")))
    nxt = p.getnext()
    if not has_content:
        p.getparent().remove(p)
    elif nxt is not None and _local(nxt) == "p":
        anchor = nxt.find(qn("w:pPr"))
        for c in reversed([c for c in p if _local(c) != "pPr"]):
            (anchor.addnext(c) if anchor is not None else nxt.insert(0, c))
        p.getparent().remove(p)


def _revert_props(change, kind: str) -> None:
    holder = change.getparent()
    old = change.find(qn("w:rPr" if kind == "r" else "w:pPr"))
    if old is not None and holder is not None:
        keep = {qn("w:rPr"), qn("w:sectPr")} if kind == "p" else set()
        for c in list(holder):
            if c is not change and c.tag not in keep:
                holder.remove(c)
        for c in reversed(list(old)):
            holder.insert(0, c)
    if holder is not None:
        holder.remove(change)


def _drop_marks_and_changes(body) -> None:
    for name in _RANGE_MARKS:
        for el in list(body.iter(qn(name))):
            el.getparent().remove(el)
    for name in _TABLE_CHANGES:
        for el in list(body.iter(qn(name))):
            el.getparent().remove(el)


def accept_all(body) -> None:
    """Every tracked change accepted: insertions stay, deletions go, moved text is where it was moved to."""
    marked = []
    for el in list(body.iter(qn("w:del"), qn("w:moveFrom"))):
        parent = el.getparent()
        if parent is None:
            continue
        if _local(parent) == "trPr":                                         # a deleted table row
            tr = parent.getparent()
            tr.getparent().remove(tr)
        elif _local(parent) == "rPr" and _local(parent.getparent()) == "pPr":  # a deleted paragraph mark
            marked.append(parent.getparent().getparent())
            parent.remove(el)
        else:
            parent.remove(el)
    for el in list(body.iter(qn("w:ins"), qn("w:moveTo"))):
        parent = el.getparent()
        if parent is None:
            continue
        (parent.remove(el) if _local(parent) in ("trPr", "rPr") else _unwrap(el))
    for p in marked:
        if p.getparent() is not None:
            _join_into_next(p)
    for el in list(body.iter(qn("w:rPrChange"), qn("w:pPrChange"))):
        el.getparent().remove(el)
    _drop_marks_and_changes(body)


def reject_all(body) -> None:
    """Every tracked change rejected: the text as it was before the reviewers."""
    marked = []
    for el in list(body.iter(qn("w:ins"), qn("w:moveTo"))):
        parent = el.getparent()
        if parent is None:
            continue
        if _local(parent) == "trPr":
            tr = parent.getparent()
            tr.getparent().remove(tr)
        elif _local(parent) == "rPr" and _local(parent.getparent()) == "pPr":  # an inserted paragraph mark
            marked.append(parent.getparent().getparent())
            parent.remove(el)
        else:
            parent.remove(el)
    for el in list(body.iter(qn("w:del"), qn("w:moveFrom"))):
        parent = el.getparent()
        if parent is None:
            continue
        if _local(parent) in ("trPr", "rPr"):
            parent.remove(el)
        else:
            for dt in el.iter(qn("w:delText")):
                dt.tag = qn("w:t")
            for dt in el.iter(qn("w:delInstrText")):
                dt.tag = qn("w:instrText")
            _unwrap(el)
    for p in marked:
        if p.getparent() is not None:
            _join_into_next(p)
    for el in list(body.iter(qn("w:rPrChange"))):
        _revert_props(el, "r")
    for el in list(body.iter(qn("w:pPrChange"))):
        _revert_props(el, "p")
    _drop_marks_and_changes(body)


# ── reading a paragraph ──────────────────────────────────────────────────────
@dataclass
class Piece:
    kind: str                 # t | tab | br
    text: str
    run: Any
    el: Any
    link: Any = None


@dataclass
class Para:
    el: Any
    pieces: List[Piece]
    text: str
    style_id: str
    base_b: bool = False
    base_i: bool = False


# the characters symbol fonts (Symbol, Wingdings) keep in the private-use area, as the Unicode characters they draw
SYMBOL_GLYPHS = {"": "•", "": "▪", "": "➢", "": "❖", "": "✓", "": "□", "": "■", "": "➢",
                 "o": "◦", "§": "▪", "": "•", "": "•", "": "✗", "": "☑", "": "➔", "": "⇨", "": "•",
                 "": "❑"}


def _symbol(c) -> str:
    """The character a ``w:sym`` draws (its code in the symbol font, F0xx or 00xx), as Unicode where it has a twin."""
    try:
        code = int(c.get(qn("w:char")) or "", 16)
    except ValueError:
        return "□"
    ch = chr(code + 0xF000 if code < 0x100 else code)
    return SYMBOL_GLYPHS.get(ch, ch)


def read_para(ctx: Ctx, p) -> Para:
    """What a paragraph shows, piece by piece, as Word shows it: its text, tabs, line breaks and symbols, and the results of its
    fields (a table of contents, page numbers, document properties, cross-references: an edit there goes into the field's
    result, which Word may refresh). Field codes, hidden text, images, footnote marks and other things that are not shown text
    are skipped here and left alone when the paragraph is edited."""
    pieces: List[Piece] = []
    fields: List[Dict[str, Any]] = []
    ppr = p.find(qn("w:pPr"))
    ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
    sid = ps.get(qn("w:val")) if ps is not None else ctx.styles.default_p
    hidden_style = bool(ctx.styles.flag(sid, "vanish"))

    def run(r, link) -> None:
        hidden = _style_flag(ctx, r, "vanish", hidden_style)
        for c in r:
            t = _local(c)
            if t == "fldChar":
                kind = c.get(qn("w:fldCharType"))
                if kind == "begin":
                    fields.append({"instr": "", "phase": "instr"})
                elif kind == "separate" and fields:
                    fields[-1]["phase"] = "result"
                elif kind == "end" and fields:
                    fields.pop()
                continue
            if t == "instrText":
                if fields and fields[-1]["phase"] == "instr":
                    fields[-1]["instr"] += c.text or ""
                continue
            if hidden or any(f["phase"] == "instr" for f in fields):
                continue
            if t == "t":
                pieces.append(Piece("t", c.text or "", r, c, link))
            elif t == "tab":
                pieces.append(Piece("tab", "\t", r, c, link))
            elif t in ("br", "cr"):                  # a line, page or column break: what follows starts on a new line
                pieces.append(Piece("br", "\n", r, c, link))
            elif t == "noBreakHyphen":
                pieces.append(Piece("t", "‑", r, c, link))
            elif t == "sym":
                pieces.append(Piece("sym", _symbol(c), r, c, link))

    def visit(container, link) -> None:
        for ch in container:
            t = _local(ch)
            if t == "r":
                run(ch, link)
            elif t == "hyperlink":
                visit(ch, ch)
            elif t in ("smartTag", "customXml", "ins", "moveTo", "dir", "bdo", "fldSimple"):
                visit(ch, link)
            elif t == "sdt":
                content = ch.find(qn("w:sdtContent"))
                if content is not None:
                    visit(content, link)

    visit(p, None)
    return Para(p, pieces, "".join(pc.text for pc in pieces), sid, bool(ctx.styles.flag(sid, "b")), bool(ctx.styles.flag(sid, "i")))


def _direct(r, name: str) -> bool:
    rpr = r.find(qn("w:rPr"))
    e = rpr.find(qn(f"w:{name}")) if rpr is not None else None
    if e is None:
        return False
    return e.get(qn("w:val")) != "none" if name == "u" else _on(e)


def _style_flag(ctx: Ctx, r, name: str, base: bool) -> bool:
    rpr = r.find(qn("w:rPr"))
    if rpr is not None:
        e = rpr.find(qn(f"w:{name}"))
        if e is not None:
            return _on(e)
        rs = rpr.find(qn("w:rStyle"))
        if rs is not None:
            v = ctx.styles.flag(rs.get(qn("w:val")), name)
            if v is not None:
                return v
    return base


def piece_fmt(ctx: Ctx, para: Para, pc: Piece) -> Fmt:
    """The formatting of a piece as the Markdown shows it: bold and italic that the paragraph's own style supplies are not
    shown (a heading is not ``**bold**``); underline and strike are those set on the text itself."""
    b = _style_flag(ctx, pc.run, "b", para.base_b) and not para.base_b
    i = _style_flag(ctx, pc.run, "i", para.base_i) and not para.base_i
    underline = _direct(pc.run, "u") and pc.link is None                   # a link is underlined because it is a link
    return (b, i, underline, _direct(pc.run, "strike") or _direct(pc.run, "dstrike"))


def para_segs(ctx: Ctx, para: Para) -> List[B.Seg]:
    segs = []
    for pc in para.pieces:
        b, i, u, s = piece_fmt(ctx, para, pc)
        segs.append(B.Seg(pc.text, b, i, u, s, False, ctx.href(pc.link) if pc.link is not None else None))
    return B._merge(segs)


def char_fmts(ctx: Ctx, para: Para) -> List[Fmt]:
    out: List[Fmt] = []
    for pc in para.pieces:
        out.extend([piece_fmt(ctx, para, pc)] * len(pc.text))
    return out


# ── Markdown for a paragraph ─────────────────────────────────────────────────
_LINE_START = ((re.compile(r"^(\s*)(#{1,6}\s)"), r"\1\\\2"), (re.compile(r"^(\s*)([-+]\s)"), r"\1\\\2"), (re.compile(r"^(\s*\d+)([.)]\s)"), r"\1\\\2"))


def _protect_lines(md: str) -> str:
    """Escape what would make a line start a heading, a bullet or a numbered item."""
    out = []
    for line in md.split("\n"):
        for rx, repl in _LINE_START:
            if rx.match(line):
                line = rx.sub(repl, line, count=1)
                break
        out.append(line)
    return "\n".join(out)


def _punct(ch: str) -> bool:
    """A punctuation mark or a symbol, as CommonMark counts them."""
    return bool(ch) and unicodedata.category(ch)[0] in "PS"


def _apart(ch: str) -> bool:
    """Nothing, a space or punctuation: a ``*`` next to it opens or closes emphasis whatever is on its other side."""
    return not ch or ch.isspace() or _punct(ch)


def _mark(text: str, seg: B.Seg, before: str = "", after: str = "") -> str:
    """``text`` with the Markdown of ``seg``'s formatting. ``before`` and ``after`` are the characters next to it: a ``*`` or ``~~``
    between a letter and a punctuation mark cannot open or close emphasis (it would read as a literal star), so punctuation at the
    edge of bold, italic or struck text goes outside the markers when a letter touches it."""
    m = re.match(r"^(\s*)(.*?)(\s*)$", text, re.S)
    lead, core, trail = m.group(1), m.group(2), m.group(3)
    if not core:
        return text
    head = tail = ""
    if (seg.bold or seg.italic or seg.strike) and not (seg.underline or seg.href):      # the markers are the outermost ones
        if not lead and not _apart(before) and _punct(core[0]):
            k = 0
            while k < len(core) and (_punct(core[k]) or core[k].isspace()):
                k += 1
            head, core = core[:k], core[k:]
        if core and not trail and not _apart(after) and _punct(core[-1]):
            k = len(core)
            while k > 0 and (_punct(core[k - 1]) or core[k - 1].isspace()):
                k -= 1
            core, tail = core[:k], core[k:]
        if not core:
            return lead + _ESC.sub(r"\\\1", head + tail) + trail
    out = _ESC.sub(r"\\\1", core)
    if seg.italic:
        out = f"*{out}*"
    if seg.bold:
        out = f"**{out}**"
    if seg.strike:
        out = f"~~{out}~~"
    if seg.underline:
        out = f"<u>{out}</u>"
    if seg.href:
        url = seg.href.replace("<", "%3C").replace(">", "%3E")
        out = f"[{out}](<{url}>)" if re.search(r"[\s()]", url) else f"[{out}]({url})"
    return lead + _ESC.sub(r"\\\1", head) + out + _ESC.sub(r"\\\1", tail) + trail


def segs_markdown(segs: Sequence[B.Seg], *, cell: bool = False, protect: bool = True) -> str:
    """Markdown for formatted text. A line break is a hard break (or ``<br>`` in a table cell); spaces and breaks at the very
    start or end are left out (Markdown cannot hold them). ``protect`` escapes a line that would start a heading, bullet or
    numbered item (not needed after a ``## `` marker or inside a cell)."""
    segs = [s for s in segs if s.text]
    while segs and not segs[0].text.strip():
        segs.pop(0)
    while segs and not segs[-1].text.strip():
        segs.pop()
    if segs:
        segs[0] = B.Seg(segs[0].text.lstrip(), *segs[0].flags(), href=segs[0].href)
        segs[-1] = B.Seg(segs[-1].text.rstrip(), *segs[-1].flags(), href=segs[-1].href)
    parts = []
    for k, s in enumerate(segs):
        before = segs[k - 1].text[-1:] if k else ""
        after = segs[k + 1].text[:1] if k + 1 < len(segs) else ""
        pieces = s.text.split("\n")
        parts.append(("<br>" if cell else "\\\n").join(
            _mark(p, s, before if j == 0 else "", after if j == len(pieces) - 1 else "") if p.strip() else p for j, p in enumerate(pieces)))
    md = "".join(parts)
    return md if cell or not protect else _protect_lines(md)


# ── walking the document ─────────────────────────────────────────────────────
@dataclass
class CellRef:
    tc: Any                                   # the w:tc (None where a merged cell covers this column)
    paras: List[Para]                         # the paragraphs with text
    text: str = ""                            # the cell as Markdown
    plain: str = ""                           # the cell's words, a line for each paragraph


@dataclass
class TableInfo:
    el: Any
    rows: List[List[CellRef]]
    trs: List[Any]
    complex: bool = False

    def plain(self) -> List[List[str]]:
        return [[c.plain for c in row] for row in self.rows]


@dataclass
class Node:
    kind: str                                 # heading | paragraph | list_item | table
    el: Any
    md: str
    text: str = ""
    level: int = 0
    ordered: bool = False
    number: Optional[int] = None
    para: Optional[Para] = None
    table: Optional[TableInfo] = None
    num_id: Optional[str] = None
    ilvl: int = 0                             # a list item's level in Word (``level`` is its nesting in the Markdown, counted from its list)
    tpl_ppr: Any = None                       # copies taken before any edit, to give new paragraphs the look of this one
    tpl_rpr: Any = None
    generated: bool = False                   # Word makes it (a table of contents) or draws it (a text box): read, not rewritten by the AI


def _top_blocks(parent) -> Iterator[Any]:
    for ch in parent:
        t = _local(ch)
        if t in ("p", "tbl"):
            yield ch
        elif t == "sdt":
            content = ch.find(qn("w:sdtContent"))
            if content is not None:
                yield from _top_blocks(content)
        elif t == "customXml":
            yield from _top_blocks(ch)


def _box_paras(p) -> Iterator[Any]:
    """The paragraphs of the text boxes drawn in paragraph ``p``, each once: a box drawn for both new and old versions of Word is in
    the file twice, and the copy for old versions (``mc:Fallback``) is skipped."""
    for box in p.iter(qn("w:txbxContent")):
        if any(_local(a) == "Fallback" for a in box.iterancestors()):
            continue
        yield from box.findall(qn("w:p"))


_DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
_DML = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _diagrams(ctx: Ctx, p) -> Iterator[Tuple[Any, List[str]]]:
    """The SmartArt diagrams drawn in paragraph ``p``: the run holding each, and the text of its shapes in reading order (a shape,
    then the shapes under it). The text is in the diagram's own part of the file, not in the paragraph."""
    for rel in p.iter(f"{_DGM}relIds"):
        run = next(rel.iterancestors(qn("w:r")), None)
        part = ctx.doc.part.related_parts.get(rel.get(f"{_REL}dm"))
        if run is None or part is None:
            continue
        try:
            root = parse_xml(part.blob)
        except Exception:       # noqa: BLE001 - a damaged diagram is left to Word to draw
            continue
        texts: Dict[str, List[str]] = {}
        top = None
        for pt in root.iter(f"{_DGM}pt"):
            kind = pt.get("type") or "node"
            if kind == "doc":
                top = pt.get("modelId")
            elif kind in ("node", "asst"):
                t = pt.find(f"{_DGM}t")
                lines = ("".join(r.text or "" for r in ap.iter(f"{_DML}t")).strip() for ap in (t.findall(f"{_DML}p") if t is not None else []))
                texts[pt.get("modelId")] = [line for line in lines if line]
        under: Dict[str, List[Tuple[int, str]]] = {}
        for c in root.iter(f"{_DGM}cxn"):
            if (c.get("type") or "parOf") == "parOf":
                order = c.get("srcOrd") or "0"
                under.setdefault(c.get("srcId"), []).append((int(order) if order.isdigit() else 0, c.get("destId")))
        out: List[str] = []
        seen: set = set()
        stack = [top]
        while stack:
            mid = stack.pop()
            if mid in seen:
                continue
            seen.add(mid)
            out += texts.get(mid, [])
            stack += [child for _, child in sorted(under.get(mid, []), reverse=True)]
        out += [line for mid, lines in texts.items() if mid not in seen for line in lines]
        if out:
            yield run, out


def _outer(el):
    """The body paragraph or table ``el`` is part of (a text box's paragraph and a diagram are drawn inside one): new blocks go
    next to that, never into the box."""
    for a in reversed(list(el.iterancestors())):
        if _local(a) in ("p", "tbl"):
            return a
    return el


def _generated(ctx: Ctx, el, para: Para) -> bool:
    """A paragraph Word writes itself: a table of contents or of figures (its entries, its heading, anything in its content control)."""
    name = ctx.styles.lname(para.style_id)
    if name.startswith("toc") or name.startswith("table of figures"):
        return True
    for sdt in el.iterancestors(qn("w:sdt")):
        gal = sdt.find(qn("w:sdtPr") + "/" + qn("w:docPartObj") + "/" + qn("w:docPartGallery"))
        if gal is not None and (gal.get(qn("w:val")) or "").lower().startswith("table of contents"):
            return True
    return False


def _wrapped(parent, name: str) -> Iterator[Any]:
    """The ``name`` children of ``parent``, also those inside content controls and custom XML (a repeating section's rows)."""
    for ch in parent:
        t = _local(ch)
        if t == name:
            yield ch
        elif t == "sdt":
            content = ch.find(qn("w:sdtContent"))
            if content is not None:
                yield from _wrapped(content, name)
        elif t == "customXml":
            yield from _wrapped(ch, name)


def _cell_paras(tc) -> Iterator[Any]:
    """A cell's paragraphs in reading order: its own and its content controls', its text boxes', and those of a table nested in it,
    row by row."""
    for ch in tc:
        t = _local(ch)
        if t == "p":
            yield ch
            yield from _box_paras(ch)
        elif t == "tbl":
            for tr in _wrapped(ch, "tr"):
                for inner in _wrapped(tr, "tc"):
                    yield from _cell_paras(inner)
        elif t == "sdt":
            content = ch.find(qn("w:sdtContent"))
            if content is not None:
                yield from _cell_paras(content)
        elif t == "customXml":
            yield from _cell_paras(ch)


def _span(tc) -> int:
    g = tc.find(qn("w:tcPr") + "/" + qn("w:gridSpan"))
    try:
        return max(1, int(g.get(qn("w:val")))) if g is not None else 1
    except (TypeError, ValueError):
        return 1


def _vcontinue(tc) -> bool:
    v = tc.find(qn("w:tcPr") + "/" + qn("w:vMerge"))
    return v is not None and v.get(qn("w:val")) in (None, "continue")


def _read_table(ctx: Ctx, tbl) -> Optional[TableInfo]:
    trs = list(_wrapped(tbl, "tr"))
    rows: List[List[CellRef]] = []
    complex_ = False
    for tr in trs:
        row: List[CellRef] = []
        pr = tr.find(qn("w:trPr"))
        before = pr.find(qn("w:gridBefore")) if pr is not None else None
        if before is not None:
            row += [CellRef(None, []) for _ in range(int(before.get(qn("w:val")) or 0))]
            complex_ = True
        for tc in _wrapped(tr, "tc"):
            span = _span(tc)
            if span > 1 or tc.find(qn("w:tcPr") + "/" + qn("w:vMerge")) is not None:
                complex_ = True
            if _vcontinue(tc):
                row += [CellRef(None, []) for _ in range(span)]
                continue
            paras = [p for p in (read_para(ctx, p) for p in _cell_paras(tc)) if p.text.strip()]
            row.append(CellRef(tc, paras, "<br>".join(segs_markdown(para_segs(ctx, p), cell=True) for p in paras),
                               "\n".join(p.text.strip() for p in paras)))
            row += [CellRef(None, []) for _ in range(span - 1)]
        if tr.find(".//" + qn("w:tbl")) is not None:            # a nested table reads as lines of its cell; its grid is not rebuilt
            complex_ = True
        rows.append(row)
    if not rows:
        return None
    width = max(len(r) for r in rows)
    if len({len(r) for r in rows}) > 1:
        complex_ = True
    for r in rows:
        r.extend(CellRef(None, []) for _ in range(width - len(r)))
    if all(not c.text for r in rows for c in r):
        return None
    return TableInfo(tbl, rows, trs, complex_)


def _table_md(info: TableInfo) -> str:
    lines = ["| " + " | ".join(c.text for c in row) + " |" for row in info.rows]
    lines.insert(1, "| " + " | ".join(["---"] * len(info.rows[0])) + " |")
    return "\n".join(lines)


def _heading_level(ctx: Ctx, para: Para) -> Optional[int]:
    """0 for a Title, else the heading number 1-9; None when the paragraph is not a heading."""
    name = ctx.styles.lname(para.style_id)
    if name == "title":
        return 0
    m = _HEADING.match(name)
    if m:
        return int(m.group(1))
    if name.startswith("toc") or name.startswith("table of"):
        return None
    ppr = para.el.find(qn("w:pPr"))
    direct = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
    if direct is not None and (direct.get(qn("w:val")) or "").isdigit():
        lvl = int(direct.get(qn("w:val")))
    else:
        lvl = ctx.styles.outline(para.style_id)
    return lvl + 1 if lvl is not None and lvl < 9 else None


def _list_info(ctx: Ctx, para: Para) -> Tuple[Optional[str], int]:
    ppr = para.el.find(qn("w:pPr"))
    np_ = ppr.find(qn("w:numPr")) if ppr is not None else None
    num = ilvl = None
    if np_ is not None:
        n, l = np_.find(qn("w:numId")), np_.find(qn("w:ilvl"))
        num = n.get(qn("w:val")) if n is not None else None
        ilvl = int(l.get(qn("w:val")) or 0) if l is not None else None
    s_num, s_lvl = ctx.styles.numpr(para.style_id)
    num = num if num is not None else s_num
    ilvl = ilvl if ilvl is not None else s_lvl
    return (None if num in (None, "0") else num), ilvl or 0


def walk(doc, ctx: Optional[Ctx] = None) -> Tuple[List[Node], Dict[str, Any]]:
    """The blocks of a document in order, each tied to its element, and a few facts about the file."""
    ctx = ctx or Ctx(doc)
    raw: List[Any] = []                       # (element, paragraph, table, generated) or a diagram's finished node
    for el in _top_blocks(doc.element.body):
        if _local(el) == "tbl":
            info = _read_table(ctx, el)
            if info is not None:
                raw.append((el, None, info, False))
            continue
        para = read_para(ctx, el)
        if para.text.strip():
            raw.append((el, para, None, _generated(ctx, el, para)))
        for box in _box_paras(el):                                     # the text in its text boxes, after it
            inside = read_para(ctx, box)
            if inside.text.strip():
                raw.append((box, inside, None, True))
        for drawn, lines in _diagrams(ctx, el):                        # and in its SmartArt diagrams, a line for each shape
            text = "\n".join(lines)
            raw.append(Node("paragraph", drawn, segs_markdown([B.Seg(text)]), text, generated=True))
    has_title = any(isinstance(e, tuple) and e[1] is not None and ctx.styles.lname(e[1].style_id) == "title" for e in raw)
    ctx.has_title = has_title
    nodes: List[Node] = []
    counters: Dict[Tuple[str, int], int] = {}
    run: List[int] = []                       # the Word levels of the list being read, one per Markdown nesting level
    auto_numbered = False
    for entry in raw:
        if isinstance(entry, Node):
            nodes.append(entry)
            continue
        el, para, info, generated = entry
        if info is not None:
            nodes.append(Node("table", el, _table_md(info), "\n".join(" ".join(r) for r in info.plain()), table=info))
            continue
        if generated:                                                  # as it reads: never a heading or a list of the document's own
            nodes.append(Node("paragraph", el, segs_markdown(para_segs(ctx, para)), para.text.strip(), para=para, generated=True))
            continue
        segs = para_segs(ctx, para)
        ppr = el.find(qn("w:pPr"))
        level = _heading_level(ctx, para)
        num_id, ilvl = _list_info(ctx, para)
        ordered = ctx.numbering.ordered(num_id, ilvl) if num_id is not None else None
        if level is not None:
            md_level = 1 if level == 0 else min(6, level + (1 if has_title else 0))
            body = segs_markdown(segs, protect=False).replace("\\\n", " ")
            trailing = re.search(r"#+$", body)
            if trailing:
                body = body[:trailing.start()] + "\\#" * len(trailing.group())
            node = Node("heading", el, "#" * md_level + " " + body, para.text.strip(), level=md_level, para=para)
            auto_numbered = auto_numbered or num_id is not None
        elif ordered is not None:
            number = None
            if ordered:
                key = (num_id, ilvl)
                counters[key] = counters.get(key, ctx.numbering.start(num_id, ilvl) - 1) + 1
                for k in [k for k in counters if k[0] == num_id and k[1] > ilvl]:
                    del counters[k]
                number = counters[key]
            # Markdown nests a list from its first item, one level at a time: a list that starts at Word's third level, under a
            # heading, is the first level here (indented, it would read back as code)
            if not nodes or nodes[-1].kind != "list_item":
                run = [ilvl]
            else:
                while len(run) > 1 and ilvl < run[-1]:
                    run.pop()
                if ilvl > run[-1] and len(run) < 6:
                    run.append(ilvl)
                elif ilvl < run[-1]:
                    run[-1] = ilvl
            depth = len(run) - 1
            node = Node("list_item", el, "    " * depth + (f"{number}. " if ordered else "- ") + segs_markdown(segs), para.text.strip(),
                        level=depth, ordered=ordered, number=number, para=para, num_id=num_id, ilvl=ilvl)
        else:
            node = Node("paragraph", el, segs_markdown(segs), para.text.strip(), para=para)
        node.tpl_ppr = copy.deepcopy(ppr) if ppr is not None else None
        first = next((pc.run for pc in para.pieces if pc.kind == "t"), None)
        rpr = first.find(qn("w:rPr")) if first is not None else None
        node.tpl_rpr = copy.deepcopy(rpr) if rpr is not None else None
        nodes.append(node)
    return nodes, {"has_title": has_title, "auto_numbered": auto_numbered}


def render_nodes(nodes: Sequence[Node]) -> str:
    out: List[str] = []
    prev: Optional[Node] = None
    for n in nodes:
        if prev is not None:
            out.append("\n" if prev.kind == "list_item" and n.kind == "list_item" else "\n\n")
        out.append(n.md)
        prev = n
    return "".join(out) + ("\n" if out else "")


def _cells_norm(rows: Sequence[Sequence[str]]) -> List[List[str]]:
    return [[B.normalize(c) for c in row] for row in rows]


def _block_problem(i: int, node: Node, kind: str, b: B.Block) -> Optional[str]:
    if node.kind != kind:
        return f"block {i + 1} should be a {kind} but is a {node.kind}"
    if kind == "heading" and node.level != b.level:
        return f"block {i + 1}: heading level {node.level} instead of {b.level}"
    if kind == "list_item" and (node.level != b.level or node.ordered != b.ordered):
        return f"block {i + 1}: the list item differs"
    if kind == "table":
        return None if _cells_norm(node.table.plain()) == _cells_norm([[c.text for c in r] for r in b.rows]) else f"block {i + 1}: the table cells differ"
    if B.normalize(node.text) == B.normalize(b.text):
        return None
    have, want = B.normalize(node.text), B.normalize(b.text)
    at = next((k for k, (x, y) in enumerate(zip(have, want)) if x != y), min(len(have), len(want)))
    return f"block {i + 1} reads …{have[max(0, at - 20):at + 30]!r} instead of …{want[max(0, at - 20):at + 30]!r}"


def check_blocks(nodes: Sequence[Node], blocks: Sequence[B.Block], *, lenient: bool = False) -> Optional[str]:
    """None when the Markdown of ``nodes`` reads back as ``blocks``, else what differs. Leniently, rules are skipped and quote,
    code and unknown blocks count as paragraphs (that is what Word holds them as)."""
    want = [b for b in blocks if not (lenient and b.kind == "hr")]
    if len(nodes) != len(want):
        return f"{len(want)} blocks were expected and there are {len(nodes)}"
    for i, (n, b) in enumerate(zip(nodes, want)):
        kind = b.kind if not lenient or b.kind in ("heading", "list_item", "table") else "paragraph"
        problem = _block_problem(i, n, kind, b)
        if problem:
            return problem
    return None


# ── import ───────────────────────────────────────────────────────────────────
def import_docx(path: str, *, twin: bool = False) -> Imported:
    """The Markdown of a Word file. With ``twin`` the result also says how to write edits back into the file (``native``)."""
    try:
        doc = Document(path)
    except Exception as exc:       # noqa: BLE001 - python-docx raises many kinds of errors for a damaged file
        raise ImportProblem("This Word file could not be opened. It may be damaged, password protected, or an older .doc file saved with the wrong extension.") from exc
    notes: List[str] = []
    tracked = count_tracked(doc.element.body)
    base_md: Optional[str] = None
    if tracked:
        before = Document(path)
        reject_all(before.element.body)
        base_md = render_nodes(walk(before)[0])
        accept_all(doc.element.body)
        notes.append(f"The file has {tracked} tracked change{'s' if tracked != 1 else ''}. They are accepted in the working text; the Changes view shows them against the text before them.")
    nodes, meta = walk(doc)
    if not nodes:
        raise ImportProblem("No text was found in this Word file (it may be only images, or a scan). Upload a PDF of it to read it with text recognition.")
    md = render_nodes(nodes)
    problem = check_blocks(nodes, B.parse(md))
    native: Optional[Dict[str, Any]] = None
    if twin:
        if problem is None:
            twin_path = None
            if tracked:
                twin_path = re.sub(r"(\.docx)?$", ".twin.docx", path, count=1, flags=re.I)
                doc.save(twin_path)
            native = {"map": {"v": 1, "blocks": len(nodes), "sha": hashlib.sha1(md.encode("utf-8")).hexdigest(), "tracked": tracked,
                              "auto_numbered": meta["auto_numbered"], "has_title": meta["has_title"]}, "path": twin_path}
            notes.append("Edits are written into the original file, so its fonts, styles, header and footer stay.")
        else:
            notes.append("The formatting of this Word file cannot be written back (" + problem + "); a clean Word copy is made when you export.")
    if meta["auto_numbered"]:
        notes.append("Headings are numbered by Word itself, so their numbers are not part of the text.")
    return Imported(md, "docx", notes, base_md, native)


# ── writing edits into the file: text ────────────────────────────────────────
def _tokens(text: str) -> List[Tuple[str, int, int]]:
    return [(m.group(), m.start(), m.end()) for m in _TOKEN.finditer(text)]


def _key(tok: str) -> str:
    if tok.isspace():
        return "\n" if "\n" in tok else ("\t" if "\t" in tok else " ")
    return tok


def _trim(toks: List[Tuple[str, int, int]]) -> List[Tuple[str, int, int]]:
    lo, hi = 0, len(toks)
    while lo < hi and toks[lo][0].isspace():
        lo += 1
    while hi > lo and toks[hi - 1][0].isspace():
        hi -= 1
    return toks[lo:hi]


def diff_ops(old: str, new: str) -> List[Tuple[str, int, int, int, int]]:
    """``[(tag, old start, old end, new start, new end)]`` in character offsets. ``equal`` comes once per token (white space of
    any width is one token, so a double space is not an edit); other tags cover runs of tokens."""
    a, b = _trim(_tokens(old)), _trim(_tokens(new))
    if not a or not b:
        return [("replace", 0, len(old), 0, len(new))] if old != new else []
    sm = SequenceMatcher(None, [_key(t[0]) for t in a], [_key(t[0]) for t in b], autojunk=False)
    words_a, words_b = [t[0] for t in a if not t[0].isspace()], [t[0] for t in b if not t[0].isspace()]
    if len(words_a) > 3 and SequenceMatcher(None, words_a, words_b, autojunk=False).ratio() < 0.3:   # few words in common: one clean replacement
        return [("replace", a[0][1], a[-1][2], b[0][1], b[-1][2])]
    ops: List[Tuple[str, int, int, int, int]] = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            ops += [("equal", a[i1 + n][1], a[i1 + n][2], b[j1 + n][1], b[j1 + n][2]) for n in range(i2 - i1)]
            continue
        a1, a2 = (a[i1][1], a[i2 - 1][2]) if i2 > i1 else (a[i1][1] if i1 < len(a) else a[-1][2],) * 2
        b1, b2 = (b[j1][1], b[j2 - 1][2]) if j2 > j1 else (b[j1][1] if j1 < len(b) else b[-1][2],) * 2
        ops.append(("replace" if i2 > i1 and j2 > j1 else ("delete" if i2 > i1 else "insert"), a1, a2, b1, b2))
    return ops


def _set_text(t, text: str) -> None:
    t.text = text
    t.set(XML_SPACE, "preserve")


def _content(r) -> List[Any]:
    return [c for c in r if isinstance(c.tag, str) and _local(c) != "rPr"]


def _new_run(rpr, child) -> Any:
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    r.append(child)
    return r


def _explode_run(r) -> None:
    """A run holding several things (text, a tab, an image ...) becomes one run for each, with the same properties."""
    kids = _content(r)
    if len(kids) <= 1:
        return
    rpr = r.find(qn("w:rPr"))
    prev = r
    for child in kids[1:]:
        nr = _new_run(rpr, child)
        prev.addnext(nr)
        prev = nr


def _split_run_text(pc: Piece, k: int) -> None:
    text = pc.el.text or ""
    tail = copy.deepcopy(pc.run)
    _set_text(pc.el, text[:k])
    _set_text(next(c for c in _content(tail) if _local(c) == "t"), text[k:])
    pc.run.addnext(tail)


def _pieces(ctx: Ctx, p) -> List[Piece]:
    return read_para(ctx, p).pieces


def _isolate(ctx: Ctx, p, a: int, b: int) -> List[Piece]:
    """Cut the runs so that ``[a, b)`` is made of whole runs (one piece each); returns those pieces."""
    pos = 0
    for pc in _pieces(ctx, p):
        s, e = pos, pos + len(pc.text)
        pos = e
        if e >= a and s <= b:
            _explode_run(pc.run)
    for boundary in sorted({a, b}, reverse=True):
        pos = 0
        for pc in _pieces(ctx, p):
            s, e = pos, pos + len(pc.text)
            pos = e
            if pc.kind == "t" and s < boundary < e:
                _split_run_text(pc, boundary - s)
                break
    out, pos = [], 0
    for pc in _pieces(ctx, p):
        s, e = pos, pos + len(pc.text)
        pos = e
        if s >= a and e <= b and e > s:
            out.append(pc)
    return out


def _delete(ctx: Ctx, p, a: int, b: int) -> None:
    for pc in _isolate(ctx, p, a, b):
        parent = pc.run.getparent()
        parent.remove(pc.run)
        if _local(parent) == "hyperlink" and not any(_local(c) == "r" for c in parent):
            parent.getparent().remove(parent)


def _flag_run(run_el, want: Fmt, have: Fmt) -> None:
    f = Run(run_el, None).font
    if want[0] != have[0]:
        f.bold = want[0]
    if want[1] != have[1]:
        f.italic = want[1]
    if want[2] != have[2]:
        f.underline = want[2]
    if want[3] != have[3]:
        f.strike = want[3]


def _atoms(text: str) -> List[Tuple[str, str]]:
    """``text`` as ``('t', text)``, ``('tab', '')`` and ``('br', '')`` parts."""
    out: List[Tuple[str, str]] = []
    for part in re.split(r"(\n|\t)", text):
        if part == "\n":
            out.append(("br", ""))
        elif part == "\t":
            out.append(("tab", ""))
        elif part:
            out.append(("t", part))
    return out


def _highlighted(rpr):
    from docx.enum.text import WD_COLOR_INDEX

    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    Run(r, None).font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r.find(qn("w:rPr"))


def _base_rpr(rpr):
    """A copy of a run's properties for new text: no link style, no highlight."""
    out = copy.deepcopy(rpr) if rpr is not None else None
    if out is not None:
        for e in out.findall(qn("w:rStyle")):
            if (e.get(qn("w:val")) or "").lower() == "hyperlink":
                out.remove(e)
        for e in out.findall(qn("w:highlight")):
            out.remove(e)
    return out


def _style_link(ctx: Ctx, r) -> None:
    sid = ctx.styles.id_for("Hyperlink")
    if sid:
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        rs = OxmlElement("w:rStyle")
        rs.set(qn("w:val"), sid)
        rpr.insert(0, rs)
        return
    from docx.shared import RGBColor

    f = Run(r, None).font
    f.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    f.underline = True


def _make_runs(ctx: Ctx, groups: Sequence[Tuple[str, NewFmt]], rpr, host: Fmt, *, wrap_links: bool = True) -> List[Any]:
    """Run elements for new text. ``groups`` is ``[(text, (bold, italic, underline, strike, href))]``; ``rpr`` is the look of
    the text around it and ``host`` the formatting that look stands for in Markdown: only differences are applied. Text with a
    link gets a new link, unless ``wrap_links`` is off (it is going into a link that is already there)."""
    from .exporters import _split_placeholders

    out: List[Any] = []
    for text, fmt in groups:
        wrap = None
        link = bool(fmt[4]) and wrap_links
        if link:
            wrap = OxmlElement("w:hyperlink")
            wrap.set(qn("r:id"), ctx.doc.part.relate_to(fmt[4], RT.HYPERLINK, is_external=True))
        for piece, is_placeholder in _split_placeholders(text):
            for kind, value in _atoms(piece):
                child = OxmlElement({"t": "w:t", "tab": "w:tab", "br": "w:br"}[kind])
                if kind == "t":
                    _set_text(child, value)
                r = _new_run(_highlighted(rpr) if is_placeholder else rpr, child)
                _flag_run(r, fmt[:4], host)
                if link:
                    _style_link(ctx, r)
                    wrap.append(r)
                else:
                    out.append(r)
        if wrap is not None:
            out.append(wrap)
    return out


def _insert(ctx: Ctx, p, at: int, groups: Sequence[Tuple[str, NewFmt]], rpr, host: Fmt) -> None:
    """New text at character offset ``at``: after the run before it (inside a link only when it stays in that link)."""
    _isolate(ctx, p, at, at)
    prev = nxt = None
    pos = 0
    for pc in _pieces(ctx, p):
        s, e = pos, pos + len(pc.text)
        pos = e
        if e <= at and e > s:
            prev = pc
        elif s >= at and nxt is None and e > s:
            nxt = pc
    hrefs = {g[1][4] for g in groups}
    into = None                                                    # text with the link of a neighbouring link goes inside that link
    if len(hrefs) == 1 and next(iter(hrefs)):
        into = next((c for c in (prev.link if prev else None, nxt.link if nxt else None) if c is not None and ctx.href(c) == next(iter(hrefs))), None)
    if into is not None:
        base = prev if prev is not None and prev.link is into else nxt
        new = _make_runs(ctx, groups, copy.deepcopy(base.run.find(qn("w:rPr"))), host, wrap_links=False)
        anchor = base.run
        for el in new:
            (anchor.addnext if base is prev else anchor.addprevious)(el)
            anchor = el if base is prev else anchor
        return
    new = _make_runs(ctx, groups, rpr, host)
    if not new:
        return
    wants_link = any(g[1][4] for g in groups)
    if prev is not None:
        anchor = prev.run
        if prev.link is not None and (nxt is None or nxt.link is not prev.link or wants_link):
            anchor = prev.link
        for el in new:
            anchor.addnext(el)
            anchor = el
    elif nxt is not None:
        anchor = nxt.link if nxt.link is not None else nxt.run
        for el in new:
            anchor.addprevious(el)
    else:
        anchor = p.find(qn("w:pPr"))
        for el in new:
            (p.insert(0, el) if anchor is None else anchor.addnext(el))
            anchor = el


def _same_rpr(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return etree.tostring(a, method="c14n") == etree.tostring(b, method="c14n")


def _merge_runs(p) -> None:
    """Join neighbouring runs that look the same (editing splits runs; Word would join them, so do we)."""
    for container in [p] + list(p.iter(qn("w:hyperlink"))):
        prev = None
        for ch in list(container):
            if _local(ch) != "r":
                prev = None
                continue
            kids = _content(ch)
            simple = len(kids) == 1 and _local(kids[0]) == "t"
            if prev is not None and simple and _same_rpr(prev.find(qn("w:rPr")), ch.find(qn("w:rPr"))):
                a = _content(prev)[0]
                _set_text(a, (a.text or "") + (kids[0].text or ""))
                container.remove(ch)
                continue
            prev = ch if simple else None


def _fmts_for(text: str, segs: Sequence[B.Seg]) -> List[NewFmt]:
    out: List[NewFmt] = []
    for s in segs:
        out.extend([(s.bold, s.italic, s.underline, s.strike, s.href)] * len(s.text))
    return out if len(out) == len(text) else [NO_NEW] * len(text)


def _groups(text: str, fmts: Sequence[NewFmt]) -> List[Tuple[str, NewFmt]]:
    out: List[Tuple[str, NewFmt]] = []
    for ch, f in zip(text, fmts):
        if out and out[-1][1] == f:
            out[-1] = (out[-1][0] + ch, f)
        else:
            out.append((ch, f))
    return out


def _host_run(ctx: Ctx, p, offset: int):
    pos, last = 0, None
    for pc in _pieces(ctx, p):
        s, e = pos, pos + len(pc.text)
        pos = e
        if e > s:
            last = pc.run
            if s <= offset < e:
                return pc.run
    return last


def patch_paragraph(ctx: Ctx, para: Para, new_text: str, new_segs: Sequence[B.Seg]) -> None:
    """Make the paragraph say ``new_text`` with ``new_segs``' formatting, changing only what differs."""
    p = para.el
    old_text, old_fmt = para.text, char_fmts(ctx, para)
    new_fmt = _fmts_for(new_text, new_segs)
    for tag, a1, a2, b1, b2 in reversed(diff_ops(old_text, new_text)):
        if tag == "equal":                                                   # same words: only their formatting may differ
            whitespace = old_text[a1:a2].isspace()
            spans: List[Tuple[int, int, Fmt, Fmt]] = []
            for k in range(a1, a2):
                want, have = new_fmt[b1 if whitespace else b1 + (k - a1)][:4], old_fmt[k]
                if want == have:
                    continue
                if spans and spans[-1][1] == k and spans[-1][2:] == (have, want):
                    spans[-1] = (spans[-1][0], k + 1, have, want)
                else:
                    spans.append((k, k + 1, have, want))
            for lo, hi, have, want in reversed(spans):
                for pc in _isolate(ctx, p, lo, hi):
                    _flag_run(pc.run, want, have)
            continue
        if tag == "delete":
            _delete(ctx, p, a1, a2)
            continue
        groups = _groups(new_text[b1:b2], new_fmt[b1:b2])
        if tag == "insert":
            at = a1 - 1 if a1 > 0 else 0
            host = _host_run(ctx, p, at)
            _insert(ctx, p, a1, groups, _base_rpr(host.find(qn("w:rPr")) if host is not None else None), old_fmt[at] if old_fmt else NO_FMT)
        else:                                                                # replace: the new words look like the ones they replace
            host = _host_run(ctx, p, a1)
            rpr = _base_rpr(host.find(qn("w:rPr")) if host is not None else None)
            _delete(ctx, p, a1, a2)
            _insert(ctx, p, a1, groups, rpr, old_fmt[a1] if old_fmt else NO_FMT)
    _merge_runs(p)


# ── writing edits into the file: new blocks ──────────────────────────────────
class Templates:
    """The look of the file's own paragraphs, for paragraphs the edit adds."""

    def __init__(self, nodes: Sequence[Node]):
        self.nodes = nodes

    def nearest(self, near: int, pred) -> Optional[Node]:
        best, dist = None, None
        for i, n in enumerate(self.nodes):
            if n.kind != "table" and not n.generated and pred(n) and (dist is None or abs(i - near) < dist):
                best, dist = n, abs(i - near)
        return best

    def for_block(self, kind: str, level: int, ordered: bool, near: int) -> Optional[Node]:
        if kind == "heading":
            return self.nearest(near, lambda n: n.kind == "heading" and n.level == level)
        if kind == "list_item":
            return (self.nearest(near, lambda n: n.kind == "list_item" and n.ordered == ordered and n.level == level)
                    or self.nearest(near, lambda n: n.kind == "list_item" and n.ordered == ordered)
                    or self.nearest(near, lambda n: n.kind == "list_item"))
        return self.nearest(near, lambda n: n.kind == "paragraph" and len(n.text) > 40) or self.nearest(near, lambda n: n.kind == "paragraph")


def _strip_emphasis(rpr, keep_bold: bool):
    out = copy.deepcopy(rpr) if rpr is not None else None
    if out is not None:
        for tag in EMPHASIS:
            if not (keep_bold and tag in ("b", "bCs")):
                for e in out.findall(qn(f"w:{tag}")):
                    out.remove(e)
    return out


def _set_numpr(ppr, num_id: str, ilvl: int) -> None:
    for old in ppr.findall(qn("w:numPr")):
        ppr.remove(old)
    np_ = OxmlElement("w:numPr")
    for tag, val in (("ilvl", ilvl), ("numId", num_id)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:val"), str(val))
        np_.append(e)
    _insert_after_any(ppr, np_, _PPR_BEFORE_NUMPR)


def _set_pstyle(ppr, style_id: str) -> None:
    for old in ppr.findall(qn("w:pStyle")):
        ppr.remove(old)
    ps = OxmlElement("w:pStyle")
    ps.set(qn("w:val"), style_id)
    ppr.insert(0, ps)


@dataclass
class Placed:
    kind: str
    el: Any
    ordered: bool = False
    level: int = 0
    num_id: Optional[str] = None


def build_paragraph(ctx: Ctx, tpls: Templates, block: B.Block, prev: Optional[Placed], near: int) -> Placed:
    """A new paragraph for ``block`` (heading, paragraph or list item) with the look of the file's neighbours."""
    kind = block.kind if block.kind in ("heading", "list_item") else "paragraph"
    tpl = tpls.for_block(kind, block.level, block.ordered, near)
    p = OxmlElement("w:p")
    ppr = copy.deepcopy(tpl.tpl_ppr) if tpl is not None and tpl.tpl_ppr is not None else OxmlElement("w:pPr")
    for tag in ("w:sectPr", "w:pageBreakBefore", "w:pPrChange"):
        for e in ppr.findall(qn(tag)):
            ppr.remove(e)
    if kind == "heading" and tpl is None:                                    # no heading of this level yet: the file's style for it
        sid = ctx.heading_style(block.level)
        if sid:
            _set_pstyle(ppr, sid)
    elif kind == "list_item" and tpl is None:
        sid = ctx.styles.id_for("list paragraph")
        if sid:
            _set_pstyle(ppr, sid)
    num_id = None
    if kind == "list_item":
        if prev is not None and prev.kind == "list_item" and prev.ordered == block.ordered and prev.num_id:
            num_id = prev.num_id
        elif tpl is not None and tpl.num_id and not block.ordered:
            num_id = tpl.num_id
        elif tpl is not None and tpl.num_id and ctx.numbering.abstract_of(tpl.num_id):
            num_id = ctx.new_num(ctx.numbering.abstract_of(tpl.num_id), block.number or 1)
        else:
            num_id = ctx.ensure_list(block.ordered)
        # Word's level: the item it copies, moved by the difference in nesting (a list may start deeper than Word's first level)
        ilvl = tpl.ilvl + block.level - tpl.level if tpl is not None and tpl.kind == "list_item" else block.level
        _set_numpr(ppr, num_id, max(0, min(ilvl, 8)))
    elif kind == "paragraph":
        for e in ppr.findall(qn("w:numPr")):
            ppr.remove(e)
    if len(ppr):
        p.append(ppr)
    rpr = _strip_emphasis(tpl.tpl_rpr if tpl is not None else None, keep_bold=(kind == "heading"))
    text = block.text
    segs = block.segments or [B.Seg(text)]
    for r in _make_runs(ctx, _groups(text, _fmts_for(text, segs)), rpr, NO_FMT):
        p.append(r)
    return Placed(kind, p, block.ordered, block.level, num_id)


def _clear_text(p) -> None:
    for ch in list(p):
        if _local(ch) in ("r", "hyperlink", "smartTag", "sdt", "fldSimple", "bookmarkStart", "bookmarkEnd", "proofErr"):
            p.remove(ch)


def fill_cell(ctx: Ctx, tc, lines: Sequence[Sequence[B.Seg]], rpr=None) -> None:
    """Put the lines of a Markdown cell into ``tc``: the first paragraph holds the first line, further lines get paragraphs
    of their own with the same paragraph properties."""
    paras = tc.findall(qn("w:p"))
    first = paras[0] if paras else None
    if first is None:
        first = OxmlElement("w:p")
        tc.append(first)
    for extra in paras[1:] + [ch for ch in tc if _local(ch) in ("tbl", "sdt", "customXml")]:     # the lines are the whole cell
        tc.remove(extra)
    _clear_text(first)
    prev = first
    for n, segs in enumerate(lines):
        text = "".join(s.text for s in segs)
        target = first
        if n:
            target = copy.deepcopy(first)
            _clear_text(target)
            prev.addnext(target)
            prev = target
        for r in _make_runs(ctx, _groups(text, _fmts_for(text, segs)), rpr, NO_FMT):
            target.append(r)


def _cell_lines(cell: B.Cell) -> List[List[B.Seg]]:
    lines: List[List[B.Seg]] = [[]]
    for s in cell.segments:
        for k, part in enumerate(s.text.split("\n")):
            if k:
                lines.append([])
            if part:
                lines[-1].append(B.Seg(part, s.bold, s.italic, s.underline, s.strike, s.code, s.href))
    return lines


def _cell_rpr(ref: CellRef):
    first = next((pc.run for pc in (ref.paras[0].pieces if ref.paras else []) if pc.kind == "t"), None)
    return _strip_emphasis(_base_rpr(first.find(qn("w:rPr")) if first is not None else None), False)


def _patch_cell(ctx: Ctx, ref: CellRef, cell: B.Cell) -> None:
    lines = _cell_lines(cell)
    if ref.paras and len(lines) == sum(len(p.text.strip().split("\n")) for p in ref.paras):     # same lines: edit each paragraph in place
        at = 0
        for para in ref.paras:
            n = len(para.text.strip().split("\n"))
            segs: List[B.Seg] = []
            for k, line in enumerate(lines[at:at + n]):
                if k:
                    segs.append(B.Seg("\n"))
                segs.extend(line)
            patch_paragraph(ctx, para, "".join(s.text for s in segs), B._merge(segs))
            at += n
        return
    fill_cell(ctx, ref.tc, lines, _cell_rpr(ref))                                       # the lines changed: one paragraph each


def _build_table(ctx: Ctx, block: B.Block) -> Any:
    """A new table for ``block`` with borders and a shaded header row."""
    rows, width = block.rows, max(len(r) for r in block.rows)
    table = ctx.doc.add_table(rows=len(rows), cols=width)
    tbl = table._tbl
    tbl.getparent().remove(tbl)
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>' + "".join(
        f'<w:{side} w:val="single" w:sz="4" w:space="0" w:color="808080"/>' for side in ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>")
    after = [c for c in tbl.tblPr if c.tag in {qn(f"w:{n}") for n in ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd")}]
    (after[-1].addnext(borders) if after else tbl.tblPr.insert(0, borders))        # before tblLayout / tblLook, as Word's schema has it
    for r, row in enumerate(rows):
        for c in range(width):
            tc = table.cell(r, c)._tc
            fill_cell(ctx, tc, _cell_lines(row[c]) if c < len(row) else [[]])
            if r == 0:
                shd = OxmlElement("w:shd")
                for k, v in (("val", "clear"), ("color", "auto"), ("fill", "F2F2F2")):         # light grey, as the exported documents have
                    shd.set(qn(f"w:{k}"), v)
                tc.get_or_add_tcPr().append(shd)
                for run in tc.iter(qn("w:r")):
                    Run(run, None).font.bold = True
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(h)
    return tbl


def _patch_table(ctx: Ctx, info: TableInfo, old: B.Block, new: B.Block) -> Optional[Any]:
    """Edit the table in place. Returns a replacement element when the table had to be rebuilt (the column count changed, a
    merged table changed shape, or text was typed where a merged cell covers), else None."""
    old_rows = [[c.text for c in r] for r in old.rows]
    new_rows = [[c.text for c in r] for r in new.rows]
    width = max(len(r) for r in new_rows)
    if max(len(r) for r in old_rows) != width:
        return _build_table(ctx, new)
    pairs = B.pair_rows(old_rows, new_rows)
    structural = any(s in ("inserted", "deleted") for s, _, _ in pairs)
    if info.complex and structural:
        return _build_table(ctx, new)
    for s, i, j in pairs:
        if s != "modified":
            continue
        for c in range(width):
            ref = info.rows[i][c]
            cell = new.rows[j][c] if c < len(new.rows[j]) else B.Cell("", [])
            if old_rows[i][c] == new_rows[j][c] and [x.flags() for x in old.rows[i][c].segments] == [x.flags() for x in cell.segments]:
                continue
            if ref.tc is None:
                if cell.text.strip():
                    return _build_table(ctx, new)
                continue
            _patch_cell(ctx, ref, cell)
    if not structural:
        return None
    t = 1 if len(info.trs) > 1 else 0                                       # a data row is the model for new rows, not the header
    anchor = None
    for s, i, j in pairs:
        if s in ("equal", "modified"):
            anchor = info.trs[i]
        elif s == "deleted":
            info.trs[i].getparent().remove(info.trs[i])
        else:
            clone = copy.deepcopy(info.trs[t])
            for c, (tc, cell) in enumerate(zip(clone.findall(qn("w:tc")), new.rows[j])):
                fill_cell(ctx, tc, _cell_lines(cell), _cell_rpr(info.rows[t][c]))
            (anchor.addnext(clone) if anchor is not None else info.trs[0].addprevious(clone))
            anchor = clone
    return None


# ── writing edits into the file: the whole document ──────────────────────────
def _remove_block(el) -> None:
    ppr = el.find(qn("w:pPr")) if _local(el) == "p" else None
    parent = el.getparent()
    if ppr is not None and ppr.find(qn("w:sectPr")) is not None:             # a section break must survive: keep the paragraph, empty
        _clear_text(el)
    elif _local(parent) == "txbxContent" and len(parent.findall(qn("w:p"))) == 1:   # a text box must keep one paragraph
        _clear_text(el)
    else:
        parent.remove(el)


_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"


def _sync_fallbacks(body) -> None:
    """A text box is in the file twice, the second copy for old versions of Word: after an edit it gets the edited text too."""
    for alt in body.iter(f"{_MC}AlternateContent"):
        choice, fallback = alt.find(f"{_MC}Choice"), alt.find(f"{_MC}Fallback")
        if choice is None or fallback is None:
            continue
        new, old = list(choice.iter(qn("w:txbxContent"))), list(fallback.iter(qn("w:txbxContent")))
        if len(new) != len(old):
            continue
        for a, b in zip(new, old):
            if [t.text for t in a.iter(qn("w:t"))] != [t.text for t in b.iter(qn("w:t"))]:
                for ch in list(b):
                    b.remove(ch)
                b.extend(copy.deepcopy(ch) for ch in a)


def _place(anchor, el, first_old) -> None:
    if anchor is not None:
        _outer(anchor).addnext(el)
    elif first_old is not None:
        _outer(first_old).addprevious(el)
    else:
        raise PatchProblem("There is nowhere in the file to put new text.")


def export_in_place(twin_path: str, new_markdown: str) -> Tuple[bytes, List[str]]:
    """The file at ``twin_path`` with its text changed to ``new_markdown``: the bytes of a .docx and notes about anything that
    could not be kept as it was. Raises ``PatchProblem`` when the result would not read back as ``new_markdown``."""
    doc = Document(twin_path)
    ctx = Ctx(doc)
    nodes, _meta = walk(doc, ctx)
    old_blocks = B.parse(render_nodes(nodes))
    problem = check_blocks(nodes, old_blocks)
    if problem:
        raise PatchProblem("The file does not read back as it did: " + problem)
    new_blocks = B.parse(new_markdown)
    tpls = Templates(nodes)
    first_old = nodes[0].el if nodes else None
    anchor: Optional[Any] = None
    prev: Optional[Placed] = None
    near = 0
    for pr in B.align(old_blocks, new_blocks):
        if pr.status == "deleted":
            _remove_block(nodes[pr.base].el)
            continue
        block = new_blocks[pr.cur]
        if pr.status == "equal":
            node = nodes[pr.base]
            anchor, near, prev = node.el, pr.base, Placed(node.kind, node.el, node.ordered, node.level, node.num_id)
            continue
        if pr.status == "modified":
            node, old = nodes[pr.base], old_blocks[pr.base]
            near = pr.base
            if node.kind != "table" and node.para is None:
                raise PatchProblem("The text of a diagram (SmartArt) can only be changed in Word.")
            if node.kind == "table" and block.kind == "table":
                replacement = _patch_table(ctx, node.table, old, block)
                if replacement is not None:
                    ctx.warn("A table was rebuilt because its shape changed, so it has a plain grid instead of its original look.")
                    node.el.addprevious(replacement)
                    node.el.getparent().remove(node.el)
                anchor = replacement if replacement is not None else node.el
                prev = Placed("table", anchor)
                continue
            same_kind = node.kind == block.kind and (node.kind not in ("heading", "list_item") or (old.level == block.level and old.ordered == block.ordered))
            if same_kind and node.para is not None:
                patch_paragraph(ctx, node.para, block.text, block.segments)
                anchor, prev = node.el, Placed(node.kind, node.el, node.ordered, node.level, node.num_id)
                continue
            if node.kind == "table" or block.kind == "table":                # a table became text or the other way round
                placed = Placed("table", _build_table(ctx, block)) if block.kind == "table" else build_paragraph(ctx, tpls, block, prev, near)
            else:                                                            # its kind changed: a new paragraph of the new kind takes its place
                placed = build_paragraph(ctx, tpls, block, prev, near)
            sect = node.el.find(qn("w:pPr") + "/" + qn("w:sectPr")) if _local(node.el) == "p" else None
            if sect is not None and _local(placed.el) == "p":
                ppr = placed.el.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr")
                    placed.el.insert(0, ppr)
                ppr.append(copy.deepcopy(sect))
            node.el.addprevious(placed.el)
            node.el.getparent().remove(node.el)
            anchor, prev = placed.el, placed
            continue
        if block.kind == "hr":                                              # an inserted block: a rule has no paragraph in Word
            continue
        if block.kind == "table":
            el = _build_table(ctx, block)
            if anchor is not None and _local(_outer(anchor)) == "tbl":       # two tables side by side would merge into one
                gap = OxmlElement("w:p")
                _outer(anchor).addnext(gap)
                anchor = gap
            _place(anchor, el, first_old)
            anchor, prev = el, Placed("table", el)
        else:
            placed = build_paragraph(ctx, tpls, block, prev, near)
            _place(anchor, placed.el, first_old)
            anchor, prev = placed.el, placed
    _sync_fallbacks(doc.element.body)
    out = io.BytesIO()
    doc.save(out)
    data = out.getvalue()
    problem = check_blocks(walk(Document(io.BytesIO(data)))[0], new_blocks, lenient=True)
    if problem:
        raise PatchProblem("The edited file does not read back as the edited text: " + problem)
    return data, ctx.warnings


# ── tidying the spacing of a Word file ───────────────────────────────────────
_KEEPS_PARAGRAPH = {"drawing", "pict", "object", "AlternateContent", "fldChar", "fldSimple", "bookmarkStart", "bookmarkEnd", "commentRangeStart",
                    "commentReference", "footnoteReference", "endnoteReference", "sym", "ptab", "sdt", "oMath", "oMathPara"}


def _blank_line(p) -> bool:
    """A paragraph typed only to make space: no text and nothing else that shows or matters (a picture, a field, a bookmark,
    a page or section break, a frame)."""
    if any((t.text or "").strip() for t in p.iter(qn("w:t"))):
        return False
    ppr = p.find(qn("w:pPr"))
    if ppr is not None and any(ppr.find(qn(f"w:{n}")) is not None for n in ("sectPr", "pageBreakBefore", "framePr")):
        return False
    for el in p.iter():
        name = _local(el)
        if name in _KEEPS_PARAGRAPH or (name in ("br", "cr") and el.get(qn("w:type")) in ("page", "column")):
            return False
    return True


def _tidy_text(ctx: Ctx, p) -> Tuple[int, int]:
    """Paragraph ``p`` without spaces, tabs and line breaks at its start and end, and with one space where it had two or three in
    a row. Returns (spaces taken out, line breaks taken out)."""
    pieces = read_para(ctx, p).pieces
    spaces = breaks = 0

    def drop(pc) -> None:
        nonlocal spaces, breaks
        if pc.el.getparent() is not None:
            pc.el.getparent().remove(pc.el)
            if pc.kind == "br":
                breaks += 1
            else:
                spaces += 1

    for order in (pieces, pieces[::-1]):                    # the start, then the end
        for pc in order:
            if pc.el.getparent() is None:
                continue
            if pc.kind == "tab" or (pc.kind == "br" and pc.el.get(qn("w:type")) in (None, "textWrapping")):
                drop(pc)
                continue
            if pc.kind != "t":
                break
            text = pc.el.text or ""
            kept = text.lstrip() if order is pieces else text.rstrip()
            spaces += len(text) - len(kept)
            if kept != text:
                _set_text(pc.el, kept)
            if kept:
                break
    live = [pc for pc in pieces if pc.el.getparent() is not None]          # inside: one space for two or three, also across runs
    flat: List[str] = []
    owner: List[Optional[Tuple[int, int]]] = []
    for k, pc in enumerate(live):
        text = (pc.el.text or "") if pc.kind == "t" else "\x00"
        flat += list(text)
        owner += [(k, j) for j in range(len(text))] if pc.kind == "t" else [None]
    drop: Dict[int, set] = {}
    for m in re.finditer(r" {2,}", "".join(flat)):
        if m.end() - m.start() <= 3:            # ponytail: four or more spaces line up a form, so they stay; tabs would be the fix
            for i in range(m.start() + 1, m.end()):
                k, j = owner[i]
                drop.setdefault(k, set()).add(j)
    for k, offsets in drop.items():
        text = live[k].el.text or ""
        new = "".join(ch for j, ch in enumerate(text) if j not in offsets)
        spaces += len(text) - len(new)
        _set_text(live[k].el, new)
    for t in list(p.iter(qn("w:t"))):                       # text emptied by the tidy
        if not t.text and t.getparent() is not None:
            t.getparent().remove(t)
    return spaces, breaks


def tidy(path: str, *, write: bool = True) -> Tuple[Optional[bytes], Dict[str, int]]:
    """The Word file at ``path`` with the spacing a professional document does not have taken out: spaces, tabs and line breaks
    at the start and end of a paragraph, two or three spaces between words (four or more line up a form, and stay), and empty
    lines at the end of the document (they can leave a blank last page). Empty lines anywhere else lay the pages out (a cover
    page, room for a signature) and stay, as do what Word makes or draws (its table of contents), pictures, fields, bookmarks,
    page and section breaks, and the paragraph after a table.
    Returns the new file and what was taken out (``spaces``, ``line_breaks``, ``empty_lines``); no file when there is nothing to
    tidy, or when ``write`` is off (the counts only)."""
    doc = Document(path)
    ctx = Ctx(doc)
    body = doc.element.body
    counts = {"spaces": 0, "line_breaks": 0, "empty_lines": 0}
    for p in list(body.iter(qn("w:p"))):
        if any(_local(a) == "Fallback" for a in p.iterancestors()):          # the copy for old Words follows the tidied box
            continue
        if _generated(ctx, p, read_para(ctx, p)):
            continue                                                         # a table of contents is Word's own
        s, b = _tidy_text(ctx, p)
        counts["spaces"] += s
        counts["line_breaks"] += b
    seq = list(_top_blocks(body))
    blank = [_local(el) == "p" and el.getparent() is body and _blank_line(el) for el in seq]
    content = [i for i, x in enumerate(blank) if not x]
    if content:
        tail = list(range(content[-1] + 1, len(seq)))
        for i in tail[1:] if _local(seq[content[-1]]) == "tbl" else tail:        # a table cannot end a Word document
            if blank[i]:
                body.remove(seq[i])
                counts["empty_lines"] += 1
    _sync_fallbacks(body)
    if not write or not any(counts.values()):
        return None, counts
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), counts
