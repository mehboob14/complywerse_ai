"""Version numbers of governance documents: the one a file states, and whether a new version is a major or a minor change.

* ``detect`` - the version an uploaded Word or PDF file states: the latest row of its revision table (a table with a Version
  or Revision column), else a "Version" / "Revision" entry in a document-control table, its header or footer, or a line of its
  opening pages ("Revision: 4.0 (Q1-2013)"), else its file name ("... v1 1.pdf", "... 4.0 - Approved.docx"); 1.0 when
  nothing says. The document's numbering carries on from there.
* ``decide`` - whether a new version is a major or a minor change. Comply Agent reads what changed (the changed blocks
  before and after, never the whole document) and says which, and why in one sentence. When it cannot answer - not
  configured, refused, or not within ``AI_SECONDS`` - a rule decides: a section added or removed, or more than a fifth of the
  words changed, is major; anything else is minor.
* ``next_version`` - the number after the current one for that change.
"""
from __future__ import annotations

import concurrent.futures
import contextvars
import logging
import os
import re
from typing import Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

AI_SECONDS = 25                 # a save waits this long for the decision, then the rule decides
MAX_CHANGES_CHARS = 14000       # what changed, as Comply Agent reads it
FIRST_PAGES = 6                 # the pages of a PDF where its version is written (cover, document control)
MAX_TEXT_PAGES = 120            # the pages of a PDF read to see what changed

_HEAD = re.compile(r"\b(?:version|revision|ver|rev|issue)\b", re.I)
_LABEL = re.compile(r"^\W*(?:current\s+|document\s+)?(?:version|revision|ver|rev|issue)\.?(?:\s*(?:no\.?|number|#))?\W*$", re.I)
_CELL = re.compile(r"^\s*(?:v|ver\.?|version|rev\.?|revision)?\s*(\d{1,2}(?:\.\d{1,2}){1,2})\s*$", re.I)
_LINE = re.compile(r"^\s*(?:current\s+|document\s+)?(?:version|revision|ver|rev)\b\.?\s*(?:no\.?|number|#)?\s*[:\-–]?\s*v?(\d{1,2}(?:\.\d{1,2}){0,2})(?![\d.]*\d)",
                   re.I | re.M)
_SHORT_V = re.compile(r"(?:^|[\s·|(\-–])v(\d{1,2}\.\d{1,2}(?:\.\d{1,2})?)(?![\d.]*\d)", re.I)     # a footer's "POL-9 · v2.3"
_NAME_V = re.compile(r"(?:^|[\s_\-(\[])v(?:er(?:sion)?)?\.?\s*(\d{1,2})(?:[.\s_](\d{1,2}))?(?=$|[\s_\-)\].])", re.I)
_NAME_BARE = re.compile(r"(?<![\d.])(\d{1,2}\.\d{1,2})(?![\d.])")


def _key(version: str) -> Tuple[int, ...]:
    return tuple(int(p) for p in version.split("."))


def _norm(version: str) -> str:
    return version if "." in version else f"{version}.0"


# ── the version a file states ────────────────────────────────────────────────
def _read_docx(path: str) -> Tuple[List[List[List[str]]], List[str]]:
    from docx import Document

    doc = Document(path)
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables[:30]]
    lines: List[str] = []
    for section in doc.sections:
        for part in (section.first_page_header, section.header, section.first_page_footer, section.footer):
            try:
                lines += [p.text for p in part.paragraphs]
                tables += [[[c.text for c in row.cells] for row in t.rows] for t in part.tables]
            except Exception:       # noqa: BLE001 - a header python-docx cannot read says nothing
                continue
    lines += [p.text for p in doc.paragraphs[:400]]
    return tables, lines


def _read_pdf(path: str) -> Tuple[List[List[List[str]]], List[str]]:
    import fitz

    doc = fitz.open(path)
    try:
        tables: List[List[List[str]]] = []
        lines: List[str] = []
        for n in range(min(doc.page_count, FIRST_PAGES)):
            page = doc.load_page(n)
            lines += page.get_text().splitlines()
            try:
                tables += [[[re.sub(r"\s+", " ", c or "").strip() for c in row] for row in t.extract()] for t in page.find_tables().tables]
            except Exception:       # noqa: BLE001 - table detection is a bonus
                pass
        return tables, lines
    finally:
        doc.close()


def _from_tables(tables: List[List[List[str]]]) -> Tuple[Optional[str], Optional[str]]:
    """(the latest version in a revision table, a version a document-control table gives as "Version | 4.0")."""
    history: List[str] = []
    stated: Optional[str] = None
    for rows in tables:
        rows = [r for r in rows if any((c or "").strip() for c in r)]          # a PDF's ruled grid reads with empty rows
        if not rows:
            continue
        cols = [k for k, h in enumerate(rows[0]) if _HEAD.search(h or "")]
        for k in cols:
            # the version column; a PDF's table can shift a cell one column over, so the columns beside it are read when it is empty
            for at in (k, k - 1, k + 1):
                found = [m.group(1) for row in rows[1:] if 0 <= at < len(row) for m in [_CELL.match(row[at] or "")] if m]
                if found:
                    history += found
                    break
        for row in rows:
            if stated is None and len(row) >= 2 and _LABEL.match(row[0] or ""):
                stated = next((m.group(1) for c in row[1:] for m in [_CELL.match(c or "")] if m), None)
    return (max(history, key=_key) if history else None), stated


def from_name(name: str) -> Optional[str]:
    base = os.path.splitext(os.path.basename(name or ""))[0]
    m = _NAME_V.search(base)
    if m:
        return f"{int(m.group(1))}.{int(m.group(2) or 0)}"
    m = _NAME_BARE.search(base)
    return m.group(1) if m else None


def detect(path: Optional[str], name: str) -> Tuple[str, str]:
    """``(version, where it was found)``: "revision table", "document", "file name" or "default" (1.0)."""
    ext = os.path.splitext(name or path or "")[1].lower()
    tables: List[List[List[str]]] = []
    lines: List[str] = []
    if path and os.path.isfile(path) and ext in (".docx", ".pdf"):
        try:
            tables, lines = _read_docx(path) if ext == ".docx" else _read_pdf(path)
        except Exception:       # noqa: BLE001 - a file that cannot be read states nothing; its name may
            logger.info("could not read the version of %s", name, exc_info=True)
    latest, stated = _from_tables(tables)
    if latest:
        return _norm(latest), "revision table"
    if stated:
        return _norm(stated), "document"
    m = _LINE.search("\n".join(lines))
    if m:
        return _norm(m.group(1)), "document"
    for line in lines:                  # a short line (a header or footer) with "v2.3" in it
        m = _SHORT_V.search(line) if len(line.strip()) <= 60 else None
        if m:
            return m.group(1), "document"
    named = from_name(name)
    if named:
        return named, "file name"
    return "1.0", "default"


# ── major or minor ───────────────────────────────────────────────────────────
SYSTEM = (
    "You are Comply Agent. You decide whether a new version of a governance document (a policy, standard, procedure, guideline or "
    "charter) is a major or a minor change, as a document controller does before numbering it.\n"
    "Major: it changes what the document requires or whom it binds - a requirement, obligation, control or prohibition added, removed, "
    "or made stronger or weaker; its scope or audience; roles and responsibilities; who owns or approves it; a threshold, frequency or "
    "time limit; a section added or removed.\n"
    "Minor: the meaning stays - wording, spelling, grammar, formatting, numbering, clarifications, examples, references to other "
    "documents or standards, names, titles, dates and contact details.\n"
    "Answer with the JSON asked for: change (major or minor), and reason - one sentence a reviewer reads in the version history, "
    "naming the change that decided it."
)
SCHEMA = {"type": "object", "additionalProperties": False, "required": ["change", "reason"],
          "properties": {"change": {"type": "string", "enum": ["major", "minor"]}, "reason": {"type": "string"}}}


def changes(old_text: str, new_text: str) -> Tuple[str, Dict[str, int]]:
    """What changed, block by block (before and after), and how much: blocks and words, and headings added or removed."""
    from .studio import blocks as B

    a, b = B.parse(old_text or ""), B.parse(new_text or "")
    out: List[str] = []
    n = {"words": sum(x.words for x in a) or 1, "changed_words": 0, "sections": 0, "blocks": 0}
    for p in B.align(a, b):
        if p.status == "equal":
            continue
        n["blocks"] += 1
        if p.status == "modified":
            before, after = a[p.base], b[p.cur]
            n["changed_words"] += sum(len(re.findall(r"\w+", t)) for op, t in B.word_diff(before.text, after.text) if op != "equal")
            out.append(f"Changed from:\n{before.source}\nto:\n{after.source}")
        elif p.status == "inserted":
            n["changed_words"] += b[p.cur].words
            n["sections"] += b[p.cur].kind == "heading"
            out.append(f"Added:\n{b[p.cur].source}")
        else:
            n["changed_words"] += a[p.base].words
            n["sections"] += a[p.base].kind == "heading"
            out.append(f"Removed:\n{a[p.base].source}")
    text = "\n\n".join(out)
    if len(text) > MAX_CHANGES_CHARS:
        text = text[:MAX_CHANGES_CHARS].rstrip() + "\n\n(more changes, not shown)"
    return text, n


def _rule(n: Dict[str, int]) -> Dict[str, str]:
    if n["sections"]:
        return {"change": "major", "reason": "A section was added or removed.", "by": "rule"}
    if n["changed_words"] > n["words"] / 5:
        return {"change": "major", "reason": "More than a fifth of the document's words changed.", "by": "rule"}
    return {"change": "minor", "reason": "A limited change to the wording.", "by": "rule"}


def decide(old_text: str, new_text: str, *, title: str = "") -> Dict[str, str]:
    """``{"change": "major" | "minor", "reason": one sentence, "by": "ai" | "rule"}`` for the version ``new_text`` makes of ``old_text``."""
    from ...services.ai_gateway import claude_client

    text, n = changes(old_text, new_text)
    if not n["blocks"]:
        return {"change": "minor", "reason": "Only the file changed, not its text.", "by": "rule"}
    if not claude_client.configured():
        return _rule(n)
    user = (f"<document title=\"{(title or 'Untitled').replace(chr(34), chr(39))}\" words=\"{n['words']}\">\n"
            f"{n['blocks']} block(s) changed, about {n['changed_words']} word(s).\n\n{text}\n</document>")
    pool = concurrent.futures.ThreadPoolExecutor(max_workers=1)
    try:
        # on a thread of its own, with the request's context (who and which organisation the AI usage is for), so a save never
        # waits longer than AI_SECONDS
        future = pool.submit(contextvars.copy_context().run, claude_client.complete, use_case="governance.version_change", system=SYSTEM,
                             messages=[{"role": "user", "content": user}], schema=SCHEMA, max_tokens=2000, effort="low")
        data = future.result(timeout=AI_SECONDS).data or {}
        if data.get("change") in ("major", "minor") and str(data.get("reason") or "").strip():
            return {"change": data["change"], "reason": str(data["reason"]).strip()[:500], "by": "ai"}
    except concurrent.futures.TimeoutError:
        logger.warning("the major/minor decision for %r took longer than %ss; the rule decided", title, AI_SECONDS)
    except Exception:       # noqa: BLE001 - not configured, refused, blocked or failed: the rule decides, the save goes on
        logger.warning("the major/minor decision for %r failed; the rule decided", title, exc_info=True)
    finally:
        pool.shutdown(wait=False)
    return _rule(n)


def next_version(current: str, change: str) -> str:
    from .routers.versions import increment_version

    return increment_version(current or "1.0", "major" if change == "major" else "minor")


def verdict(decision: Dict[str, str]) -> str:
    """The decision as the version history says it: "Major change: <why>"."""
    by = "" if decision.get("by") == "ai" else " (decided by rule: Comply Agent was not available)"
    return f"{decision['change'].capitalize()} change: {decision['reason']}{by}"


def summary_of(done: str, decision: Dict[str, str]) -> str:
    return f"{done}. {verdict(decision)}"


_FOUND = {"revision table": "from its revision table", "document": "as the document states it", "file name": "from its file name"}


def first_summary(version: str, found_in: str) -> str:
    """What the history says about a document's first version: the uploaded file, and where its number came from."""
    if found_in in _FOUND:
        return f"Initial version: the uploaded file, version {version} {_FOUND[found_in]}"
    return "Initial version: the uploaded file (it states no version, so it starts at 1.0)"


def _safe_key(version: str) -> Tuple[int, ...]:
    try:
        return _key(version or "1.0")
    except ValueError:
        return (0,)


def for_new_file(current: str, old_path: Optional[str], old_name: str, new_path: str, new_name: str, *,
                 old_content: Optional[str] = None, title: str = "") -> Tuple[str, str, str]:
    """``(number, "major" | "minor", what the history says)`` for a new file of a document at version ``current``: the version the
    file states when it is later than ``current`` (the organisation's own numbering), else the next number by what changed."""
    stated, found_in = detect(new_path, new_name)
    if found_in in _FOUND and _safe_key(stated) > _safe_key(current):
        change = "major" if _safe_key(stated)[0] > _safe_key(current)[0] else "minor"
        return stated, change, f"The file states version {stated} ({_FOUND[found_in]})."
    new_text = file_text(new_path, new_name)
    if not new_text.strip():
        return next_version(current, "minor"), "minor", "Minor change: the new file's text could not be read to compare it with the last version."
    decision = decide(file_text(old_path, old_name) or old_content or "", new_text, title=title)
    return next_version(current, decision["change"]), decision["change"], verdict(decision)


def file_text(path: Optional[str], name: str) -> str:
    """The text of a Word or PDF file, read to see what changed between two of its versions ("" for any other file)."""
    ext = os.path.splitext(name or path or "")[1].lower()
    if not path or not os.path.isfile(path):
        return ""
    try:
        if ext == ".docx":
            from .studio.docx_io import import_docx

            return import_docx(path).markdown
        if ext == ".pdf":
            import fitz

            doc = fitz.open(path)
            try:
                return "\n\n".join(doc.load_page(i).get_text().strip() for i in range(min(doc.page_count, MAX_TEXT_PAGES)))
            finally:
                doc.close()
    except Exception:       # noqa: BLE001 - an unreadable file: the change is decided from what can be read
        logger.info("could not read the text of %s", name, exc_info=True)
    return ""
