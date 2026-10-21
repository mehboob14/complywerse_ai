"""Policy statements <-> the original document file (tenants in ``config.DOCUMENT_SYNC_TENANTS``, cfsb).

Two promises:

1. A statement is the document's own wording. After the AI extracts statements, each one is looked up in the file's text.
   Found as written -> ``exact``; found with small differences (a paraphrase, a lost line break) -> ``corrected`` to the
   file's wording; not found -> ``not_found`` (kept, flagged, and never written back). The wording as the file has it is
   kept in ``source_text`` with its offset ``source_start``.
2. An edit to a statement reaches the file. ``statement_text != source_text`` means "edited, not yet in the document";
   ``apply`` writes every such edit into a new copy of the file, creates the next document version and sends an approved
   document back to review. A Word file is edited in place (formatting stays); a PDF cannot be, so the new version is a Word
   copy built from its text, and the PDF stays as the earlier version.

Everything works on the file's text as one string (paragraphs joined by newlines) so a sentence wrapped over two lines or
two paragraphs is still one statement. Matching ignores differences in white space, quotation marks and dashes.
"""
from __future__ import annotations

import logging
import os
import re
import shutil
import uuid
from bisect import bisect_left
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

FUZZY_MIN = 0.90          # how close a statement must be to a sentence of the file to be corrected to it
_WRITABLE = ("exact", "corrected")
_REVIEW_FROM = ("approved", "published", "exception_applied")
_QUOTES = str.maketrans({"‘": "'", "’": "'", "‚": "'", "“": '"', "”": '"', "„": '"',
                         "–": "-", "—": "-", "−": "-", "‐": "-", "‑": "-"})
_DROP = {"­", "​", "‌", "‍", "﻿"}


def enabled(db) -> bool:
    from ...config import DOCUMENT_SYNC_TENANTS
    from ...models import Tenant

    tenant = db.query(Tenant).first()
    return bool(tenant and ("*" in DOCUMENT_SYNC_TENANTS or (tenant.slug or "").lower() in DOCUMENT_SYNC_TENANTS))


# ── text of a file ───────────────────────────────────────────────────────────
def _norm(text: str) -> Tuple[str, List[int]]:
    """``(normalised text, index of each of its characters in the original)``."""
    out: List[str] = []
    idx: List[int] = []
    for i, ch in enumerate(text):
        if ch in _DROP:
            continue
        if ch.isspace() or ch == " ":
            if out and out[-1] == " ":
                continue
            out.append(" ")
        else:
            out.append(ch.translate(_QUOTES))
        idx.append(i)
    return "".join(out), idx


def _norm_plain(text: str) -> str:
    return _norm(text or "")[0].strip()


@dataclass
class Paragraph:
    text: str
    ref: Any = None        # the python-docx paragraph, when there is one


def _runs_text(p) -> str:
    return "".join(r.text for r in p.runs)


def docx_paragraphs(path: str) -> Tuple[Any, List[Paragraph]]:
    """The editable paragraphs of a Word file: body paragraphs, then table cells (merged cells once). Text is that of the
    runs, which is what can be rewritten."""
    from docx import Document

    doc = Document(path)
    paragraphs = [Paragraph(_runs_text(p), p) for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                paragraphs.extend(Paragraph(_runs_text(p), p) for p in cell.paragraphs)
    return doc, [p for p in paragraphs if p.text.strip()]


def file_paragraphs(path: str, file_type: str) -> Tuple[Any, List[Paragraph], bool]:
    """``(docx document or None, paragraphs, editable)``: a Word file as it is, any other file as the text it extracts to."""
    ft = (file_type or "").lower().lstrip(".")
    if ft == "docx":
        doc, paragraphs = docx_paragraphs(path)
        return doc, paragraphs, True
    from .routers.policy_parser import extract_text_from_file

    text = extract_text_from_file(path, ft)
    return None, [Paragraph(line) for line in text.split("\n") if line.strip()], False


def flatten(paragraphs: List[Paragraph]) -> Tuple[str, List[Tuple[int, int]]]:
    """The text as one string and each paragraph's ``(start, end)`` in it."""
    spans, pos, parts = [], 0, []
    for p in paragraphs:
        spans.append((pos, pos + len(p.text)))
        parts.append(p.text)
        pos += len(p.text) + 1
    return "\n".join(parts), spans


# ── finding a statement in the text ──────────────────────────────────────────
def _overlaps(claimed: List[Tuple[int, int]], a: int, b: int) -> bool:
    return any(a < y and x < b for x, y in claimed)


def _sentences(norm: str, idx: List[int], paragraphs: List[Tuple[int, int]]) -> List[Tuple[int, int]]:
    """Sentence windows of the normalised text (cut after . ; : ! ? and at each paragraph), plus pairs of neighbours (a
    statement the AI merged)."""
    cuts = {0, len(norm)} | {m.end() for m in re.finditer(r"(?<=[.;:!?])\s+", norm)}
    cuts |= {bisect_left(idx, start) for start, _end in paragraphs}
    ordered = sorted(c for c in cuts if 0 <= c <= len(norm))
    spans = [(a, b) for a, b in zip(ordered, ordered[1:]) if norm[a:b].strip()]
    return spans + [(spans[i][0], spans[i + 1][1]) for i in range(len(spans) - 1)]


def locate(flat: str, text: str, claimed: List[Tuple[int, int]], _cache: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """Where ``text`` is in ``flat``: ``{"start", "end", "match"}`` in ``flat`` coordinates, ``match`` exact or corrected."""
    if "n" not in _cache:
        _cache["n"] = _norm(flat)
    norm, idx = _cache["n"]
    wanted = _norm_plain(text)
    if not wanted:
        return None

    def to_flat(a: int, b: int) -> Tuple[int, int]:
        start, end = idx[a], idx[b - 1] + 1
        return start, end

    at = norm.find(wanted)
    while at != -1:
        if not _overlaps(claimed, *to_flat(at, at + len(wanted))):
            start, end = to_flat(at, at + len(wanted))
            return {"start": start, "end": end, "match": "exact"}
        at = norm.find(wanted, at + 1)

    if "s" not in _cache:
        _cache["s"] = _sentences(norm, idx, _cache.get("spans", []))
    lowered = wanted.lower()
    best, best_ratio = None, FUZZY_MIN
    for a, b in _cache["s"]:
        window = norm[a:b].strip()
        if not window or not 0.7 <= len(window) / len(wanted) <= 1.4:
            continue
        matcher = SequenceMatcher(None, window.lower(), lowered, autojunk=False)
        if matcher.real_quick_ratio() < best_ratio or matcher.quick_ratio() < best_ratio:
            continue
        ratio = matcher.ratio()
        if ratio > best_ratio:
            start, end = to_flat(a + (len(norm[a:b]) - len(norm[a:b].lstrip())), a + len(norm[a:b].rstrip()))
            if not _overlaps(claimed, start, end):
                best, best_ratio = {"start": start, "end": end, "match": "corrected"}, ratio
    return best


# ── anchoring statements ─────────────────────────────────────────────────────
def _statements(db, document_id: int):
    from ...models import PolicyStatement

    return (db.query(PolicyStatement).filter(PolicyStatement.document_id == document_id, PolicyStatement.status == "active")
            .order_by(PolicyStatement.statement_code, PolicyStatement.id).all())


def anchor(db, document, user_id: Optional[int] = None, *, again: bool = False) -> Dict[str, int]:
    """Match the document's statements to its file. Statements never matched are matched; with ``again`` the ones already
    matched are re-located too (after the file changed). A statement with an edit not yet in the file is left alone."""
    counts = {"exact": 0, "corrected": 0, "not_found": 0, "total": 0}
    if not document.file_path or not os.path.exists(document.file_path):
        return counts
    _doc, paragraphs, _editable = file_paragraphs(document.file_path, document.file_type)
    flat, _spans = flatten(paragraphs)
    cache: Dict[str, Any] = {"spans": _spans}
    claimed: List[Tuple[int, int]] = []
    from .routers.policy_parser import _create_version_snapshot

    for s in _statements(db, document.id):
        counts["total"] += 1
        pending = s.source_text is not None and (s.statement_text or "").strip() != s.source_text.strip()
        if pending or (s.source_match is not None and not again):
            if s.source_start is not None and s.source_text:
                claimed.append((s.source_start, s.source_start + len(s.source_text)))
            counts[s.source_match or "not_found"] = counts.get(s.source_match or "not_found", 0) + 1
            continue
        found = locate(flat, s.statement_text or "", claimed, cache)
        if not found:
            s.source_match, s.source_text, s.source_start = "not_found", None, None
            counts["not_found"] += 1
            continue
        original = flat[found["start"]:found["end"]]
        claimed.append((found["start"], found["end"]))
        if (s.statement_text or "").strip() != original.strip():       # the statement is the document's own wording, character for character
            s.statement_text = original[:10000]
            _create_version_snapshot(db, s, "verbatim_correction", user_id or document.owner_id,
                                     "Wording corrected to match the document")
        s.source_text, s.source_start, s.source_match = original, found["start"], found["match"]
        counts[found["match"]] += 1
    db.commit()
    return counts


def anchor_after_parse(db, document_id: int, user_id: Optional[int] = None) -> None:
    """Called when statements were extracted or re-extracted: never lets a failure fail the parse."""
    try:
        from ...models import GovernanceDocument

        if not enabled(db):
            return
        document = db.query(GovernanceDocument).filter(GovernanceDocument.id == document_id).first()
        if document:
            logger.info("document %s: statements matched to the file: %s", document_id, anchor(db, document, user_id))
    except Exception:       # noqa: BLE001
        db.rollback()
        logger.exception("could not match the statements of document %s to its file", document_id)


# ── what is waiting to be written ────────────────────────────────────────────
def state(db, document) -> Dict[str, Any]:
    stmts = _statements(db, document.id)
    pending = [{"id": s.id, "code": s.statement_code, "was": s.source_text, "now": s.statement_text} for s in stmts
               if s.source_match in _WRITABLE and s.source_text is not None and (s.statement_text or "").strip() != s.source_text.strip()]
    blocked = [{"id": s.id, "code": s.statement_code} for s in stmts if s.source_match == "not_found"]
    ft = (document.file_type or "").lower()
    has_file = bool(document.file_path and os.path.exists(document.file_path))
    reason = None
    if not has_file:
        reason = "This document has no file to update."
    elif ft not in ("docx", "pdf"):
        reason = "Only Word and PDF documents can be updated from their statements."
    return {
        "enabled": True, "file_type": ft, "version": document.current_version, "status": document.status,
        "can_apply": bool(pending) and reason is None, "reason": reason,
        "creates_word_copy": ft != "docx",
        "pending": pending, "not_found": blocked,
        "counts": {"statements": len(stmts), "checked": sum(1 for s in stmts if s.source_match),
                   "exact": sum(1 for s in stmts if s.source_match == "exact"),
                   "corrected": sum(1 for s in stmts if s.source_match == "corrected"),
                   "not_found": sum(1 for s in stmts if s.source_match == "not_found")},
    }


# ── writing edits into the file ──────────────────────────────────────────────
def _set_range(paragraph, start: int, end: int, new: str) -> None:
    """Replace characters ``start:end`` of the paragraph's run text with ``new``, inside the run where it begins, so the
    formatting of the surrounding text is untouched."""
    pos, first = 0, True
    for run in paragraph.runs:
        text = run.text
        a, b = pos, pos + len(text)
        pos = b
        if b <= start or a >= end:
            continue
        lo, hi = max(start, a) - a, min(end, b) - a
        if first:
            run.text = text[:lo] + new + text[hi:]
            first = False
        else:
            run.text = text[:lo] + text[hi:]


def rewrite(paragraphs: List[Paragraph], spans: List[Tuple[int, int]], edits: List[Tuple[int, int, str]]) -> None:
    """Apply ``(start, end, new)`` edits (flat coordinates) to the paragraphs. A range over several paragraphs puts the new
    text where it starts and removes the rest of it from the others."""
    for start, end, new in sorted(edits, key=lambda e: -e[0]):
        first = True
        for i in (i for i, (a, b) in enumerate(spans) if a < end and start < b):
            a, b = spans[i]
            lo, hi = max(start, a) - a, min(end, b) - a
            _set_range(paragraphs[i].ref, lo, hi, new if first else "")
            first = False


def _word_copy(paragraphs: List[Paragraph], title: str, path: str) -> None:
    from docx import Document

    doc = Document()
    doc.core_properties.title = title or ""
    for p in paragraphs:
        doc.add_paragraph(p.text)
    doc.save(path)


def apply(db, document, user, change_reason: Optional[str] = None) -> Dict[str, Any]:
    """Write every pending statement edit into a new copy of the document's file and record the next version."""
    from ...models import GovernanceDocumentVersion, PolicyStatement
    from .routers.documents import GOVERNANCE_UPLOAD_DIR
    from .routers.versions import create_audit_log
    from .permissions import ensure_not_archived

    ensure_not_archived(document)
    info = state(db, document)
    if info["reason"]:
        raise ValueError(info["reason"])
    if not info["pending"]:
        raise LookupError("No statement has been changed since the document was last written.")

    old_path, old_name, old_type = document.file_path, document.file_name, (document.file_type or "").lower()
    stem = os.path.splitext(old_name or "document")[0]
    new_name = f"{stem}.docx"
    new_path = os.path.join(GOVERNANCE_UPLOAD_DIR, f"{uuid.uuid4()}_{new_name}")

    _doc, paragraphs, editable = file_paragraphs(old_path, old_type)
    if editable:
        shutil.copyfile(old_path, new_path)
        doc, paragraphs = docx_paragraphs(new_path)
    else:
        _word_copy(paragraphs, document.title, new_path)
        doc, paragraphs = docx_paragraphs(new_path)
    flat, spans = flatten(paragraphs)

    by_id = {s.id: s for s in db.query(PolicyStatement).filter(PolicyStatement.id.in_([p["id"] for p in info["pending"]])).all()}
    edits: List[Tuple[int, int, str]] = []
    conflicts: List[str] = []
    for item in info["pending"]:
        s = by_id[item["id"]]
        # the wording must still be where it was recorded; if the file moved on, look for it again
        text = s.source_text
        at = s.source_start if s.source_start is not None and flat[s.source_start:s.source_start + len(text)] == text else flat.find(text)
        if at < 0:
            conflicts.append(s.statement_code or str(s.id))
            continue
        edits.append((at, at + len(s.source_text), s.statement_text.strip()))
    if not edits:
        raise LookupError("The edited wording could not be found in the file any more: " + ", ".join(conflicts))
    if len({e[0] for e in edits}) != len(edits):
        raise ValueError("Two statements point at the same place in the file.")
    before = "\n\n".join(p.text for p in paragraphs)
    rewrite(paragraphs, spans, edits)
    doc.save(new_path)

    # ── the version ──
    old_version = document.current_version or "1.0"
    versions = db.query(GovernanceDocumentVersion).filter(GovernanceDocumentVersion.document_id == document.id)
    if versions.count() == 0:
        db.add(GovernanceDocumentVersion(
            document_id=document.id, version_number=old_version, change_type="baseline", title=document.title,
            content=document.content, file_name=old_name, file_path=old_path, file_size=document.file_size,
            file_type=old_type, change_summary="Original document, before statement edits were written back",
            status="superseded", created_by=user.id, created_at=datetime.utcnow()))
    else:
        replaced = versions.filter(GovernanceDocumentVersion.version_number == old_version).order_by(GovernanceDocumentVersion.id.desc()).first()
        if replaced is not None and not replaced.file_path and old_path:      # its file was put in place before its row held it
            replaced.file_name, replaced.file_path, replaced.file_size, replaced.file_type = old_name, old_path, document.file_size, old_type
        versions.filter(GovernanceDocumentVersion.status == "current").update({"status": "superseded"})
    # major or minor, as Comply Agent reads what changed in the document's text (doc_version)
    from .doc_version import decide, next_version, verdict

    decision = decide(before, "\n\n".join(p.text for p in paragraphs), title=document.title)
    new_version = next_version(old_version, decision["change"])
    codes = [p["code"] for p in info["pending"] if p["code"] not in conflicts]
    content = document.content
    if content:
        for item in info["pending"]:
            if item["code"] not in conflicts and content.count(item["was"]) == 1:
                content = content.replace(item["was"], item["now"].strip(), 1)
    row = GovernanceDocumentVersion(
        document_id=document.id, version_number=new_version, change_type=decision["change"], title=document.title, content=content,
        file_name=new_name, file_path=new_path, file_size=os.path.getsize(new_path), file_type="docx",
        change_summary=f"{len(edits)} policy statement(s) updated in the document: " + ", ".join(codes) + ". " + verdict(decision),
        change_reason=change_reason, status="current", created_by=user.id, created_at=datetime.utcnow())
    db.add(row)
    db.flush()

    old_status = document.status
    document.file_name, document.file_path, document.file_size, document.file_type = new_name, new_path, row.file_size, "docx"
    document.current_version, document.content, document.updated_at = new_version, content, datetime.utcnow()
    from . import doc_lifecycle

    sent = doc_lifecycle.changed_after_review(db, document, old_status, user.id,
                                              f"Statements written into the document (version {new_version}), so it needs approving again")

    for item in info["pending"]:
        if item["code"] in conflicts:
            continue
        s = by_id[item["id"]]
        s.source_text, s.document_version_id = s.statement_text.strip(), row.id
    db.flush()
    # offsets moved: find every statement again in the new file
    anchor(db, document, user.id, again=True)

    create_audit_log(db=db, document_id=document.id, tenant_id=document.tenant_id, user_id=user.id,
                     action="statements_written_to_document",
                     action_details=f"{len(edits)} statement edit(s) written into the document; version {new_version} created"
                                    + doc_lifecycle.outcome(old_status, document.status, document),
                     field_changed="current_version", old_value=old_version, new_value=new_version)
    db.commit()
    return {"version": new_version, "version_id": row.id, "updated": len(edits), "conflicts": conflicts,
            "status": document.status, "sent_to_review": sent,
            "file_type": "docx", "word_copy": not editable}
