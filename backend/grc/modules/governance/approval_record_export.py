"""The document control pages (the approval record, approval_record.py, laid out as the pages of a bank's own policy: Document
Description, Review and Approval, Revision History and Approval History) as files - on their own, or placed inside the document they
belong to (the user, 2026-10-25: the section is part of the document, kept by the platform from the workflow, not a card or an
appendix). Built from the record only - no database - so the functions are the same whatever the document is; the drawn signatures
come in as ``images`` ({e-signature id: PNG bytes}), a signature without ink is left out.

* ``to_pdf`` / ``to_docx``: the control pages on their own (the standalone export; ``to_pdf`` also makes each part of a
  platform-written document's PDF that stands for one of its sections, controlled_copy.pdf).
* ``place_in_docx``: the pages inserted after the cover page of a Word file (else at its end); ``place_in_pdf``: after page 1 of a PDF
  that looks like a cover (else at its end). The file given is read and a copy returned: the stored file is never modified, and every
  original element of a Word file stays as it was.

Every time is printed as ``YYYY-MM-DD HH:MM UTC`` - the date AND the time, with the meaning of each signature beside it (what a
signed record must show in every human-readable form, 21 CFR 11.50). The pages of a PDF carry "<title> · Document control · generated
by the platform <time> · <reference>" and "Document control page X of Y" at the foot (X of Y counting only these pages); the pages
inside a Word file are in the document's own header and footer, with the same line as their last paragraph.
"""
from __future__ import annotations

import html
import io
import re
from datetime import datetime
from typing import Any, List, Mapping, Optional, Sequence, Tuple
from urllib.parse import quote

from . import routes

TITLE = "Document control"
SLOTS = ("description", "signoff", "revision")                  # the sections, in the order they are printed
SLOT_TITLE = {"description": "Document Description", "signoff": "Review and Approval", "revision": "Revision History"}
APPROVAL_TITLE = "Approval History"                              # printed with the revision history
_STATUS = {"complete": "Complete", "pending": "Pending", "changes_requested": "Changes requested", "rejected": "Rejected"}
_RULE = {"all": "everyone signs", "any": "any one signs"}
_DECIDED = {"signed": "Signed", "rejected": "Rejected", "changes_requested": "Asked for changes"}
SIGN_BOX = (5.2, 1.6)                          # cm: the room a drawn signature gets above its line
NAMELESS = "Unknown"                           # a person is never printed without a name (the record never lists one, but a name can scrub to nothing)
COMMENT_CHARS = 220
VERSION_MIN, DATE_MIN = 1.9, 2.2               # cm: the narrowest a version or a date column gets on a narrow page, so "Version" and "2026-10-28" stay whole
COVER_CHARS = 700                              # a first page with fewer characters of text than this is a cover
COVER_SCAN = 40                                # a Word file's cover ends at a page boundary among its first elements
MIME = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
A4 = (595.2756, 841.8898)                      # points


_UNFIT = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]")         # not allowed in a Word file (XML), and nothing to print in a PDF


def scrub(value: Any) -> Any:
    """The record with every text cleaned of characters a document cannot hold (a comment may carry anything a person pasted)."""
    if isinstance(value, str):
        return _UNFIT.sub("", value)
    if isinstance(value, Mapping):
        return {k: scrub(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [scrub(v) for v in value]
    return value


# ── what the pages say, as text (shared by the PDF and the Word file) ───────────
def when(iso: Optional[str]) -> str:
    """``2026-10-21 10:15 UTC`` (a record's times are naive UTC); a dash when there is none."""
    if not iso:
        return "-"
    try:
        return datetime.fromisoformat(iso).strftime("%Y-%m-%d %H:%M UTC")
    except ValueError:
        return str(iso)


def day(iso: Optional[str]) -> str:
    return str(iso)[:10] if iso else "-"


def trim(text: Optional[str], n: int = COMMENT_CHARS) -> str:
    text = (text or "").strip()
    return text if len(text) <= n else text[:n].rstrip() + " …"


def rule_text(stage: Mapping[str, Any]) -> str:
    if stage["rule"] == "at_least":
        return f"at least {stage['min_count'] or 1} sign"
    return _RULE.get(stage["rule"], "everyone signs")


def stage_title(stage: Mapping[str, Any]) -> str:
    """``STAGE 2 — CONCURRED BY · Risk concurrence`` (a stage nobody named has no name after the bar's words: ``routes.UNNAMED``)."""
    name = (stage.get("name") or "").strip()
    return f"STAGE {stage['position']} — {str(stage['label']).upper()}" + (f" · {name}" if name and name not in routes.UNNAMED else "")


def stage_note(stage: Mapping[str, Any]) -> str:
    """``Complete · 2 of 2 signed`` (and the rule, when it is not that everyone signs)."""
    signed = sum(1 for p in stage["people"] if p["decision"] == "signed")
    note = f"{_STATUS.get(stage['status'], stage['status'])} · {signed} of {len(stage['people'])} signed"
    return note + (f" · {rule_text(stage)}" if stage["rule"] != "all" else "")


def person_lines(stage: Mapping[str, Any], p: Mapping[str, Any]) -> List[Tuple[str, str]]:
    """What is printed under a person's signature line: the name, designation, what they did and when, the e-signature's seal and
    their comment. (text, style): name, text, small, quote."""
    lines = [(p["name"] or NAMELESS, "name")]
    if p.get("on_behalf_of"):
        lines.append((f"for {p['on_behalf_of']}", "small"))
    if p.get("designation"):
        lines.append((p["designation"], "text"))
    decided = _DECIDED.get(p["decision"])
    if decided:
        lines.append((f"{decided} {when(p['signed_at'])}" if p.get("signed_at") else decided, "text"))
    else:
        lines.append(("Handed over" if p["decision"] == "delegated" else "Waiting", "text"))
    if decided and p.get("meaning") and p["meaning"] != stage["label"]:
        lines.append((f"as {p['meaning']}", "small"))
    sig = p.get("signature")
    if sig:
        lines.append((f"Drawn e-signature · seal {sig['seal_short']}" if sig.get("drawn") else f"E-signature · seal {sig['seal_short']}", "small"))
    if p.get("comment"):
        lines.append((f"“{trim(p['comment'])}”", "quote"))
    return lines


def description_rows(record: Mapping[str, Any]) -> List[Tuple[str, str]]:
    """The Document Description: label and value for each row that has one (the title, version and status always)."""
    c = record["control"]
    date = lambda v: day(v) if v else None          # noqa: E731
    rows = [("Title", c["title"], True), ("Document code", c["code"], False), ("Version", c["version"], True), ("Status", c["status"], True),
            ("Document type", c["type_label"], False), ("Classification", (c["classification"] or "").title(), False),
            ("Ownership", c["ownership"], False), ("Approving authority", c["approving_authority"], False),
            ("Date of approval", date(c["date_of_approval"]), False), ("Date of implementation", date(c["date_of_implementation"]), False),
            ("Review frequency", c["review_frequency"], False), ("Last approval date", date(c["last_approval_date"]), False),
            ("Next review date", date(c["next_review_date"]), False)]
    return [(label, value or "-") for label, value, always in rows if value or always]


def revision_rows(record: Mapping[str, Any]) -> List[Tuple[str, bool, str, str, str, str]]:
    """(version, released, date, change type, summary, updated by) of the Revision History, oldest first."""
    return [(r["version"], bool(r["released"]), day(r["date"]), r["change_type"], r["summary"] or "-", r["updated_by"] or "-")
            for r in record["revision_history"]]


def approvers_text(row: Mapping[str, Any]) -> str:
    """The approvers of a version as names, each stage's names followed by the stage's name when it has one: ``Mina Manager (Board)``."""
    groups: List[Tuple[Optional[str], List[str]]] = []
    for a in row["approved_by"]:
        if groups and groups[-1][0] == a.get("stage"):
            groups[-1][1].append(a["name"] or NAMELESS)
        else:
            groups.append((a.get("stage"), [a["name"] or NAMELESS]))
    return ", ".join(", ".join(names) + (f" ({stage})" if stage else "") for stage, names in groups) or "-"


def approval_rows(record: Mapping[str, Any]) -> List[Tuple[str, str, str, str]]:
    """(version, reviewed by, approved by, approval date) of the Approval History, oldest first."""
    return [(r["version"], ", ".join(x["name"] or NAMELESS for x in r["reviewed_by"]) or "-", approvers_text(r), day(r["approval_date"]))
            for r in record["approval_history"]]


def foot_parts(record: Mapping[str, Any]) -> Tuple[str, str]:
    """The line that closes the pages - ``<title> · Document control · generated by the platform <time> · <reference>`` - as the
    title (which gives way when the page is narrow) and the rest."""
    return record["document"]["title"], f" · {TITLE} · generated by the platform {when(record['generated_at'])} · {record['reference']}"


def foot_text(record: Mapping[str, Any]) -> str:
    return "".join(foot_parts(record))


def content_disposition(filename: str) -> str:
    """An attachment header that survives any file name: an ASCII fallback, and the real name encoded (RFC 5987) when it differs."""
    plain = re.sub(r"[^A-Za-z0-9._ -]+", "_", filename).strip() or "document"
    return f'attachment; filename="{plain}"' + ("" if plain == filename else f"; filename*=UTF-8''{quote(filename)}")


def stem(record: Mapping[str, Any]) -> str:
    """The first part of the file name: the document's code, else its title."""
    d = record["document"]
    return re.sub(r"[^\w\-]+", "-", d["code"] or d["title"] or "document", flags=re.UNICODE).strip("-")[:70] or "document"


def fit(png: bytes, box: Tuple[float, float] = SIGN_BOX) -> Optional[Tuple[float, float]]:
    """The size, in the box's unit, a drawn signature takes in a cell (its own shape, as large as the box allows); None when unreadable."""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(png)) as im:
            w, h = im.size
        scale = min(box[0] / w, box[1] / h)
        return w * scale, h * scale
    except Exception:       # noqa: BLE001 - a picture that cannot be read is left out, the pages still print
        return None


def _ink(images: Mapping[int, bytes], p: Mapping[str, Any]) -> Optional[bytes]:
    sig = p.get("signature")
    return images.get(sig["id"]) if sig and p["decision"] == "signed" else None


def _grid_pairs(people: Sequence[Any]) -> List[List[Any]]:
    return [list(people[i:i + 2]) for i in range(0, len(people), 2)]


def _columns(total: float, parts: Sequence[float], floors: Sequence[float]) -> List[float]:
    """The widths of a table's columns: ``total`` shared in the proportions of ``parts``, but no column narrower than its floor (a
    version number and a date must not be cut in two on a narrow page); the other columns give way. Same unit in and out."""
    held: set = set()
    while True:
        free = [i for i in range(len(parts)) if i not in held]
        rest = total - sum(floors[i] for i in held)
        if not free or rest <= 0:                                           # no room to hold them all: as proportioned
            return [total * p / sum(parts) for p in parts]
        widths = [floors[i] if i in held else rest * parts[i] / sum(parts[j] for j in free) for i in range(len(parts))]
        short = [i for i in free if widths[i] < floors[i]]
        if not short:
            return widths
        held.update(short)


# ── PDF ──────────────────────────────────────────────────────────────────────────
def to_pdf(record: Mapping[str, Any], images: Optional[Mapping[int, bytes]] = None, *, slots: Sequence[str] = SLOTS,
           titles: Optional[Mapping[str, Optional[str]]] = None, size: Optional[Tuple[float, float]] = None, first: int = 1,
           of: Optional[int] = None, lead: Optional[str] = None, notes: bool = True) -> bytes:
    """The control pages as a PDF (A4 unless ``size``): the ``slots`` asked for in the order given, each under its title
    (``titles``: a slot's heading text, None for no heading line; the platform's titles by default); ``lead`` is a line of its own
    above them (the document's title, when the pages stand at the top of a document that has none). The foot of every page says
    "Document control page X of Y": X counts from ``first`` and Y is ``of`` (when this PDF is one of several parts) or its own pages.
    ``notes``: what the record cannot tie down, and the line that says the pages are produced by the platform, close the last page."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import CondPageBreak, Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .studio.exporters import _pdf_font

    images, record = images or {}, scrub(record)
    size = size or A4
    margin = min(2.0 * cm, size[0] * 0.09)
    avail = size[0] - 2 * margin
    regular, bold, italic, _bi = _pdf_font()
    ink, grey, light, rule, fill, bar = (colors.HexColor(c) for c in ("#000000", "#404040", "#8C8C8C", "#BFBFBF", "#F2F2F2", "#E0E0E0"))
    base = ParagraphStyle("base", fontName=regular, fontSize=9, leading=12)
    cell = ParagraphStyle("cell", parent=base, fontSize=8.5, leading=10.8)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, leading=9.5, textColor=grey)
    head = ParagraphStyle("head", parent=cell, fontName=bold)
    h1 = ParagraphStyle("h1", parent=base, fontName=bold, fontSize=14, leading=18, spaceBefore=14, spaceAfter=7, keepWithNext=1)
    lead_style = ParagraphStyle("lead", parent=base, fontName=bold, fontSize=18, leading=22, spaceAfter=6)
    centred = {"name": ParagraphStyle("name", parent=cell, fontName=bold, alignment=1, spaceBefore=2),
               "text": ParagraphStyle("ctext", parent=cell, alignment=1), "small": ParagraphStyle("csmall", parent=small, alignment=1),
               "quote": ParagraphStyle("cquote", parent=small, fontName=italic, alignment=1, textColor=ink, spaceBefore=2)}
    title_style = ParagraphStyle("bar", parent=base, fontName=bold, fontSize=9, leading=11.5)
    note_style = ParagraphStyle("note", parent=small, alignment=2, leading=11.5)

    def esc(value: Any) -> str:
        return html.escape(str(value if value is not None else "")).replace("\n", "<br/>")

    def grid(data: List[List[Any]], widths: Sequence[float], floors: Sequence[float], header: bool = True) -> Table:
        t = Table(data, colWidths=_columns(avail, widths, floors), repeatRows=1 if header else 0)
        style = [("GRID", (0, 0), (-1, -1), 0.4, ink), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
        t.setStyle(TableStyle(style + ([("BACKGROUND", (0, 0), (-1, 0), fill)] if header else [])))
        return t

    def title(text: Optional[str]) -> List[Any]:
        return [CondPageBreak(5 * cm), Paragraph(esc(text), h1)] if text else []

    def heading(slot: str) -> List[Any]:
        return title(SLOT_TITLE[slot] if titles is None else titles.get(slot))         # None: the part that goes under a heading the document has

    def description() -> List[Any]:
        t = Table([[Paragraph(esc(k), head), Paragraph(esc(v), cell)] for k, v in description_rows(record)], colWidths=[5.0 * avail / 17.0, 12.0 * avail / 17.0])
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, ink), ("BACKGROUND", (0, 0), (0, -1), fill), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
                               ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]))
        return heading("description") + [t]

    inside = avail / 2 - 12                                            # the room in a cell of the signature grid

    def person_cell(stage: Mapping[str, Any], p: Mapping[str, Any], room_wide: float = inside) -> List[Any]:
        """The drawn signature (or the empty line, with "Awaiting signature" on it) above its rule, then the printed lines;
        ``room_wide``: the width of the cell (a person alone in a row has the whole row)."""
        png = _ink(images, p)
        room = fit(png, SIGN_BOX) if png else None
        if png and room:
            mark: Any = Image(io.BytesIO(png), width=room[0] * cm, height=room[1] * cm)
        else:
            wait = p["decision"] in ("pending", "delegated")
            mark = Paragraph(esc("Awaiting signature") if wait else "&nbsp;", ParagraphStyle("wait", parent=centred["small"], textColor=light))
        line = Table([["", mark, ""]], colWidths=[room_wide * 0.14, room_wide * 0.72, room_wide * 0.14], rowHeights=[1.75 * cm])
        line.setStyle(TableStyle([("ALIGN", (1, 0), (1, 0), "CENTER"), ("VALIGN", (1, 0), (1, 0), "BOTTOM"), ("LINEBELOW", (1, 0), (1, 0), 0.6, ink),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0),
                                  ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        return [Spacer(1, 0.1 * cm), line] + [Paragraph(esc(t), centred[style]) for t, style in person_lines(stage, p)] + [Spacer(1, 0.2 * cm)]

    def stage_table(stage: Mapping[str, Any]) -> Table:
        """A title bar over a grid of two columns of bordered cells, one for each person; the bar is repeated when the grid runs over a page."""
        data: List[List[Any]] = [[Paragraph(esc(stage_title(stage)), title_style), "", "", Paragraph(esc(stage_note(stage)), note_style)]]
        boxes: List[Any] = [("BOX", (0, 0), (-1, 0), 0.6, ink), ("BACKGROUND", (0, 0), (-1, 0), bar), ("SPAN", (0, 0), (2, 0))]
        for r, pair in enumerate(_grid_pairs(stage["people"]), start=1):
            if len(pair) == 1:                                           # alone in the last row: the whole row, not half of it and an empty half
                data.append([person_cell(stage, pair[0], avail - 12), "", "", ""])
                boxes += [("SPAN", (0, r), (3, r)), ("BOX", (0, r), (3, r), 0.6, ink)]
                continue
            data.append([person_cell(stage, pair[0]), "", person_cell(stage, pair[1]), ""])
            boxes += [("SPAN", (0, r), (1, r)), ("SPAN", (2, r), (3, r)), ("BOX", (0, r), (1, r), 0.6, ink), ("BOX", (2, r), (3, r), 0.6, ink)]
        t = Table(data, colWidths=[avail * f for f in (0.25, 0.25, 0.12, 0.38)], repeatRows=1)
        t.setStyle(TableStyle(boxes + [("VALIGN", (0, 0), (-1, -1), "TOP"), ("VALIGN", (0, 0), (-1, 0), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3),
                                       ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
        return t

    def signoff() -> List[Any]:
        out = heading("signoff")
        prep = record["prepared_by"]
        made = [(prep["name"] or "-", "name")] + ([(prep["designation"], "text")] if prep.get("designation") else []) + [
            (f"Version {record['document']['version']} prepared {when(prep['at'])}" if prep.get("at") else f"Version {record['document']['version']}", "small")]
        t = Table([[Paragraph("PREPARED BY", title_style)], [[Paragraph(esc(x), centred[s]) for x, s in made]]], colWidths=[avail], repeatRows=1)
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, ink), ("LINEBELOW", (0, 0), (-1, 0), 0.6, ink), ("BACKGROUND", (0, 0), (-1, 0), bar),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("LEFTPADDING", (0, 0), (-1, -1), 6)]))
        out += [t, Spacer(1, 0.45 * cm)]
        if not record["stages"]:
            out.append(Paragraph("No review or approval is recorded for this version.", small))
        for stage in record["stages"]:
            out += [stage_table(stage), Spacer(1, 0.45 * cm)]
        return out

    def revision() -> List[Any]:
        out = heading("revision")
        data = [[Paragraph(esc(h), head) for h in ("Version", "Date", "Change type", "Summary of change", "Updated by")]]
        for version, released, date, kind, summary, by in revision_rows(record):
            mark = [Paragraph(esc(version), cell)] + ([Paragraph("In progress", small)] if not released else [])
            said = [Paragraph(esc(line), cell) for line in summary.split("\n")]                        # a line for each change
            data.append([mark, Paragraph(esc(date), cell), Paragraph(esc(kind), cell), said, Paragraph(esc(by), cell)])
        out.append(grid(data, [2.2, 2.4, 2.8, 6.6, 3.0], [VERSION_MIN * cm, DATE_MIN * cm, 0, 0, 0]))
        out += title(APPROVAL_TITLE)
        rows = approval_rows(record)
        data = [[Paragraph(esc(h), head) for h in ("Version", "Reviewed by", "Approved by", "Approval date")]]
        data += [[Paragraph(esc(c), cell) for c in row] for row in rows] or [[Paragraph("No version has been approved yet.", cell), "", "", ""]]
        t = grid(data, [2.2, 5.6, 6.2, 3.0], [VERSION_MIN * cm, 0, 0, DATE_MIN * cm])
        if not rows:
            t.setStyle(TableStyle([("SPAN", (0, 1), (-1, 1))]))
        out.append(t)
        return out

    parts = {"description": description, "signoff": signoff, "revision": revision}
    story: List[Any] = [Paragraph(esc(lead), lead_style)] if lead else []
    for slot in slots:
        story += parts[slot]()
    if notes:
        story.append(Spacer(1, 6))
        story += [Paragraph(esc("- " + n), small) for n in record["notes"]]
        story.append(KeepTogether(Paragraph(esc("These pages are produced by the platform from the document's sign-off, signature and audit data. They are not edited by hand."),
                                            ParagraphStyle("end", parent=small, fontName=italic, spaceBefore=4))))
    foot_title, foot_tail = foot_parts(record)

    class Canvas(rl_canvas.Canvas):
        """Two passes, so every page can say how many there are."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._pages: List[dict] = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._chrome(total)
                super().showPage()
            super().save()

        def _chrome(self, total: int) -> None:
            self.setStrokeColor(rule)
            self.setLineWidth(0.5)
            self.setFillColor(grey)
            self.line(margin, 1.7 * cm, size[0] - margin, 1.7 * cm)
            right = f"{TITLE} page {first + self._pageNumber - 1} of {of or first + total - 1}"
            self.setFont(regular, 7.5)

            def wide(s: str) -> float:
                return pdfmetrics.stringWidth(s, regular, 7.5)

            one_line = wide(foot_title + foot_tail) + wide(right) + 12 <= avail
            shown = foot_title
            while len(shown) > 1 and wide(shown + foot_tail) > avail:          # a title too long even for a line of its own is cut
                shown = shown[:-1]
            self.drawString(margin, 1.2 * cm if one_line else 1.4 * cm, (shown if shown == foot_title else shown.rstrip() + "…") + foot_tail)
            self.drawRightString(size[0] - margin, 1.2 * cm if one_line else 1.0 * cm, right)       # else the count goes on the line under it

    out = io.BytesIO()
    d = record["document"]
    SimpleDocTemplate(out, pagesize=size, leftMargin=margin, rightMargin=margin, topMargin=2.0 * cm, bottomMargin=2.2 * cm,
                      title=f"{TITLE} - {d['title']} v{d['version']}", author="").build(story, canvasmaker=Canvas)
    return out.getvalue()


def _is_cover(page) -> bool:
    """Page 1 looks like a cover: little text on it (fewer than COVER_CHARS characters), or it is mostly a picture."""
    import fitz

    if len("".join((page.get_text("text") or "").split())) < COVER_CHARS:
        return True
    covered = 0.0
    for info in page.get_image_info():
        r = fitz.Rect(info["bbox"]) & page.rect
        covered += max(r.width, 0) * max(r.height, 0)
    return covered >= 0.5 * page.rect.width * page.rect.height


def place_in_pdf(data: bytes, record: Mapping[str, Any], images: Optional[Mapping[int, bytes]] = None) -> bytes:
    """A PDF's bytes with the control pages after its cover - page 1, when it looks like one and the file has more pages - else after
    its last page; made at the size of page 1 (the file given is read, never changed). ValueError when it cannot be read or is protected."""
    import fitz

    base = fitz.open(stream=data, filetype="pdf")
    try:
        # a protected file is not rewritten, whether it asks for a password or opens without one (its owner's restrictions would be dropped)
        if base.needs_pass or base.is_encrypted or (base.metadata or {}).get("encryption"):
            raise ValueError("The PDF is protected.")
        first = base[0].rect
        at = 1 if base.page_count > 1 and _is_cover(base[0]) else base.page_count
        size = (first.width, first.height) if first.width >= 380 and first.height >= 500 else None        # a tiny page keeps the pages A4
        try:
            made = to_pdf(record, images, size=size)
        except Exception:       # noqa: BLE001 - a page size the layout cannot use: A4
            made = to_pdf(record, images)
        extra = fitz.open(stream=made, filetype="pdf")
        try:
            base.insert_pdf(extra, start_at=at)
        finally:
            extra.close()
        return base.tobytes(deflate=True)
    finally:
        base.close()


# ── Word ─────────────────────────────────────────────────────────────────────────
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_FONT = {"name": (True, None, 0.0, False), "text": (None, None, 0.0, False), "small": (None, None, -1.5, True), "quote": (None, True, -1.5, False),
         "bold": (True, None, 0.0, False), "plain": (None, None, 0.0, False)}        # style -> (bold, italic, size added, grey)
_GRID = (0.25, 0.25, 0.12, 0.38)           # the signature grid: two cells of two columns each, and the title bar over the first three


def _run(paragraph, text: str, style: str = "plain", size: float = 8.5):
    from docx.shared import Pt, RGBColor

    bold, italic, delta, grey = _FONT.get(style, _FONT["plain"])
    run = paragraph.add_run(text)
    run.font.size = Pt(size + delta)
    run.bold, run.italic = bold, italic
    if grey:
        run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return run


def _title(doc, text: str, size: float = 14, before: float = 14):
    """A heading by hand (never a Heading style: the document's own headings may be numbered, and its contents list reads them)."""
    from docx.shared import Pt

    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.font.size = True, Pt(size)
    p.paragraph_format.space_before, p.paragraph_format.space_after, p.paragraph_format.keep_with_next = Pt(before), Pt(6), True
    return p


def _gap(doc, points: float = 6) -> None:
    """A little empty space between two tables (an empty paragraph with a small mark: Word puts two tables together otherwise)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    props, size = OxmlElement("w:rPr"), OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(points * 2)))
    props.append(size)
    p._p.get_or_add_pPr().append(props)


def _grid(table) -> None:
    """Thin borders on every cell, set on the table itself so a file without a "Table Grid" style still shows them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "808080")):
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    after = {qn("w:shd"), qn("w:tblLayout"), qn("w:tblCellMar"), qn("w:tblLook")}          # Word's schema puts them before these
    nxt = next((c for c in pr if c.tag in after), None)
    (nxt.addprevious(borders) if nxt is not None else pr.append(borders))


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _no_edge(cell, *edges: str) -> None:
    """No line on these sides of a cell ("left", "right"): a title bar and its note read as one band. Before ``_fill``/``_shade``, so the
    cell's properties stay in the order Word's schema has them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:tcBorders")
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    cell._tc.get_or_add_tcPr().append(borders)


def _fill(cell, lines: Sequence[Tuple[str, str]], size: float = 8.5, shade: Optional[str] = None, align: Optional[str] = None, keep: bool = False) -> None:
    """The lines of a cell, each a paragraph (text, style): the style is one of _FONT; ``align``: "centre" or "right"."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    cell.text = ""
    for stray in cell.paragraphs[0].runs:               # `cell.text = ""` leaves an empty run behind
        stray._r.getparent().remove(stray._r)
    for i, (text, style) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(0), Pt(1)
        p.paragraph_format.keep_with_next = True if keep else None
        if align:
            p.alignment = {"centre": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        _run(p, text, style, size)
    if shade:
        _shade(cell, shade)


def _no_split(table, header: bool = False) -> None:
    """A row (a person with a drawn signature) is never cut across two pages; the first row repeats when ``header``."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for r, tr in enumerate(table.rows):
        props = tr._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if header and r == 0:
            mark = OxmlElement("w:tblHeader")
            mark.set(qn("w:val"), "true")
            props.append(mark)


def _widths(table, fractions: Sequence[float], width) -> None:
    """The grid, the cells (a merged one as wide as the columns it covers) and the table all say the same widths (Word and
    LibreOffice read different ones)."""
    from docx.oxml.ns import qn

    for c, fraction in enumerate(fractions):
        table.columns[c].width = int(width * fraction)
    for tr in table._tbl.tr_lst:
        column = 0
        for tc in tr.tc_lst:
            span = tc.grid_span
            tc.width = int(width * sum(fractions[column:column + span]))
            column += span
    total = table._tbl.tblPr.find(qn("w:tblW"))
    if total is not None:
        total.set(qn("w:type"), "dxa")
        total.set(qn("w:w"), str(int(width / 635)))      # twips: 635 EMU each


def _table(doc, headers: Optional[Sequence[str]], rows: Sequence[Sequence[Any]], fractions: Sequence[float], width, shade_first: bool = False,
           whole_row: bool = False) -> None:
    """A table of ``rows`` (each cell a list of (text, style)) as wide as ``width``, under a shaded header row when ``headers`` are
    given (it repeats on every page); the first column is shaded when ``shade_first``; ``whole_row``: its one row is a single cell
    across the table (a line that says there is nothing to list, not a sentence in the narrow first column)."""
    table = doc.add_table(rows=(1 if headers else 0) + len(rows), cols=len(fractions))
    table.autofit = False
    _grid(table)
    top = 1 if headers else 0
    for c, h in enumerate(headers or ()):
        _fill(table.rows[0].cells[c], [(h, "bold")], shade="F2F2F2")
    for r, row in enumerate(rows, start=top):
        for c, value in enumerate(row):
            _fill(table.rows[r].cells[c], value, shade="F2F2F2" if shade_first and c == 0 else None)
    if whole_row:
        cells = table.rows[top].cells
        cells[0].merge(cells[-1])
    _no_split(table, header=bool(headers))
    _widths(table, fractions, width)


def _person_cell(cell, stage: Mapping[str, Any], p: Mapping[str, Any], images: Mapping[int, bytes]) -> None:
    """A person's cell: the drawn signature - or room for it, "Awaiting signature" on it when they have not signed yet - over a
    rule, then their printed lines."""
    from io import BytesIO

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    from .studio.exporters import _border

    png = _ink(images, p)
    cell.text = ""
    first = cell.paragraphs[0]
    for stray in first.runs:                            # `cell.text = ""` leaves an empty run behind
        stray._r.getparent().remove(stray._r)
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.space_before, first.paragraph_format.space_after, first.paragraph_format.keep_with_next = Pt(4), Pt(2), True
    room = fit(png) if png else None
    placed = False
    if room:
        try:
            first.add_run().add_picture(BytesIO(png), width=Cm(room[0]), height=Cm(room[1]))
            placed = True
        except Exception:       # noqa: BLE001 - a picture Word cannot take is left out, the pages still print
            pass
    if not placed:
        first.paragraph_format.space_before = Pt(26)
        if p["decision"] in ("pending", "delegated"):
            run = first.add_run("Awaiting signature")
            run.font.size, run.font.color.rgb = Pt(7.5), RGBColor(0x8C, 0x8C, 0x8C)
    _border(first, "bottom", 6, "000000", 1)
    for text, style in person_lines(stage, p):
        q = cell.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_before, q.paragraph_format.space_after = Pt(0), Pt(1)
        _run(q, text, style)


def _signature_grid(doc, label: str, note: str, people: Sequence[Any], fill_person, width) -> None:
    """A title bar (``label``, and ``note`` at its right when there is one) over a grid of bordered cells, two to a row, one for each
    of ``people``; a person alone in the last row has the whole row (as on the screen), not half of it and an empty half."""
    pairs = _grid_pairs(people)
    table = doc.add_table(rows=1 + len(pairs), cols=4)
    table.autofit = False
    _grid(table)
    bar = table.rows[0].cells
    if note:
        title = bar[0].merge(bar[2])
        _no_edge(title, "right")
        _no_edge(bar[3], "left")
        _fill(title, [(label, "bold")], size=9, shade="E0E0E0", keep=True)
        _fill(bar[3], [(note, "small")], shade="E0E0E0", align="right", keep=True)
    else:
        _fill(bar[0].merge(bar[3]), [(label, "bold")], size=9, shade="E0E0E0", keep=True)
    for r, pair in enumerate(pairs, start=1):
        row = table.rows[r].cells
        if len(pair) == 1:
            fill_person(row[0].merge(row[3]), pair[0])
            continue
        for c, (a, b) in enumerate(((0, 1), (2, 3))):
            fill_person(row[a].merge(row[b]), pair[c])
    _no_split(table, header=True)
    _widths(table, _GRID, width)
    _gap(doc)


def _add_block(doc, record: Mapping[str, Any], images: Mapping[int, bytes], width, *, lead_break: bool = False, trailing_break: bool = False,
               stamp: bool = True) -> None:
    """The control pages' paragraphs and tables, added at the end of ``doc``: the Document Description, the Review and Approval
    grids (a title bar over two columns of bordered cells, each a person: the drawn signature over a rule, then the name,
    designation, what they did and when) and the Revision and Approval History tables. ``lead_break``: they begin on a new page
    of their own; ``trailing_break``: what follows them does; ``stamp``: their last line says by what and when they were made."""
    from docx.shared import Cm

    record = scrub(record)
    first = [True]

    def heading(text: str) -> None:
        p = _title(doc, text)
        if first[0] and lead_break:
            p.paragraph_format.page_break_before = True
        first[0] = False

    def fractions(parts: Sequence[float], floors: Sequence[float]) -> List[float]:
        """The columns of a table as parts of the text width, a version and a date never narrower than they can be read."""
        return [c / width for c in _columns(float(width), parts, [Cm(f) if f else 0 for f in floors])]

    heading(SLOT_TITLE["description"])
    _table(doc, None, [[[(k, "bold")], [(v, "plain")]] for k, v in description_rows(record)], (5 / 17, 12 / 17), width, shade_first=True)
    heading(SLOT_TITLE["signoff"])
    prep = record["prepared_by"]
    made = [(prep["name"] or "-", "name")] + ([(prep["designation"], "text")] if prep.get("designation") else []) + [
        (f"Version {record['document']['version']} prepared {when(prep['at'])}" if prep.get("at") else f"Version {record['document']['version']}", "small")]
    _signature_grid(doc, "PREPARED BY", "", [made], lambda cell, lines: _fill(cell, lines, align="centre"), width)
    for stage in record["stages"]:
        _signature_grid(doc, stage_title(stage), stage_note(stage), stage["people"], lambda cell, p, s=stage: _person_cell(cell, s, p, images), width)
    if not record["stages"]:
        _run(doc.add_paragraph(), "No review or approval is recorded for this version.", "small")
    heading(SLOT_TITLE["revision"])
    rows = [[[(version, "plain")] + ([] if released else [("In progress", "small")]), [(date, "plain")], [(kind, "plain")], [(line, "plain") for line in summary.split("\n")], [(by, "plain")]]
            for version, released, date, kind, summary, by in revision_rows(record)]
    _table(doc, ("Version", "Date", "Change type", "Summary of change", "Updated by"), rows, fractions((2.2, 2.4, 2.8, 6.6, 3.0), (VERSION_MIN, DATE_MIN, 0, 0, 0)), width)
    heading(APPROVAL_TITLE)
    rows = [[[(c, "plain")] for c in row] for row in approval_rows(record)]
    _table(doc, ("Version", "Reviewed by", "Approved by", "Approval date"), rows or [[[("No version has been approved yet.", "plain")], [], [], []]],
           fractions((2.2, 5.6, 6.2, 3.0), (VERSION_MIN, 0, 0, DATE_MIN)), width, whole_row=not rows)
    _gap(doc)
    for note in record["notes"]:
        _run(doc.add_paragraph(), "- " + note, "small", 8.5)
    _run(doc.add_paragraph(), "These pages are produced by the platform from the document's sign-off, signature and audit data. They are not edited by hand.", "quote", 8.5)
    if stamp:
        _run(doc.add_paragraph(), foot_text(record), "small", 8.5)
    if trailing_break:
        doc.add_page_break()


def to_docx(record: Mapping[str, Any], images: Optional[Mapping[int, bytes]] = None) -> bytes:
    """The control pages on their own as a Word document: A4, black headings, the foot says "Document control page X of Y"."""
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Pt, RGBColor

    from .studio.exporters import GREY, _border, _field, _style_doc

    record = scrub(record)
    doc = Document()
    _style_doc(doc)
    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    foot = section.footer.paragraphs[0]
    foot.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    run = foot.add_run(foot_text(record))
    run.font.size, run.font.color.rgb = Pt(7.5), RGBColor.from_string(GREY)
    foot.add_run(f"\t{TITLE} page ").font.size = Pt(7.5)
    _field(foot, "PAGE")
    foot.add_run(" of ").font.size = Pt(7.5)
    _field(foot, "NUMPAGES")
    _border(foot, "top", 6, "BFBFBF", 6)
    doc.core_properties.title = f"{TITLE} - {record['document']['title']} v{record['document']['version']}"
    _add_block(doc, record, images or {}, width, stamp=False)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ── placing the pages in a Word file ─────────────────────────────────────────────
def _tokens(el) -> List[str]:
    """What a paragraph shows, in order: "text" (a word, a picture) and "break" (a page break)."""
    out: List[str] = []
    for r in el.iter(_W + "r"):
        for c in r:
            if c.tag == _W + "br" and c.get(_W + "type") == "page":
                out.append("break")
            elif (c.tag == _W + "t" and (c.text or "").strip()) or c.tag in (_W + "drawing", _W + "pict", _W + "object"):
                out.append("text")
    return out


def _next_section_starts_page(el) -> bool:
    """A paragraph holds the settings of the section that ends with it: the next section's own settings (its type) say whether it
    begins on a page of its own - "continuous" does not."""
    for sibling in el.itersiblings():
        sect = sibling.find(f"{_W}pPr/{_W}sectPr") if sibling.tag == _W + "p" else sibling if sibling.tag == _W + "sectPr" else None
        if sect is not None:
            kind = sect.find(_W + "type")
            return kind is None or kind.get(_W + "val") != "continuous"
    return True


def _boundary(el) -> Optional[str]:
    """"after" when the page ends with this element (a page break at the end of a paragraph, a section break to a new page), "before"
    when a new page begins with it (``pageBreakBefore``, or a paragraph that opens with a page break), else None."""
    if el.tag == _W + "sdt":                    # Word's own cover pages are content controls that end with the page break
        paragraphs = el.findall(f"{_W}sdtContent/{_W}p")
        return "after" if paragraphs and _tokens(paragraphs[-1])[-1:] == ["break"] else None
    if el.tag != _W + "p":
        return None
    props = el.find(_W + "pPr")
    if props is not None:
        if props.find(_W + "sectPr") is not None and _next_section_starts_page(el):
            return "after"
        flag = props.find(_W + "pageBreakBefore")
        if flag is not None and flag.get(_W + "val") not in ("0", "false", "off"):
            return "before"
    tokens = _tokens(el)
    if "break" not in tokens:
        return None
    if tokens[-1] == "break":
        return "after"
    return "before" if "text" not in tokens[:tokens.index("break")] else None


def cover_end(body) -> Optional[Tuple[Any, str]]:
    """Where a Word file's cover page ends: the first page boundary among its first COVER_SCAN elements, as (the element, "after" |
    "before" it), or None when it has none (then it is taken to have no cover)."""
    for i, el in enumerate([c for c in body if c.tag != _W + "sectPr"][:COVER_SCAN]):
        where = _boundary(el)
        if where == "before" and i == 0:                    # an empty first page is not a cover
            continue
        if where:
            return el, where
    return None


def _width_at(doc, following) -> int:
    """The width of the text area (EMU) of the section the pages land in: the one whose settings come next in the file."""
    el = following
    while el is not None:
        sect = el.find(f"{_W}pPr/{_W}sectPr") if el.tag == _W + "p" else el if el.tag == _W + "sectPr" else None
        if sect is not None:
            try:
                size, margin = sect.find(_W + "pgSz"), sect.find(_W + "pgMar")
                return (int(size.get(_W + "w")) - int(margin.get(_W + "left")) - int(margin.get(_W + "right"))) * 635
            except (AttributeError, TypeError, ValueError):
                break
        el = el.getnext()
    section = doc.sections[-1]
    return section.page_width - section.left_margin - section.right_margin


def place_in_docx(data: bytes, record: Mapping[str, Any], images: Optional[Mapping[int, bytes]] = None) -> bytes:
    """A Word file's bytes with the control pages after its cover page, on pages of their own - else, when it has no cover (no page
    boundary among its first elements), after its content. The file given is read, never changed; every element of the original is as
    it was, and the pages take the document's own header, footer and page numbers."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    body = doc.element.body
    spot = cover_end(body)
    if spot is None:
        doc.add_page_break()
        _add_block(doc, record, images or {}, _width_at(doc, None))
    else:
        target, where = spot
        following = target.getnext() if where == "after" else target
        sect = body.find(f"{_W}sectPr")
        last = sect.getprevious() if sect is not None else (body[-1] if len(body) else None)
        # the original goes on from a new page: after the pages, unless what follows them begins one itself
        trailing = where == "after" and following is not None and following.tag != _W + "sectPr" and _boundary(following) != "before"
        _add_block(doc, record, images or {}, _width_at(doc, following), lead_break=where == "before", trailing_break=trailing)
        new, el = [], (last.getnext() if last is not None else body[0])
        while el is not None and el is not sect:
            new.append(el)
            el = el.getnext()
        cursor = target if where == "after" else target.getprevious()
        for el in new:
            cursor.addnext(el)
            cursor = el
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
