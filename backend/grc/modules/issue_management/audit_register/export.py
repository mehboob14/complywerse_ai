"""The monthly pack as the workbook the committee is used to: their SUMMARY
sheet, section by section and row label by row label, filled from `pack()`.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_BOLD = Font(bold=True)
_HEAD = PatternFill("solid", fgColor="D9E1F2")
_THIN = Side(style="thin", color="A6A6A6")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header(sheet, row: int, columns: List[Dict[str, str]], check: bool = True) -> None:
    titles = ["Source "] + [c["title"] for c in columns] + ["Total"]
    for col, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=col, value=title)
        cell.font, cell.fill, cell.border = _BOLD, _HEAD, _BOX
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    if check:
        sheet.cell(row=row, column=len(titles) + 2, value="check").font = _BOLD


def _line(sheet, row: int, label: str, cells: Dict[str, Any], columns, total=None) -> None:
    sheet.cell(row=row, column=1, value=label).border = _BOX
    for col, column in enumerate(columns, start=2):
        sheet.cell(row=row, column=col, value=cells.get(column["source"], 0)).border = _BOX
    sheet.cell(row=row, column=len(columns) + 2,
               value=sum(cells.get(c["source"], 0) for c in columns) if total is None else total
               ).border = _BOX


def build_workbook(data: Dict[str, Any], organisation: str) -> bytes:
    columns = data["columns"]
    width = len(columns) + 2                      # label + sources + total
    sheet = Workbook().active
    sheet.title = "SUMMARY"
    sheet.column_dimensions["A"].width = 46
    for col in range(2, width + 3):
        sheet.column_dimensions[sheet.cell(row=1, column=col).column_letter].width = 16

    as_of = date.fromisoformat(data["as_of"])
    sheet["A1"] = f"{organisation.upper()} AUDIT SERVICES"
    sheet["A2"] = "Summary: Issue & Status     Recommendation & Status"
    sheet["A4"] = data["title"]
    for ref in ("A1", "A2", "A4"):
        sheet[ref].font = _BOLD

    # Roll-forward
    rf = data["roll_forward"]
    by = {m: {r["source"]: r[m] for r in rf["rows"]} for m in ("beginning", "opened", "closed", "ending")}
    _header(sheet, 5, columns)
    _line(sheet, 6, "Beginning Number of Issues", by["beginning"], columns)
    _line(sheet, 7, "Add: Issues opened this month", by["opened"], columns)
    _line(sheet, 8, "Less: Issues closed this month", by["closed"], columns)
    _line(sheet, 9, "Ending Number of Issues ", by["ending"], columns)
    sheet.cell(row=9, column=width + 2,
               value=rf["totals"]["ending"] - (rf["totals"]["beginning"] + rf["totals"]["opened"]
                                               - rf["totals"]["closed"]))

    # Status of issue resolution progress
    res = data["resolution"]
    sheet["A12"] = f"Open Issue status as of {as_of:%B} {as_of.day} {as_of.year}"
    sheet["A13"] = "STATUS OF ISSUE RESOLUTION PROGRESS:"
    sheet["A12"].font = sheet["A13"].font = _BOLD
    lines = [("Not Started", res["not_started"]), ("In Progress", res["in_progress"]),
             ("Delayed /Past Due", res["delayed_past_due"]),
             ("Extended (Audit Committee Approval Required)", res["extended"])]
    if res["unstated"]:
        lines.append(("Other or no status (RM, N/A, blank)", res["unstated"]))
    lines.append(("Closed this month", res["closed_this_month"]))
    row = 14
    for label, value in lines:
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=3, value=value)
        row += 1
    sheet.cell(row=row, column=2, value="Subtotal of Issues ").font = _BOLD
    sheet.cell(row=row, column=3, value=res["subtotal"]).font = _BOLD
    sheet.cell(row=row + 1, column=1, value="LESS:")
    sheet.cell(row=row + 1, column=2, value="Issues Closed this Month:")
    sheet.cell(row=row + 1, column=3, value=-res["less_closed"])
    sheet.cell(row=row + 2, column=1, value="Total Issues as of:").font = _BOLD
    sheet.cell(row=row + 2, column=2, value=as_of)
    sheet.cell(row=row + 2, column=3, value=res["total"]).font = _BOLD

    # Aging
    row += 5
    sheet.cell(row=row, column=1, value="ISSUE AGING ANALYSIS").font = _BOLD
    for offset, bucket in enumerate(data["aging"], start=1):
        if offset == 1:
            sheet.cell(row=row + offset, column=1, value="Days Past Due")
        sheet.cell(row=row + offset, column=2, value=bucket["bucket"])
        sheet.cell(row=row + offset, column=3, value=bucket["total"])

    # Status × source
    row += len(data["aging"]) + 3
    _header(sheet, row, columns)
    for status in data["status_rows"]:
        row += 1
        _line(sheet, row, status["label"], status["by_source"], columns)
    row += 1
    sheet.cell(row=row, column=1, value="List of issues closed this month: ").font = _BOLD
    sheet.cell(row=row, column=3, value="Number of Issue Closed")
    sheet.cell(row=row, column=width, value=len(data["closed"]))
    row += 1
    sheet.cell(row=row, column=1, value="Source").font = _BOLD
    sheet.cell(row=row, column=2, value="Issue #").font = _BOLD
    sheet.cell(row=row, column=3, value="Issue").font = _BOLD
    sheet.cell(row=row, column=width, value="Closed on").font = _BOLD
    titles = {c["source"]: c["title"] for c in columns}
    for item in data["closed"]:
        row += 1
        sheet.cell(row=row, column=1, value=titles.get(item["source"], item["source"]))
        sheet.cell(row=row, column=2, value=item["reference"])
        sheet.cell(row=row, column=3, value=item["title"])
        sheet.cell(row=row, column=width, value=date.fromisoformat(item["closed_on"]))
    row += 1
    total_open = data["total_open"]
    _line(sheet, row, "Total Open Issues", total_open["by_source"], columns)
    sheet.cell(row=row, column=width + 2, value=total_open["check"])
    for col in range(1, width + 1):
        sheet.cell(row=row, column=col).font = _BOLD

    # Recommendations
    row += 4
    _header(sheet, row, columns, check=False)
    _line(sheet, row + 1, "Recommendations (not tracked)", data["recommendations"]["by_source"], columns)

    out = BytesIO()
    sheet.parent.save(out)
    return out.getvalue()


def _summary_skeleton(sheet, organisation: str) -> None:
    """Their SUMMARY sheet's labels, figures left blank to fill in."""
    from .template import SOURCE_BY_SUMMARY_COLUMN, SOURCE_TITLES

    titles = [SOURCE_TITLES[s] for s in dict.fromkeys(SOURCE_BY_SUMMARY_COLUMN.values())]
    header = ["Source "] + titles + ["Total"]
    lines = [
        [f"{organisation.upper()} AUDIT SERVICES".strip()],
        ["Summary: Issue & Status     Recommendation & Status"], [],
        [f"{date.today():%B %Y} Issue Summary"], header,
        ["Beginning Number of Issues"], ["Add: Issues opened this month"],
        ["Less: Issues closed this month"], ["Ending Number of Issues "], [], [],
        ["Open Issue status as of"], ["STATUS OF ISSUE RESOLUTION PROGRESS:"],
        ["Not Started"], ["In Progress"], ["Delayed /Past Due"],
        ["Extended (Audit Committee Approval Required)"], ["", "Subtotal of Issues "],
        ["LESS:", "Issues Closed this Month:"], ["Total Issues as of:"], [], [],
        ["ISSUE AGING ANALYSIS"], ["Days Past Due", "0-30"], ["", "30-60"], ["", "60-90"],
        ["", "90-180"], [], [], header,
        ["Not Started"], ["In Progress"], ["Delayed"], ["Past Due"], ["Extension"],
        ["List of issues closed this month: ", "", "Number of Issue Closed"], [], ["Source"], [], [],
        ["Total Open Issues"], [], [], [], header, ["Recommendations (not tracked)"],
    ]
    for line in lines:
        sheet.append(line)
    for row in sheet.iter_rows():
        if row and row[0].value in ("Source ",):
            for cell in row:
                cell.font, cell.fill, cell.border = _BOLD, _HEAD, _BOX
    sheet.column_dimensions["A"].width = 46


def build_template(sheet_headers: Dict[str, List[str]], definitions: Dict[str, Dict[str, str]],
                   organisation: str) -> bytes:
    """The workbook to fill in: their sheets, their header rows, no findings.

    Each sheet keeps the status definitions it states above its header, so an
    upload of the filled template carries them in again.
    """
    book = Workbook()
    _summary_skeleton(book.active, organisation)
    book.active.title = "SUMMARY"
    for name, headers in sheet_headers.items():
        sheet = book.create_sheet(name[:31])
        sheet.append(["AS OF", None])
        for text in (definitions.get(name) or {}).values():
            sheet.append([text])
        sheet.append([])
        sheet.append(headers)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font, cell.fill, cell.border = _BOLD, _HEAD, _BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.column_dimensions[cell.column_letter].width = max(12, min(40, len(str(cell.value)) + 4))
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    out = BytesIO()
    book.save(out)
    return out.getvalue()
