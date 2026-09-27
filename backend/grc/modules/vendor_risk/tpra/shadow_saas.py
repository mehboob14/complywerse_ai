"""Shadow SaaS: software our people use that no one has assessed, and what to do about it.

Apps arrive from a discovery tool's export (CSV or Excel, in whatever column
names it uses) or straight from Grip Security when its token is connected. Each
is matched to what we already hold — an app already listed, a supplier already
on record — and scored 0–100 from the facts that make it risky here: how many
of our people use it, whether it has MFA and single sign-on, whether files are
shared or uploaded to it, the OAuth permissions it holds, AI processing, and
past breaches. A discovery tool's own score is kept beside ours, not mixed in.

Each app is decided once: onboarded (it becomes a supplier request, its intake
pre-answered only where discovery gives a reason), denied (and, with a Zscaler
connection, added to the URL category the web gateway blocks), or dismissed as
noise. Suppliers named in our own records but missing from the register are
listed alongside, and can be brought in for the same decision.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session

from ....models import GRCUser, IntegrationConnection, TPRAShadowApp, Vendor, get_db
from ....routers.auth_router import get_user_tenants, require_auth
from . import graph, intake, rbac, service

router = APIRouter(tags=["Vendor shadow SaaS"])

STATUSES = ("pending", "onboarded", "denied", "dismissed")
MAX_FILE = 5 * 1024 * 1024
MAX_ROWS = 5000
TEMPLATE = ["name", "domain", "category", "users", "business owner", "owner email", "risk score", "mfa",
            "sso %", "file sharing", "upload bytes", "breaches (3 years)", "risk types", "description"]
# What discovery tools call each field, lower-cased with punctuation as spaces.
COLUMNS: Dict[str, Tuple[str, ...]] = {
    "name": ("name", "application", "app", "app name", "application name", "vendor", "vendor name", "saas", "saas name"),
    "domain": ("domain", "url", "website", "vendor domain", "app url", "application url"),
    "category": ("category", "application category", "app category"),
    "users": ("users", "user count", "number of users", "numberofusers", "active users", "users count"),
    "owner_name": ("business owner", "owner", "relationship manager", "primary contact", "app owner"),
    "owner_email": ("owner email", "business owner email", "contact email", "primary contact email"),
    "source_risk": ("risk score", "risk", "riskscore", "grip risk score"),
    "mfa": ("mfa", "mfa support", "mfasupport", "mfa supported"),
    "sso": ("sso", "sso %", "sso percentage", "ssopercentage", "sso pct"),
    "file_sharing": ("file sharing", "filesharing", "file share"),
    "upload_bytes": ("upload bytes", "uploadbytes", "upload", "uploaded"),
    "breaches": ("breaches (3 years)", "breaches", "breaches in three years", "breaches in 3 years"),
    "risk_types": ("risk types", "risk type", "risks"),
    "description": ("description", "notes", "purpose"),
}
_LOOKUP = {alias: key for key, aliases in COLUMNS.items() for alias in aliases}


def _norm(text) -> str:
    return " ".join(re.sub(r"[^a-z0-9%()]+", " ", str(text or "").lower()).split())


def _bool(value) -> Optional[bool]:
    text = str(value if value is not None else "").strip().lower()
    return True if text in ("yes", "y", "true", "1", "supported", "enabled") else \
        False if text in ("no", "n", "false", "0", "not supported", "disabled", "none") else None


def _int(value) -> Optional[int]:
    found = re.search(r"-?\d[\d,]*", str(value if value is not None else ""))
    return int(found.group(0).replace(",", "")) if found else None


def _bytes(value) -> Optional[int]:
    text = str(value if value is not None else "").strip().lower().replace(",", "")
    found = re.match(r"^(\d+(?:\.\d+)?)\s*(b|kb|mb|gb|tb)?$", text)
    if not found:
        return None
    scale = {"b": 1, "kb": 1024, "mb": 1024 ** 2, "gb": 1024 ** 3, "tb": 1024 ** 4}[found.group(2) or "b"]
    return int(float(found.group(1)) * scale)


def _host(value) -> Optional[str]:
    host = intake.host(str(value or "")).split(":")[0].strip(".")
    return host if "." in host and " " not in host else None


# ── from a row, from Grip ────────────────────────────────────────────────────

def from_row(raw: dict) -> Optional[dict]:
    """A row of a discovery export, in our shape; None when it names no app."""
    row = {_LOOKUP[_norm(k)]: v for k, v in raw.items() if _norm(k) in _LOOKUP and v not in (None, "")}
    name = " ".join(str(row.get("name") or "").split())[:255]
    if not name:
        return None
    risk_types = [t.strip() for t in re.split(r"[;,|]", str(row.get("risk_types") or "")) if t.strip()][:8]
    sso = _int(row.get("sso"))
    return {
        "name": name, "domain": _host(row.get("domain")), "category": (str(row.get("category") or "").strip() or None),
        "description": (str(row.get("description") or "").strip()[:2000] or None), "users": _int(row.get("users")),
        "owner_name": (str(row.get("owner_name") or "").strip()[:255] or None),
        "owner_email": (str(row.get("owner_email") or "").strip().lower()[:255] or None),
        "source_risk": _clamp(_int(row.get("source_risk"))),
        "facts": {k: v for k, v in {
            "mfa": _bool(row.get("mfa")), "sso_pct": sso if sso is not None and 0 <= sso <= 100 else None,
            "file_sharing": _bool(row.get("file_sharing")), "upload_bytes": _bytes(row.get("upload_bytes")),
            "breaches_3y": _int(row.get("breaches")), "risk_types": risk_types or None}.items() if v is not None},
    }


def from_grip(app: dict) -> Optional[dict]:
    g = app.get("gripData") if isinstance(app.get("gripData"), dict) else {}
    name = " ".join(str(app.get("name") or "").split())[:255]
    if not name:
        return None
    oauth = g.get("oauthScopes") if isinstance(g.get("oauthScopes"), dict) else {}
    contact = g.get("primaryContact") if isinstance(g.get("primaryContact"), dict) else {}
    ai = g.get("aiDepth") if isinstance(g.get("aiDepth"), dict) else {}
    mfa = g.get("mfaSupported")
    return {
        "name": name, "domain": _host(app.get("url")), "category": g.get("category"), "external_id": str(app.get("id") or "") or None,
        "description": (ai.get("description") or None), "users": _int(g.get("numberOfUsers")),
        "owner_name": contact.get("name") or g.get("businessOwner") or g.get("buisnessOwner"),
        "owner_email": (contact.get("email") or "").lower() or None, "source_risk": _clamp(_int(g.get("riskScore"))),
        "facts": {k: v for k, v in {
            "mfa": mfa if isinstance(mfa, bool) else _bool(mfa), "sso_pct": _int(g.get("SSOPercentage")),
            "oauth_high": _int((oauth.get("high") or {}).get("count")) if isinstance(oauth.get("high"), dict) else None,
            "ai": (str(ai.get("level")) if ai.get("level") and str(ai.get("level")).lower() != "none" else None),
            "compliances": [str(c) for c in g.get("compliances") or []][:10] or None}.items() if v is not None},
    }


def _clamp(value: Optional[int]) -> Optional[int]:
    return None if value is None else max(0, min(100, value))


# ── our score ────────────────────────────────────────────────────────────────

def score(users: Optional[int], facts: dict) -> Tuple[int, List[str]]:
    """0–100 and the reasons, from the facts that make an unassessed app risky here."""
    points, why = 0, []

    def add(n: int, reason: str):
        nonlocal points
        points += n
        why.append(reason)

    if users:
        add(20 if users >= 100 else 12 if users >= 20 else 6 if users >= 5 else 2,
            f"{users} of our people use it")
    if facts.get("mfa") is False:
        add(15, "No multi-factor sign-in")
    elif facts.get("mfa") is None:
        add(5, "Multi-factor sign-in not known")
    if facts.get("sso_pct") is not None and facts["sso_pct"] < 50:
        add(5, f"Only {facts['sso_pct']}% sign in through single sign-on")
    if facts.get("file_sharing"):
        add(10, "Files are shared through it")
    uploaded = facts.get("upload_bytes") or 0
    if uploaded >= 1024 ** 3:
        add(15, f"{uploaded / 1024 ** 3:.1f} GB uploaded to it")
    elif uploaded >= 100 * 1024 ** 2:
        add(8, f"{uploaded / 1024 ** 2:.0f} MB uploaded to it")
    if facts.get("oauth_high"):
        add(15, f"Holds {facts['oauth_high']} high-risk OAuth permission{'s' if facts['oauth_high'] != 1 else ''}")
    if facts.get("ai"):
        add(10, f"Puts data through AI ({facts['ai']})")
    if facts.get("breaches_3y"):
        add(min(20, 10 * facts["breaches_3y"]), f"{facts['breaches_3y']} breach{'es' if facts['breaches_3y'] != 1 else ''} in three years")
    if facts.get("risk_types"):
        add(min(20, 5 * len(facts["risk_types"])), "Flagged by discovery: " + ", ".join(facts["risk_types"]))
    return min(100, points), why


def level(risk: Optional[int]) -> str:
    return "high" if (risk or 0) >= 60 else "medium" if (risk or 0) >= 35 else "low"


# ── bringing apps in ─────────────────────────────────────────────────────────

def _known(db: Session, tenant_id: int) -> Tuple[Dict[str, Vendor], Dict[str, Vendor]]:
    by_host, by_name = {}, {}
    for v in db.query(Vendor).filter(Vendor.tenant_id == tenant_id, Vendor.deleted_at.is_(None)):
        by_name[v.name.strip().lower()] = v
        for d in [v.website, *(v.domains or [])]:
            host = _host(d)
            if host:
                by_host[host] = v
    return by_host, by_name


def bring_in(db: Session, tenant_id: int, rows: List[dict], source: str, dry_run: bool,
             now: Optional[datetime] = None) -> dict:
    """Add new apps and refresh known ones. Apps already on the supplier register are
    reported, not added; a decision already taken is kept."""
    now = now or datetime.utcnow()
    by_host, by_name = _known(db, tenant_id)
    apps = db.query(TPRAShadowApp).filter(TPRAShadowApp.tenant_id == tenant_id).all()
    by_ext = {(a.source, a.external_id): a for a in apps if a.external_id}
    by_domain = {a.domain: a for a in apps if a.domain}
    by_app_name = {a.name.lower(): a for a in apps}
    added, refreshed, suppliers = [], [], []
    for row in rows:
        vendor = by_host.get(row["domain"] or "") or by_name.get(row["name"].lower())
        if vendor is not None:
            suppliers.append({"name": row["name"], "supplier": vendor.name, "vendor_id": vendor.id})
            continue
        app = (by_ext.get((source, row.get("external_id"))) if row.get("external_id") else None) \
            or by_domain.get(row["domain"] or "") or by_app_name.get(row["name"].lower())
        risk, _ = score(row.get("users"), row.get("facts") or {})
        if app is None:
            added.append(row["name"])
            if not dry_run:
                app = TPRAShadowApp(tenant_id=tenant_id, source=source, first_seen=now, status="pending")
                db.add(app)
                by_app_name[row["name"].lower()] = app
                if row["domain"]:
                    by_domain[row["domain"]] = app
        else:
            refreshed.append(row["name"])
        if dry_run or app is None:
            continue
        for key in ("name", "domain", "category", "description", "users", "owner_name", "owner_email", "source_risk",
                    "external_id"):
            if row.get(key) is not None:
                setattr(app, key, row[key])
        app.facts = {**(app.facts or {}), **(row.get("facts") or {})}
        app.risk = score(app.users, app.facts)[0]
        app.last_seen = now
    return {"new": len(added), "updated": len(refreshed), "suppliers": suppliers[:200],
            "names": {"new": added[:200], "updated": refreshed[:200]}}


def _read(filename: str, content: bytes) -> List[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        values = list(sheet.iter_rows(values_only=True))
        header = [str(h or "") for h in (values[0] if values else [])]
        return [{header[i]: r[i] for i in range(min(len(header), len(r)))} for r in values[1:]
                if any(c not in (None, "") for c in r)]
    return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))


# ── connections ──────────────────────────────────────────────────────────────

def connection(db: Session, tenant_id: int, category: str, integration: str) -> Optional[dict]:
    """A connected integration's settings and secrets together, or None."""
    from ....services.connector_credentials import decrypt_credentials

    row = (db.query(IntegrationConnection).filter(
        IntegrationConnection.tenant_id == tenant_id, IntegrationConnection.category == category,
        IntegrationConnection.integration_type == integration, IntegrationConnection.is_active.is_(True)).first())
    if row is None:
        return None
    return {**(row.provider_config or {}), **(decrypt_credentials(row.encrypted_credentials) or {}),
            **({"base_url": row.console_url} if row.console_url and not (row.provider_config or {}).get("base_url") else {})}


def _gateway(db: Session, tenant_id: int):
    from ...connectors.providers.saas_governance import ZscalerSession

    creds = connection(db, tenant_id, "web_gateway", "zscaler")
    return (lambda: ZscalerSession(creds)) if creds else None


# ── REST ─────────────────────────────────────────────────────────────────────

def _tids(user: GRCUser, db: Session) -> List[int]:
    tids = get_user_tenants(user, db)
    if not tids:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "No tenant context")
    return tids


def _app(db: Session, app_id: int, tids: List[int]) -> TPRAShadowApp:
    app = db.query(TPRAShadowApp).filter(TPRAShadowApp.id == app_id, TPRAShadowApp.tenant_id.in_(tids)).first()
    if app is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "App not found")
    return app


def _row(a: TPRAShadowApp, names: Dict[int, str]) -> dict:
    risk, why = score(a.users, a.facts or {})
    return {"id": a.id, "name": a.name, "domain": a.domain, "category": a.category, "description": a.description,
            "users": a.users, "owner_name": a.owner_name, "owner_email": a.owner_email, "source": a.source,
            "source_risk": a.source_risk, "risk": risk, "level": level(risk), "why": why, "facts": a.facts or {},
            "status": a.status, "vendor_id": a.vendor_id, "blocked": bool(a.blocked), "decided_by": names.get(a.decided_by),
            "decided_at": a.decided_at, "decision_note": a.decision_note, "first_seen": a.first_seen, "last_seen": a.last_seen}


def _names(db: Session, ids) -> Dict[int, str]:
    ids = {i for i in ids if i}
    if not ids:
        return {}
    return {u.id: u.display_name or u.username or u.email for u in db.query(GRCUser).filter(GRCUser.id.in_(ids))}


@router.get("/shadow-saas")
def board(status_: Optional[str] = Query(None, alias="status"), db: Session = Depends(get_db),
          user: GRCUser = Depends(require_auth)):
    tid = _tids(user, db)[0]
    apps = db.query(TPRAShadowApp).filter(TPRAShadowApp.tenant_id == tid).all()
    counts = {k: 0 for k in STATUSES}
    for a in apps:
        counts[a.status] = counts.get(a.status, 0) + 1
    kept = [a for a in apps if not status_ or a.status == status_]
    names = _names(db, [a.decided_by for a in kept])
    rows = sorted((_row(a, names) for a in kept), key=lambda r: (-r["risk"], -(r["users"] or 0), r["name"].lower()))
    pending = [r for r in rows if r["status"] == "pending"] if not status_ or status_ == "pending" else []
    listed = {a.name.lower() for a in apps}
    return {
        "items": rows, "counts": counts,
        "pending_high": sum(1 for r in pending if r["level"] == "high"),
        "people_on_pending": sum(r["users"] or 0 for r in pending),
        "blocked": sum(1 for a in apps if a.blocked),
        "records": [r for r in graph.shadow_suppliers(db, tid, 50) if r["name"].lower() not in listed],
        "connected": {"grip": connection(db, tid, "saas_discovery", "grip") is not None,
                      "zscaler": connection(db, tid, "web_gateway", "zscaler") is not None},
    }


@router.get("/shadow-saas/import/template")
def template(user: GRCUser = Depends(require_auth)):
    return Response(",".join(TEMPLATE) + "\r\n", media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="shadow-saas-template.csv"'})


@router.post("/shadow-saas/import")
async def import_apps(file: UploadFile = File(...), dry_run: bool = Query(True), db: Session = Depends(get_db),
                      user: GRCUser = Depends(require_auth)):
    tid = _tids(user, db)[0]
    rbac.require_write(db, user, "vendors", "edit")
    content = await file.read(MAX_FILE + 1)
    if len(content) > MAX_FILE:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "The file is larger than 5 MB")
    try:
        raw = _read(file.filename or "", content)
    except Exception:  # noqa: BLE001 — a file that will not parse is the user's to fix
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The file could not be read as CSV or Excel")
    if len(raw) > MAX_ROWS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Split the file: at most {MAX_ROWS} rows at a time")
    rows = [r for r in (from_row(x) for x in raw) if r]
    if raw and not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No row names an app. The first row must be the column names, "
                                                         "with one called name or application.")
    result = bring_in(db, tid, rows, "csv", dry_run)
    if not dry_run:
        service.write_audit(db, tid, entity="shadow_app", action="import", actor_id=user.id,
                            to_value=f"{result['new']} new, {result['updated']} updated", reason=file.filename)
        db.commit()
    return {**result, "rows": len(raw), "skipped": len(raw) - len(rows), "dry_run": dry_run}


@router.post("/shadow-saas/sync")
def sync_grip(db: Session = Depends(get_db), user: GRCUser = Depends(require_auth)):
    from ...connectors.providers.saas_governance import grip_apps

    tid = _tids(user, db)[0]
    rbac.require_write(db, user, "vendors", "edit")
    creds = connection(db, tid, "saas_discovery", "grip")
    if not creds or not creds.get("api_token") or not creds.get("base_url"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Connect Grip Security under Admin → Connectors first")
    try:
        found = grip_apps(creds["base_url"], creds["api_token"])
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status.HTTP_502_BAD_GATEWAY, f"Grip could not be read: {exc}"[:300])
    result = bring_in(db, tid, [r for r in (from_grip(a) for a in found) if r], "grip", dry_run=False)
    service.write_audit(db, tid, entity="shadow_app", action="sync", actor_id=user.id,
                        to_value=f"{result['new']} new, {result['updated']} updated")
    db.commit()
    return result


class FromRecordsIn(BaseModel):
    name: str = Field(..., min_length=1, max_length=255)


@router.post("/shadow-saas/from-records", status_code=status.HTTP_201_CREATED)
def from_records(body: FromRecordsIn, db: Session = Depends(get_db), user: GRCUser = Depends(require_auth)):
    """A supplier our own records name but the register does not, brought in for a decision."""
    tid = _tids(user, db)[0]
    rbac.require_write(db, user, "vendors", "edit")
    match = next((r for r in graph.shadow_suppliers(db, tid, 500) if r["name"].lower() == body.name.strip().lower()), None)
    if match is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Our records do not name that supplier")
    app = TPRAShadowApp(tenant_id=tid, name=match["name"], source="records", status="pending",
                        description=f"Named in our {', '.join(match['sources'])} on {match['asset_count']} asset(s)"
                                    + (f": {', '.join(match['products'])}" if match["products"] else ""),
                        facts={}, risk=score(None, {})[0])
    db.add(app)
    db.flush()
    service.write_audit(db, tid, entity="shadow_app", action="create", entity_id=app.id, actor_id=user.id, to_value=app.name)
    db.commit()
    return _row(app, {})


class DecisionIn(BaseModel):
    note: Optional[str] = Field(None, max_length=2000)
    block: bool = False


def _decide(db: Session, app: TPRAShadowApp, user: GRCUser, to: str, note: str) -> None:
    before = app.status
    app.status, app.decided_by, app.decided_at, app.decision_note = to, user.id, datetime.utcnow(), note or None
    service.write_audit(db, app.tenant_id, entity="shadow_app", action="decide", entity_id=app.id, vendor_id=app.vendor_id,
                        actor_id=user.id, from_value=before, to_value=to, reason=note or None)


def _intake(app: TPRAShadowApp) -> dict:
    """Answers discovery gives a reason for; everything else is left for the requester."""
    facts, answers, why = app.facts or {}, {"engagement_type": "saas"}, {}
    purpose = app.description or (f"{app.category} tool" if app.category else None)
    if purpose:
        answers["purpose"] = f"Already in use without assessment: {purpose}"[:500]
    if app.users:
        answers["users_count"] = app.users
    if facts.get("file_sharing") or (facts.get("upload_bytes") or 0) > 0:
        answers["hosts_data"] = "yes"
        why["hosts_data"] = "Discovery saw our people upload or share files through it."
    if facts.get("ai"):
        answers["uses_ai"] = "yes"
        why["uses_ai"] = f"Discovery reports it processes data with AI ({facts['ai']})."
    if facts.get("sso_pct") is not None:
        answers["sso"] = "yes" if facts["sso_pct"] > 0 else "no"
    return {"answers": answers, "justifications": why}


@router.post("/shadow-saas/{app_id}/onboard")
def onboard(app_id: int, body: DecisionIn = DecisionIn(), db: Session = Depends(get_db), user: GRCUser = Depends(require_auth)):
    """The app becomes a supplier request, pre-answered where discovery gives a reason."""
    tids = _tids(user, db)
    app = _app(db, app_id, tids)
    rbac.require_write(db, user, "intake", "create")
    if app.status != "pending":
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"This app was already {app.status}")
    by_host, by_name = _known(db, app.tenant_id)
    existing = by_host.get(app.domain or "") or by_name.get(app.name.lower())
    if existing is not None:
        raise HTTPException(status.HTTP_409_CONFLICT, f"{existing.name} is already on record (ID {existing.id})")
    owner = None
    if app.owner_email:
        owner = db.query(GRCUser.id).filter(func.lower(GRCUser.email) == app.owner_email, GRCUser.is_active.is_(True)).scalar()
    answers = _intake(app)
    v = Vendor(tenant_id=app.tenant_id, name=app.name, website=f"https://{app.domain}" if app.domain else None, tier="medium",
               status="requested", intake_status="draft", requested_by=user.id, owner_id=owner or user.id,
               stakeholder_ids=[owner] if owner and owner != user.id else None,
               description=app.description, intake=intake.clean(answers["answers"], answers["justifications"]))
    db.add(v)
    db.flush()
    app.vendor_id = v.id
    service.write_audit(db, v.tenant_id, entity="request", action="create", vendor_id=v.id, actor_id=user.id,
                        to_value="draft", reason=f"From shadow SaaS: {app.name}")
    _decide(db, app, user, "onboarded", " ".join((body.note or "").split()))
    db.commit()
    return {**_row(app, _names(db, [user.id])), "request_id": v.id}


@router.post("/shadow-saas/{app_id}/deny")
def deny(app_id: int, body: DecisionIn, db: Session = Depends(get_db), user: GRCUser = Depends(require_auth)):
    """Not to be used. With a Zscaler connection, its domain can be added to the blocked category."""
    tids = _tids(user, db)
    app = _app(db, app_id, tids)
    rbac.require_write(db, user, "vendors", "edit")
    note = " ".join((body.note or "").split())
    if app.status != "pending":
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"This app was already {app.status}")
    if len(note) < 5:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Say why it is denied")
    _decide(db, app, user, "denied", note)
    block_error = None
    if body.block:
        gateway = _gateway(db, app.tenant_id)
        if not app.domain:
            block_error = "It has no domain to block"
        elif gateway is None:
            block_error = "Connect Zscaler under Admin → Connectors to block it"
        else:
            try:
                with gateway() as z:
                    z.block(app.domain)
                app.blocked, app.blocked_at = True, datetime.utcnow()
                service.write_audit(db, app.tenant_id, entity="shadow_app", action="block", entity_id=app.id,
                                    actor_id=user.id, to_value=app.domain)
            except Exception as exc:  # noqa: BLE001 — the decision stands; the block is reported as not done
                block_error = f"Zscaler did not take it: {exc}"[:300]
    db.commit()
    return {**_row(app, _names(db, [user.id])), "block_error": block_error}


@router.post("/shadow-saas/{app_id}/dismiss")
def dismiss(app_id: int, body: DecisionIn, db: Session = Depends(get_db), user: GRCUser = Depends(require_auth)):
    tids = _tids(user, db)
    app = _app(db, app_id, tids)
    rbac.require_write(db, user, "vendors", "edit")
    note = " ".join((body.note or "").split())
    if app.status != "pending":
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"This app was already {app.status}")
    if len(note) < 5:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Say why it can be ignored")
    _decide(db, app, user, "dismissed", note)
    db.commit()
    return _row(app, _names(db, [user.id]))


@router.post("/shadow-saas/{app_id}/reopen")
def reopen(app_id: int, db: Session = Depends(get_db), user: GRCUser = Depends(require_auth)):
    """Back to be decided; a blocked domain is taken off the gateway's list."""
    tids = _tids(user, db)
    app = _app(db, app_id, tids)
    rbac.require_write(db, user, "vendors", "edit")
    if app.status in ("pending", "onboarded"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Only a denied or dismissed app can be reopened")
    if app.blocked:
        gateway = _gateway(db, app.tenant_id)
        if gateway is None:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "It is blocked in Zscaler; connect Zscaler to unblock it first")
        try:
            with gateway() as z:
                z.allow(app.domain)
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status.HTTP_502_BAD_GATEWAY, f"Zscaler did not unblock it: {exc}"[:300])
        app.blocked, app.blocked_at = False, None
        service.write_audit(db, app.tenant_id, entity="shadow_app", action="unblock", entity_id=app.id,
                            actor_id=user.id, from_value=app.domain)
    before = app.status
    app.status, app.decided_by, app.decided_at, app.decision_note = "pending", None, None, None
    service.write_audit(db, app.tenant_id, entity="shadow_app", action="reopen", entity_id=app.id, actor_id=user.id,
                        from_value=before, to_value="pending")
    db.commit()
    return _row(app, {})
