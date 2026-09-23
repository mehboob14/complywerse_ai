"""A questionnaire as a workbook, for vendors who answer offline.

The workbook says which questionnaire and which template version it was made
from, so an upload is matched to the right questionnaire and refused if it was
made for another one. Answers are read by the labels the vendor saw; one that is
not among a question's answers is refused with the row it is on, rather than
guessed at. The attestation travels on its own sheet.
"""
from __future__ import annotations

import io
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

FORMAT = "complyverse-questionnaire-1"
FIRST_ROW = 5
MAX_BYTES = 5 * 1024 * 1024
_CONFIRM = "I confirm these answers are accurate and complete, and that I am authorised to give them"


def _options(question: dict) -> List[dict]:
    if (question.get("type") or "yes_no") == "rating":
        return [{"value": str(n), "label": str(n)} for n in range(1, 6)]
    return [o if isinstance(o, dict) else {"value": str(o), "label": str(o)} for o in question.get("options") or []]


def _label(question: dict, value) -> str:
    if value is None:
        return ""
    for option in _options(question):
        if str(option.get("value")).lower() == str(value).lower():
            return str(option.get("label") or option.get("value"))
    return str(value)


def build(qr, questions: List[dict], vendor_name: str, locked: Optional[Dict[str, str]] = None) -> bytes:
    """The workbook for one questionnaire, prefilled with the answers so far."""
    locked = locked or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Questionnaire"
    ws["A1"] = f"Questionnaire for {vendor_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Answer in the Answer column, choosing from the list where there is one. "
                "Change only the Answer and Comment columns. Fill in the Attestation sheet, then upload the file.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 32
    for col, title in enumerate(["Key", "#", "Question", "Answer", "Comment"], start=1):
        cell = ws.cell(row=FIRST_ROW - 1, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF4")
    answers, comments = qr.responses or {}, qr.vendor_comments or {}
    for i, q in enumerate(questions):
        row = FIRST_ROW + i
        key = str(q.get("id"))
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=i + 1)
        text = str(q.get("text") or key)
        if key in locked:
            text += "  [answered by your certificate]"
        ws.cell(row=row, column=3, value=text).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=4, value=_label(q, answers.get(key)))
        ws.cell(row=row, column=5, value=comments.get(key))
        labels = [str(o.get("label") or o.get("value")) for o in _options(q)]
        if labels and (q.get("type") or "yes_no") != "text":
            listing = ",".join(label.replace(",", " ") for label in labels)
            if len(listing) < 250:                     # Excel's limit on an inline list
                dv = DataValidation(type="list", formula1=f'"{listing}"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(ws.cell(row=row, column=4))
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 40

    att = wb.create_sheet("Attestation")
    rows = [("Name", qr.attested_name), ("Job title", qr.attested_title), ("Email", qr.attested_email),
            (_CONFIRM, "No")]
    for r, (label, value) in enumerate(rows, start=1):
        att.cell(row=r, column=1, value=label).font = Font(bold=True)
        att.cell(row=r, column=2, value=value)
    confirm = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    att.add_data_validation(confirm)
    confirm.add(att.cell(row=4, column=2))
    att.column_dimensions["A"].width = 60
    att.column_dimensions["B"].width = 40

    meta = wb.create_sheet("_meta")
    for r, (k, v) in enumerate([("format", FORMAT), ("questionnaire_id", qr.id),
                                ("template_version_id", qr.template_version_id)], start=1):
        meta.cell(row=r, column=1, value=k)
        meta.cell(row=r, column=2, value=v)
    meta.sheet_state = "hidden"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def read(content: bytes, qr, questions: List[dict]) -> dict:
    """Answers, comments and attestation from a completed workbook. Raises
    ValueError saying what is wrong, row by row where it can."""
    if len(content) > MAX_BYTES:
        raise ValueError("The workbook is larger than 5 MB")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:  # noqa: BLE001 — anything openpyxl cannot open
        raise ValueError("This is not a workbook made from the questionnaire")
    if "_meta" not in wb.sheetnames or "Questionnaire" not in wb.sheetnames:
        raise ValueError("This is not a workbook made from the questionnaire")
    meta = {str(r[0]): r[1] for r in wb["_meta"].iter_rows(values_only=True) if r and r[0]}
    if meta.get("format") != FORMAT or str(meta.get("questionnaire_id")) != str(qr.id):
        raise ValueError("This workbook was made for a different questionnaire")
    if str(meta.get("template_version_id")) != str(qr.template_version_id):
        raise ValueError("This workbook was made from a different version of the questionnaire. "
                         "Download it again and copy your answers across.")

    by_key = {str(q.get("id")): q for q in questions}
    answers: Dict[str, str] = {}
    comments: Dict[str, str] = {}
    problems: List[str] = []
    for row_no, row in enumerate(wb["Questionnaire"].iter_rows(min_row=FIRST_ROW, values_only=True), start=FIRST_ROW):
        key = str(row[0]).strip() if row and row[0] is not None else ""
        if key not in by_key:
            continue
        q = by_key[key]
        raw = "" if len(row) < 4 or row[3] is None else str(row[3]).strip()
        comment = "" if len(row) < 5 or row[4] is None else str(row[4]).strip()
        if comment:
            comments[key] = comment[:4000]
        if not raw:
            continue
        if (q.get("type") or "yes_no") == "text":
            answers[key] = raw[:10000]
            continue
        match = next((o for o in _options(q)
                      if raw.lower() in (str(o.get("label") or "").lower(), str(o.get("value")).lower())), None)
        if match is None:
            problems.append(f"Row {row_no}: '{raw[:60]}' is not one of the answers to question {row[1] or key}")
        else:
            answers[key] = str(match.get("value"))
    if problems:
        raise ValueError("; ".join(problems[:10]))

    attestation = {}
    if "Attestation" in wb.sheetnames:
        values = [r[1] if r and len(r) > 1 else None for r in wb["Attestation"].iter_rows(max_row=4, values_only=True)]
        values += [None] * (4 - len(values))
        attestation = {"name": values[0] or "", "title": values[1] or "", "email": values[2] or "",
                       "confirm": str(values[3] or "").strip().lower() == "yes"}
    return {"responses": answers, "comments": comments, "attestation": attestation}
