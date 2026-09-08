"""Build the NIST CSF 2.0 framework library from NIST's own CPRT export.

Our library was CSF v1.1 (ID.BE-2, PR.AC-1). SCF crosswalks CSF 2.0 only
(GV.OC-01, DE.AE-02), and 2.0 renumbered the core and dissolved whole v1.1
categories into new ones — so the old library resolved to zero against
`nist_csf_20` and could not be machine-migrated. This rebuilds it from source.

Source: the CSF 2.0 Reference Tool export, a NIST publication and therefore a
US Government work in the public domain (17 U.S.C. §105) — unlike ISO/PCI/CIS,
the requirement text here is ours to reproduce verbatim.

    curl -L -o csf20.xlsx \
      "https://csrc.nist.gov/extensions/nudp/services/json/csf/download?olirids=all"
    python -m grc.tools.build_nist_csf_library --source csf20.xlsx

The export carries the v1.1 core alongside 2.0: retired categories and
subcategories sit inline under their function, each tagged by NIST with
"[Withdrawn: Incorporated into ...]". We filter on that marker rather than on
row position, because withdrawal is per-subcategory — ID.AM-06 is withdrawn
while the rest of the live ID.AM category is not. 185 rows − 79 withdrawn = 106.

Nothing here is paraphrased or AI-authored: subcategory statements and
implementation examples are copied verbatim from the export.
"""
from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "seed_data" / "frameworks" / "nist_csf.json"
SCF_MAPPINGS = ROOT / "seed_data" / "scf" / "mappings.csv.gz"

# CSF 2.0 core shape, from NIST CSWP 29. These are assertions, not guesses: if
# a future export disagrees, the build should fail loudly rather than ship a
# library that silently gained or lost requirements.
EXPECT_FUNCTIONS = 6
EXPECT_CATEGORIES = 22
EXPECT_SUBCATEGORIES = 106

# The trailing description is optional — each function block is closed by a
# bare repeat of its own header ("GOVERN (GV)") with no description after it.
FUNCTION_RE = re.compile(r"^(.+?)\s*\(([A-Z]{2})\)\s*:?", re.S)
CATEGORY_RE = re.compile(r"^(.+?)\s*\(([A-Z]{2}\.[A-Z]{2})\)\s*:?", re.S)
SUBCAT_RE = re.compile(r"^([A-Z]{2}\.[A-Z]{2}-\d{2})\s*:\s*(.+)$", re.S)
EXAMPLE_RE = re.compile(r"(?:^|\n)\s*(Ex\d+)\s*:\s*", re.S)

# NIST's own retirement marker, carried on both category and subcategory cells.
WITHDRAWN = "[Withdrawn"


def _clean(text: Any) -> str:
    """Collapse the export's soft-wrapped whitespace without joining words."""
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _split_examples(cell: Any) -> List[str]:
    """'Ex1: do a thing\nEx2: do another' -> ['Ex1: do a thing', 'Ex2: ...'].

    Split on the Ex-N markers themselves so a newline *inside* one example
    doesn't fracture it — several examples wrap across three or four lines.
    """
    raw = str(cell or "").strip()
    if not raw:
        return []
    marks = list(EXAMPLE_RE.finditer(raw))
    if not marks:
        return [_clean(raw)]
    out = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(raw)
        body = _clean(raw[m.end():end])
        if body:
            out.append(f"{m.group(1)}: {body}")
    return out


def parse(path: Path) -> List[Dict[str, Any]]:
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "CSF 2.0" not in wb.sheetnames:
        sys.exit(f"{path} has no 'CSF 2.0' sheet (found: {wb.sheetnames})")
    rows = list(wb["CSF 2.0"].iter_rows(values_only=True))[2:]  # banner + header

    controls: List[Dict[str, Any]] = []
    withdrawn: List[str] = []
    fn_name = cat_code = cat_name = ""
    cat_withdrawn = False

    for row in rows:
        fn_cell, cat_cell, sub_cell = row[0], row[1], row[2]
        ex_cell = row[3] if len(row) > 3 else None

        if fn_cell:
            m = FUNCTION_RE.match(str(fn_cell))
            if not m:
                sys.exit(f"unparseable function cell: {str(fn_cell)[:120]!r}")
            fn_name = _clean(m.group(1))
            continue

        if cat_cell:
            m = CATEGORY_RE.match(str(cat_cell))
            if not m:
                sys.exit(f"unparseable category cell: {str(cat_cell)[:120]!r}")
            cat_name, cat_code = _clean(m.group(1)), m.group(2)
            cat_withdrawn = WITHDRAWN in str(cat_cell)
            continue

        if not sub_cell:
            continue

        m = SUBCAT_RE.match(str(sub_cell).strip())
        if not m:
            sys.exit(f"unparseable subcategory cell: {str(sub_cell)[:120]!r}")
        code, statement = m.group(1), _clean(m.group(2))
        if not code.startswith(cat_code + "-"):
            sys.exit(f"{code} does not belong to category {cat_code}")

        if statement.startswith(WITHDRAWN):
            withdrawn.append(code)
            continue
        if cat_withdrawn:
            # A live subcategory under a retired category would mean the marker
            # no longer means what we think it does.
            sys.exit(f"{code} is live but its category {cat_code} is withdrawn")

        examples = _split_examples(ex_cell)
        full_text = statement
        if examples:
            full_text += "\n\nImplementation examples:\n" + "\n".join(f"- {e}" for e in examples)

        controls.append({
            "control_id": code,
            "original_reference": code,
            # CSF publishes no subcategory titles — the statement IS the
            # requirement, so it stands as the title rather than an invented one.
            "title": statement,
            "description": statement,
            "full_text": full_text,
            "domain": fn_name,
            "category": cat_name,
            # Verbatim NIST implementation examples. Not evidence requirements:
            # CSF publishes none, and inventing them is what we are avoiding.
            "assessment_criteria": examples,
            "section_number": code,
            "parent_section": cat_code,
            "is_mandatory": True,
            "priority": "medium",
            "evidence_requirements": [],
        })

    print(f"  {len(withdrawn)} withdrawn v1.1 subcategories skipped")
    return controls


def scf_codes() -> set:
    if not SCF_MAPPINGS.exists():
        return set()
    out = set()
    with gzip.open(SCF_MAPPINGS, "rt", encoding="utf-8", newline="") as gz:
        for row in csv.DictReader(gz):
            if row["source_slug"] == "nist_csf_20":
                out.add(row["requirement_code"])
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", required=True, type=Path, help="CSF 2.0 Reference Tool .xlsx export")
    ap.add_argument("--emit", action="store_true", help="write seed_data/frameworks/nist_csf.json")
    args = ap.parse_args()

    controls = parse(args.source)

    codes = [c["control_id"] for c in controls]
    functions = {c["control_id"][:2] for c in controls}
    categories = {c["parent_section"] for c in controls}

    print(f"parsed {len(controls)} subcategories in {len(categories)} categories, {len(functions)} functions")

    dupes = [k for k, v in Counter(codes).items() if v > 1]
    if dupes:
        sys.exit(f"duplicate subcategory codes: {dupes}")
    if len(functions) != EXPECT_FUNCTIONS:
        sys.exit(f"expected {EXPECT_FUNCTIONS} functions, got {len(functions)}: {sorted(functions)}")
    if len(categories) != EXPECT_CATEGORIES:
        sys.exit(f"expected {EXPECT_CATEGORIES} categories, got {len(categories)}: {sorted(categories)}")
    if len(controls) != EXPECT_SUBCATEGORIES:
        sys.exit(f"expected {EXPECT_SUBCATEGORIES} subcategories, got {len(controls)}")
    if any(not c["title"] for c in controls):
        sys.exit("some subcategories parsed with empty statement text")

    by_fn: Dict[str, int] = defaultdict(int)
    for c in controls:
        by_fn[c["control_id"][:2]] += 1
    print("  " + "  ".join(f"{k}:{v}" for k, v in sorted(by_fn.items())))

    no_examples = [c["control_id"] for c in controls if not c["assessment_criteria"]]
    print(f"implementation examples on {len(controls) - len(no_examples)}/{len(controls)}"
          + (f"; none for {no_examples}" if no_examples else ""))

    # Independent cross-check: SCF resolved CSF 2.0 from its own copy of the
    # standard. Agreement between two unrelated readings is the strongest
    # signal we can get that the parse is right.
    scf = scf_codes()
    if scf:
        scf_subs = {c for c in scf if re.fullmatch(r"[A-Z]{2}\.[A-Z]{2}-\d{2}", c)}
        ours = set(codes)
        print(f"\ncross-check vs SCF nist_csf_20: {len(scf_subs)} subcategory codes")
        only_scf, only_ours = sorted(scf_subs - ours), sorted(ours - scf_subs)
        print(f"  will resolve : {len(ours & scf_subs)}")
        if only_scf:
            print(f"  ! in SCF, not in our parse ({len(only_scf)}): {only_scf}")
        if only_ours:
            print(f"  unmapped by SCF ({len(only_ours)}): {only_ours}")
        if only_scf:
            sys.exit("SCF names CSF 2.0 codes our parse missed — fix the parser before emitting")

    if args.emit:
        payload = {
            "metadata": {
                "name": "NIST Cybersecurity Framework",
                "description": (
                    "The NIST Cybersecurity Framework (CSF) 2.0 organizes cybersecurity outcomes "
                    "into 6 Functions, 22 Categories and 106 Subcategories. Each Subcategory states "
                    "an outcome to be achieved rather than a prescriptive control."
                ),
                "version": "2.0",
                "framework_type": "standard",
                "classification": "compliance",
                "source_organization": "National Institute of Standards and Technology (NIST)",
                "regulatory_authority": "NIST",
                "adoption_approach": "voluntary",
                "hierarchy_structure": "Function > Category > Subcategory",
                "source_note": (
                    "Built from the NIST CSF 2.0 Reference Tool export "
                    "(csrc.nist.gov). NIST publications are US Government works in the public "
                    "domain (17 U.S.C. §105); subcategory statements and implementation "
                    "examples are reproduced verbatim."
                ),
            },
            "controls": controls,
        }
        OUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"\nwrote {OUT} ({OUT.stat().st_size:,} bytes, {len(controls)} controls)")
    else:
        print("\n(dry run — pass --emit to write the library)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
