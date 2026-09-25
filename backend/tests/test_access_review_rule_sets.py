"""A review runs the rules someone chose, and reports against exactly those.

The rules used to be the whole library's enabled set, read again every time the
review was shown — so switching a rule off rewrote finished reviews. Now a review
runs every enabled rule, the rules that evidence one framework, or a hand-picked
set; what ran is frozen when its checks run; and the report and every export
show each rule's result, its category, and the frameworks it evidences.
"""
import importlib
import io
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.db import get_tenant_db

# the module itself: grc.routers re-exports its APIRouter under the same name
ar = importlib.import_module("grc.routers.access_review_router")

SOC2, ECC = "aicpa_tsc_soc2", "emea_saudi_arabia_ecc_1_2018"


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(m.Tenant(id=1, name="Demo Bank", slug="demo"))
        session.add_all([
            m.SCFSource(release_id=1, source_key="k1", source_slug=SOC2, display_name="AICPA TSC 2017:2022 (used for SOC 2)"),
            m.SCFSource(release_id=1, source_key="k2", source_slug=ECC, display_name="EMEA Saudi Arabia ECC-1 2018"),
            # AUTH-01 (IAC-06) answers both; PRIV-05 (IAC-06.1) answers SOC 2 only
            m.SCFMapping(release_id=1, scf_id="IAC-06", source_slug=SOC2, requirement_code="CC6.6"),
            m.SCFMapping(release_id=1, scf_id="IAC-06", source_slug=ECC, requirement_code="2-2-3-2"),
            m.SCFMapping(release_id=1, scf_id="IAC-06.1", source_slug=SOC2, requirement_code="CC6.1"),
        ])
        session.commit()
        yield session


@pytest.fixture
def api(db, monkeypatch):
    monkeypatch.setattr(ar, "_require_admin", lambda *a, **k: SimpleNamespace(id=7))
    app = FastAPI()
    app.include_router(ar.router)
    app.dependency_overrides[get_tenant_db] = lambda: db
    return TestClient(app)


def _sample(db, campaign_id):
    """Two sampled people: one without MFA (fails AUTH-01), one fine."""
    for i, mfa in ((1, False), (2, True)):
        db.add(m.AccessReviewItem(tenant_id=1, campaign_id=campaign_id, email=f"u{i}@bank.example",
                                  display_name=f"User {i}", department="IT", account_enabled=True,
                                  mfa_enabled=mfa, is_privileged=False, decision="pending",
                                  last_sign_in=None, roles_snapshot=["Staff"]))
    campaign = db.get(m.AccessReviewCampaign, campaign_id)
    campaign.status = "sampled"
    db.commit()


def test_a_framework_review_runs_that_frameworks_rules_and_keeps_them_when_the_library_changes(db, api):
    made = api.post("/access-reviews", json={"name": "ECC review", "rule_scope": "framework", "rule_framework": ECC})
    assert made.status_code == 200, made.text
    review = made.json()
    assert review["rule_scope"] == "framework" and review["rule_framework_name"] == "EMEA Saudi Arabia ECC-1 2018"
    # a framework no runnable rule evidences can't make a review
    assert api.post("/access-reviews", json={"name": "x", "rule_scope": "framework", "rule_framework": "nope"}).status_code == 400

    _sample(db, review["id"])
    assert api.post(f"/access-reviews/{review['id']}/run-checks").status_code == 200
    detail = api.get(f"/access-reviews/{review['id']}").json()
    assert detail["campaign"]["rules_run"] == ["AUTH-01"]               # ECC's runnable rules, nothing else
    [mfa] = detail["rule_results"]
    assert (mfa["id"], mfa["domain"], mfa["failed"], mfa["passed"]) == ("AUTH-01", "Authentication", 1, 1)
    assert mfa["frameworks"][0]["slug"] == ECC                           # the review's own framework first
    assert {r["status"] for it in detail["items"] for r in it["rules"]} == {"pass", "fail"}

    # switching the rule off in the library later doesn't rewrite this review
    assert api.patch("/access-reviews/rules/AUTH-01", json={"enabled": False}).status_code == 200
    again = api.get(f"/access-reviews/{review['id']}").json()
    assert [r["id"] for r in again["rule_results"]] == ["AUTH-01"]


def test_a_custom_review_runs_only_the_rules_picked_and_can_be_changed_until_sealed(db, api):
    made = api.post("/access-reviews", json={"name": "Picked", "rule_scope": "custom",
                                             "rule_ids": ["IDM-04", "AUTH-01", "NET-01", "NOPE"]})
    assert made.status_code == 200, made.text
    review = made.json()
    assert review["rule_ids"] == ["IDM-04", "AUTH-01"]                   # rules that can't run are dropped
    assert api.post("/access-reviews", json={"name": "x", "rule_scope": "custom", "rule_ids": ["NET-01"]}).status_code == 400

    _sample(db, review["id"])
    api.post(f"/access-reviews/{review['id']}/run-checks")
    assert sorted(api.get(f"/access-reviews/{review['id']}").json()["campaign"]["rules_run"]) == ["AUTH-01", "IDM-04"]

    changed = api.put(f"/access-reviews/{review['id']}/rules", json={"rule_scope": "framework", "rule_framework": SOC2})
    assert changed.status_code == 200 and changed.json()["rule_scope"] == "framework"
    api.post(f"/access-reviews/{review['id']}/run-checks")               # run again to apply it
    ran = api.get(f"/access-reviews/{review['id']}").json()["campaign"]["rules_run"]
    assert "PRIV-05" in ran and "IDM-04" not in ran                      # SOC 2's rules now

    api.post(f"/access-reviews/{review['id']}/close")
    assert api.put(f"/access-reviews/{review['id']}/rules", json={"rule_scope": "enabled"}).status_code == 400


def test_the_report_and_every_export_show_what_each_rule_found(db, api):
    review = api.post("/access-reviews", json={"name": "Q3 & Q4 review", "rule_scope": "framework",
                                               "rule_framework": SOC2}).json()
    _sample(db, review["id"])
    api.post(f"/access-reviews/{review['id']}/run-checks")

    report = api.get(f"/access-reviews/{review['id']}/report").json()
    assert report["rule_scope"] == "framework" and report["rule_framework_name"].startswith("AICPA")
    assert report["provisional"] is True and report["pending"] == 2
    assert {r["id"] for r in report["rule_results"]} >= {"AUTH-01", "PRIV-05"}

    csv = api.get(f"/access-reviews/{review['id']}/report/export?format=csv")
    assert csv.status_code == 200 and "Rules failed" in csv.text and "AUTH-01 No MFA" in csv.text
    xlsx = api.get(f"/access-reviews/{review['id']}/report/export?format=xlsx")
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Users", "Rules"]
    rules_sheet = [[c.value for c in row] for row in book["Rules"].iter_rows()]
    assert rules_sheet[0][:3] == ["Rule", "Name", "Category"]
    assert any(r[0] == "AUTH-01" and r[4] == "Fail" and "SOC 2" in r[7] for r in rules_sheet[1:])
    pdf = api.get(f"/access-reviews/{review['id']}/report/export?format=pdf")     # '&' in names and categories
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
