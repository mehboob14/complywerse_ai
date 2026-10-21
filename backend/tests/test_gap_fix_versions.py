"""A gap fix is saved as every other change is: the document's next version, the versions before it kept as they were (the user,
2026-10-21). Override takes one fix back out on its own - every other change kept - as a version of its own, and the finding is
open again. A document that is its Word file gets the fix, and its take back, written into a new file."""
import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.governance.routers import documents
from grc.modules.governance.routers import gap_analysis as gaps

TEXT = ("# Access policy\n\n## Purpose\n\nAccess is granted on need to know. This policy sets how the bank grants, reviews and removes "
        "access to its systems, its data and its premises, for staff, contractors and service providers alike.\n\n"
        "## Rules\n\nAccounts are reviewed every year.")
YEARLY = "Accounts are reviewed every year."
SHARPER = "Accounts are reviewed every year and removed when their owner leaves."
QUARTERLY = "Access rights are reviewed every quarter by their owners."


@pytest.fixture
def db(monkeypatch, tmp_path):
    monkeypatch.setattr(gaps, "get_user_tenants", lambda user, session: [1])
    monkeypatch.setattr(documents, "GOVERNANCE_UPLOAD_DIR", str(tmp_path))
    for key in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN"):          # major or minor: the rule decides, no AI call
        monkeypatch.delenv(key, raising=False)
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(m.Tenant(id=1, name="Demo Bank", slug="demo"))
        session.add(m.GRCUser(id=3, username="sara", email="sara@bank.example", display_name="Sara Khan"))
        session.add(m.GovernanceDocument(id=5, tenant_id=1, title="Access policy", doc_type="policy", status="draft", content=TEXT,
                                         current_version="1.0", owner_id=3))
        session.add(m.GovernanceDocumentVersion(document_id=5, version_number="1.0", change_type="major", title="Access policy", content=TEXT,
                                                status="current", created_by=3))
        session.add(m.PolicyGapAnalysisRun(id=1, tenant_id=1, document_id=5))
        session.add(m.PolicyGapFinding(id=1, tenant_id=1, analysis_run_id=1, document_id=5, framework_name="ISO/IEC 27001", clause_reference="A.5.18",
                                       compliance_status="not_addressed", remediation_status="open"))
        session.add(m.PolicyGapFinding(id=2, tenant_id=1, analysis_run_id=1, document_id=5, framework_name="ISO/IEC 27001", clause_reference="A.5.15",
                                       compliance_status="partially_compliant", remediation_status="open"))
        session.commit()
        yield session


def _user(db):
    return db.get(m.GRCUser, 3)


def _versions(db):
    rows = db.query(m.GovernanceDocumentVersion).filter_by(document_id=5).order_by(m.GovernanceDocumentVersion.id).all()
    return [(v.version_number, v.status) for v in rows]


def _add(db, finding_id=1, text=QUARTERLY):
    return gaps.apply_finding_fix(finding_id=finding_id, body=gaps.ApplyFixRequest(proposed_text=text), db=db, current_user=_user(db))


def _replace(db, finding_id=2):
    return gaps.apply_finding_fix(finding_id=finding_id, body=gaps.ApplyFixRequest(mode="replace", current_text=YEARLY, proposed_text=SHARPER),
                                  db=db, current_user=_user(db))


def _undo(db, finding_id=1):
    return gaps.override_finding(finding_id=finding_id, override=gaps.FindingOverrideRequest(override_status="not_addressed", override_justification="Not this way"),
                                 db=db, current_user=_user(db))


def _row(db, number):
    return db.query(m.GovernanceDocumentVersion).filter_by(document_id=5, version_number=number).one()


def test_each_fix_on_a_draft_makes_its_own_version_and_undo_takes_back_only_its_own_fix(db):
    first, second = _add(db), _replace(db)
    assert (first["in_place"], first["applied_version_number"]) == (False, "2.0")               # a section added: major
    assert (second["in_place"], second["applied_version_number"]) == (False, "2.1")             # a sentence sharpened: minor
    assert second["applied_version_id"] != first["applied_version_id"]
    assert _versions(db) == [("1.0", "superseded"), ("2.0", "superseded"), ("2.1", "current")]
    assert QUARTERLY in _row(db, "2.0").content and SHARPER not in _row(db, "2.0").content          # 2.0 keeps its own text
    doc = db.get(m.GovernanceDocument, 5)
    assert QUARTERLY in doc.content and SHARPER in doc.content and doc.current_version == "2.1"
    added_statement = db.get(m.PolicyGapFinding, 1).applied_statement_id

    out = _undo(db)
    doc = db.get(m.GovernanceDocument, 5)
    assert doc.content == TEXT.replace(YEARLY, SHARPER)                          # only the first fix (and its heading) went
    assert _versions(db) == [("1.0", "superseded"), ("2.0", "superseded"), ("2.1", "superseded"), ("3.0", "current")]
    row = _row(db, "3.0")                                                        # a section taken out: major, a version of its own
    assert row.content == doc.content and out["taken_back_version"] == "3.0" == doc.current_version
    assert {i["origin"] for i in row.change_set["items"]} == {"revert"}
    assert QUARTERLY in _row(db, "2.1").content                                  # the fix is still in the versions that had it
    assert out["remediation_status"] == "open" and out["compliance_status"] == "not_addressed" and out["applied_at"] is None
    assert db.get(m.PolicyStatement, added_statement).status == "superseded"     # the statement the fix added is retracted
    kept = db.get(m.PolicyGapFinding, 2)
    assert kept.applied_at is not None and kept.remediation_status == "in_progress"

    _undo(db, 2)                                                                  # the other one too: the text of 1.0 again
    assert db.get(m.GovernanceDocument, 5).content == TEXT and db.get(m.GovernanceDocument, 5).current_version == "3.1"


def test_after_a_publication_each_fix_makes_the_next_number(db):
    doc = db.get(m.GovernanceDocument, 5)
    doc.status, doc.last_published_version, doc.published_version = "published", "1.0", "1.0"
    db.commit()
    first = _replace(db)
    assert first["in_place"] is False and first["applied_version_number"] == "1.1"     # a reworded rule: minor
    assert first["document_status"] == "draft"                                        # nobody to sign it: a draft again
    second = _add(db)
    assert second["applied_version_number"] == "2.0"                                  # a section added: major
    assert _versions(db) == [("1.0", "superseded"), ("1.1", "superseded"), ("2.0", "current")]


def test_an_undo_leaves_a_fix_edited_since_in_the_document_and_still_overrides_the_finding(db):
    _add(db)
    doc = db.get(m.GovernanceDocument, 5)
    doc.content = doc.content.replace("every quarter", "every month")
    db.commit()
    edited, versions = doc.content, _versions(db)
    out = _undo(db)                                    # no dead end: the person's edit stands, the reply says so
    assert out["fix_left_in_document"] is True and out["taken_back_version"] is None
    assert db.get(m.GovernanceDocument, 5).content == edited and _versions(db) == versions
    assert out["applied_at"] is None and out["remediation_status"] == "open" and out["compliance_status"] == "not_addressed"
    finding = db.get(m.PolicyGapFinding, 1)
    assert finding.is_overridden and finding.applied_statement_id is None


def test_undoing_an_added_clause_takes_back_its_own_copy_not_an_earlier_one(db):
    _add(db, text=YEARLY)                              # the same words as the rule already in the document
    out = _undo(db)
    assert db.get(m.GovernanceDocument, 5).content == TEXT and out["fix_left_in_document"] is False


def test_a_document_that_is_its_word_file_gets_the_fix_and_its_take_back_in_a_new_file(db, tmp_path):
    path = tmp_path / "access.docx"
    word = Document()
    word.add_heading("Access policy", level=1)
    word.add_paragraph("Access is granted on need to know.")
    word.save(path)
    doc = db.get(m.GovernanceDocument, 5)
    doc.content, doc.file_path, doc.file_name, doc.file_type = "", str(path), "access.docx", "docx"
    db.query(m.GovernanceDocumentVersion).update({"content": None, "file_path": str(path), "file_name": "access.docx", "file_type": "docx"})
    db.commit()

    out = _add(db)
    doc = db.get(m.GovernanceDocument, 5)
    assert doc.file_path != str(path) and not (doc.content or "").strip()        # still its file: a new one
    text = "\n".join(p.text for p in Document(doc.file_path).paragraphs)
    assert QUARTERLY in text and "Access is granted on need to know." in text
    assert _versions(db) == [("1.0", "superseded"), (out["applied_version_number"], "current")]

    _undo(db)
    doc = db.get(m.GovernanceDocument, 5)
    text = "\n".join(p.text for p in Document(doc.file_path).paragraphs)
    assert QUARTERLY not in text and "A.5.18" not in text and "Access is granted on need to know." in text
    rows = db.query(m.GovernanceDocumentVersion).filter_by(document_id=5).order_by(m.GovernanceDocumentVersion.id).all()
    assert [r.version_number for r in rows] == ["1.0", out["applied_version_number"], "3.0"]     # the take back is a version of its own
    assert QUARTERLY in "\n".join(p.text for p in Document(rows[1].file_path).paragraphs)          # and the fix's file stays as it was
