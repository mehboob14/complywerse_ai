"""Follow-up questionnaires, and answering offline in a workbook.

An answer can call for another questionnaire: when the questionnaire is accepted
the follow-up goes out once, and its trigger is what stops a second. A vendor
can answer in a workbook instead of the portal: it carries the questionnaire and
version it was made from, so it is refused anywhere else, and an answer that is
not among a question's answers is refused with its row rather than guessed.
"""
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from grc.models import (
    Base, Evidence, EvidenceQualityCheck, GRCUser, Tenant, TPRAAuditLog, TPRAEvidenceLink, TPRAFinding,
    TPRAQuestion, TPRAQuestionResponse, TPRATemplateVersion, TPRATieringConfig, Vendor, VendorAssessment,
    VendorQuestionnaireEvidence, VendorQuestionnaireResponse, VendorQuestionnaireTemplate, get_db,
)
from grc.modules.vendor_risk.routers import questionnaires as routes
from grc.modules.vendor_risk.tpra import follow_ups, portal, rbac, versions
from grc.routers.auth_router import require_auth

_TABLES = [Tenant, GRCUser, Vendor, VendorAssessment, VendorQuestionnaireTemplate, VendorQuestionnaireResponse,
           VendorQuestionnaireEvidence, TPRATemplateVersion, TPRAQuestion, TPRAQuestionResponse, TPRAFinding,
           TPRATieringConfig, TPRAAuditLog, Evidence, EvidenceQualityCheck, TPRAEvidenceLink]
ATTEST = {"name": "Sam Vendor", "email": "sam@vendor.test", "confirm": True}


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[m.__table__ for m in _TABLES])
    s = sessionmaker(bind=engine)()
    s.add(Tenant(id=1, name="Acme", slug="acme"))
    s.add(GRCUser(id=7, username="analyst", email="a@x.test", is_active=True))
    s.commit()
    yield s
    s.close()


@pytest.fixture()
def http(db, monkeypatch):
    monkeypatch.setattr(rbac, "require_write", lambda *a, **k: None)
    monkeypatch.setattr(routes, "_email_link", lambda *a, **k: None)
    monkeypatch.setattr(portal, "throttle", portal.Throttle())
    app = FastAPI()
    app.include_router(routes.router, prefix="/vendor-risk")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: db.get(GRCUser, 7)
    return TestClient(app)


def _templates(db):
    fourth = VendorQuestionnaireTemplate(tenant_id=1, name="Fourth parties", questions=[
        {"id": "fp1", "text": "List your subprocessors with access to our data.", "type": "text"}])
    db.add(fourth)
    db.flush()
    main = VendorQuestionnaireTemplate(tenant_id=1, name="Security", questions=[
        {"id": "subs", "text": "Do you use subprocessors?", "type": "yes_no",
         "follow_up": {"when": ["yes"], "template_id": fourth.id}},
        {"id": "mfa", "text": "Is MFA enforced?", "type": "yes_no"},
        {"id": "pen", "text": "How often are you tested?", "type": "multiple_choice",
         "options": [{"value": "annual", "label": "Annually", "score": 1}, {"value": "never", "label": "Never", "score": 0}]},
        {"id": "notes", "text": "Anything else?", "type": "text"},
    ])
    vendor = Vendor(tenant_id=1, name="Acme Cloud", status="active")
    db.add_all([main, vendor])
    db.commit()
    return vendor, main, fourth


def _send(http, vendor, template):
    body = http.post("/vendor-risk/questionnaires/send",
                     json={"vendor_id": vendor.id, "template_id": template.id,
                           "respondent_email": "sam@vendor.test"}).json()
    return body["token"], body["questionnaire_response"]["id"]


# ── follow-ups ───────────────────────────────────────────────────────────────

def test_an_answer_that_calls_for_a_follow_up_sends_it_once_on_acceptance(db, http):
    vendor, main, fourth = _templates(db)
    token, qr_id = _send(http, vendor, main)
    http.post(f"/vendor-risk/questionnaires/external/{token}",
              json={"responses": {"subs": "yes", "mfa": "yes"}, "attestation": ATTEST})
    accepted = http.post(f"/vendor-risk/questionnaire-responses/{qr_id}/accept").json()

    assert len(accepted["follow_ups"]) == 1
    child = db.get(VendorQuestionnaireResponse, accepted["follow_ups"][0])
    assert child.parent_response_id == qr_id and child.trigger_key == f"{qr_id}:subs:{fourth.id}"
    assert child.template_id == fourth.id and child.template_version_id is not None
    assert (child.respondent_email, child.status) == ("sam@vendor.test", "pending")

    parent = db.get(VendorQuestionnaireResponse, qr_id)
    again = follow_ups.send(db, parent, [dict(q) for q in db.get(TPRATemplateVersion, parent.template_version_id).questions], 7)
    assert again == [] and db.query(VendorQuestionnaireResponse).count() == 2


def test_no_follow_up_when_the_answer_does_not_call_for_one(db, http):
    vendor, main, _ = _templates(db)
    token, qr_id = _send(http, vendor, main)
    http.post(f"/vendor-risk/questionnaires/external/{token}",
              json={"responses": {"subs": "no", "mfa": "yes"}, "attestation": ATTEST})
    assert http.post(f"/vendor-risk/questionnaire-responses/{qr_id}/accept").json()["follow_ups"] == []


# ── the workbook ─────────────────────────────────────────────────────────────

def _fill(content: bytes, answers: dict, comments: dict = None, attestation=None) -> bytes:
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Questionnaire"]
    for row in ws.iter_rows(min_row=5):
        key = row[0].value
        if key in answers:
            row[3].value = answers[key]
        if comments and key in comments:
            row[4].value = comments[key]
    if attestation:
        att = wb["Attestation"]
        for r, value in enumerate(attestation, start=1):
            att.cell(row=r, column=2, value=value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx(content: bytes):
    return {"file": ("answers.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_a_workbook_goes_out_and_comes_back_as_answers(db, http):
    vendor, main, _ = _templates(db)
    token, qr_id = _send(http, vendor, main)
    blank = http.get(f"/vendor-risk/questionnaires/external/{token}/workbook")
    assert blank.status_code == 200 and blank.content[:2] == b"PK"

    filled = _fill(blank.content, {"subs": "No", "mfa": "Partial", "pen": "Annually", "notes": "We are ISO certified."},
                   comments={"mfa": "Rolling out to contractors."},
                   attestation=["Sam Vendor", "CISO", "sam@vendor.test", "Yes"])
    back = http.post(f"/vendor-risk/questionnaires/external/{token}/workbook", files=_xlsx(filled))
    assert back.status_code == 200, back.text
    assert back.json()["status"] == "in_progress" and back.json()["attestation"]["name"] == "Sam Vendor"

    qr = db.get(VendorQuestionnaireResponse, qr_id)
    assert qr.responses == {"subs": "no", "mfa": "partial", "pen": "annual", "notes": "We are ISO certified."}
    assert qr.vendor_comments == {"mfa": "Rolling out to contractors."}
    assert qr.attested_name is None                         # the vendor still attests in the portal


def test_a_workbook_answer_that_is_not_an_option_is_refused_by_row(db, http):
    vendor, main, _ = _templates(db)
    token, _ = _send(http, vendor, main)
    blank = http.get(f"/vendor-risk/questionnaires/external/{token}/workbook").content
    refused = http.post(f"/vendor-risk/questionnaires/external/{token}/workbook",
                        files=_xlsx(_fill(blank, {"mfa": "Mostly"})))
    assert refused.status_code == 400 and "Row 6" in refused.json()["detail"] and "Mostly" in refused.json()["detail"]


def test_a_workbook_made_for_another_questionnaire_or_version_is_refused(db, http):
    vendor, main, _ = _templates(db)
    token_a, _ = _send(http, vendor, main)
    token_b, qr_b = _send(http, vendor, main)
    from_a = http.get(f"/vendor-risk/questionnaires/external/{token_a}/workbook").content
    wrong = http.post(f"/vendor-risk/questionnaires/external/{token_b}/workbook", files=_xlsx(from_a))
    assert wrong.status_code == 400 and "different questionnaire" in wrong.json()["detail"]

    from_b = http.get(f"/vendor-risk/questionnaires/external/{token_b}/workbook").content
    main.questions = [{**main.questions[0], "text": "Edited"}] + list(main.questions[1:])
    db.commit()
    newer = versions.publish(db, main)
    db.get(VendorQuestionnaireResponse, qr_b).template_version_id = newer.id    # as if pinned to another
    db.commit()
    stale = http.post(f"/vendor-risk/questionnaires/external/{token_b}/workbook", files=_xlsx(from_b))
    assert stale.status_code == 400 and "different version" in stale.json()["detail"]


def test_an_analyst_imports_a_returned_workbook_on_the_vendors_attestation(db, http):
    vendor, main, _ = _templates(db)
    token, qr_id = _send(http, vendor, main)
    blank = http.get(f"/vendor-risk/questionnaire-responses/{qr_id}/workbook").content

    unattested = _fill(blank, {"subs": "No", "mfa": "Yes"})
    assert http.post(f"/vendor-risk/questionnaire-responses/{qr_id}/workbook", files=_xlsx(unattested),
                     data={"submit": "true"}).status_code == 400

    attested = _fill(blank, {"subs": "No", "mfa": "Yes"}, attestation=["Sam Vendor", "CISO", "sam@vendor.test", "Yes"])
    done = http.post(f"/vendor-risk/questionnaire-responses/{qr_id}/workbook", files=_xlsx(attested),
                     data={"submit": "true"})
    assert done.status_code == 200, done.text
    qr = db.get(VendorQuestionnaireResponse, qr_id)
    assert (qr.status, qr.attested_name, qr.attested_title) == ("submitted", "Sam Vendor", "CISO")
    assert [a.action for a in db.query(TPRAAuditLog).filter(TPRAAuditLog.entity == "questionnaire")] == ["import"]
