"""The Change log records every change to a document's statements (the user, 2026-10-21: "change log should also contain changes
done through statements ... if any statement changed applied to document, even reverted etc all changes should be logged and
documented there"): added, edited, removed, rolled back, re-read from the document, matched to the document's wording, changed by a
gap fix and that fix undone, written into the document (that version's changes say they came from the statements) and taken back
again (the statement then shows its edit as not in the document). A statement linked to controls and issues can be removed.

The database is SQLite with foreign keys on and a session that does not autoflush, as a tenant's Postgres session (PITFALLS 74)."""
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, inspect
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.governance import change_chain as C
from grc.modules.governance import doc_sync
from grc.modules.governance.routers import change_chain as chain_routes
from grc.modules.governance.routers import documents, policy_parser
from grc.modules.governance.routers import gap_analysis as gaps
from grc.routers.auth_router import require_auth

S1 = "All administrators shall use multi-factor authentication."
S2 = "Passwords must be changed every 90 days."
S2_NEW = "Passwords must be changed every 60 days."
S3 = "Backups shall be tested “quarterly”."          # curly quotes in the file, straight ones in the statement
LONG = "Access rights must be reviewed at least once every year by the system owner."
TEXT = "# Access policy\n\n## Rules\n\nAccounts are reviewed every year.\n"
YEARLY = "Accounts are reviewed every year."
SHARPER = "Accounts are reviewed every year and removed when their owner leaves."
QUARTERLY = "Access rights are reviewed every quarter by their owners."


def make_docx(path):
    doc = Document()
    doc.add_paragraph("Information Security Policy")
    p = doc.add_paragraph()
    p.add_run("1. ").bold = True
    p.add_run(S1)
    doc.add_paragraph(S2 + " Accounts that are idle are disabled.")
    doc.add_paragraph(S3)
    doc.add_paragraph(LONG)
    doc.save(path)


@pytest.fixture
def world(tmp_path, monkeypatch):
    monkeypatch.setattr(documents, "GOVERNANCE_UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr(policy_parser, "_reparse_proposals", {})
    monkeypatch.setattr(policy_parser, "_parsing_status", {})
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    event.listen(engine, "connect", lambda conn, _: conn.execute("PRAGMA foreign_keys=ON"))
    m.Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, autoflush=False)()
    path = str(tmp_path / "policy.docx")
    make_docx(path)
    db.add(m.Tenant(id=1, name="CFSB", slug="cfsb"))                  # cfsb: statements are held to the document's file
    db.add(m.GRCUser(id=1, username="olivia", email="o@cfsb.test", display_name="Olivia Owner"))
    db.flush()
    db.add(m.GovernanceDocument(id=5, tenant_id=1, title="InfoSec policy", doc_type="policy", status="approved", current_version="1.0",
                                file_name="policy.docx", file_path=path, file_type="docx", file_size=1, content=None, owner_id=1))
    db.add(m.GovernanceDocument(id=6, tenant_id=1, title="Access policy", doc_type="policy", status="draft", content=TEXT, current_version="1.0",
                                owner_id=1))
    db.flush()
    db.add(m.GovernanceDocumentVersion(document_id=6, version_number="1.0", change_type="major", title="Access policy", content=TEXT,
                                       status="current", created_by=1))
    statements = [(5, f"PS-0005-{i:03d}", t) for i, t in enumerate([S1, S2, S3.replace("“", '"').replace("”", '"'), LONG], start=1)]
    for doc_id, code, text in statements + [(6, "PS-0006-001", YEARLY)]:
        s = m.PolicyStatement(tenant_id=1, document_id=doc_id, statement_code=code, statement_text=text, status="active")
        db.add(s)
        db.flush()
        policy_parser._create_version_snapshot(db, s, "initial_parse", 1)       # as reading the document into statements does
    db.add(m.PolicyGapAnalysisRun(id=1, tenant_id=1, document_id=6))
    db.flush()
    db.add(m.PolicyGapFinding(id=1, tenant_id=1, analysis_run_id=1, document_id=6, framework_name="ISO/IEC 27001", clause_reference="A.5.15",
                              compliance_status="partially_compliant", remediation_status="open"))
    db.add(m.PolicyGapFinding(id=2, tenant_id=1, analysis_run_id=1, document_id=6, framework_name="ISO/IEC 27001", clause_reference="A.5.18",
                              compliance_status="not_addressed", remediation_status="open"))
    db.commit()

    app = FastAPI()
    app.include_router(policy_parser.router)
    app.include_router(chain_routes.router)
    app.dependency_overrides[m.get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: db.get(m.GRCUser, 1)
    yield SimpleNamespace(db=db, http=TestClient(app), path=path)
    db.close()


def statement(db, code):
    return db.query(m.PolicyStatement).filter_by(statement_code=code).one()


def rows(db, doc_id=5, action=None):
    q = db.query(m.DocumentAuditLog).filter_by(document_id=doc_id)
    return (q.filter_by(action=action) if action else q.filter(m.DocumentAuditLog.action.like("statement%"))).order_by(m.DocumentAuditLog.id).all()


def edit(world, code, **body):
    out = world.http.put(f"/documents/5/statements/{statement(world.db, code).id}", json=body)
    assert out.status_code == 200, out.text


# ── each change is one row ───────────────────────────────────────────────────
def test_adding_editing_and_rolling_back_a_statement_each_write_one_row_with_its_wording(world):
    db = world.db
    added = world.http.post("/documents/5/statements", json={"statement_text": "Visitors shall be escorted at all times."})
    assert added.status_code == 200, added.text
    (row,) = rows(db, action="statement_added")
    assert (row.field_changed, row.old_value, row.new_value) == ("statement", None, "Visitors shall be escorted at all times.")
    assert row.action_details == added.json()["statement_code"] and row.performed_by == 1

    edit(world, "PS-0005-002", statement_text=S2_NEW, change_reason="Tighter rule")
    (row,) = rows(db, action="statement_edited")
    assert (row.field_changed, row.old_value, row.new_value, row.action_details) == ("statement", S2, S2_NEW, "PS-0005-002 · Tighter rule")
    edit(world, "PS-0005-002", statement_text=S2_NEW)                       # nothing changed: nothing to log
    edit(world, "PS-0005-002", priority="critical")                          # the wording stays: the row says what changed
    assert [(r.old_value, r.new_value, r.action_details) for r in rows(db, action="statement_edited")][1:] == [
        (None, None, "PS-0005-002 · changed: priority")]

    first = world.http.get(f"/documents/5/statements/{statement(db, 'PS-0005-002').id}/versions").json()["versions"][-1]
    back = world.http.post(f"/documents/5/statements/{statement(db, 'PS-0005-002').id}/rollback", json={"version_id": first["id"]})
    assert back.status_code == 200, back.text
    (row,) = rows(db, action="statement_rolled_back")
    assert (row.old_value, row.new_value, row.action_details) == (S2_NEW, S2, "PS-0005-002 · back to version 0")
    assert db.query(m.AuditLog).filter_by(action="statement_rollback").count() == 1          # the organisation's audit log keeps its row
    assert [r.action for r in rows(db)] == ["statement_added", "statement_edited", "statement_edited", "statement_rolled_back"]


def test_a_long_statement_keeps_the_part_that_changed_in_the_log():
    head = "Every system owner shall " + "review the access rights of every account on the system, " * 12
    before, after = policy_parser._wording(head + "once a year.", head + "twice a year.")
    assert len(before) < 320 and "once a year." in before and "twice a year." in after and before.startswith("… ")
    assert policy_parser._wording("  Short\n text ", None) == ("Short text", None)


def test_the_change_log_and_its_export_show_a_statement_change_with_its_wording(world):
    from openpyxl import load_workbook

    edit(world, "PS-0005-002", statement_text=S2_NEW)
    chain = world.http.get("/documents/5/change-chain").json()
    (e,) = [e for e in chain["events"] if e["kind"] == "statement_edited"]
    assert (e["actor"], e["detail"], e["before"], e["after"], e["from"], e["to"]) == ("Olivia Owner", "PS-0005-002", S2, S2_NEW, None, None)
    assert C.EVENT_LABEL["statement_edited"] == "Statement edited"
    (row,) = [r for r in C.catalog(chain)["approvals"] if r[1] == "Statement edited"]
    assert row[2:] == ["Olivia Owner", f"PS-0005-002 · “{S2}” → “{S2_NEW}”"]
    xlsx = world.http.get("/documents/5/change-chain/export?format=xlsx")
    assert xlsx.status_code == 200
    sheet = [r for r in load_workbook(BytesIO(xlsx.content))["Approvals"].iter_rows(values_only=True)]
    assert ("Statement edited", "Olivia Owner") in {(r[1], r[2]) for r in sheet}


# ── matched to the document, written into it, and taken back ────────────────
def test_a_statement_matched_to_the_documents_wording_is_logged_only_when_its_words_change(world):
    db = world.db
    statement(db, "PS-0005-001").statement_text = "All administrators must use multi-factor authentication."      # the AI's paraphrase
    db.commit()
    assert world.http.post("/documents/5/document-sync/check?again=true").status_code == 200
    # its first match is part of reading the document into statements (a new document's parse corrects many): its history keeps
    # it, the Change log does not
    assert statement(db, "PS-0005-001").statement_text == S1 and rows(db, action="statement_corrected_to_file") == []
    assert statement(db, "PS-0005-003").statement_text == S3                  # its quotation marks follow the file
    word = Document(world.path)                                               # the file's wording changes afterwards (a new file)
    next(p for p in word.paragraphs if p.text == LONG).text = LONG.replace("system owner", "system's owner")
    word.save(world.path)
    assert world.http.post("/documents/5/document-sync/check?again=true").status_code == 200
    (row,) = rows(db, action="statement_corrected_to_file")
    assert (row.old_value, row.new_value, row.action_details) == (LONG, LONG.replace("system owner", "system's owner"), "PS-0005-004")
    world.http.post("/documents/5/document-sync/check?again=true")
    assert len(rows(db, action="statement_corrected_to_file")) == 1


def _write_back(world):
    """PS-0005-002 edited (90 -> 60 days) and written into the Word file: version 1.1."""
    assert world.http.post("/documents/5/document-sync/check?again=true").status_code == 200
    edit(world, "PS-0005-002", statement_text=S2_NEW)
    out = world.http.post("/documents/5/document-sync/apply", json={"change_reason": "Tighter rule"})
    assert out.status_code == 200, out.text
    return out.json()


def test_statements_written_into_the_document_make_a_version_whose_changes_come_from_the_statements(world):
    out = _write_back(world)
    db = world.db
    row = db.get(m.GovernanceDocumentVersion, out["version_id"])
    assert row.change_set and {i["origin"] for i in row.change_set["items"]} == {"statement"}
    (item,) = [i for i in row.change_set["items"] if "60 days" in i["after_text"]]
    assert "90 days" in item["before_text"] and item["reason"] == "From statement PS-0005-002" and item["by"] == "Olivia Owner"

    chain = world.http.get("/documents/5/change-chain").json()
    version = next(v for v in chain["versions"] if v["id"] == row.id)
    assert (version["origin"], version["origin_label"], version["changes"]) == ("statement", "From statements", len(row.change_set["items"]))
    (written,) = [e for e in chain["events"] if e["kind"] == "statements_written_to_document"]
    assert "(PS-0005-002)" in written["detail"] and C.EVENT_LABEL[written["kind"]] == "Statements written into the document"
    assert [e["kind"] for e in chain["events"] if e["kind"].startswith("statement")] == ["statement_edited", "statements_written_to_document"]


def test_taking_back_a_write_back_is_logged_and_the_statement_shows_its_edit_as_not_in_the_document_again(world):
    out = _write_back(world)
    db = world.db
    doc = db.get(m.GovernanceDocument, 5)
    doc.status, doc.last_published_version, doc.published_version = "published", out["version"], out["version"]     # approved since
    db.commit()
    before = db.query(m.GovernanceDocumentVersion).filter_by(document_id=5, version_number="1.0").one()
    red = world.http.get(f"/documents/5/change-chain/redline?from={before.id}").json()
    (hunk,) = [h for h in red["hunks"] if "60 days" in h["after_text"]]
    back = world.http.post("/documents/5/change-chain/revert", json={"from_version_id": before.id, "hunks": [hunk["id"]], "hash": red["hash"],
                                                                     "reason": "The board kept 90 days"})
    assert back.status_code == 200, back.text
    db.expire_all()
    assert any(S2 in p.text for p in Document(db.get(m.GovernanceDocument, 5).file_path).paragraphs)      # the file says 90 days again

    ps2 = statement(db, "PS-0005-002")
    assert (ps2.statement_text, ps2.source_text, ps2.source_match) == (S2_NEW, S2, "exact")          # its edit, not in the document
    listed = {s["statement_code"]: s for s in world.http.get("/documents/5/policy-statements").json()["statements"]}
    assert listed["PS-0005-002"]["pending_in_document"] is True
    assert [p["code"] for p in doc_sync.state(db, db.get(m.GovernanceDocument, 5))["pending"]] == ["PS-0005-002"]     # it can be written again

    chain = world.http.get("/documents/5/change-chain").json()
    assert chain["versions"][-1]["origin_label"] == "Taken back" and chain["versions"][-1]["reverts"]["since_version"] == "1.0"
    (gone,) = [e for e in chain["events"] if e["kind"] == "statement_out_of_document"]
    assert (gone["before"], gone["after"], gone["detail"]) == (S2, S2_NEW, "PS-0005-002")
    assert C.EVENT_LABEL[gone["kind"]] == "Statement edit no longer in the document"


# ── gap fixes ────────────────────────────────────────────────────────────────
def test_a_gap_fix_that_changes_or_adds_a_statement_and_its_undo_are_logged(world, monkeypatch):
    monkeypatch.setattr(gaps, "get_user_tenants", lambda user, session: [1])
    db, me = world.db, world.db.get(m.GRCUser, 1)
    gaps.apply_finding_fix(finding_id=1, body=gaps.ApplyFixRequest(mode="replace", current_text=YEARLY, proposed_text=SHARPER), db=db, current_user=me)
    gaps.apply_finding_fix(finding_id=2, body=gaps.ApplyFixRequest(proposed_text=QUARTERLY), db=db, current_user=me)
    for finding_id in (1, 2):
        gaps.override_finding(finding_id=finding_id, override=gaps.FindingOverrideRequest(override_status="not_addressed", override_justification="No"),
                              db=db, current_user=me)
    added = db.query(m.PolicyStatement).filter_by(document_id=6, statement_text=QUARTERLY).one()
    assert [(r.action, r.action_details, r.old_value, r.new_value) for r in rows(db, 6)] == [
        ("statement_changed_by_fix", "PS-0006-001 · the fix for ISO/IEC 27001 A.5.15", YEARLY, SHARPER),
        ("statement_changed_by_fix", f"{added.statement_code} · added by the fix for ISO/IEC 27001 A.5.18", None, QUARTERLY),
        ("statement_fix_undone", "PS-0006-001 · the fix for ISO/IEC 27001 A.5.15 taken back", SHARPER, YEARLY),
        ("statement_fix_undone", f"{added.statement_code} · retracted: the fix for ISO/IEC 27001 A.5.18 that added it was taken back", QUARTERLY, None)]
    kinds = [e["kind"] for e in world.http.get("/documents/6/change-chain").json()["events"]]
    assert kinds == ["statement_changed_by_fix", "statement_changed_by_fix", "statement_fix_undone", "statement_fix_undone"]


def test_a_gap_fix_in_a_word_document_and_its_undo_leave_the_statement_in_step_with_the_file(world, monkeypatch):
    """The fix writes the new wording into the file and into the statement: nothing waits to be written (it used to, and "Update
    document" then failed with 409), and after its undo the statement and the file say the old wording again."""
    monkeypatch.setattr(gaps, "get_user_tenants", lambda user, session: [1])
    db, me = world.db, world.db.get(m.GRCUser, 1)
    assert world.http.post("/documents/5/document-sync/check?again=true").status_code == 200      # the statements are matched to the file
    db.add(m.PolicyGapAnalysisRun(id=2, tenant_id=1, document_id=5))
    db.flush()
    db.add(m.PolicyGapFinding(id=3, tenant_id=1, analysis_run_id=2, document_id=5, framework_name="ISO/IEC 27001", clause_reference="A.5.17",
                              compliance_status="partially_compliant", remediation_status="open"))
    db.commit()
    gaps.apply_finding_fix(finding_id=3, body=gaps.ApplyFixRequest(mode="replace", current_text=S2, proposed_text=S2_NEW), db=db, current_user=me)
    db.expire_all()
    doc = db.get(m.GovernanceDocument, 5)
    assert any(S2_NEW in p.text for p in Document(doc.file_path).paragraphs)
    ps2 = statement(db, "PS-0005-002")
    assert (ps2.statement_text, ps2.source_text, ps2.source_match) == (S2_NEW, S2_NEW, "exact")
    assert doc_sync.state(db, doc)["pending"] == []

    gaps.override_finding(finding_id=3, override=gaps.FindingOverrideRequest(override_status="not_addressed", override_justification="No"),
                          db=db, current_user=me)
    db.expire_all()
    doc = db.get(m.GovernanceDocument, 5)
    assert any(S2 in p.text for p in Document(doc.file_path).paragraphs)
    ps2 = statement(db, "PS-0005-002")
    assert (ps2.statement_text, ps2.source_text, ps2.source_match) == (S2, S2, "exact")
    assert doc_sync.state(db, doc)["pending"] == []


# ── re-read from the document ────────────────────────────────────────────────
def test_applying_what_reading_the_document_again_proposes_is_one_row_with_the_statements_codes(world):
    db = world.db
    ps2 = statement(db, "PS-0005-002")
    policy_parser._save_proposals(None, 5, {"proposals": [
        {"type": "update", "existing_statement_id": ps2.id, "status": "pending", "new_statement": {"statement_text": S2, "category": "access"}},
        {"type": "new", "status": "pending", "new_statement": {"statement_text": "Accounts that are idle are disabled."}},
        {"type": "new", "status": "pending", "new_statement": {"statement_text": "Something the reader did not want."}}]})
    out = world.http.post("/documents/5/reparse-proposals/apply", json={"decisions": [
        {"index": 0, "action": "accept"}, {"index": 1, "action": "accept"}, {"index": 2, "action": "reject"}]})
    assert out.status_code == 200, out.text
    (row,) = rows(db, action="statements_reparsed")
    assert (row.field_changed, row.action_details) == ("statement", "1 changed (PS-0005-002) · 1 added (PS-0005-005) · 1 not taken")
    assert db.query(m.AuditLog).filter_by(action="reparse_proposals_applied").count() == 1
    assert C.EVENT_LABEL["statements_reparsed"] == "Statements re-read from the document"


# ── removing a statement ─────────────────────────────────────────────────────
def test_a_statement_linked_to_controls_issues_and_a_gap_fix_is_removed_and_the_records_from_it_are_kept(world):
    db = world.db
    ps4 = statement(db, "PS-0005-004")
    db.add(m.StatementControlMapping(tenant_id=1, statement_id=ps4.id, control_kind="normalized", control_code="IAC-17"))
    db.add(m.Issue(id=3, tenant_id=1, title="Reviews were late"))
    db.add(m.InternalControl(id=9, tenant_id=1, control_id="IC-9", name="Yearly access review", source_document_id=5, source_statement_id=ps4.id))
    db.flush()
    db.add(m.IssueGovernanceLink(issue_id=3, policy_statement_id=ps4.id))
    db.query(m.PolicyGapFinding).filter_by(id=1).update({"applied_statement_id": ps4.id})
    db.commit()

    out = world.http.delete(f"/documents/5/statements/{ps4.id}")
    assert out.status_code == 200, out.text
    assert db.query(m.PolicyStatement).filter_by(statement_code="PS-0005-004").count() == 0
    assert db.query(m.StatementControlMapping).count() == 0 and db.query(m.IssueGovernanceLink).count() == 0     # links that only pointed at it
    assert db.get(m.Issue, 3) is not None and db.get(m.InternalControl, 9).source_statement_id is None              # the records stay
    assert db.get(m.PolicyGapFinding, 1).applied_statement_id is None
    (row,) = rows(db, action="statement_deleted")
    assert (row.field_changed, row.old_value, row.new_value, row.action_details) == ("statement", LONG, None, "PS-0005-004")


def test_every_reference_to_a_statement_is_cleared_or_cascaded_when_it_is_removed():
    cascaded = {rel.mapper.local_table.name for rel in inspect(m.PolicyStatement).relationships if rel.cascade.delete}
    listed = {f"{a.class_.__tablename__}.{a.key}" for a in policy_parser._DELETE_WITH_STATEMENT + policy_parser._UNLINK_FROM_STATEMENT}
    missing = [f"{t.name}.{fk.parent.name}" for t in m.Base.metadata.sorted_tables for fk in t.foreign_keys
               if fk.column.table.name == m.PolicyStatement.__tablename__ and t.name not in cascaded and f"{t.name}.{fk.parent.name}" not in listed]
    assert not missing, missing
