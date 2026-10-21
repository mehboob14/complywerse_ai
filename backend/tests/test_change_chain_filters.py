"""The Change log's filters on its export ("Export filtered", the user 2026-10-21: "filter on changes made by particular user, date, or
source of changes"): people, a time span, the source of a change, which entries (changes to the text, reviews and approvals,
statements, other activity) and a search narrow the chain as the screen does (grc-frontend/src/lib/changeLog.ts); the counts are of
what is left and the file says it is filtered. Without filters the export is what it was."""
import csv
import io
from datetime import datetime, timedelta

import pytest

import grc.models as m
from grc.modules.governance import change_chain as C

from test_change_chain import POLICY, _audit, _edit, library        # noqa: F401  (library is a fixture)
from test_studio_flow import world                                    # noqa: F401  (world is a fixture, library's own)

T0 = datetime(2026, 10, 1, 9, 0)


def _change(n, status, section, before, after, origin, by, by_id, **extra):
    return {"n": n, "status": status, "section": section, "locked": False, "blocks": 1, "words_added": len(after.split()) if after else 0,
            "words_removed": len(before.split()) if before else 0, "format_only": False, "before_text": before, "after_text": after, "truncated": False,
            "before_digests": [], "after_digests": [], "origin": origin, "by": by, "by_id": by_id, "at": (T0 + timedelta(days=2)).isoformat(),
            "request": extra.get("request", ""), "reason": extra.get("reason", ""), "sources": extra.get("sources", []), "ref": {}, "coverage": 1.0}


def _log(library):
    """1.0 the original, 1.1 Olivia by hand, 1.2 saved by Sam with a Comply Agent change Olivia accepted and his own edit; a review
    step, a statement edited by Sam and an owner change between them."""
    db = library.db
    doc = db.get(m.GovernanceDocument, 10)
    doc.content, doc.status = POLICY, "draft"
    db.commit()
    _edit(db, 10, POLICY.replace("every year", "every six months"), user_id=1, reason="Review frequency", at=T0)
    _audit(db, 10, 1, "sent_for_review", T0 + timedelta(hours=1), status="pending_review", old="draft")
    row = _edit(db, 10, db.get(m.GovernanceDocument, 10).content.replace("by the CISO.", "by the CISO and logged."), user_id=2,
                reason="Évidence of approval", at=T0 + timedelta(days=2))
    iso = {"ref": "R1", "kind": "framework", "title": "ISO/IEC 27001", "locator": "A.8.2", "quote": "privileged access rights", "basis": "reference", "verified": True}
    row.change_set = {"v": 1, "items": [
        _change(1, "modified", "3. Access Control", "Privileged access shall be approved by the CISO.", "Privileged access shall be approved by the CISO and logged.",
                "ai", "Olivia Owner", 1, request="Ask for proof", reason="An approval leaves a record", sources=[iso]),
        _change(2, "inserted", "2. Scope", "", "Third parties with access are in scope.", "manual", "Sam Other", 2)], "sources": [iso]}
    db.add(m.DocumentAuditLog(document_id=10, tenant_id=1, action="statement_edited", field_changed="statement", action_details="PS-0010-001 edited",
                              old_value="Accounts are reviewed every year.", new_value="Accounts are reviewed every quarter.", performed_by=2,
                              performed_at=T0 + timedelta(days=3)))
    db.commit()
    _audit(db, 10, 1, "owner_assigned", T0 + timedelta(days=4), details="Owner: Sam Other")
    library.me["id"] = 1
    return library.chain()


def _shape(c):
    return [(v["number"], [i["n"] for i in v["items"]]) for v in c["versions"]], [e["kind"] for e in c["events"]]


def test_without_filters_the_chain_is_kept_as_it_is_and_unknown_values_are_refused(library):
    c = _log(library)
    assert _shape(c) == ([("1.0", []), ("1.1", [1]), ("1.2", [1, 2])], ["sent_for_review", "statement_edited", "owner_assigned"])
    assert C.filter_chain(c) is c and C.filter_chain(c, people=[], sources=[], kinds=[], search="  ", date_from="", date_to=None) is c
    for bad in ({"kinds": ["approvals"]}, {"sources": ["robot"]}, {"date_from": "last week"}):
        with pytest.raises(ValueError):
            C.filter_chain(c, **bad)
    assert library.lib.get("/governance/documents/10/change-chain/export?format=csv&kinds=approvals").status_code == 400
    assert library.lib.get("/governance/documents/10/change-chain/export?format=csv&date_to=tomorrow").status_code == 400


def test_a_person_keeps_their_own_changes_in_each_version_and_their_own_activity(library):
    c = _log(library)
    sam = C.filter_chain(c, people=["2"])
    assert _shape(sam) == ([("1.2", [2])], ["statement_edited"])          # his edit in 1.2, not Olivia's accepted change in it
    assert sam["summary"] == {"versions": 1, "changes": 1, "people": 1, "ai_assisted": 0, "sources": 0, "words_added": 7, "words_removed": 0}
    assert sam["versions"][0]["changes"] == 1 and sam["filters"] == [("People", "Sam Other")]
    olivia = C.filter_chain(c, people=["1"])
    assert _shape(olivia) == ([("1.0", []), ("1.1", [1]), ("1.2", [1])], ["sent_for_review", "owner_assigned"])     # 1.0 is hers too
    assert olivia["summary"]["ai_assisted"] == 1 and olivia["sources"][0]["title"] == "ISO/IEC 27001"
    both = C.filter_chain(c, people=["1", "2"])
    assert _shape(both) == _shape(c) and both["filters"] == [("People", "Olivia Owner, Sam Other")]
    # someone known only by name in an older change is the person with that name
    c["versions"][2]["items"][1]["by_id"] = None
    assert _shape(C.filter_chain(c, people=["2"]))[0] == [("1.2", [2])]


def test_a_source_narrows_the_changes_and_never_removes_activity(library):
    c = _log(library)
    ai = C.filter_chain(c, sources=["ai"])
    assert _shape(ai) == ([("1.2", [1])], ["sent_for_review", "statement_edited", "owner_assigned"])
    assert ai["filters"] == [("Source", "Comply Agent")]
    assert _shape(C.filter_chain(c, sources=["manual"]))[0] == [("1.1", [1]), ("1.2", [2])]
    assert _shape(C.filter_chain(c, sources=["initial"]))[0] == [("1.0", [])]             # the original matches by itself


def test_the_kinds_choose_the_entries_and_each_event_is_a_review_a_statement_or_other_activity(library):
    c = _log(library)
    assert _shape(C.filter_chain(c, kinds=["text"])) == ([("1.0", []), ("1.1", [1]), ("1.2", [1, 2])], [])
    assert _shape(C.filter_chain(c, kinds=["review"])) == ([], ["sent_for_review"])
    assert _shape(C.filter_chain(c, kinds=["statement", "other"])) == ([], ["statement_edited", "owner_assigned"])
    group = {k: C.event_group({"kind": k, "to": None}) for k in C.EVENT_LABEL}
    assert {k for k, g in group.items() if g == "other"} == {"created", "archived", "restored", "status_changed", "version_rollback", "owner_assigned"}
    assert {k for k, g in group.items() if g == "statement"} == set(C.STATEMENT_EVENT_LABEL) | {"statements_updated_from_document"}
    assert {"sent_for_review", "signed", "approved", "acknowledged", "changes_requested", "route_set"} <= {k for k, g in group.items() if g == "review"}
    assert C.event_group({"kind": "status_changed", "to": "published"}) == "review"             # set to published by hand: an approval
    assert C.event_group({"kind": "reviewers_asked_again", "to": None}) == "review"             # a new step of the trail


def test_a_time_span_keeps_what_happened_in_it_in_utc(library):
    c = _log(library)
    later = C.filter_chain(c, date_from=(T0 + timedelta(days=1)).isoformat())
    assert _shape(later) == ([("1.2", [1, 2])], ["statement_edited", "owner_assigned"])
    assert later["filters"] == [("From", "2026-10-02 09:00 UTC")]
    karachi = C.filter_chain(c, date_from="2026-10-02T14:00:00+05:00")                           # the same moment, told in another zone
    assert _shape(karachi) == _shape(later)
    before = C.filter_chain(c, date_to=(T0 + timedelta(days=1)).isoformat() + "Z")
    assert _shape(before) == ([("1.0", []), ("1.1", [1])], ["sent_for_review"])                # the end is not included
    assert _shape(C.filter_chain(c, date_from="2026-10-04T00:00:00", date_to="2026-10-05T00:00:00")) == ([], ["statement_edited"])


def test_a_search_finds_words_in_the_changes_reasons_and_details_without_accents(library):
    c = _log(library)
    assert _shape(C.filter_chain(c, search="third PARTIES")) == ([("1.2", [2])], [])            # only the change that says it
    assert _shape(C.filter_chain(c, search="iso a.8.2")) == ([("1.2", [1])], [])               # the source a change rests on
    assert _shape(C.filter_chain(c, search="evidence")) == ([("1.2", [1, 2])], [])             # "Évidence": the version's own reason
    assert _shape(C.filter_chain(c, search="v1.1")) == ([("1.1", [1])], [])                    # a version by its number, all of it
    assert _shape(C.filter_chain(c, search="quarter")) == ([], ["statement_edited"])          # a statement's new wording
    assert _shape(C.filter_chain(c, search="owner sam")) == ([], ["owner_assigned"])            # every word, in one entry


def test_export_filtered_holds_what_the_screen_shows_and_says_it_is_filtered(library):
    from openpyxl import load_workbook

    _log(library)
    url = "/governance/documents/10/change-chain/export"
    whole = library.lib.get(f"{url}?format=csv")
    rows = list(csv.reader(io.StringIO(whole.content.decode("utf-8-sig"))))
    assert [r[2] for r in rows[1:]] == ["1.1.1", "1.2.1", "1.2.2"] and "filtered" not in whole.headers["content-disposition"]

    sam = library.lib.get(f"{url}?format=csv&people=2")
    rows = list(csv.reader(io.StringIO(sam.content.decode("utf-8-sig"))))
    assert [r[2] for r in rows[1:]] == ["1.2.2"] and sam.headers["content-disposition"].endswith('-change-log-filtered.csv"')

    params = [("format", "xlsx"), ("kinds", "text"), ("kinds", "review"), ("sources", "ai"), ("date_from", "2026-10-01T00:00:00")]
    xlsx = library.lib.get(url, params=params)
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Filtered"] == "From: 2026-10-01 00:00 UTC; Source: Comply Agent; Showing: Changes to the text, Reviews & approvals"
    assert summary["Versions"] == "1" and summary["Changes recorded"] == "1"
    assert [r[0] for r in wb["Changes"].iter_rows(min_row=2, values_only=True)] == ["1.2.1"]
    assert [r[1] for r in wb["Approvals"].iter_rows(min_row=2, values_only=True)] == ["Sent for review"]
    pdf = library.lib.get(url, params=[("format", "pdf"), ("search", "quarter")])
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
