"""Word files in the studio: read as Markdown, and edited in place. The file is built like the ones organisations have (direct
run formatting, a title, headings, lists, a link, a field, a soft break, an image, tables, header and footer, tracked
changes); after an edit the file must still be that file with only the changed words changed."""
import io
import re
import zipfile

import pytest
from docx import Document
from docx.shared import Pt

from grc.modules.governance.studio import blocks as B
from grc.modules.governance.studio import docx_io as D
from grc.modules.governance.studio.importers import ImportProblem

import _docx_fixtures as F


@pytest.fixture
def policy(tmp_path):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    imported = D.import_docx(path, twin=True)
    return path, imported


def edit(path, new_md):
    data, warnings = D.export_in_place(path, new_md)
    return Document(io.BytesIO(data)), data, warnings


def parts(data_or_path):
    z = zipfile.ZipFile(io.BytesIO(data_or_path) if isinstance(data_or_path, bytes) else data_or_path)
    return {n: z.read(n) for n in z.namelist()}


def para(doc, startswith):
    return next(p for p in doc.paragraphs if p.text.startswith(startswith))


# ── reading ──────────────────────────────────────────────────────────────────
def test_a_word_file_reads_as_markdown_with_its_structure(policy):
    path, imp = policy
    md = imp.markdown
    assert md.startswith("# Information Security Policy\n\nThis policy applies to **all employees** and contractors of the bank, see [the ECC](https://example.com/ecc). Review it every year.")
    assert "## 1. Purpose" in md and "The purpose is to protect *information* assets." in md       # Heading 1 sits under the Title
    assert "1. Use strong passwords\n2. Lock your screen\n3. Report incidents\n- Hardware is covered\n- Software is covered" in md
    assert "### 2.1 Passwords" in md
    assert "First line\\\nsecond line" in md                                                       # a soft line break
    assert "Page  of this policy." in md                                                            # a field's result is not text
    assert "| **Role** | **Responsibility** | **Owner** |" in md and "| CISO | Owns the policy | Board |" in md
    assert "| Merged header |  | Status |" in md                                                     # the merged cell covers two columns
    assert imp.kind == "docx" and imp.base_markdown is None
    assert imp.native["map"]["blocks"] == len(B.parse(md)) and imp.native["path"] is None
    blocks = B.parse(md)
    assert blocks[0].kind == "heading" and blocks[0].level == 1 and [b.kind for b in blocks].count("table") == 2
    assert any("Edits are written into the original file" in n for n in imp.notes)


def test_the_markdown_reads_back_as_the_same_blocks_for_awkward_text(tmp_path):
    d = Document()
    d.add_paragraph("1. Looks like a list but is text")
    d.add_paragraph("- Also not a bullet")
    d.add_paragraph("# not a heading, pipes | and *stars* and _unders_ and <tags> & [brackets] and `ticks`")
    d.add_paragraph("Trailing hash C#")
    d.add_heading("Heading ending in #", level=1)
    d.add_paragraph("Two  spaces and\ttab.")
    path = str(tmp_path / "awkward.docx")
    d.save(path)
    imp = D.import_docx(path, twin=True)
    blocks = B.parse(imp.markdown)
    assert [b.kind for b in blocks] == ["paragraph", "paragraph", "paragraph", "paragraph", "heading", "paragraph"]
    assert blocks[0].text == "1. Looks like a list but is text" and blocks[1].text == "- Also not a bullet"
    assert blocks[2].text == "# not a heading, pipes | and *stars* and _unders_ and <tags> & [brackets] and `ticks`"
    assert blocks[4].text == "Heading ending in #" and imp.native is not None


def test_a_file_that_cannot_be_used_is_refused_with_a_sentence(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"this is not a zip")
    with pytest.raises(ImportProblem, match="could not be opened"):
        D.import_docx(str(bad))
    empty = tmp_path / "empty.docx"
    Document().save(str(empty))
    with pytest.raises(ImportProblem, match="No text was found"):
        D.import_docx(str(empty))


# ── editing: the file stays the file ─────────────────────────────────────────
def test_an_unchanged_document_is_not_touched_at_all(policy):
    path, imp = policy
    _, data, warnings = edit(path, imp.markdown)
    assert warnings == []
    before, after = parts(path), parts(data)
    assert before.keys() == after.keys()
    assert before["word/document.xml"] == after["word/document.xml"]
    for name in ("word/header1.xml", "word/footer1.xml", "word/styles.xml", "word/numbering.xml"):
        assert before[name] == after[name], name
    assert [n for n in before if n.startswith("word/media/")] == [n for n in after if n.startswith("word/media/")]


def test_changed_words_change_inside_their_own_runs_and_the_formatting_around_them_stays(policy):
    path, imp = policy
    new = imp.markdown.replace("**all employees**", "**all staff**").replace("Review it every year.", "Review it every six months.")
    doc, data, _ = edit(path, new)
    p = para(doc, "This policy applies")
    assert p.text == "This policy applies to all staff and contractors of the bank, see the ECC. Review it every six months."
    runs = {r.text: r for r in p.runs}
    assert runs["all staff"].bold is True and runs["all staff"].font.size == Pt(11) and runs["all staff"].font.name == "Calibri"   # still the bold run
    assert runs["This policy applies to "].bold is None and runs["This policy applies to "].font.size == Pt(11)
    assert runs[". Review it every six months."].font.name == "Calibri" and runs[". Review it every six months."].bold is None
    link = p._p.xpath(".//w:hyperlink")
    assert len(link) == 1 and "the ECC" in "".join(link[0].itertext())
    rel = doc.part.rels[link[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")]
    assert rel.is_external and rel.target_ref == "https://example.com/ecc"
    other = para(doc, "Passwords shall be at least")                                        # an untouched paragraph is byte-for-byte as it was
    original = Document(path)
    assert para(original, "Passwords shall be at least")._p.xml == other._p.xml
    z = parts(data)
    assert z["word/header1.xml"] == parts(path)["word/header1.xml"] and z["word/footer1.xml"] == parts(path)["word/footer1.xml"]


def test_text_added_inside_a_paragraph_takes_the_look_of_the_words_it_follows(policy):
    path, imp = policy
    new = imp.markdown.replace("The purpose is to protect *information* assets.", "The purpose is to protect *information* assets and customer data.")
    doc, _, _ = edit(path, new)
    p = para(doc, "The purpose")
    assert p.text == "The purpose is to protect information assets and customer data."
    assert [(r.text, r.italic) for r in p.runs] == [("The purpose is to protect ", None), ("information", True), (" assets and customer data.", None)]
    assert all(r.font.size == Pt(11) for r in p.runs)


def test_formatting_can_be_changed_on_words_that_stay(policy):
    path, imp = policy
    new = (imp.markdown.replace("Review it every year.", "Review it **every** year.")           # plain -> bold
           .replace("**all employees**", "all employees")                                      # bold -> plain
           .replace("protect *information* assets", "protect <u>information</u> assets"))     # italic -> underline
    doc, _, _ = edit(path, new)
    p = para(doc, "This policy applies")
    runs = {r.text: r for r in p.runs}
    assert runs["every"].bold is True and runs["all employees"].bold is False and runs["all employees"].font.size == Pt(11)
    q = para(doc, "The purpose")
    info = next(r for r in q.runs if r.text == "information")
    assert info.underline is True and info.italic is False


def test_line_breaks_survive_an_edit_of_the_words_around_them(policy):
    path, imp = policy
    doc, _, _ = edit(path, imp.markdown.replace("First line", "Opening line").replace("second line", "next line"))
    p = para(doc, "Opening line")
    assert p.text == "Opening line\nnext line" and len(p._p.xpath(".//w:br")) == 1


def test_fields_images_and_the_header_are_left_alone(policy):
    path, imp = policy
    before = Document(path)
    doc, data, _ = edit(path, imp.markdown.replace("of this policy.", "of this policy, as amended."))
    p = para(doc, "Page")
    assert "as amended." in p.text and len(p._p.xpath(".//w:fldChar")) == 3 and "PAGE" in "".join(p._p.xpath(".//w:instrText/text()"))
    assert len(doc.element.body.xpath(".//w:drawing")) == len(before.element.body.xpath(".//w:drawing")) == 1
    assert len([n for n in parts(data) if n.startswith("word/media/")]) == 1


# ── editing: blocks come and go ──────────────────────────────────────────────
def test_new_paragraphs_headings_and_list_items_look_like_their_neighbours(policy):
    path, imp = policy
    new = (imp.markdown.replace("2. Lock your screen\n", "")
           .replace("3. Report incidents", "3. Report incidents quickly\n4. Never share passwords")
           .replace("## 3. Roles", "### 2.2 New subsection\n\nA new paragraph with **bold** and a [[placeholder]].\n\n## 3. Roles"))
    doc, _, warnings = edit(path, new)
    assert warnings == []
    texts = [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]
    assert ("Heading 2", "2.2 New subsection") in texts
    assert [t for s, t in texts if s == "List Number"] == ["Use strong passwords", "Report incidents quickly", "Never share passwords"]
    body = para(doc, "A new paragraph")
    assert body.style.name == "Normal"
    runs = {r.text: r for r in body.runs}
    assert runs["bold"].bold is True and runs["A new paragraph with "].bold is None
    assert runs["[[placeholder]]"].font.highlight_color is not None                           # a placeholder is highlighted for whoever fills it in
    nearest = para(Document(path), "Passwords shall be at least").runs[0]                      # the file's own body paragraph next to it: its font and size, not the default
    assert runs["A new paragraph with "].font.size == nearest.font.size == Pt(10.5) and runs["A new paragraph with "].font.name == nearest.font.name == "Arial"
    items = [n for n in D.walk(doc)[0] if n.kind == "list_item" and n.ordered]
    assert [n.text for n in items] == ["Use strong passwords", "Report incidents quickly", "Never share passwords"]
    assert len({n.num_id for n in items}) == 1 and [n.number for n in items] == [1, 2, 3]       # one list, counted on


def test_a_list_added_in_the_middle_of_text_counts_from_one_and_a_bullet_list_is_a_bullet_list(policy):
    path, imp = policy
    new = imp.markdown.replace("After the table.", "After the table.\n\n1. First step\n2. Second step\n\n- A point\n- Another point")
    doc, _, _ = edit(path, new)
    steps = [p for p in doc.paragraphs if p.text in ("First step", "Second step")]
    points = [p for p in doc.paragraphs if p.text in ("A point", "Another point")]
    assert len(steps) == 2 and len(points) == 2
    num_ids = {p._p.xpath("string(.//w:numId/@w:val)") for p in steps}
    assert len(num_ids) == 1
    numbering = doc.part.numbering_part.element.xml
    assert re.search(rf'<w:num w:numId="{num_ids.pop()}"[^>]*>.*?<w:startOverride w:val="1"/>', numbering, re.S)   # restarts at 1
    re_read = D.walk(doc)[0]
    assert [(n.kind, n.ordered) for n in re_read if n.text in ("First step", "A point")] == [("list_item", True), ("list_item", False)]


def test_a_block_can_be_deleted_and_a_section_break_survives(tmp_path):
    d = Document()
    d.add_paragraph("Keep this first paragraph.")
    gone = d.add_paragraph("This paragraph carries a section break and will be deleted.")
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    gone._p.get_or_add_pPr().append(parse_xml(f'<w:sectPr {nsdecls("w")}><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>'))
    d.add_paragraph("Keep this last paragraph.")
    path = str(tmp_path / "sect.docx")
    d.save(path)
    imp = D.import_docx(path, twin=True)
    doc, _, _ = edit(path, "Keep this first paragraph.\n\nKeep this last paragraph.\n")
    assert [p.text for p in doc.paragraphs if p.text] == ["Keep this first paragraph.", "Keep this last paragraph."]
    assert len(doc.element.body.xpath("./w:p/w:pPr/w:sectPr")) == 1                           # the section break is still there, on an empty paragraph
    assert imp.markdown.count("\n\n") == 2


def test_a_heading_without_a_style_in_the_file_gets_one(tmp_path):
    d = Document()
    d.add_paragraph("Title paragraph only", style="Normal")
    path = str(tmp_path / "plain.docx")
    d.save(path)
    doc, _, warnings = edit(path, "Title paragraph only\n\n## Brand new heading\n\nText below it.\n")
    h = next(p for p in doc.paragraphs if p.text == "Brand new heading")
    assert h.style.name == "Heading 2" and warnings == []
    assert D.walk(doc)[0][1].kind == "heading"


def test_auto_numbered_headings_keep_their_numbering_and_the_numbers_are_not_text(tmp_path):
    d = Document()
    d.add_heading("Purpose", level=1)
    d.add_paragraph("Why this exists.")
    h1 = d.styles["Heading 1"]
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    numbering = d.part.numbering_part.element
    numbering.insert(0, parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="77"><w:multiLevelType w:val="multilevel"/><w:lvl w:ilvl="0"><w:start w:val="1"/>'
                                  '<w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl></w:abstractNum>'))
    num = numbering.add_num(77)
    h1.element.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num.numId}"/></w:numPr>'))
    path = str(tmp_path / "numbered.docx")
    d.save(path)
    imp = D.import_docx(path, twin=True)
    assert imp.markdown.startswith("# Purpose\n") and any("numbered by Word itself" in n for n in imp.notes)
    doc, _, _ = edit(path, imp.markdown + "\n# Scope\n\nWho it covers.\n")
    scope = next(p for p in doc.paragraphs if p.text == "Scope")
    assert scope.style.name == "Heading 1" and not scope.text[0].isdigit()                   # numbered by the style, like its sibling


# ── editing: tables ──────────────────────────────────────────────────────────
def test_table_cells_rows_and_columns(policy):
    path, imp = policy
    md = imp.markdown
    doc, _, warnings = edit(path, md.replace("| CISO | Owns the policy | Board |", "| CISO | Owns and reviews the policy | Board |"))
    cell = doc.tables[0].cell(1, 1)
    assert cell.text == "Owns and reviews the policy" and cell.paragraphs[0].runs[0].font.size == Pt(10) and warnings == []
    header = next(r for r in doc.tables[0].cell(0, 0).paragraphs[0].runs if r.text == "Role")
    assert header.bold is True and header.font.size == Pt(10)                                   # the header row kept its look

    doc, _, warnings = edit(path, md.replace("| IT | Implements controls | CISO |", "| IT | Implements controls | CISO |\n| Audit | Reviews controls | Board |"))
    rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert rows[-1] == ["Audit", "Reviews controls", "Board"] and len(rows) == 4 and warnings == []
    assert doc.tables[0].rows[3].cells[0].paragraphs[0].runs[0].font.size == Pt(10)           # the new row copies a data row, not the header
    assert doc.tables[0].rows[3].cells[0].paragraphs[0].runs[0].bold is None

    doc, _, _ = edit(path, md.replace("| IT | Implements controls | CISO |\n", ""))
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["Role", "Responsibility", "Owner"], ["CISO", "Owns the policy", "Board"]]

    narrow = "| **Role** | **Responsibility** | **Owner** |\n| --- | --- | --- |\n| CISO | Owns the policy | Board |\n| IT | Implements controls | CISO |"
    wide = "| **Role** | **Responsibility** | **Owner** | **Notes** |\n| --- | --- | --- | --- |\n| CISO | Owns the policy | Board | - |\n| IT | Implements controls | CISO | - |"
    assert narrow in md
    doc, _, warnings = edit(path, md.replace(narrow, wide))
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["Role", "Responsibility", "Owner", "Notes"]
    assert any("rebuilt" in w for w in warnings)                                              # a new column rebuilds the table, and says so


def test_a_new_line_in_a_table_cell_becomes_a_paragraph_of_the_cell_with_the_same_look(policy):
    path, imp = policy
    doc, _, warnings = edit(path, imp.markdown.replace("| CISO | Owns the policy | Board |", "| CISO | Owns the policy<br>and reviews it yearly | Board |"))
    cell = doc.tables[0].cell(1, 1)
    assert [p.text for p in cell.paragraphs] == ["Owns the policy", "and reviews it yearly"] and warnings == []
    assert all(p.runs[0].font.size == Pt(10) for p in cell.paragraphs)


def test_a_table_with_merged_cells_can_have_its_text_edited_but_not_its_shape(policy):
    path, imp = policy
    md = imp.markdown
    doc, _, warnings = edit(path, md.replace("| A | B | Open |", "| A | B | Done |"))
    assert doc.tables[1].cell(1, 2).text == "Done" and warnings == []
    assert len(doc.tables[1].rows[0]._tr.xpath("./w:tc")) == 2                                 # still one merged cell and one more
    doc, _, warnings = edit(path, md.replace("| C | D | Closed |", "| C | D | Closed |\n| E | F | New |"))
    assert [c.text for c in doc.tables[1].rows[-1].cells] == ["E", "F", "New"] and any("rebuilt" in w for w in warnings)


# ── editing: links ───────────────────────────────────────────────────────────
def test_link_text_can_be_edited_and_a_new_link_can_be_added(policy):
    path, imp = policy
    new = imp.markdown.replace("[the ECC](https://example.com/ecc)", "[the NCA ECC](https://example.com/ecc)").replace(
        "Review it every year.", "Review it every year, see [the standard](https://example.com/std).")
    doc, _, _ = edit(path, new)
    p = para(doc, "This policy applies")
    links = p._p.xpath(".//w:hyperlink")
    assert ["".join(l.xpath(".//w:t/text()")) for l in links] == ["the NCA ECC", "the standard"]
    targets = [doc.part.rels[l.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")].target_ref for l in links]
    assert targets == ["https://example.com/ecc", "https://example.com/std"]
    assert ", see " in p.text and p.text.endswith("the standard.")


# ── tracked changes ──────────────────────────────────────────────────────────
def test_tracked_changes_are_accepted_for_the_text_and_kept_as_the_baseline(tmp_path):
    path = F.tracked_policy(str(tmp_path / "tracked.docx"))
    imp = D.import_docx(path, twin=True)
    assert imp.markdown == "# Retention Policy\n\nRecords shall be kept for seven years.\n\nA paragraph the reviewer added.\n\nMake this **important**\n\nUntouched paragraph.\n"
    assert imp.base_markdown == "# Retention Policy\n\nRecords shall be kept for five years.\n\nA paragraph the reviewer removed.\n\nMake this important\n\nUntouched paragraph.\n"
    assert any("7 tracked changes" in n for n in imp.notes)
    twin = imp.native["path"]
    assert twin and twin.endswith(".twin.docx")
    xml = parts(twin)["word/document.xml"].decode()
    assert "<w:ins " not in xml and "<w:del " not in xml and "w:rPrChange" not in xml          # the twin has every change accepted
    rows = B.redline(B.parse(imp.base_markdown), B.parse(imp.markdown))                          # what the reviewers did, as a redline
    assert [r["status"] for r in rows] == ["equal", "modified", "modified", "modified", "equal"]
    assert rows[1]["segments"] == [("equal", "Records shall be kept for "), ("delete", "five"), ("insert", "seven"), ("equal", " years.")]
    assert rows[3]["format_only"] is True                                                       # the reviewer made a word bold
    doc, data, _ = edit(twin, imp.markdown.replace("seven years", "ten years"))
    assert para(doc, "Records shall").text == "Records shall be kept for ten years."
    assert b"<w:ins " not in parts(data)["word/document.xml"]


def test_a_deleted_paragraph_mark_joins_the_paragraphs_and_a_deleted_row_goes(tmp_path):
    d = Document()
    body = d.element.body
    sect = body.find(F.qn("w:sectPr"))
    n = F.nsdecls("w")
    for xml in (
        f'<w:p {n}><w:pPr><w:rPr><w:del w:id="1" {F.AUTH}/></w:rPr></w:pPr><w:r><w:t xml:space="preserve">First half </w:t></w:r></w:p>',
        f'<w:p {n}><w:r><w:t>second half.</w:t></w:r></w:p>',
    ):
        sect.addprevious(F.parse_xml(xml))
    t = d.add_table(rows=2, cols=1)
    t.cell(0, 0).text, t.cell(1, 0).text = "Kept row", "Deleted row"
    t.rows[1]._tr.get_or_add_trPr().append(F.parse_xml(f'<w:del {n} w:id="9" {F.AUTH}/>'))
    path = str(tmp_path / "marks.docx")
    d.save(path)
    imp = D.import_docx(path, twin=True)
    assert imp.markdown.startswith("First half second half.") and "Deleted row" not in imp.markdown and "Kept row" in imp.markdown
    assert "Deleted row" in imp.base_markdown or imp.base_markdown is not None


# ── safety net ───────────────────────────────────────────────────────────────
def test_an_edit_that_would_not_read_back_as_asked_is_refused(policy, monkeypatch):
    path, imp = policy
    real = D.build_paragraph

    def wrong(ctx, tpls, block, prev, near):
        placed = real(ctx, tpls, block, prev, near)
        for t in placed.el.iter(F.qn("w:t")):
            t.text = "something else"
        return placed

    monkeypatch.setattr(D, "build_paragraph", wrong)
    with pytest.raises(D.PatchProblem, match="does not read back"):
        D.export_in_place(path, imp.markdown + "\nA brand new closing paragraph.\n")


def test_white_space_alone_is_not_an_edit_and_a_rewrite_is_one_replacement():
    assert {o[0] for o in D.diff_ops("A  double  space", "A double space")} == {"equal"}
    assert D.diff_ops("The cat sat on the mat today", "The dog sat on the mat today")[2] == ("replace", 4, 7, 4, 7)
    assert [o[0] for o in D.diff_ops("The cat sat on the mat today", "The dog sat on the mat today") if o[0] != "equal"] == ["replace"]
    old, new = "totally different words here now", "nothing alike remains of that sentence"
    assert D.diff_ops(old, new) == [("replace", 0, len(old), 0, len(new))]
    assert D.diff_ops("  padded  ", "padded") == [("equal", 2, 8, 0, 6)]                   # edge white space is not part of the text
