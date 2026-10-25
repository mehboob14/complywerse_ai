"""The document control pages as part of the document (modules/governance/approval_record.py keys ``control``, ``revision_history`` and
``approval_history``; approval_record_export.py, controlled_copy.py and the endpoints ``controlled-content`` and ``download-file`` /
``export`` with ``with_record``): where the places for them are in a document's Markdown, how a platform-written document's PDF is put
together, where they are inserted in a Word file or a PDF (after the cover, else at the end), the grids of signatures, and what the
endpoints answer. Fixtures and helpers are those of test_approval_record.py (Olivia owns and wrote the document, Rana and Rafay review,
Mina and Mo approve, Vera only reads the library); the files are checked on records made by hand where the database is not the point.
"""
import io
import re
import zipfile
from datetime import datetime

import fitz
import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree

import grc.models as m
from grc.modules.governance import approval_record
from grc.modules.governance import approval_record_export as X
from grc.modules.governance import controlled_copy as C
from grc.modules.governance.router import router

from _docx_fixtures import corporate_policy, png as ink, tracked_policy
from test_approval_record import (ADAM, DOCX_MEDIA, MINA, MO, OLIVIA, RAFAY, RANA, SMALL, TEXT, UMA, VERA, WHEN, approve_and_publish, ask_changes,  # noqa: F401
                                  attach, doc, edit, finished, get, names, pdf_file, pdf_text, person, publish, record, reject, role, send, set_route,
                                  sign, stage, usual, user)
from test_approval_record import _schema, _tenants, world  # noqa: F401  (the fixtures)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ══ where the places are in a document's text ════════════════════════════════════
SAMPLE = """# Access policy

Prepared for the board.

## 1. Document Description

| Field | Value |
| --- | --- |
| Title | [[Document title]] |

## 2. Purpose

Why we have it.

## 3. Approval Signoff

| Role | Name | Designation | Signature / Date |
| --- | --- | --- | --- |
| Approved by | ______ | ______ | ______ |

### 3.1 How to read this

A sub-section stays.

## 4. Revision History

| Version | Date |
| --- | --- |
| 1.0 | [[Date]] |
"""


def shape(segs):
    """The segments as (slot, has the heading) for a place, "md" for text."""
    return [(s["slot"], s["heading"]) if s["type"] == "slot" else "md" for s in segs]


def texts(segs):
    return [s["text"] for s in segs if s["type"] == "markdown"]


def test_a_document_with_the_three_sections_gets_a_place_in_each_and_keeps_every_other_line():
    segs = C.segments(SAMPLE)
    assert shape(segs) == ["md", ("description", False), "md", ("signoff", False), "md", ("revision", False)]
    assert texts(segs) == [
        "# Access policy\n\nPrepared for the board.\n\n## 1. Document Description\n",              # the heading stays: the document's numbering with it
        "## 2. Purpose\n\nWhy we have it.\n\n## 3. Approval Signoff\n",
        "### 3.1 How to read this\n\nA sub-section stays.\n\n## 4. Revision History\n"]            # a deeper heading ends the body, and survives
    assert all(set(s) == {"type", "text"} for s in segs if s["type"] == "markdown") and all(set(s) == {"type", "slot", "heading"} for s in segs if s["type"] == "slot")
    # the texts joined are the document less the bodies the places stand for
    gone = ["\n| Field | Value |\n| --- | --- |\n| Title | [[Document title]] |\n\n",
            "\n| Role | Name | Designation | Signature / Date |\n| --- | --- | --- | --- |\n| Approved by | ______ | ______ | ______ |\n\n",
            "\n| Version | Date |\n| --- | --- |\n| 1.0 | [[Date]] |\n"]
    expected = SAMPLE
    for body in gone:
        assert body in expected
        expected = expected.replace(body, "", 1)
    assert "".join(texts(segs)) == expected


@pytest.mark.parametrize("heading", [
    "Document Description", "DOCUMENT DESCRIPTION", "1. Document Description", "1.2 Document Description", "Annex A — Document Description",
    "A. Document Description", "Document Control", "Document Information", "Document details", "Document Summary", "**Document Description**",
    "Document Description:", "Document Description ##"])
def test_the_description_is_found_by_its_words_whatever_the_numbering_case_or_marks(heading):
    segs = C.segments(f"# T\n\n## {heading}\n\nold body\n\n## Next\n\ntext\n", have=["description"])
    assert shape(segs) == ["md", ("description", False), "md"] and "old body" not in "".join(texts(segs))
    assert texts(segs)[0].endswith(f"## {heading}\n") and texts(segs)[1] == "## Next\n\ntext\n"


@pytest.mark.parametrize("heading", [
    "Approval Signoff", "Approval Sign-off", "Approval sign off", "Approval", "Approvals", "Sign-off", "Signoff", "Sign Off", "Review and Approval",
    "Review & Approval", "Recommendation and Review", "Recommendation & Review", "Approvals and Sign-off", "Approvals & Sign-off", "3. APPROVAL"])
def test_the_sign_off_section_is_found_under_the_names_banks_and_the_templates_give_it(heading):
    segs = C.segments(f"# T\n\n## {heading}\n\nold body\n\n## Next\n", have=["signoff"])
    assert shape(segs) == ["md", ("signoff", False), "md"]


@pytest.mark.parametrize("heading", ["Revision History", "6 REVISION HISTORY", "Version History", "Change History", "Document History",
                                     "Document Revision History", "Annex B - Revision History"])
def test_the_revision_section_is_found_under_its_names(heading):
    assert shape(C.segments(f"# T\n\n## {heading}\n\nold\n\n## Next\n", have=["revision"])) == ["md", ("revision", False), "md"]


def test_headings_that_only_look_like_them_are_not_sections():
    for words in ("Document Descriptions of Assets", "Approval Authority Matrix", "Revision of this Policy", "History", "Approvers", "Description"):
        segs = C.segments(f"# T\n\n## {words}\n\ntext\n", have=["description", "signoff", "revision"])
        assert [s["slot"] for s in segs if s["type"] == "slot"] == ["description", "signoff", "revision"] and all(s["heading"] for s in segs if s["type"] == "slot")
        assert f"## {words}" in "".join(texts(segs))                                              # left as it is


def test_a_code_block_and_other_marks_are_not_headings():
    md = ("# T\n\nIntro.\n\n```\n## Document Description\nnot a heading\n```\n\n~~~~\n## Revision History\n~~~~\n\n    ## Approval Signoff\n\n"
          "> ## Approval\n\n#Revision History\n\nDocument Description\n====================\n\n- ## Version History\n")
    segs = C.segments(md)
    assert [(s["slot"], s["heading"]) for s in segs if s["type"] == "slot"] == [("description", True), ("signoff", True), ("revision", True)]   # all missing
    assert "".join(texts(segs)) == md                                                                    # not a character of it was replaced
    # a fence closes only on its own kind of fence, as long as the one that opened it
    md = "# T\n\n````\n```\n## Document Description\n````\n\n## Document Description\n\nbody\n"
    segs = C.segments(md, have=["description"])
    assert shape(segs) == ["md", ("description", False)] and "body" not in "".join(texts(segs))
    unclosed = "# T\n\n```\n## Document Description\ntext\n"
    assert [s["heading"] for s in C.segments(unclosed, have=["description"]) if s["type"] == "slot"] == [True]      # inside the fence to the end: no heading


def test_the_body_ends_at_the_next_heading_of_any_level_and_its_sub_sections_survive():
    md = ("# T\n\n## Approval Signoff\n\nsign-off body\n\n#### deep\n\nstays\n\n## Document Description\n\ndescription body\n\n"
          "# Another top\n\nlast words\n")
    segs = C.segments(md)
    # (the revision history has no heading: its place is after the title, before the first heading)
    assert shape(segs) == ["md", ("revision", True), "md", ("signoff", False), "md", ("description", False), "md"]
    assert texts(segs) == ["# T\n\n", "## Approval Signoff\n", "#### deep\n\nstays\n\n## Document Description\n", "# Another top\n\nlast words\n"]


def test_the_first_section_of_each_kind_wins_and_a_second_stays_as_text():
    md = "# T\n\n## Revision History\n\nfirst\n\n## Version History\n\nsecond\n\n## Document History\n\nthird\n"
    segs = C.segments(md, have=["revision"])
    assert shape(segs) == ["md", ("revision", False), "md"]
    assert "first" not in "".join(texts(segs)) and "## Version History\n\nsecond" in texts(segs)[1] and "third" in texts(segs)[1]


def test_what_the_document_does_not_have_is_a_place_with_a_title_after_its_first_heading_and_the_lines_under_it():
    segs = C.segments("# Access policy\n\nPrepared for the board.\nSecond line.\n\n## 1. Purpose\n\nWhy.\n")
    assert shape(segs) == ["md", ("description", True), ("signoff", True), ("revision", True), "md"]
    assert texts(segs) == ["# Access policy\n\nPrepared for the board.\nSecond line.\n\n", "## 1. Purpose\n\nWhy.\n"]
    # only the missing ones: the document's own section keeps its place, the others come at the top
    md = "# Access policy\n\nLead.\n\n## 1. Purpose\n\nWhy.\n\n## 2. Revision History\n\nold\n"
    segs = C.segments(md)
    assert shape(segs) == ["md", ("description", True), ("signoff", True), "md", ("revision", False)]
    assert texts(segs)[0] == "# Access policy\n\nLead.\n\n" and texts(segs)[1].startswith("## 1. Purpose") and texts(segs)[1].endswith("## 2. Revision History\n")
    # a document with nothing under its title: straight after the title
    assert shape(C.segments("# Title\n\n## One\n\ntext\n", have=["signoff"])) == ["md", ("signoff", True), "md"]
    assert texts(C.segments("# Title\n\n## One\n\ntext\n", have=["signoff"])) == ["# Title\n\n", "## One\n\ntext\n"]
    # a title and nothing else: the pages follow it
    assert shape(C.segments("# Title\n", have=["revision"])) == ["md", ("revision", True)]


def test_a_document_without_a_title_gets_the_places_at_the_very_top():
    segs = C.segments("## 1. Purpose\n\nWhy.\n\n## 2. Scope\n\nWho.\n")
    assert shape(segs) == [("description", True), ("signoff", True), ("revision", True), "md"] and texts(segs) == ["## 1. Purpose\n\nWhy.\n\n## 2. Scope\n\nWho.\n"]
    assert shape(C.segments("Just a paragraph.\n", have=["description"])) == [("description", True), "md"]
    assert shape(C.segments("")) == [("description", True), ("signoff", True), ("revision", True)]
    assert shape(C.segments(None, have=["revision"])) == [("revision", True)]


def test_a_second_level_one_heading_is_not_the_title_and_a_matched_title_keeps_its_place():
    segs = C.segments("# One\n\nlead\n\n# Two\n\nmore\n", have=["signoff"])
    assert texts(segs) == ["# One\n\nlead\n\n", "# Two\n\nmore\n"]
    segs = C.segments("# Document Description\n\nold body\n\n## Purpose\n\nWhy.\n")                      # the title is a section itself
    assert shape(segs) == ["md", ("description", False), ("signoff", True), ("revision", True), "md"]
    assert texts(segs) == ["# Document Description\n", "## Purpose\n\nWhy.\n"]


def test_windows_line_endings_make_the_same_segments():
    unix = C.segments(SAMPLE)
    assert C.segments(SAMPLE.replace("\n", "\r\n")) == unix and C.segments(SAMPLE.replace("\n", "\r")) == unix
    assert all("\r" not in t for t in texts(C.segments(SAMPLE.replace("\n", "\r\n"))))


def test_have_limits_the_places_to_the_kinds_the_platform_provides():
    assert shape(C.segments(SAMPLE, have=["signoff"])) == ["md", ("signoff", False), "md"]
    assert "Document Description" in "".join(texts(C.segments(SAMPLE, have=["signoff"]))) and "[[Document title]]" in "".join(texts(C.segments(SAMPLE, have=["signoff"])))
    assert C.segments("# T\n\ntext\n", have=[]) == [{"type": "markdown", "text": "# T\n\ntext\n"}]
    assert shape(C.segments("# T\n", have=["nonsense", "revision"])) == ["md", ("revision", True)]


def test_an_ai_draft_with_placeholders_is_split_like_any_other_and_its_other_placeholders_stay():
    draft = ("# [[Document title]]\n\n[[One sentence on what this is.]]\n\n## 1. Document Description\n\n| Field | Value |\n| --- | --- |\n| Document Owner | [[Owner name and role]] |\n"
             "| Version | 1.0 |\n\n## 2. Approval Sign-off\n\n| Role | Name | Designation | Signature / Date |\n| --- | --- | --- | --- |\n| Prepared by | [[Name]] | [[Role]] | ___ |\n\n"
             "## 3. Purpose\n\n[[Why the document exists, in two sentences.]]\n\n## 4. Revision History\n\n| Version | Date | Author | Summary of change |\n| --- | --- | --- | --- |\n"
             "| 1.0 | [[Date]] | [[Author]] | Initial version |\n\n## 5. Definitions and Acronyms\n\n| Term | Definition |\n| --- | --- |\n| [[Term]] | [[Definition]] |\n")
    segs = C.segments(draft)
    assert shape(segs) == ["md", ("description", False), "md", ("signoff", False), "md", ("revision", False), "md"]
    joined = "".join(texts(segs))
    assert "[[Owner name and role]]" not in joined and "[[Name]]" not in joined and "[[Date]]" not in joined
    assert "# [[Document title]]" in joined and "[[Why the document exists, in two sentences.]]" in joined and "[[Term]]" in joined
    assert texts(segs)[-1].startswith("## 5. Definitions and Acronyms")


def test_no_segment_is_empty_adjacent_texts_are_one_and_the_result_is_stable():
    for md in (SAMPLE, "## Document Description\n## Approval Signoff\n## Revision History\n", "# T\n## Document Description\n\n\n## Approval Signoff\n\n\n## Revision History\n\n\n\n",
               "\n\n\n", "# T\n", "text only", "# T\n\n```\nunclosed\n"):
        segs = C.segments(md)
        assert all(s["type"] == "slot" or s["text"].strip() for s in segs), md
        assert not any(a["type"] == "markdown" and b["type"] == "markdown" for a, b in zip(segs, segs[1:])), md
        assert C.segments(md) == segs
    segs = C.segments("# T\n## Document Description\n## Approval Signoff\n## Revision History\n")
    assert shape(segs) == ["md", ("description", False), "md", ("signoff", False), "md", ("revision", False)]
    assert texts(segs) == ["# T\n## Document Description\n", "## Approval Signoff\n", "## Revision History\n"]
    segs = C.segments("# Title\nOnly paragraph.\n", have=["revision"])                  # nothing after the title but its paragraph: the pages follow the paragraph
    assert shape(segs) == ["md", ("revision", True)] and texts(segs) == ["# Title\nOnly paragraph.\n"]
    segs = C.segments("# Title\n\nFirst paragraph.\nStill the first.\n\nSecond paragraph.\n\n- a list\n", have=["revision"])
    assert shape(segs) == ["md", ("revision", True), "md"] and texts(segs) == ["# Title\n\nFirst paragraph.\nStill the first.\n", "\nSecond paragraph.\n\n- a list\n"]
    assert "".join(texts(C.segments("no newline at the end", have=["revision"]))) == "no newline at the end"


# ══ the pages themselves, made from a record (no database) ═══════════════════════
def fake(stages=2, per=2, decisions=None, names_=True, notes=()):
    """A record as approval_record.build makes it (keys and shapes), with ``stages`` stages of ``per`` people each, and the pictures."""
    count = iter(range(1, 10_000))
    images, route = {}, []
    for s in range(1, stages + 1):
        people = []
        for k in range(per):
            n = next(count)
            decision = (decisions or {}).get((s, k), "signed")
            images[n] = ink((120, 40), (20, 40, 90, 255))
            signed = decision in ("signed", "rejected", "changes_requested")
            people.append({"name": f"Person {n}", "designation": f"Designation {n}", "decision": decision, "meaning": "Approved by" if s == stages else "Reviewed by",
                           "signed_at": "2026-10-22T14:05:00" if signed else None, "comment": None, "on_behalf_of": None,
                           "signature": {"id": n, "drawn": True, "seal_short": f"{n:08X}"} if decision == "signed" else None})
        route.append({"position": s, "label": "Approved by" if s == stages else "Reviewed by", "name": f"Stage {s} name" if names_ else None, "rule": "all",
                      "min_count": None, "status": "complete", "people": people})
    return {
        "reference": "DAR-9-1.0-ABCDEF12", "generated_at": "2026-10-25T09:30:00",
        "document": {"id": 9, "code": "POL-009", "title": "Access policy", "type": "policy", "classification": "internal", "version": "1.0", "status": "published",
                     "department": "Information Security", "owner": "Olivia Owner", "author": "Olivia Owner", "effective_date": None, "next_review_date": None,
                     "approved_at": "2026-10-23T10:00:00", "published_at": None, "published_by": None},
        "state": "published", "round": {"version": "1.0", "started_at": None, "completed_at": None, "complete": True},
        "prepared_by": {"name": "Olivia Owner", "designation": "Head of Policy", "at": "2026-10-01T09:00:00"},
        "stages": route, "milestones": [], "history": [], "notes": list(notes),
        "control": {"title": "Access policy", "code": "POL-009", "version": "1.0", "type": "policy", "type_label": "Policy", "classification": "internal",
                    "owner": "Olivia Owner", "department": "Information Security", "ownership": "Information Security — Olivia Owner",
                    "approving_authority": "Board of Directors", "date_of_approval": "2026-10-23T10:00:00", "date_of_implementation": "2026-11-01T00:00:00",
                    "review_frequency": "Every 12 months", "last_approval_date": "2026-10-23T10:00:00", "next_review_date": "2027-10-23T00:00:00", "status": "Published"},
        "revision_history": [{"version": "1.0", "date": "2026-10-23T10:00:00", "change_type": "Creation", "summary": "Initial version", "updated_by": "Olivia Owner",
                              "released": True}],
        "approval_history": [{"version": "1.0", "reviewed_by": [{"name": "Person 1", "designation": "Designation 1"}],
                              "approved_by": [{"name": "Person 3", "designation": "Designation 3", "stage": "Board of Directors"}], "approval_date": "2026-10-23T10:00:00"}],
    }, images


def pdf_pages(data):
    with fitz.open(stream=data, filetype="pdf") as pdf:
        return [page.get_text() for page in pdf]


def cell_texts(word):
    """The text of every cell of every table of a Word document, a merged cell once (python-docx lists it once for each column it covers)."""
    from docx.table import _Cell

    return ["\n".join(p.text for p in _Cell(tc, table).paragraphs) for table in word.tables for tr in table._tbl.tr_lst for tc in tr.tc_lst]


def drawn(data):
    with fitz.open(stream=data, filetype="pdf") as pdf:
        return sum(len(page.get_image_info()) for page in pdf)


def test_twelve_stages_with_everyone_in_them_are_all_in_the_pdf_and_the_word_file_with_every_drawn_signature():
    rec, images = fake(stages=12, per=3)
    data = X.to_pdf(rec, images)
    text = "\n".join(pdf_pages(data))
    for s in rec["stages"]:
        assert f"STAGE {s['position']} — " in text
        for p in s["people"]:
            assert p["name"] in text and p["designation"] in text
    assert drawn(data) == 36 and text.count("Signed 2026-10-22 14:05 UTC") == 36
    assert len(pdf_pages(data)) >= 4 and re.search(r"Document control page %d of %d" % (len(pdf_pages(data)), len(pdf_pages(data))), text)
    word = Document(io.BytesIO(X.to_docx(rec, images)))
    cells = "\n".join(cell_texts(word))
    assert len(word.inline_shapes) == 36 and all(f"Person {n}" in cells for n in range(1, 37))
    assert all(f"STAGE {n} — " in cells for n in range(1, 13)) and cells.count("Signed 2026-10-22 14:05 UTC") == 36


def test_a_route_of_twelve_stages_from_the_database_has_every_stage_and_every_drawn_signature_in_both_files(world):
    crew = [RANA, RAFAY, UMA, MO]
    set_route(world, *[stage("review", [person(crew[i % 4]), person(crew[(i + 1) % 4])], name=f"Step {i + 1}") for i in range(11)],
              stage("approval", [person(MINA)], name="Final approval"))
    send(world)
    for i in range(11):
        sign(world, crew[i % 4])
        sign(world, crew[(i + 1) % 4])
    sign(world, MINA)
    rec = record(world)
    assert [s["position"] for s in rec["stages"]] == list(range(1, 13)) and sum(len(s["people"]) for s in rec["stages"]) == 23
    pdf = get(world, "/approval-record/export")
    assert pdf.status_code == 200
    text = "\n".join(pdf_pages(pdf.content))
    assert all(f"STAGE {n} — " in text for n in range(1, 13)) and "STAGE 12 — APPROVED BY · Final approval" in text
    assert drawn(pdf.content) == 23 and text.count("Signed ") == 23 and "Complete · 1 of 1 signed" in text
    word = Document(io.BytesIO(get(world, "/approval-record/export", format="docx").content))
    assert len(word.inline_shapes) == 23 and "\n".join(cell_texts(word)).count("Signed ") == 23
    assert all(f"STAGE {n} — " in "\n".join(cell_texts(word)) for n in range(1, 13))
    # the same in the Word file the pages are put in, and in a PDF
    attach(world, "Data policy.docx", docx_of(lambda d: (cover(d), d.add_page_break(), content(d))), "docx")
    placed_word = Document(io.BytesIO(get(world, "/download-file", with_record="true").content))
    assert len(placed_word.inline_shapes) == 23 and "\n".join(cell_texts(placed_word)).count("Signed ") == 23
    attach(world, "Data policy.pdf", book(pages=3), "pdf")
    assert drawn(get(world, "/download-file", with_record="true").content) == 23


def test_a_stage_is_a_title_bar_over_a_grid_of_two_columns_one_cell_for_each_person():
    rec, images = fake(stages=3, per=3)
    with fitz.open(stream=X.to_pdf(rec, images), filetype="pdf") as pdf:
        page = pdf[0]
        mid = page.rect.width / 2
        at = {n: page.search_for(f"Person {n}")[0] for n in (1, 2, 3)}                 # the grid comes before the history tables on the page
        assert at[1].x1 < mid < at[2].x0 and abs(at[1].y0 - at[2].y0) < 2               # two to a row ...
        assert at[3].y0 > at[1].y0 + 40 and at[3].x1 < mid                              # ... the third begins the next
        bar = page.search_for("STAGE 1 — REVIEWED BY")[0]
        assert bar.y1 < at[1].y0
        pictures = sorted((i["bbox"] for i in page.get_image_info()), key=lambda b: (round(b[1]), b[0]))
        assert len(pictures) >= 3 and all(pictures[k][3] < at[k + 1].y0 for k in range(3))   # each signature is above the printed name
    heads = [w for page in pdf_pages(X.to_pdf(rec, images)) for w in page.splitlines() if w.startswith("STAGE ")]
    assert heads == ["STAGE 1 — REVIEWED BY · Stage 1 name", "STAGE 2 — REVIEWED BY · Stage 2 name", "STAGE 3 — APPROVED BY · Stage 3 name"]


def test_a_title_bar_says_where_the_stage_stands_and_its_rule_when_it_is_not_everyone():
    rec, _images = fake(stages=3, per=2, decisions={(1, 1): "pending", (2, 0): "rejected", (3, 0): "changes_requested", (3, 1): "delegated"})
    rec["stages"][0].update(status="pending", rule="at_least", min_count=2)
    rec["stages"][1].update(status="rejected")
    rec["stages"][2].update(status="changes_requested", rule="any")
    assert X.stage_note(rec["stages"][0]) == "Pending · 1 of 2 signed · at least 2 sign"
    assert X.stage_note(rec["stages"][1]) == "Rejected · 1 of 2 signed"
    assert X.stage_note(rec["stages"][2]) == "Changes requested · 0 of 2 signed · any one signs"
    text = "\n".join(pdf_pages(X.to_pdf(rec, {})))
    for needle in ("Pending · 1 of 2 signed · at least 2 sign", "Rejected · 1 of 2 signed", "Changes requested · 0 of 2 signed · any one signs"):
        assert needle in text


def test_each_kind_of_person_cell_says_what_happened_and_the_empty_line_says_it_waits():
    rec, images = fake(stages=1, per=6, decisions={(1, 1): "pending", (1, 2): "rejected", (1, 3): "changes_requested", (1, 4): "delegated"})
    people = rec["stages"][0]["people"]
    people[2].update(comment="Not ready: the scope is too wide.")
    people[3].update(comment="Clarify who approves an exception.")
    people[4].update(comment="Handed over to Mo Manager")
    people[5].update(on_behalf_of="Rana Reviewer", meaning="CISO")
    text = "\n".join(pdf_pages(X.to_pdf(rec, images)))
    assert text.count("Awaiting signature") == 2                                           # the one who has not signed and the one who handed over
    for needle in ("Waiting", "Rejected 2026-10-22 14:05 UTC", "Asked for changes 2026-10-22 14:05 UTC", "Handed over", "“Not ready: the scope is too wide.”",
                   "“Clarify who approves an exception.”", "“Handed over to Mo Manager”", "for Rana Reviewer", "as CISO"):
        assert needle in text, needle
    assert drawn(X.to_pdf(rec, images)) == 2                                               # only the two who signed drew
    assert X.person_lines(rec["stages"][0], people[1])[-1] == ("Waiting", "text")


def test_a_comment_is_cut_at_two_hundred_and_twenty_characters_and_a_signed_person_may_have_no_picture():
    rec, images = fake(stages=1, per=2)
    rec["stages"][0]["people"][0]["comment"] = "word " * 100
    lines = X.person_lines(rec["stages"][0], rec["stages"][0]["people"][0])
    quote = next(t for t, style in lines if style == "quote")
    assert len(quote) <= 220 + 4 and quote.endswith(" …”") and quote.startswith("“")
    text = "\n".join(pdf_pages(X.to_pdf(rec, {})))                                          # no picture available: the seal is still printed
    assert "Drawn e-signature · seal" in text and drawn(X.to_pdf(rec, {})) == 0


def test_the_description_has_a_row_for_what_is_known_and_always_the_title_version_and_status():
    rec, _images = fake()
    assert X.description_rows(rec) == [
        ("Title", "Access policy"), ("Document code", "POL-009"), ("Version", "1.0"), ("Status", "Published"), ("Document type", "Policy"),
        ("Classification", "Internal"), ("Ownership", "Information Security — Olivia Owner"), ("Approving authority", "Board of Directors"),
        ("Date of approval", "2026-10-23"), ("Date of implementation", "2026-11-01"), ("Review frequency", "Every 12 months"),
        ("Last approval date", "2026-10-23"), ("Next review date", "2027-10-23")]
    for key in rec["control"]:
        if key not in ("title", "version", "status"):
            rec["control"][key] = None
    assert X.description_rows(rec) == [("Title", "Access policy"), ("Version", "1.0"), ("Status", "Published")]
    data = X.to_pdf(rec, {}, slots=("description",))
    assert drawn(data) == 0 and len(pdf_pages(data)) == 1


def test_the_revision_and_approval_history_tables_read_as_the_banks_do():
    rec, _images = fake()
    rec["revision_history"] += [{"version": "1.1", "date": "2026-12-01T10:00:00", "change_type": "Minor revision", "summary": "Scope widened; Exceptions added",
                                 "updated_by": "Rana Reviewer", "released": True},
                                {"version": "1.2", "date": "2026-12-20T10:00:00", "change_type": "Major revision", "summary": None, "updated_by": None, "released": False}]
    rec["approval_history"] += [{"version": "1.1", "reviewed_by": [{"name": "A", "designation": None}, {"name": "B", "designation": None}],
                                 "approved_by": [{"name": "C", "designation": None, "stage": "Board"}, {"name": "D", "designation": None, "stage": "Board"},
                                                 {"name": "E", "designation": None, "stage": None}], "approval_date": "2026-12-01T10:00:00"}]
    assert X.revision_rows(rec) == [("1.0", True, "2026-10-23", "Creation", "Initial version", "Olivia Owner"),
                                    ("1.1", True, "2026-12-01", "Minor revision", "Scope widened; Exceptions added", "Rana Reviewer"),
                                    ("1.2", False, "2026-12-20", "Major revision", "-", "-")]
    assert X.approval_rows(rec) == [("1.0", "Person 1", "Person 3 (Board of Directors)", "2026-10-23"), ("1.1", "A, B", "C, D (Board), E", "2026-12-01")]
    text = "\n".join(pdf_pages(X.to_pdf(rec, {})))
    for needle in ("Version", "Date", "Change type", "Summary of change", "Updated by", "Reviewed by", "Approved by", "Approval date", "In progress",
                   "Scope widened; Exceptions added", "C, D (Board), E"):
        assert needle in text, needle
    rec["approval_history"] = []
    assert "No version has been approved yet." in "\n".join(pdf_pages(X.to_pdf(rec, {})))
    assert "No version has been approved yet." in "\n".join(cell_texts(Document(io.BytesIO(X.to_docx(rec, {})))))


def test_the_foot_of_the_pages_says_what_they_are_and_counts_only_themselves():
    rec, images = fake(stages=8, per=2)
    pages = pdf_pages(X.to_pdf(rec, images))
    assert len(pages) >= 3
    for i, text in enumerate(pages, start=1):
        assert f"Document control page {i} of {len(pages)}" in text
        assert "generated by the platform 2026-10-25 09:30 UTC" in text and "DAR-9-1.0-ABCDEF12" in text and "Access policy" in text
    later = pdf_pages(X.to_pdf(rec, images, first=5, of=12))                           # one of several parts of a document's pages
    assert "Document control page 5 of 12" in later[0] and f"Document control page {4 + len(later)} of 12" in later[-1]
    rec["document"]["title"] = "A very long title " * 8
    long = pdf_pages(X.to_pdf(rec, images))[0]
    assert "…" in long and "Document control page 1 of" in long and "DAR-9-1.0-ABCDEF12" in long          # the title gives way, nothing else


def test_whatever_is_written_in_a_name_a_comment_or_a_title_the_pages_still_open_and_say_it():
    rec, images = fake(stages=1, per=2)
    nasty = "<script>alert(1)</script> & \x00\x0b\x1f “quoted” <b>bold</b>"
    rec["stages"][0]["people"][0].update(name=nasty, comment=nasty * 3, designation=nasty)
    rec["document"]["title"] = rec["control"]["title"] = nasty
    rec["control"]["ownership"] = nasty
    rec["revision_history"][0]["summary"] = nasty
    rec["notes"] = [nasty]
    for build in (X.to_pdf, X.to_docx):
        assert build(rec, images)[:2] in (b"%P", b"PK")
    text = "\n".join(pdf_pages(X.to_pdf(rec, images)))
    assert "<script>alert(1)</script> & “quoted” <b>bold</b>" in text and "\x00" not in text and "\x0b" not in text
    cells = "\n".join(cell_texts(Document(io.BytesIO(X.to_docx(rec, images)))))
    assert "<script>alert(1)</script> & “quoted” <b>bold</b>" in " ".join(cells.split()) and "\x00" not in cells and "\x0b" not in cells       # (the controls are gone, the rest is as it was)
    rec["stages"][0]["people"][1]["name"] = "رنا الأحمد"                                  # another script: the pages still open (right-to-left is not shaped)
    assert len(pdf_pages(X.to_pdf(rec, images))) >= 1 and Document(io.BytesIO(X.to_docx(rec, images)))


def test_the_pages_are_made_at_the_size_asked_for_and_never_have_a_header_of_their_own():
    rec, images = fake(stages=2, per=2)
    letter = X.to_pdf(rec, images, size=(612, 792))
    with fitz.open(stream=letter, filetype="pdf") as pdf:
        assert abs(pdf[0].rect.width - 612) < 1 and abs(pdf[0].rect.height - 792) < 1
        top = [w for w in pdf[0].get_text("words") if w[1] < 40]
        assert top == []                                                                  # nothing in the margin above: no header that looks like the original's
    landscape = X.to_pdf(rec, images, size=(842, 595))
    assert drawn(landscape) == 4 and "Person 1" in "\n".join(pdf_pages(landscape))


# ══ the pages of a platform-written document: put in where its sections are ═══════
def page_of(pages, needle):
    return next(i for i, t in enumerate(pages) if needle in t)


def test_a_markdown_document_gets_the_control_pages_between_the_pieces_of_its_text_at_its_own_headings():
    rec, images = fake(stages=2, per=1)
    pages = pdf_pages(C.pdf(SAMPLE, rec, images, title="Access policy"))
    first = page_of(pages, "Prepared for the board")
    describe = page_of(pages, "Information Security — Olivia Owner")
    purpose = page_of(pages, "Why we have it")
    signoff = page_of(pages, "STAGE 1 — REVIEWED BY")
    sub = page_of(pages, "A sub-section stays")
    revision = page_of(pages, "Initial version")
    assert first < describe < purpose < signoff < sub < revision
    assert "1. Document Description" in pages[describe] and "3. Approval Signoff" in pages[signoff] and "4. Revision History" in pages[revision]   # the document's own headings
    assert "1. Document Description" not in pages[first] and "3. Approval Signoff" not in pages[purpose]                              # not left at the foot of the page before
    assert "[[Document title]]" not in "\n".join(pages) and "______" not in "\n".join(pages)                                         # the old bodies are not printed
    assert "3.1 How to read this" in pages[sub] and "Approval History" in pages[revision]
    counted = [t for t in pages if "Document control page" in t]
    assert len(counted) == 3 and all("of 3" in t for t in counted) and [re.search(r"page (\d) of", t).group(1) for t in counted] == ["1", "2", "3"]
    assert drawn(C.pdf(SAMPLE, rec, images, title="Access policy")) == 2


def test_a_markdown_document_without_the_sections_gets_them_with_their_titles_after_its_title():
    rec, images = fake(stages=1, per=1)
    pages = pdf_pages(C.pdf(TEXT, rec, images, title="Access policy"))
    text = "\n".join(pages)
    assert text.index("Access policy") < text.index("Document Description") < text.index("Review and Approval") < text.index("Revision History") < text.index("1. Purpose")
    counted = [t for t in pages if "Document control page" in t]
    assert counted and all(f"of {len(counted)}" in t for t in counted)
    # a document with no title heading: the pages are at the top, under the document's title
    pages = pdf_pages(C.pdf("## 1. Purpose\n\nWhy.\n", rec, images, title="Leave policy"))
    assert "Leave policy" in pages[0] and "Document Description" in pages[0] and page_of(pages, "Why.") > 0


def test_the_pdf_of_a_document_that_is_all_places_has_no_text_pages_and_the_plain_pdf_is_as_it_was(world, monkeypatch):
    from reportlab import rl_config

    monkeypatch.setattr(rl_config, "invariant", 1)
    rec, images = fake(stages=1, per=1)
    pages = pdf_pages(C.pdf("# T\n## Document Description\n## Approval Signoff\n## Revision History\n", rec, images, title="T"))
    assert [("Document control page" in t) for t in pages] == [True, True, True] or all("Document control page" in t for t in pages[1:])
    finished(world)
    d = doc(world)
    d.content = SAMPLE
    world.db.commit()
    from grc.routers._artifact_export import build_export

    plain, _media, _ext = build_export("pdf", title="Access policy", content=SAMPLE)
    assert get(world, "/export").content == plain and get(world, "/export", with_record="false").content == plain


# ══ the pages in a Word file ═════════════════════════════════════════════════════
def docx_of(build):
    d = Document()
    build(d)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def children(data):
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    return [etree.tostring(c) for c in root.find(W + "body") if c.tag != W + "sectPr"]


def placed(original, record=None, images=None):
    """(the composed bytes, where the pages begin among the original's elements, the pages' elements): the elements of the original
    must all be there, in order and untouched, with the pages in one place."""
    rec, imgs = (record, images) if record is not None else fake(stages=2, per=2)
    data = X.place_in_docx(original, rec, imgs)
    before, after = children(original), children(data)
    k = 0
    while k < len(before) and before[k] == after[k]:
        k += 1
    extra = len(after) - len(before)
    assert extra > 0 and after[k + extra:] == before[k:], "an element of the original was changed, moved or lost"
    return data, k, after[k:k + extra]


def cover(d):
    d.add_paragraph("Records Policy", style="Title")
    d.add_paragraph("Version 1.0")


def content(d):
    d.add_heading("1. Purpose", 1)
    d.add_paragraph("The purpose is to govern data.")
    d.add_heading("2. Scope", 1)
    d.add_paragraph("All data.")


def has(elements, needle):
    return any(needle.encode() in e for e in elements)


def test_the_pages_come_in_after_the_cover_page_and_the_document_goes_on_from_a_new_page():
    original = docx_of(lambda d: (cover(d), d.add_page_break(), content(d)))
    data, k, block = placed(original)
    assert k == 3                                                                       # the title, the date, the page break: then the pages
    assert has(block[:2], "Document Description") and has(block, "Review and Approval") and has(block, "Revision History") and has(block, "Approval History")
    assert b'w:br w:type="page"' in block[-1] and has([block[-1]], "w:br")              # the document continues on a new page
    word = Document(io.BytesIO(data))
    texts_ = [p.text for p in word.paragraphs]
    assert texts_.index("Document Description") < texts_.index("1. Purpose") and texts_[0] == "Records Policy"
    assert len(word.tables) == 6                                                        # the description, prepared by, the two stages, the two histories
    assert len(word.inline_shapes) == 4 and "Person 1" in "\n".join(cell_texts(word))


def test_a_page_break_at_the_end_of_the_last_cover_line_and_one_that_opens_the_next_page_are_boundaries_too():
    def inline(d):
        d.add_paragraph("Records Policy", style="Title")
        d.add_paragraph("Version 1.0").add_run().add_break(WD_BREAK.PAGE)
        content(d)

    data, k, block = placed(docx_of(inline))
    assert k == 2 and has(block, "Document Description")
    assert b'w:br w:type="page"' in block[-1]

    def opening(d):
        cover(d)
        p = d.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p.add_run("1. Purpose")                                                         # the page break comes first, its text begins the next page
        d.add_paragraph("The purpose is to govern data.")

    data, k, block = placed(docx_of(opening))
    assert k == 2 and has(block, "Document Description") and b'w:br w:type="page"' not in block[-1] and has(block[:1], 'w:pageBreakBefore')


def test_a_heading_that_begins_a_page_of_its_own_is_where_the_cover_ends_and_the_pages_begin_on_a_new_page():
    def build(d):
        cover(d)
        d.add_heading("1. Purpose", 1).paragraph_format.page_break_before = True
        d.add_paragraph("The purpose is to govern data.")

    data, k, block = placed(docx_of(build))
    assert k == 2 and has(block[:1], "w:pageBreakBefore") and has(block[:1], "Document Description")
    assert not has(block[-1:], 'w:br w:type="page"')                                    # the heading after them begins a page by itself: no empty page between
    assert [p.text for p in Document(io.BytesIO(data)).paragraphs][-2:] == ["1. Purpose", "The purpose is to govern data."]


def test_when_the_next_element_already_begins_a_page_the_pages_do_not_add_a_break_of_their_own():
    def build(d):
        cover(d)
        d.add_page_break()
        d.add_heading("1. Purpose", 1).paragraph_format.page_break_before = True

    data, k, block = placed(docx_of(build))
    assert k == 3 and not has(block[-1:], 'w:br w:type="page"')


def test_a_section_break_to_a_new_page_ends_the_cover_and_a_continuous_one_does_not():
    def new_page(d):
        cover(d)
        d.add_section(WD_SECTION.NEW_PAGE)
        content(d)

    data, k, block = placed(docx_of(new_page))
    assert k == 3 and has(block, "Document Description")                                # after the paragraph that holds the section settings
    assert Document(io.BytesIO(data)).sections.__len__() == 2

    def continuous(d):
        cover(d)
        d.add_section(WD_SECTION.CONTINUOUS)
        content(d)

    original = docx_of(continuous)
    data, k, block = placed(original)
    assert k == len(children(original))                                                   # no page boundary: no cover, so the pages are at the end
    assert has(block[:1], 'w:br w:type="page"')


def test_without_a_page_boundary_in_its_first_forty_elements_the_pages_go_at_the_end_after_a_page_break():
    original = docx_of(lambda d: (cover(d), content(d)))
    data, k, block = placed(original)
    assert k == len(children(original)) and has(block[:1], 'w:br w:type="page"') and has(block, "Document Description")
    late = docx_of(lambda d: ([d.add_paragraph(f"Line {i}") for i in range(45)], d.add_page_break(), content(d)))      # a boundary, but not among the first forty
    data, k, block = placed(late)
    assert k == len(children(late)) and has(block[:1], 'w:br w:type="page"')
    near = docx_of(lambda d: ([d.add_paragraph(f"Line {i}") for i in range(38)], d.add_page_break(), content(d)))
    assert placed(near)[1] == 39
    # the section settings stay last in the body
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    assert root.find(W + "body")[-1].tag == W + "sectPr"


def test_a_cover_made_by_a_content_control_ends_with_its_own_page_break():
    def build(d):
        body = d.element.body
        control = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Cover Pages"/></w:docPartObj></w:sdtPr><w:sdtContent>'
                            '<w:p><w:r><w:t>Policy</w:t></w:r></w:p><w:p><w:r><w:br w:type="page"/></w:r></w:p></w:sdtContent></w:sdt>')
        body.insert(0, control)
        content(d)

    data, k, block = placed(docx_of(build))
    assert k == 1 and has(block, "Document Description")                                   # after the control, never inside it


def test_a_page_break_in_the_middle_of_a_line_of_text_is_not_a_boundary():
    def build(d):
        p = d.add_paragraph("Before the break")
        p.add_run().add_break(WD_BREAK.PAGE)
        p.add_run("after the break")
        content(d)

    original = docx_of(build)
    assert placed(original)[1] == len(children(original))


def test_the_original_parts_of_a_real_word_file_are_untouched_and_the_pages_take_its_header_and_footer(tmp_path):
    for build in (corporate_policy, tracked_policy):
        original = open(build(str(tmp_path / f"{build.__name__}.docx")), "rb").read()
        word = Document(io.BytesIO(original))                                            # a cover in front of the text: the title, then a page break
        word.element.body.insert(1, parse_xml(f'<w:p {nsdecls("w")}><w:r><w:br w:type="page"/></w:r></w:p>'))
        out = io.BytesIO()
        word.save(out)
        covered = out.getvalue()
        data, k, block = placed(covered)
        assert k == 2 and has(block, "Document Description"), build.__name__
        before, after = zipfile.ZipFile(io.BytesIO(covered)), zipfile.ZipFile(io.BytesIO(data))
        assert set(before.namelist()) <= set(after.namelist())
        for part in sorted(n for n in before.namelist() if n.startswith(("word/header", "word/footer")) or n in ("word/styles.xml", "word/numbering.xml")):
            assert after.read(part) == before.read(part), part
        assert len(Document(io.BytesIO(data)).inline_shapes) >= 4
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    assert xml.count("<w:ins ") >= 2 and xml.count("<w:del ") >= 2 and "rPrChange" in xml                          # the reviewer's changes are still tracked


def test_the_pages_are_as_wide_as_the_text_of_the_section_they_land_in_and_never_use_a_heading_style():
    def build(d):
        cover(d)
        d.add_section(WD_SECTION.NEW_PAGE)
        s = d.sections[1]
        s.left_margin = s.right_margin = 914400                                           # one inch each side
        content(d)

    original = docx_of(build)
    data, _k, _block = placed(original)
    word = Document(io.BytesIO(data))
    section = word.sections[1]
    width = int(section.page_width - section.left_margin - section.right_margin)
    grid = next(t for t in word.tables if t.cell(0, 0).text.startswith("STAGE 1"))
    assert abs(sum(int(c.width) for c in grid.columns) - width) <= 4 * 635                       # the columns add up to the text's width
    assert abs(int(grid._tbl.tblPr.find(W + "tblW").get(W + "w")) * 635 - width) <= 635
    first = grid._tbl.tr_lst[0].tc_lst[0]
    assert first.grid_span == 3 and abs(int(first.width) - width * 0.62) <= 4 * 635               # the title bar covers three columns, a person's cell two
    assert grid._tbl.tr_lst[1].tc_lst[0].grid_span == 2 and grid._tbl.tr_lst[1].tc_lst[1].grid_span == 2
    styles = {p.style.name for p in word.paragraphs if p.text in ("Document Description", "Review and Approval", "Revision History", "Approval History")}
    assert styles == {"Normal"}
    for p in word.paragraphs:
        if p.text in ("Document Description", "Review and Approval"):
            assert p.runs[0].bold and p.paragraph_format.keep_with_next


def test_a_word_file_that_python_docx_cannot_open_raises_so_the_caller_serves_it_as_it_is():
    rec, images = fake()
    with pytest.raises(Exception):
        X.place_in_docx(b"this is not a zip file", rec, images)


# ══ the pages in a PDF ═══════════════════════════════════════════════════════════
def book(pages=4, cover_text="Retention Policy", size=(595, 842), body_lines=60, image=False, cover_lines=1):
    """A PDF whose first page is a cover (a title, a few words) and whose other pages are full of text."""
    pdf = fitz.open()
    first = pdf.new_page(width=size[0], height=size[1])
    for i in range(cover_lines):
        first.insert_text((72, 300 + i * 12), cover_text, fontsize=20 if cover_lines == 1 else 9)
    if image:
        first.insert_image(fitz.Rect(30, 30, size[0] - 30, size[1] - 30), stream=ink((400, 400), (200, 30, 30, 255)))
    for i in range(pages - 1):
        page = pdf.new_page(width=size[0], height=size[1])
        for j in range(body_lines):
            page.insert_text((50, 60 + j * 11.5), f"Page {i + 2} line {j}: access shall be reviewed regularly by the owner of the system.", fontsize=8)
    data = pdf.tobytes()
    pdf.close()
    return data


def test_the_pages_come_in_after_a_pdfs_cover_and_every_original_page_is_as_it_was_in_its_place():
    rec, images = fake(stages=2, per=2)
    original = book(pages=4)
    data = X.place_in_pdf(original, rec, images)
    before, after = pdf_pages(original), pdf_pages(data)
    n = len(after) - len(before)
    assert n >= 2 and after[0] == before[0] and after[1 + n:] == before[1:]            # the cover, the pages, the rest of the file as it was
    assert all("Document control page" in t for t in after[1:1 + n]) and "Document Description" in after[1]
    assert [re.search(r"page (\d+) of (\d+)", t).groups() for t in after[1:1 + n]] == [(str(i), str(n)) for i in range(1, n + 1)]
    with fitz.open(stream=data, filetype="pdf") as pdf:
        assert all(abs(pdf[i].rect.width - 595) < 1 and abs(pdf[i].rect.height - 842) < 1 for i in range(len(pdf)))
        assert sum(len(p.get_image_info()) for p in pdf) == 4


def test_a_first_page_full_of_text_is_not_a_cover_and_the_pages_go_at_the_end_so_does_a_one_page_file():
    rec, images = fake()
    original = book(pages=3)
    with fitz.open(stream=original, filetype="pdf") as pdf:                              # the same text on page 1 as on the others
        pdf.delete_page(0)
        full = pdf.tobytes()
    data = X.place_in_pdf(full, rec, images)
    before, after = pdf_pages(full), pdf_pages(data)
    assert after[:len(before)] == before and len(after) > len(before) and all("Document control page" in t for t in after[len(before):])
    single = book(pages=1)
    after = pdf_pages(X.place_in_pdf(single, rec, images))
    assert "Retention Policy" in after[0] and "Document Description" in after[1] and len(after) >= 2


def test_a_first_page_that_is_mostly_a_picture_is_a_cover_even_with_a_lot_of_text_on_it():
    rec, images = fake()
    original = book(pages=3, cover_text="word " * 12, cover_lines=20, image=True)
    assert len("".join(pdf_pages(original)[0].split())) > 700
    after = pdf_pages(X.place_in_pdf(original, rec, images))
    assert "Document Description" in after[1] and "Page 3 line 0" in after[-1]
    no_picture = book(pages=3, cover_text="word " * 12, cover_lines=20)                      # the same text without the picture: a page of text
    after = pdf_pages(X.place_in_pdf(no_picture, rec, images))
    assert "Document Description" not in after[1] and "Document Description" in after[3]


def test_the_pages_are_made_at_the_size_of_the_first_page_and_a_tiny_one_keeps_them_a_four():
    rec, images = fake()
    with fitz.open(stream=X.place_in_pdf(book(pages=3, size=(612, 792)), rec, images), filetype="pdf") as pdf:
        assert all(abs(p.rect.width - 612) < 1 and abs(p.rect.height - 792) < 1 for p in pdf) and "Document Description" in pdf[1].get_text()
    with fitz.open(stream=X.place_in_pdf(book(pages=3, size=(842, 595), body_lines=40), rec, images), filetype="pdf") as pdf:
        assert abs(pdf[1].rect.width - 842) < 1 and "Document Description" in pdf[1].get_text()
    with fitz.open(stream=X.place_in_pdf(book(pages=2, size=(300, 400), body_lines=20), rec, images), filetype="pdf") as pdf:
        assert abs(pdf[0].rect.width - 300) < 1 and abs(pdf[1].rect.width - 595) < 2 and "Document Description" in pdf[1].get_text()


def test_a_protected_or_unreadable_pdf_is_refused_so_it_is_served_as_it_is():
    rec, images = fake()
    locked = fitz.open()
    locked.new_page().insert_text((72, 72), "Protected")
    protected = locked.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="", permissions=fitz.PDF_PERM_ACCESSIBILITY)
    with pytest.raises(ValueError):
        X.place_in_pdf(protected, rec, images)
    with pytest.raises(Exception):
        X.place_in_pdf(b"%PDF-1.4 truncated", rec, images)


def test_the_stored_file_a_function_reads_is_not_changed_by_it():
    rec, images = fake()
    original = book(pages=3)
    copy = bytes(original)
    X.place_in_pdf(original, rec, images)
    word = docx_of(lambda d: (cover(d), d.add_page_break(), content(d)))
    word_copy = bytes(word)
    X.place_in_docx(word, rec, images)
    assert original == copy and word == word_copy


# ══ the keys of the record that the pages are made from ═══════════════════════════
def put(world, **fields):
    d = doc(world)
    for key, value in fields.items():
        setattr(d, key, value)
    world.db.commit()


def test_the_description_of_a_published_document_on_the_usual_route(world):
    from datetime import datetime

    world.db.add(m.GRCDepartment(id=7, tenant_id=1, name="Information Security", code="IS"))
    put(world, department_id=7, effective_date=datetime(2026, 11, 1), next_review_date=datetime(2027, 10, 23), review_cycle_months=12)
    rec = finished(world)
    c = rec["control"]
    assert c == {"title": "Access policy", "code": "POL-005", "version": "1.0", "type": "policy", "type_label": "Policy", "classification": "internal",
                 "owner": "Olivia Owner", "department": "Information Security", "ownership": "Information Security — Olivia Owner",
                 "approving_authority": "Mina Manager", "date_of_approval": rec["document"]["approved_at"], "date_of_implementation": "2026-11-01T00:00:00",
                 "review_frequency": "Every 12 months", "last_approval_date": rec["document"]["approved_at"], "next_review_date": "2027-10-23T00:00:00",
                 "status": "Published"}
    assert c["date_of_approval"] and c["date_of_approval"] == rec["history"][0]["approved_at"]
    assert [r["version"] for r in rec["revision_history"]] == ["1.0"] and rec["revision_history"][0] == {
        "version": "1.0", "date": c["date_of_approval"], "change_type": "Creation", "summary": "Initial version", "updated_by": "Olivia Owner", "released": True}
    assert rec["approval_history"] == [{"version": "1.0", "reviewed_by": [{"name": "Rana Reviewer", "designation": "Compliance Manager"},
                                                                        {"name": "Rafay Reviewer", "designation": "Risk Officer"}],
                                        "approved_by": [{"name": "Mina Manager", "designation": "CISO", "stage": None}], "approval_date": c["date_of_approval"]}]


def test_the_review_frequency_is_said_in_months_and_the_type_in_words(world):
    for months, said in ((12, "Every 12 months"), (1, "Every month"), (36, "Every 36 months"), (None, None), (0, None)):
        put(world, review_cycle_months=months)
        assert record(world)["control"]["review_frequency"] == said
    for kind, said in (("standard", "Standard"), ("charter", "Charter"), ("risk_register", "Risk Register"), ("tprm-form", "Tprm Form")):
        put(world, doc_type=kind)
        assert record(world)["control"]["type"] == kind and record(world)["control"]["type_label"] == said
    put(world, owner_id=None, classification="restricted")
    c = record(world)["control"]
    assert c["owner"] is None and c["ownership"] is None and c["classification"] == "restricted"


def test_the_approving_authority_is_the_last_stages_own_name_else_the_people_asked_at_it(world):
    set_route(world, stage("review", [person(RANA), person(RAFAY)], name="Compliance review", rule="at_least", min_count=1),
              stage("concurrence", [person(UMA)], name="Legal concurrence"), stage("approval", [role()], name="Board approval", rule="any"))
    send(world)
    assert record(world)["control"]["approving_authority"] == "Board approval" and record(world)["control"]["status"] == "In review"
    sign(world, RAFAY)
    sign(world, UMA)
    sign(world, MO)
    rec = record(world)
    assert rec["control"]["approving_authority"] == "Board approval" and rec["control"]["status"] == "Approved, not published"
    assert rec["approval_history"][0]["reviewed_by"] == [{"name": "Rafay Reviewer", "designation": "Risk Officer"}, {"name": "Uma User", "designation": "Analyst"}]
    assert rec["approval_history"][0]["approved_by"] == [{"name": "Mo Manager", "designation": None, "stage": "Board approval"}]
    assert rec["revision_history"][0]["released"] is True
    # a stage that kept the name its kind gives it ("Approval") has no name of its own: the people at it are the authority
    world.db.query(m.DocumentRouteStage).filter_by(document_id=5, position=3).update({"name": "Approval"})
    world.db.commit()
    assert record(world)["control"]["approving_authority"] == "Mo Manager"


def test_a_draft_and_a_document_in_review_have_a_description_and_a_revision_in_progress_and_nothing_approved(world):
    rec = record(world)
    assert rec["control"]["status"] == "Draft" and rec["control"]["approving_authority"] is None and rec["control"]["date_of_approval"] is None
    assert rec["control"]["last_approval_date"] is None and rec["approval_history"] == []
    assert rec["revision_history"] == [{"version": "1.0", "date": "2026-10-01T09:00:00", "change_type": "Creation", "summary": "Initial version",
                                        "updated_by": "Olivia Owner", "released": False}]
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    rec = record(world)
    assert rec["control"]["status"] == "In approval"
    assert rec["control"]["approving_authority"] == "Mina Manager" and rec["control"]["date_of_approval"] is None and rec["approval_history"] == []
    assert [(r["version"], r["released"]) for r in rec["revision_history"]] == [("1.0", False)]


def test_two_released_versions_and_one_in_progress_make_three_rows_and_the_saves_between_are_in_the_summary(world):
    approve_and_publish(world, (RANA, RAFAY), (MINA,))
    edit(world, SMALL, "Temporary staff added to the scope")
    edit(world, SMALL + "\nA clause.\n", "A clause about exceptions")
    edit(world, SMALL + "\nA clause.\nAnother.\n", "Retention period stated")
    usual(world, (RAFAY,), (MO,))
    send(world)
    sign(world, RAFAY)
    sign(world, MO)
    publish(world)
    edit(world, SMALL + "\nA clause.\nAnother.\nMore.\n", "Work in progress")
    rec = record(world)
    assert [(r["version"], r["released"], r["change_type"]) for r in rec["revision_history"]] == [
        ("1.0", True, "Creation"), ("1.3", True, "Minor revision"), ("1.4", False, "Minor revision")]
    first, second, third = rec["revision_history"]
    assert first["summary"] == "Initial version" and first["updated_by"] == "Olivia Owner"
    # the released version's own summary, then what the saves since the version released before it said
    assert second["summary"] == "Retention period stated; Temporary staff added to the scope; A clause about exceptions"
    assert third["summary"] == "Work in progress" and third["date"] != second["date"] and third["date"] < "2999"
    assert second["date"] == rec["approval_history"][1]["approval_date"]                        # a released version is dated by its approval
    assert [a["version"] for a in rec["approval_history"]] == ["1.0", "1.3"]                      # oldest first, one for each released version
    assert [x["name"] for x in rec["approval_history"][0]["reviewed_by"]] == ["Rana Reviewer", "Rafay Reviewer"]
    assert [x["name"] for x in rec["approval_history"][1]["reviewed_by"]] == ["Rafay Reviewer"] and [x["name"] for x in rec["approval_history"][1]["approved_by"]] == ["Mo Manager"]
    assert rec["control"]["last_approval_date"] == rec["approval_history"][1]["approval_date"] and rec["control"]["date_of_approval"] is None
    assert [h["version"] for h in rec["history"]] == ["1.3", "1.0"]                                # the history is as it was: newest first
    # the record of the earlier version is its own: its rows are the versions up to it
    old = record(world, version="1.0")
    assert [r["version"] for r in old["revision_history"]] == ["1.0"] and [a["version"] for a in old["approval_history"]] == ["1.0"]
    assert old["control"]["version"] == "1.0" and old["control"]["status"] == "Published" and old["revision_history"][0]["released"] is True


def test_many_saves_between_two_releases_list_four_and_say_how_many_more(world):
    approve_and_publish(world, (RANA,), (MINA,))
    for i in range(1, 8):
        edit(world, SMALL + f"\nClause {i}.\n", f"Change number {i}")
    usual(world, (RAFAY,), (MO,))
    send(world)
    sign(world, RAFAY)
    sign(world, MO)
    row = record(world)["revision_history"][-1]
    assert row["version"] == "1.7" and row["released"] is True
    assert row["summary"] == "Change number 7; Change number 1; Change number 2; Change number 3; Change number 4; and 2 more"
    long = "x" * 300
    edit(world, SMALL + "\nLong.\n", long)
    assert len(record(world)["revision_history"][-1]["summary"]) <= 305


def test_approvals_that_cannot_be_tied_to_a_version_are_one_unknown_row_in_both_histories(world):
    approve_and_publish(world, (RANA,), (MINA,))
    world.db.query(m.DocumentSignature).update({m.DocumentSignature.version_id: None})
    edit(world)
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    sign(world, MO)
    ids = [s.id for s in world.db.query(m.DocumentSignature).filter(m.DocumentSignature.signer_user_id.in_((RANA, MINA)))]
    world.db.query(m.ElectronicSignature).filter(m.ElectronicSignature.record_id.in_(ids)).delete(synchronize_session=False)
    world.db.commit()
    rec = record(world)
    assert [r["version"] for r in rec["revision_history"]] == ["unknown", "1.1"] and [a["version"] for a in rec["approval_history"]] == ["unknown", "1.1"]
    unknown = rec["revision_history"][0]
    assert (unknown["change_type"], unknown["updated_by"], unknown["released"]) == ("Creation", None, True)       # the first row of the table is the creation
    row = rec["approval_history"][0]
    assert [x["name"] for x in row["reviewed_by"]] == ["Rana Reviewer"] and [x["name"] for x in row["approved_by"]] == ["Mina Manager"]


def test_a_reader_who_may_not_see_drafts_gets_no_row_for_a_version_after_the_one_in_force(world):
    approve_and_publish(world, (RANA,), (MINA,))
    edit(world, SMALL, "Temporary staff added")
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    rec = get(world, "/approval-record", who=VERA).json()
    assert [r["version"] for r in rec["revision_history"]] == ["1.0"] and [a["version"] for a in rec["approval_history"]] == ["1.0"]
    assert "Temporary staff" not in str(rec) and rec["control"]["version"] == "1.0" and rec["control"]["status"] == "Published"
    assert rec["control"]["approving_authority"] == "Mina Manager"


# ══ what the screens rely on ═════════════════════════════════════════════════════
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(\.\d+)?$")


def released(rec):
    return [(r["version"], r["released"]) for r in rec["revision_history"]]


def test_a_version_is_released_when_it_completed_approval_and_only_the_one_being_worked_on_is_not(world):
    rec = record(world)                                                                    # a draft: the version being worked on
    assert released(rec) == [("1.0", False)]
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    assert released(record(world)) == [("1.0", False)] and record(world)["state"] == "in_approval"      # half way is not approved
    sign(world, MINA)
    assert record(world)["state"] == "approved" and released(record(world)) == [("1.0", True)]            # approved, not yet published: released
    publish(world)
    assert released(record(world)) == [("1.0", True)] and [a["version"] for a in record(world)["approval_history"]] == ["1.0"]
    edit(world, SMALL, "Temporary staff added")                                          # the next version is being worked on
    assert released(record(world)) == [("1.0", True), ("1.1", False)]
    assert [a["version"] for a in record(world)["approval_history"]] == ["1.0"]            # (the approval history has only what was approved)
    assert released(record(world, version="1.0")) == [("1.0", True)]                       # the earlier version's record has no row for a later one


def test_a_version_that_was_never_approved_and_has_been_overtaken_is_not_listed_as_in_progress(world):
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    reject(world, MINA)
    edit(world, SMALL, "Temporary staff added to the scope")
    send(world)
    sign(world, RANA)
    sign(world, MINA)
    assert record(world)["state"] == "approved"
    now = record(world)
    assert released(now) == [("1.1", True)] and now["revision_history"][0]["change_type"] == "Creation"
    assert now["revision_history"][0]["summary"] == "Temporary staff added to the scope; Initial version"
    old = record(world, version="1.0")                                                      # 1.0 was rejected: it is not "in progress" nor released
    assert old["revision_history"] == [] and old["approval_history"] == [] and old["control"]["status"] == "Draft"


def test_a_document_approved_or_published_without_a_complete_round_has_its_version_released_all_the_same(world):
    d = doc(world)
    d.status, d.published_at, d.published_by = "published", datetime(2026, 9, 3), ADAM
    world.db.commit()
    rec = record(world)                                                                    # no signature at all: published by hand
    assert released(rec) == [("1.0", True)] and rec["revision_history"][0]["date"] == "2026-09-03T00:00:00"
    assert rec["approval_history"] == [{"version": "1.0", "reviewed_by": [], "approved_by": [], "approval_date": None}]
    assert rec["control"]["status"] == "Published" and rec["control"]["date_of_approval"] is None and rec["control"]["last_approval_date"] is None
    assert any(n.startswith("No sign-off is recorded") for n in rec["notes"])
    assert X.approval_rows(rec) == [("1.0", "-", "-", "-")]                                  # the pages print what is not known as a dash


def test_an_archived_document_keeps_the_approval_it_had_in_both_histories(world):
    rec = finished(world)
    world.db.get(m.GovernanceDocument, 5).status = "archived"
    from grc.modules.governance.routers import documents

    documents.create_audit_log(world.db, 5, 1, ADAM, "archived", "Archived by Adam Admin")
    world.db.commit()
    arch = record(world)
    assert arch["state"] == "archived" and arch["control"]["status"] == "Archived" and arch["history"] == []        # (the existing keys are as they were)
    assert released(arch) == [("1.0", True)] and [a["version"] for a in arch["approval_history"]] == ["1.0"]
    assert [x["name"] for x in arch["approval_history"][0]["approved_by"]] == ["Mina Manager"]
    assert arch["control"]["date_of_approval"] and arch["control"]["date_of_approval"] == arch["control"]["last_approval_date"]
    assert arch["control"]["date_of_approval"][:16] == rec["document"]["approved_at"][:16]
    world.db.get(m.GovernanceDocument, 5).status = "archived"                                  # archived while a draft nobody approved: nothing released
    world.db.query(m.DocumentSignature).delete()
    world.db.commit()
    assert released(record(world)) == [("1.0", False)]


def test_only_a_name_the_route_gave_a_stage_is_its_name_in_the_approval_history(world):
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    sign(world, MINA)
    assert record(world)["approval_history"][0]["approved_by"] == [{"name": "Mina Manager", "designation": "CISO", "stage": None}]     # the usual route has none
    for name, stage_name in (("Board approval", "Board approval"), ("Approval", None), ("Approved by", None), (None, None)):
        world.db.query(m.GovernanceDocument).filter_by(id=5).update({"status": "draft"})
        world.db.query(m.DocumentSignature).delete()
        world.db.commit()
        world.db.expire_all()
        named = {} if name is None else {"name": name}
        set_route(world, stage("review", [person(RANA)], name="Compliance review"), stage("approval", [person(MINA)], **named))
        send(world)
        sign(world, RANA)
        sign(world, MINA)
        rec = record(world)
        last = rec["approval_history"][-1]
        assert last["approved_by"][0]["stage"] == stage_name, name
        assert rec["control"]["approving_authority"] == (name if stage_name else "Mina Manager"), name
        assert all(x["stage"] is None or x["stage"] not in ("Approved by", "Reviewed by") for a in rec["approval_history"] for x in a["approved_by"])
        assert "stage" not in rec["approval_history"][-1]["reviewed_by"][0]


def test_the_status_and_the_type_are_words_for_every_state_a_document_can_be_in(world):
    seen = {}
    for status in ("draft", "pending_review", "pending_approval", "approved", "published", "expired", "archived", "exception_applied", "review_required", None):
        world.db.get(m.GovernanceDocument, 5).status = status
        world.db.commit()
        seen[status] = record(world)["control"]["status"]
    assert seen == {"draft": "Draft", "pending_review": "In review", "pending_approval": "In approval", "approved": "Approved, not published",
                    "published": "Published", "expired": "Published", "archived": "Archived", "exception_applied": "Published",
                    "review_required": "Published", None: "Draft"}
    for kind, said in (("policy", "Policy"), ("standard", "Standard"), ("procedure", "Procedure"), ("guideline", "Guideline"), ("charter", "Charter"),
                       ("framework", "Framework"), ("incident_plan", "Incident Plan")):
        world.db.get(m.GovernanceDocument, 5).doc_type = kind
        world.db.commit()
        assert record(world)["control"]["type_label"] == said and record(world)["control"]["type"] == kind


def test_the_dates_of_both_histories_are_full_naive_utc_timestamps_or_null(world):
    approve_and_publish(world, (RANA,), (MINA,))
    edit(world, SMALL, "Temporary staff added")
    world.db.query(m.GovernanceDocumentVersion).filter_by(version_number="1.1").update({"created_at": None})      # a row without a time
    world.db.commit()
    rec = record(world)
    dates = [r["date"] for r in rec["revision_history"]] + [a["approval_date"] for a in rec["approval_history"]]
    assert len(dates) == 3 and dates[1] is None                                              # 1.1 has no time to give
    assert all(d is None or (ISO.match(d) and "Z" not in d and "+" not in d) for d in dates)
    assert rec["revision_history"][0]["date"] == rec["approval_history"][0]["approval_date"] == rec["history"][0]["approved_at"]
    assert ISO.match(rec["control"]["date_of_approval"] or rec["control"]["last_approval_date"])
    rec = get(world, "/approval-record").json()                                              # and as JSON
    assert all(d is None or (ISO.match(d) and not d.endswith("Z")) for d in [r["date"] for r in rec["revision_history"]])


def test_a_person_is_never_listed_without_a_name(world):
    approve_and_publish(world, (RANA, RAFAY), (MINA,))
    world.db.query(m.ElectronicSignature).filter_by(record_type="document_signoff").update({"signer_name": ""})        # the snapshot has no name ...
    world.db.query(m.GRCUser).filter(m.GRCUser.id == RAFAY).delete()                                                    # ... and one signer's account is gone
    world.db.get(m.GRCUser, RANA).display_name = ""
    world.db.get(m.GRCUser, RANA).username = ""
    world.db.commit()
    rec = record(world)
    people = [p for s in rec["stages"] for p in s["people"]]
    assert len(people) == 3 and all(isinstance(p["name"], str) and p["name"] for p in people)
    assert [p["name"] for p in people] == [f"User #{RANA}", f"User #{RAFAY}", "Mina Manager"]
    assert all(x["name"] for a in rec["approval_history"] for x in a["reviewed_by"] + a["approved_by"])
    assert all(p["name"] for s in record(world, version="1.0")["stages"] for p in s["people"])
    # a name that is only control characters is nothing once it is cleaned: the pages print "Unknown", not an empty line
    rec2 = fake()[0]
    rec2["stages"][0]["people"][0]["name"] = "\x00\x0b"
    rec2["approval_history"][0]["reviewed_by"][0]["name"] = "\x00"
    assert "Unknown" in "\n".join(pdf_pages(X.to_pdf(rec2, {}))) and "Unknown" in "\n".join(cell_texts(Document(io.BytesIO(X.to_docx(rec2, {})))))


# ══ the endpoints ════════════════════════════════════════════════════════════════
def test_controlled_content_of_a_platform_written_document_is_its_text_with_the_places_and_the_record(world):
    rec = finished(world)
    r = get(world, "/controlled-content")
    assert r.status_code == 200
    body = r.json()
    assert list(body) == ["kind", "segments", "record"] and body["kind"] == "markdown"
    assert body["segments"] == C.segments(TEXT) and body["record"]["reference"] == rec["reference"] and body["record"]["control"] == rec["control"]
    assert [(s["slot"], s["heading"]) for s in body["segments"] if s["type"] == "slot"] == [("description", True), ("signoff", True), ("revision", True)]
    assert r.json()["record"]["stages"][0]["people"][0]["name"] == "Rana Reviewer"


def test_controlled_content_of_a_document_that_is_a_file_says_so_and_has_no_segments(world):
    finished(world)
    attach(world, "Access policy.docx", docx_of(lambda d: (cover(d), content(d))), "docx")
    body = get(world, "/controlled-content").json()
    assert body["kind"] == "file" and body["segments"] is None and body["record"]["document"]["version"] == "1.0"
    d = doc(world)                                                                        # a file that went missing: the text is what there is
    d.file_path = str(world.tmp / "gone.docx")
    world.db.commit()
    assert get(world, "/controlled-content").json()["kind"] == "markdown"
    d.content = None
    world.db.commit()
    body = get(world, "/controlled-content").json()
    assert body["kind"] == "file" and body["segments"] is None                           # neither: nothing to put the pages in


def test_controlled_content_follows_the_version_asked_for_and_the_one_in_force_for_a_reader_who_may_not_see_drafts(world):
    approve_and_publish(world, (RANA,), (MINA,))
    edit(world, SMALL, "Temporary staff added")
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    now = get(world, "/controlled-content").json()
    assert now["record"]["document"]["version"] == "1.1" and "temporary staff" in "".join(s["text"] for s in now["segments"] if s["type"] == "markdown")
    old = get(world, "/controlled-content", version="1.0").json()                         # its text, its record
    joined = "".join(s["text"] for s in old["segments"] if s["type"] == "markdown")
    assert old["record"]["document"]["version"] == "1.0" and "temporary staff" not in joined and "contractors, vendors" in joined
    assert [x["name"] for x in old["record"]["stages"][0]["people"]] == ["Rana Reviewer"]
    vera = get(world, "/controlled-content", who=VERA)                                    # reads what is in force, and not a word of what is on its way
    assert vera.status_code == 200 and vera.json()["record"]["document"]["version"] == "1.0"
    assert vera.json()["segments"] == old["segments"] and "temporary staff" not in vera.text and "Rafay" not in vera.text
    assert get(world, "/controlled-content", who=VERA, version="1.1").status_code == 404
    assert get(world, "/controlled-content", version="9.9").status_code == 404 and get(world, "/controlled-content", version="9.9").json()["detail"] == "That version of the document was not found."


def test_controlled_content_of_a_draft_is_for_the_people_who_may_read_the_draft(world):
    body = get(world, "/controlled-content").json()
    assert body["record"]["state"] == "draft" and body["kind"] == "markdown" and body["record"]["stages"] == []
    assert get(world, "/controlled-content", who=VERA).status_code == 403                       # a library reader does not see a draft
    world.db.add(m.GRCUser(id=9, username="una", email="una@cfsb.test", display_name="Una Unrelated"))
    world.db.commit()
    assert get(world, "/controlled-content", who=9).status_code == 403
    assert world.http.get("/governance/documents/99/controlled-content").status_code == 404


def test_download_file_with_record_gives_a_word_file_its_pages_after_the_cover_and_says_so(world):
    rec = finished(world)
    original = docx_of(lambda d: (cover(d), d.add_page_break(), content(d)))
    attach(world, "Data policy.docx", original, "docx")
    plain = get(world, "/download-file")
    assert plain.status_code == 200 and plain.content == original and plain.headers["x-controlled-copy"] == "false"
    assert get(world, "/download-file", with_record="false").content == original
    r = get(world, "/download-file", with_record="true")
    assert r.status_code == 200 and r.headers["x-controlled-copy"] == "true" and r.headers["content-type"] == DOCX_MEDIA
    assert r.headers["content-disposition"] == 'attachment; filename="Data policy.docx"'
    word = Document(io.BytesIO(r.content))
    paragraphs = [p.text for p in word.paragraphs]
    assert paragraphs[0] == "Records Policy" and paragraphs.index("Document Description") < paragraphs.index("1. Purpose")
    cells = "\n".join(cell_texts(word))
    assert "Rana Reviewer" in cells and "Mina Manager" in cells and rec["reference"] in "\n".join(paragraphs) and len(word.inline_shapes) == 3
    assert (world.tmp / "Data policy.docx").read_bytes() == original                         # the stored file is as it was


def test_download_file_with_record_gives_a_pdf_its_pages_after_its_cover(world):
    rec = finished(world)
    original = book(pages=3)
    attach(world, "Data policy.pdf", original, "pdf")
    r = get(world, "/download-file", with_record="true")
    assert r.status_code == 200 and r.headers["x-controlled-copy"] == "true" and r.headers["content-type"] == "application/pdf"
    before, after = pdf_pages(original), pdf_pages(r.content)
    n = len(after) - len(before)
    assert n >= 2 and after[0] == before[0] and after[1 + n:] == before[1:] and rec["reference"] in after[1] + after[2]
    assert get(world, "/download-file").content == original


def test_download_file_serves_any_other_file_and_one_it_cannot_extend_as_it_is_with_the_header_false(world):
    finished(world)
    attach(world, "Register.xlsx", b"PK\x03\x04 pretend spreadsheet", "xlsx")
    r = get(world, "/download-file", with_record="true")
    assert r.status_code == 200 and r.content == b"PK\x03\x04 pretend spreadsheet" and r.headers["x-controlled-copy"] == "false"
    attach(world, "broken.docx", b"this is not a zip file", "docx")
    r = get(world, "/download-file", with_record="true")
    assert r.content == b"this is not a zip file" and r.headers["x-controlled-copy"] == "false"
    locked = fitz.open()
    locked.new_page().insert_text((72, 72), "Protected")
    protected = locked.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="", permissions=fitz.PDF_PERM_ACCESSIBILITY)
    attach(world, "locked.pdf", protected, "pdf")
    r = get(world, "/download-file", with_record="true")
    assert r.content == protected and r.headers["x-controlled-copy"] == "false"
    assert get(world, "/download-file", with_record="true", who=VERA).status_code == 200      # published: everybody asked reads it


def test_the_file_of_an_older_version_has_that_versions_pages(world):
    approve_and_publish(world, (RANA,), (MINA,))
    row = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=5, version_number="1.0").one()
    original = docx_of(lambda d: (cover(d), d.add_page_break(), content(d)))
    (world.tmp / "v10.docx").write_bytes(original)
    row.file_name, row.file_path, row.file_type, row.file_size = "Policy v1.0.docx", str(world.tmp / "v10.docx"), "docx", len(original)
    world.db.commit()
    edit(world, SMALL, "Temporary staff added")
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    newest = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=5, version_number="1.1").one()
    for who in (OLIVIA, VERA):                                                             # the owner, and a reader who reads what is in force
        r = get(world, "/download-file", who=who, version_id=row.id, with_record="true")
        assert r.status_code == 200 and r.headers["x-controlled-copy"] == "true", who
        cells = "\n".join(cell_texts(Document(io.BytesIO(r.content))))
        assert "Rana Reviewer" in cells and "Mina Manager" in cells and "Rafay Reviewer" not in cells and "Mo Manager" not in cells
        assert "STAGE 2 — APPROVED BY" in cells
    raw = get(world, "/download-file", version_id=row.id)
    assert raw.content == original and raw.headers["x-controlled-copy"] == "false"
    assert get(world, "/download-file", version_id=row.id + 99, with_record="true").status_code == 404
    assert get(world, "/download-file", who=VERA, version_id=newest.id, with_record="true").status_code == 404      # never a version after the one in force


def test_the_export_of_a_platform_written_document_puts_the_pages_at_the_places_of_its_sections(world):
    rec = finished(world)
    d = doc(world)
    d.content = SAMPLE
    world.db.commit()
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.headers["content-disposition"] == 'attachment; filename="Access_policy-with-approval-record.pdf"'
    assert "x-approval-record" not in r.headers
    pages = pdf_pages(r.content)
    assert page_of(pages, "Prepared for the board") < page_of(pages, "Olivia Owner") < page_of(pages, "Why we have it")
    assert page_of(pages, "STAGE 1 — REVIEWED BY") < page_of(pages, "A sub-section stays") < page_of(pages, "Initial version")


def test_every_answer_of_download_file_says_whether_it_is_a_controlled_copy_and_the_export_keeps_its_own_header(world, monkeypatch):
    finished(world)
    d = doc(world)
    d.file_name = d.file_path = d.file_type = None
    world.db.commit()
    for params in ({}, {"with_record": "true"}, {"version_id": 999}, {"version_id": 999, "with_record": "true"}):
        r = get(world, "/download-file", **params)
        assert r.status_code == 404 and r.headers["x-controlled-copy"] == "false", params                    # nothing to serve: still says it is not a copy
    row = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=5).first()
    r = get(world, "/download-file", version_id=row.id, with_record="true")
    assert r.status_code == 404 and r.json()["detail"] == "File not found for this version" and r.headers["x-controlled-copy"] == "false"
    original = docx_of(lambda d_: (cover(d_), d_.add_page_break(), content(d_)))
    attach(world, "Data policy.docx", original, "docx")
    for params, said in (({}, "false"), ({"with_record": "false"}, "false"), ({"with_record": "true"}, "true")):
        r = get(world, "/download-file", **params)
        assert r.status_code == 200 and r.headers["x-controlled-copy"] == said, params
    # the export keeps what it says: a file it cannot extend is served as it is, "separate"; a copy it makes has no such header
    assert "x-approval-record" not in get(world, "/export", with_record="true").headers
    attach(world, "Register.xlsx", b"PK\x03\x04 pretend spreadsheet", "xlsx")
    assert get(world, "/export", with_record="true").headers["x-approval-record"] == "separate"
    d = doc(world)
    d.file_name = d.file_path = d.file_type = None
    world.db.commit()

    def boom(*args, **kwargs):
        raise RuntimeError("the pages could not be put together")

    monkeypatch.setattr(C, "pdf", boom)
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.headers["x-approval-record"] == "separate" and r.content[:5] == b"%PDF-"
    assert r.headers["content-disposition"] == 'attachment; filename="Access_policy.pdf"'


def test_the_routes_are_guarded_like_their_neighbours_and_say_what_they_take():
    from fastapi.routing import APIRoute

    from grc.modules.governance import permissions

    def calls(dependant):
        for d in dependant.dependencies:
            yield d.call
            yield from calls(d)

    routes_ = {r.path: r for r in router.routes if isinstance(r, APIRoute) and "GET" in r.methods}
    content_ = routes_["/governance/documents/{document_id}/controlled-content"]
    assert permissions.require_doc_reader in set(calls(content_.dependant)) and [q.name for q in content_.dependant.query_params] == ["version"]
    download = routes_["/governance/documents/{document_id}/download-file"]
    assert permissions.require_doc_reader in set(calls(download.dependant))
    assert {q.name: q.default for q in download.dependant.query_params} == {"version_id": None, "with_record": False}
    export = routes_["/governance/documents/{document_id}/export"]
    assert {q.name: q.default for q in export.dependant.query_params} == {"with_record": False}


def test_reading_the_controlled_copies_writes_nothing(world, monkeypatch):
    from sqlalchemy import event
    from sqlalchemy.orm import Session

    finished(world)
    attach(world, "Data policy.docx", docx_of(lambda d: (cover(d), d.add_page_break(), content(d))), "docx")
    flushed, commits = [], []
    event.listen(world.db, "after_flush", lambda session, ctx: flushed.append(1))
    monkeypatch.setattr(Session, "commit", lambda self: commits.append(1))
    tables = {t: world.db.execute(m.Base.metadata.tables[t].select()).fetchall() for t in
              ("grc_document_signatures", "grc_electronic_signatures", "grc_document_audit_logs", "grc_document_obligations",
               "grc_document_route_holds", "grc_governance_documents", "grc_governance_document_versions")}
    for path, params in (("/controlled-content", {}), ("/controlled-content", {"version": "1.0"}), ("/download-file", {"with_record": "true"}),
                         ("/export", {"with_record": "true"}), ("/approval-record", {}), ("/approval-record/export", {"format": "docx"})):
        assert get(world, path, **params).status_code == 200, path
    d = doc(world)
    d.file_path, d.file_type = None, None
    world.db.expire(d)
    for path, params in (("/controlled-content", {}), ("/export", {"with_record": "true"})):
        assert get(world, path, **params).status_code == 200, path
    assert not world.db.new and not world.db.dirty and not world.db.deleted
    assert flushed == [] and commits == []
    assert {t: world.db.execute(m.Base.metadata.tables[t].select()).fetchall() for t in tables} == tables
