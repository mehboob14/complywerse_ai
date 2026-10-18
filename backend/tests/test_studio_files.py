"""The studio with real files: a Word file uploaded or opened from the library is edited in place and comes back as the file;
a PDF is read from its layout and leaves as a Word copy; moving a Word file into the library keeps it a Word file."""
import io
import json
import os
import zipfile

import pytest
from docx import Document

import grc.models as m
from grc.modules.governance.studio import docx_io, exporters

import _docx_fixtures as F
from test_studio_flow import answer, change, reply, world          # noqa: F401  (world is a fixture)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def lib(world, tmp_path, monkeypatch):
    """The library's own folder for files, away from the real one."""
    folder = tmp_path / "gov"
    folder.mkdir()
    monkeypatch.setattr("grc.modules.governance.routers.documents.GOVERNANCE_UPLOAD_DIR", str(folder))
    return folder


def upload(world, path, name="policy.docx", kind=DOCX):
    with open(path, "rb") as f:
        return world.http.post(world.api("/sessions/upload"), files={"file": (name, f.read(), kind)}, data={"doc_type": "policy"}).json()


def export(world, sid, fmt="docx", **body):
    return world.http.post(world.api(f"/sessions/{sid}/export"), json={"format": fmt, **body})


def docx_of(response):
    assert response.status_code == 200, response.text
    return Document(io.BytesIO(response.content)), zipfile.ZipFile(io.BytesIO(response.content))


def test_an_uploaded_word_file_is_edited_by_chat_and_leaves_as_the_same_file(world, tmp_path):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    s = upload(world, path)
    assert s["source"] == {"kind": "docx", "name": "policy.docx", "in_place": True} and s["kind"] == "library" and s["document_id"]
    assert s["content"].startswith("# Information Security Policy") and "**all employees**" in s["content"]

    paragraph = s["content"].split("\n\n")[1]
    world.respond(lambda body: reply(answer([change("replace", "B2", markdown=paragraph.replace("**all employees**", "**all staff**"), reason="wording")], "Changed the audience.")))
    msg = world.http.post(world.api(f"/sessions/{s['id']}/messages"), json={"text": "Say staff instead of employees."}).json()
    accepted = world.http.post(world.api(f"/sessions/{s['id']}/messages/{msg['id']}/accept"), json={}).json()
    assert accepted["accepted"] == 1 and "**all staff**" in accepted["content"]

    check = world.http.post(world.api(f"/sessions/{s['id']}/export/check"), json={"format": "docx"}).json()
    assert any(n["rule"] == "word_in_place" for n in check["notes"]) and check["can_export"] is True
    assert not any(n["rule"] in ("entity", "logo") for n in check["notes"])                    # its own header is kept: the studio's is not added
    doc, z = docx_of(export(world, s["id"]))
    p = next(p for p in doc.paragraphs if p.text.startswith("This policy applies"))
    assert "all staff" in p.text and "ACME BANK - CONFIDENTIAL" in z.read("word/header1.xml").decode()
    original = zipfile.ZipFile(path)
    assert z.read("word/header1.xml") == original.read("word/header1.xml") and z.read("word/footer1.xml") == original.read("word/footer1.xml")
    assert len([n for n in z.namelist() if n.startswith("word/media/")]) == 1                    # the picture is still there


def test_a_word_file_edited_in_place_gets_the_watermark_in_its_own_header_when_asked(world, tmp_path):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    s = upload(world, path)
    world.http.patch(world.api(f"/sessions/{s['id']}"), json={"instructions": {"watermark": True, "watermark_text": "CONFIDENTIAL"}})
    check = world.http.post(world.api(f"/sessions/{s['id']}/export/check"), json={"format": "docx"}).json()
    assert any(n["rule"] == "watermark" and "CONFIDENTIAL" in n["text"] for n in check["notes"])
    _doc, z = docx_of(export(world, s["id"]))
    header = z.read("word/header1.xml").decode()
    assert "ACME BANK - CONFIDENTIAL" in header and 'string="CONFIDENTIAL"' in header               # its own header, with the watermark added
    assert z.read("word/footer1.xml") == zipfile.ZipFile(path).read("word/footer1.xml")
    world.http.patch(world.api(f"/sessions/{s['id']}"), json={"instructions": {"watermark": False}})
    _doc, z = docx_of(export(world, s["id"]))
    assert z.read("word/header1.xml") == zipfile.ZipFile(path).read("word/header1.xml")              # off again: the header is the file's own


def test_when_the_changes_cannot_be_written_into_the_word_file_a_clean_copy_is_made_and_the_person_is_told(world, tmp_path, monkeypatch):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    s = upload(world, path)
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"] + "\nA closing paragraph.\n", "base_revision": s["revision"]})

    def refuse(*args, **kwargs):
        raise docx_io.PatchProblem("block 9 reads differently")

    monkeypatch.setattr(docx_io, "export_in_place", refuse)
    check = world.http.post(world.api(f"/sessions/{s['id']}/export/check"), json={"format": "docx"}).json()
    assert any(n["rule"] == "word_rebuilt" and n["level"] == "warning" for n in check["notes"])
    doc, z = docx_of(export(world, s["id"]))
    assert "A closing paragraph." in "\n".join(p.text for p in doc.paragraphs) and "ACME BANK" not in z.read("word/header1.xml").decode()


def test_a_pdf_is_read_from_its_layout_and_leaves_as_a_word_copy(world, tmp_path):
    md = ("# Retention Policy\n\n## 1. Purpose\n\n" + "Records shall be kept as long as the law requires and no longer. " * 12 + "\n\n## 2. Rules\n\n1. Keep invoices for ten years.\n2. Keep emails for two years.\n\n"
          "| Record | Years |\n|--------|-------|\n| Invoice | 10 |\n| Email | 2 |\n")
    path = tmp_path / "retention.pdf"
    path.write_bytes(exporters.to_pdf(md, title="Retention", entity_name="Capital Financial Services Bank", doc_code="POL-9", version="1.0"))
    s = upload(world, str(path), "retention.pdf", "application/pdf")
    assert s["source"]["kind"] == "pdf" and s["source"]["in_place"] is False
    content = s["content"]
    assert "# Retention Policy" in content and "## 1. Purpose" in content and "1. Keep invoices for ten years." in content and "| Invoice | 10 |" in content
    assert "POL-9" not in content and "Capital Financial Services Bank" not in content           # the running header and footer are dropped
    doc, _ = docx_of(export(world, s["id"]))
    assert any(p.text == "Retention Policy" and p.style.name == "Title" for p in doc.paragraphs) and len(doc.tables) == 1
    assert world.http.post(world.api("/sessions/upload"), files={"file": ("x.pdf", b"%PDF-1.4 broken", "application/pdf")}).status_code == 400


def make_library_doc(world, folder, *, status="published", file_fixture=F.corporate_policy):
    path = file_fixture(str(folder / "orig.docx"))
    size = os.path.getsize(path)
    world.db.add(m.GovernanceDocument(id=20, tenant_id=1, title="Uploaded Policy", doc_type="policy", status=status, current_version="1.0", file_name="orig.docx",
                                      file_path=path, file_size=size, file_type="docx", owner_id=1, content=None))
    world.db.add(m.GovernanceDocumentVersion(document_id=20, version_number="1.0", change_type="major", title="Uploaded Policy", file_name="orig.docx", file_path=path,
                                             file_size=size, file_type="docx", change_summary="Initial version", status="current", created_by=1))
    world.db.commit()
    return path


def test_a_library_word_file_is_opened_edited_and_saved_as_the_next_version_file(world, lib):
    original = make_library_doc(world, lib)
    world.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=20, role_type="reviewer", target_type="user", target_id=1, added_by=1))
    world.db.commit()                                                          # it has a reviewer, so the save sends it back to review
    before = open(original, "rb").read()
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 20}).json()
    assert s["source"]["in_place"] is True and s["document_id"] == 20 and "**all employees**" in s["content"] and s["base_label"] == "Uploaded Policy v1.0"
    paragraph = s["content"].split("\n\n")[1]
    new = s["content"].replace(paragraph, paragraph.replace("Review it every year.", "Review it every six months."))
    assert world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": new, "base_revision": s["revision"]}).status_code == 200

    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={"change_reason": "Review cycle"}).json()
    assert saved["version"] == "1.1" and saved["status"] == "pending_review"                   # the approved text changed: back to review
    doc = world.db.get(m.GovernanceDocument, 20)
    world.db.refresh(doc)
    assert doc.file_path != original and doc.file_type == "docx" and not (doc.content or "").strip()     # the file is still the document
    assert open(original, "rb").read() == before                                                  # the earlier version's file is untouched
    versions = {v.version_number: v for v in world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=20)}
    assert versions["1.0"].status == "superseded" and versions["1.0"].file_path == original
    assert versions["1.1"].status == "current" and versions["1.1"].file_path == doc.file_path and versions["1.1"].change_reason == "Review cycle"
    new_file = Document(doc.file_path)
    assert "Review it every six months." in next(p for p in new_file.paragraphs if p.text.startswith("This policy applies")).text
    assert zipfile.ZipFile(doc.file_path).read("word/header1.xml") == zipfile.ZipFile(original).read("word/header1.xml")

    assert world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={}).status_code == 400          # nothing new since
    again = world.http.get(world.api(f"/sessions/{s['id']}")).json()
    assert again["base_label"] == "Uploaded Policy v1.1"


def test_a_word_file_saved_while_in_review_is_signed_again_from_the_start_by_its_reviewers(world, lib):
    from datetime import datetime

    make_library_doc(world, lib, status="pending_review")
    world.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=20, role_type="reviewer", target_type="user", target_id=1, added_by=1))
    world.db.get(m.GovernanceDocument, 20).signoff_round_at = datetime(2026, 1, 1)
    world.db.commit()
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 20}).json()
    assert (s["document_status"], s["unsaved_in_library"]) == ("pending_review", False)
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"].replace("Review it every year.", "Review it every six months."),
                                                                     "base_revision": s["revision"]})
    assert world.http.get(world.api(f"/sessions/{s['id']}")).json()["unsaved_in_library"] is True
    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={}).json()
    assert (saved["version"], saved["status"], saved["change"]) == ("1.1", "pending_review", "minor")
    doc = world.db.get(m.GovernanceDocument, 20)
    world.db.refresh(doc)
    assert doc.signoff_round_at > datetime(2026, 1, 1)
    assert world.db.query(m.DocumentObligation).filter_by(document_id=20, kind="review", user_id=1, status="pending").count() == 1
    row = world.db.query(m.DocumentAuditLog).filter_by(document_id=20, action="content_edited").one()
    assert row.action_details.endswith("; signed again from the start by its reviewers")
    assert world.http.get(world.api(f"/sessions/{s['id']}")).json()["unsaved_in_library"] is False


def test_saving_over_a_library_file_that_was_replaced_is_a_conflict_unless_forced(world, lib):
    make_library_doc(world, lib, status="draft")
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 20}).json()
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"] + "\nAn addition.\n", "base_revision": s["revision"]})
    other = F.corporate_policy(str(lib / "someone_elses.docx"))
    world.db.get(m.GovernanceDocument, 20).file_path = other
    world.db.commit()
    assert world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={}).status_code == 409
    ok = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={"force": True})
    assert ok.status_code == 200 and ok.json()["version"] == "1.1" and ok.json()["status"] == "draft"


def test_a_library_word_file_with_tracked_changes_shows_them_as_the_redline(world, lib, tmp_path):
    make_library_doc(world, lib, file_fixture=F.tracked_policy)
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 20}).json()
    assert "seven years" in s["content"] and s["base_label"] == "Uploaded Policy v1.0 before its tracked changes"
    diff = world.http.get(world.api(f"/sessions/{s['id']}/diff"), params={"from": 0}).json()
    assert diff["stats"]["modified"] == 3 and any(("delete", "five") in [tuple(x) for x in (r.get("segments") or [])] for r in diff["rows"])


def test_an_uploaded_word_file_is_in_the_library_at_once_as_the_file_itself_numbered_as_it_states(world, lib, tmp_path):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    s = upload(world, path, name="Information Security Policy v3.2.docx")
    doc = world.db.get(m.GovernanceDocument, s["document_id"])
    only = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id).one()
    # its footer says "POL-001 v1.0": what the document states comes before what its file name says
    assert (s["kind"], doc.status, doc.current_version, only.version_number, only.status) == ("library", "draft", "1.0", "1.0", "current")
    assert only.change_summary == "Initial version: the uploaded file, version 1.0 as the document states it" and s["base_label"].endswith(" v1.0")
    assert open(doc.file_path, "rb").read() == open(path, "rb").read() and only.file_path == doc.file_path        # the file itself, not a re-saved copy
    assert doc.file_path.startswith(str(lib)) and doc.file_type == "docx" and not (doc.content or "").strip()
    assert "uploaded into the Comply AI Studio" in world.db.query(m.DocumentAuditLog).filter_by(document_id=doc.id, action="created").one().action_details
    assert world.http.post(world.api(f"/sessions/{s['id']}/move-to-library"), json={}).status_code == 400            # it is there already


def test_edits_to_an_uploaded_word_file_save_as_its_next_version_major_or_minor_as_comply_agent_decides(world, lib, tmp_path):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    s = upload(world, path)
    paragraph = s["content"].split("\n\n")[1]
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"].replace(paragraph, paragraph.replace("every year", "every quarter")), "base_revision": s["revision"]})
    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={}).json()
    doc = world.db.get(m.GovernanceDocument, s["document_id"])
    assert saved["version"] == "1.1" and doc.file_type == "docx" and doc.file_path.startswith(str(lib))      # two words: minor (no AI answer here: the rule)
    assert "every quarter" in "\n".join(p.text for p in Document(doc.file_path).paragraphs)                    # the edit is in the file, in the original's formatting
    assert zipfile.ZipFile(doc.file_path).read("word/header1.xml") == zipfile.ZipFile(path).read("word/header1.xml")
    # the uploaded file stays as the first version, as it was
    first, second = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id).order_by(m.GovernanceDocumentVersion.id).all()
    assert (first.version_number, first.status, first.file_name) == ("1.0", "superseded", "policy.docx") and open(first.file_path, "rb").read() == open(path, "rb").read()
    assert (second.version_number, second.status, second.file_path, second.change_type) == ("1.1", "current", doc.file_path, "minor")
    assert second.change_set["items"] and "every year" in second.change_set["items"][0]["before_text"] and "every quarter" in second.change_set["items"][0]["after_text"]
    # the next one adds an obligation: Comply Agent calls it a major change, so it is version 2.0 of the same file
    world.respond(lambda body: reply(json.dumps({"change": "major", "reason": "Staff must now report incidents within one day."})))
    again = world.http.get(world.api(f"/sessions/{s['id']}")).json()
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": again["content"] + "\nStaff shall report incidents within one day.\n", "base_revision": again["revision"]})
    saved2 = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={"change_reason": "Incident reporting"}).json()
    assert saved2["version"] == "2.0" and "Staff shall report incidents within one day." in "\n".join(p.text for p in Document(world.db.get(m.GovernanceDocument, doc.id).file_path).paragraphs)
    third = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id, version_number="2.0").one()
    assert third.change_type == "major" and third.change_summary == "Edited in the Comply AI Studio: Incident reporting. Major change: Staff must now report incidents within one day."
    sent = world.wire.requests[-1].body
    assert sent["output_config"]["effort"] == "low" and "Staff shall report incidents within one day." in sent["messages"][0]["content"]
    assert "Review it every quarter" not in sent["messages"][0]["content"]                       # only what changed is read, not the whole document


def test_an_uploaded_pdf_is_in_the_library_at_once_as_the_pdf_and_its_edits_save_as_a_word_copy(world, lib, tmp_path):
    md = "# Retention Policy\n\n## 1. Purpose\n\n" + "Records shall be kept as long as the law requires and no longer. " * 6 + "\n\n## 2. Rules\n\n1. Keep invoices for ten years.\n2. Keep emails for two years.\n"
    pdf = tmp_path / "retention.pdf"
    pdf.write_bytes(exporters.to_pdf(md, title="Retention", entity_name="Capital Financial Services Bank", doc_code="POL-9", version="2.3"))
    s = upload(world, str(pdf), "retention.pdf", "application/pdf")
    doc = world.db.get(m.GovernanceDocument, s["document_id"])
    assert (doc.file_type, doc.current_version, doc.status) == ("pdf", "2.3", "draft") and open(doc.file_path, "rb").read() == pdf.read_bytes()   # its footer says v2.3
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"].replace("two years", "three years"), "base_revision": s["revision"]})
    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={}).json()
    world.db.refresh(doc)
    assert saved["version"] == "2.4" and doc.file_type == "docx" and "three years" in "\n".join(p.text for p in Document(doc.file_path).paragraphs)
    first, second = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id).order_by(m.GovernanceDocumentVersion.id).all()
    assert (first.version_number, first.file_name, first.file_type, first.status) == ("2.3", "retention.pdf", "pdf", "superseded")    # the PDF is kept as it was
    assert (second.version_number, second.status, second.change_type, second.file_type) == ("2.4", "current", "minor", "docx")
    assert "three years" in second.change_set["items"][0]["after_text"]


def test_an_upload_session_from_before_moves_to_the_library_numbered_as_its_file_states(world, lib, tmp_path):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    s = upload(world, path, name="Information Security Policy v2.0.docx")
    row = world.db.get(m.StudioSession, s["id"])                 # a session uploaded before an upload went into the library at once
    row.kind, row.document_id, row.settings = "upload", None, {}
    world.db.commit()
    paragraph = s["content"].split("\n\n")[1]
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"].replace(paragraph, paragraph.replace("every year", "every quarter")), "base_revision": s["revision"]})
    moved = world.http.post(world.api(f"/sessions/{s['id']}/move-to-library"), json={"title": "Information Security Policy"}).json()
    doc = world.db.get(m.GovernanceDocument, moved["document_id"])
    first, second = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id).order_by(m.GovernanceDocumentVersion.id).all()
    assert (first.version_number, first.status, first.change_summary) == ("1.0", "superseded", "Initial version: the uploaded file, version 1.0 as the document states it")
    assert (second.version_number, second.change_type, doc.current_version, doc.file_type) == ("1.1", "minor", "1.1", "docx")
    assert open(first.file_path, "rb").read() == open(path, "rb").read() and "every quarter" in "\n".join(p.text for p in Document(doc.file_path).paragraphs)


def test_a_document_can_be_compared_with_a_file_the_person_uploads(world):
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 10}).json()
    other = b"# Information Security Policy\n\n## 1. Purpose\n\nA quite different purpose statement.\n\n## 2. Scope\n\nAll employees and contractors.\n"
    ref = world.http.post(world.api(f"/sessions/{s['id']}/references/upload"), files={"file": ("other.md", other, "text/markdown")}).json()
    out = world.http.post(world.api(f"/sessions/{s['id']}/compare"), json={"reference_id": ref["id"]}).json()
    assert out["other_title"] == "other.md" and out["stats"]["inserted"] + out["stats"]["deleted"] + out["stats"]["modified"] >= 1
    assert any(r["status"] != "equal" for r in out["rows"]) and any(r["status"] == "equal" for r in out["rows"])
    assert world.http.post(world.api(f"/sessions/{s['id']}/compare"), json={"reference_id": 9999}).status_code == 404
    mine = world.http.post(world.api("/sessions"), json={"kind": "new", "title": "Other", "template_id": "builtin:blank", "doc_type": "other"}).json()
    assert world.http.post(world.api(f"/sessions/{mine['id']}/compare"), json={"reference_id": ref["id"]}).status_code == 404          # not another session's file
    assert world.http.post(world.api(f"/sessions/{s['id']}/compare"), json={}).status_code == 400


def test_a_draft_whose_library_pdf_was_deleted_can_be_moved_back_as_text_and_saves_as_text(world, lib, tmp_path):
    from grc.modules.governance.routers.documents import delete_document

    pdf = tmp_path / "retention.pdf"
    pdf.write_bytes(exporters.to_pdf("# Retention Policy\n\n## 1. Purpose\n\n" + "Records shall be kept as long as the law requires and no longer. " * 6 + "\n", title="Retention"))
    world.db.add(m.GovernanceDocument(id=21, tenant_id=1, title="Retention Policy", doc_type="policy", status="draft", current_version="1.0", file_name="retention.pdf",
                                      file_path=str(pdf), file_size=pdf.stat().st_size, file_type="pdf", owner_id=1, content=None))
    world.db.commit()
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 21}).json()
    assert s["settings"]["file_master"] is True and s["document_id"] == 21

    delete_document(document_id=21, db=world.db, current_user=world.db.get(m.GRCUser, 1))
    detached = world.http.get(world.api(f"/sessions/{s['id']}")).json()
    assert detached["document_id"] is None and "# Retention Policy" in detached["content"]           # the draft is kept

    moved = world.http.post(world.api(f"/sessions/{s['id']}/move-to-library"), json={"title": "Retention Policy"}).json()
    assert world.http.get(world.api(f"/sessions/{s['id']}")).json()["settings"].get("file_master") is None
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": detached["content"] + "\nMore.\n", "base_revision": detached["revision"]})
    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={})
    doc = world.db.get(m.GovernanceDocument, moved["document_id"])
    world.db.refresh(doc)
    assert saved.status_code == 200 and saved.json()["version"] == "1.1" and "More." in doc.content and not doc.file_path


def test_a_text_document_in_the_library_still_saves_as_text(world, lib):
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 11}).json()
    assert s["source"]["in_place"] is False
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"] + "\nMore.\n", "base_revision": s["revision"]})
    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={}).json()
    doc = world.db.get(m.GovernanceDocument, 11)
    world.db.refresh(doc)
    assert saved["version"] == "1.1" and "More." in doc.content and not doc.file_path
