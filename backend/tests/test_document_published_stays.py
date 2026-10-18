"""The routes that serve a governance document (the user's decisions of 2026-10-17): who may read it at all, and which version they
read. While a new version of a published document is on its way, the version published stays in force: staff (who may not change,
publish or sign it off) are served that one everywhere - the document, the lists and the tree, its file, its export, its versions
and its change chain - and never the newer text; owners, editors, administrators and the people taking part in its sign-off read
the newest. Over HTTP through the whole governance router, on in-memory SQLite, on a session that does not flush by itself."""
import base64
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image, ImageDraw
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.governance import permissions as P
from grc.modules.governance.router import router
from grc.modules.governance.routers import documents
from grc.modules.governance.studio import service
from grc.routers.auth_router import require_auth

OWNER, REVIEWER, STAFF, EDITOR, ADMIN, AUDITOR, OUTSIDER, PREPARER, ACKER = range(1, 10)
TEXT, DRAFT, WORD = 5, 6, 7                    # a text in review (1.0 in force) | a draft never published, under 5 | a Word file in review
V10 = "# Access policy\n\n## 1. Rules\n\nAccess is reviewed every year.\n"
V11 = V10.replace("every year", "every quarter")
SIGNED = ("# Leave policy\n\n## 1. Rules\n\nLeave is booked ahead.\n\n## 2. Approval Signoff\n\n| Role | Name | Designation | Signature / Date |\n"
          "| --- | --- | --- | --- |\n| Approved by | | | |\n")
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _word(path, text):
    d = Document()
    d.add_paragraph(text)
    d.save(path)
    return path.read_bytes()


def _published(db, did, title, **held):
    """A document published at version 1.0 (the listener records the number in force)."""
    d = m.GovernanceDocument(id=did, tenant_id=1, title=title, doc_type="policy", status="draft", current_version="1.0", owner_id=OWNER, **held)
    db.add(d)
    db.flush()
    d.status = "published"
    return d


def _in_review(db, did, title, old, new):
    """Published at 1.0, its version 1.1 in review: the row 1.0 holds the published text or file."""
    d = _published(db, did, title, **old)
    db.add(m.GovernanceDocumentVersion(document_id=did, version_number="1.0", change_type="major", title=title, status="superseded", **old))
    db.add(m.GovernanceDocumentVersion(document_id=did, version_number="1.1", change_type="minor", title=title, status="current", **new))
    for key, value in new.items():
        setattr(d, key, value)
    d.current_version, d.status = "1.1", "pending_review"
    for kind, uid in (("reviewer", REVIEWER), ("prepared_by", PREPARER)):
        db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=did, role_type=kind, target_type="user", target_id=uid))
    db.add(m.DocumentObligation(tenant_id=1, document_id=did, kind="acknowledgement", user_id=ACKER, document_version="1.0"))


@pytest.fixture
def world(monkeypatch, tmp_path):
    monkeypatch.setenv("SESSION_SECRET", "unit-test-secret")
    monkeypatch.setattr(service, "UPLOAD_DIR", str(tmp_path / "studio"))
    monkeypatch.setattr(documents, "GOVERNANCE_UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr(documents, "_resolve_tenant_slug_for_autoparse", lambda tenant_id: None)  # never the real master database
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, autoflush=False)()
    db.add(m.Tenant(id=1, name="CFSB", slug="cfsb"))
    for uid in range(1, 10):
        db.add(m.GRCUser(id=uid, username=f"u{uid}", email=f"u{uid}@cfsb.test", display_name=f"User {uid}"))
    for rid, role, perm in [(20, "Document editors", "governance:policies:edit"), (22, "Administrator", None),
                            (23, "Auditors", "compliance:frameworks:view"), (24, "Document readers", "governance:policies:view")]:
        db.add(m.Role(id=rid, tenant_id=1, name=role))
        if perm:
            db.add(m.Permission(id=rid, name=perm, resource=perm.rsplit(":", 1)[0], action=perm.rsplit(":", 1)[1]))
            db.add(m.RolePermission(role_id=rid, permission_id=rid))
    for uid, rid in [(EDITOR, 20), (ADMIN, 22), (AUDITOR, 23), (STAFF, 24)]:
        db.add(m.UserRole(user_id=uid, role_id=rid, tenant_id=1))
    files = {"a": _word(tmp_path / "a.docx", "Published wording."), "b": _word(tmp_path / "b.docx", "New wording.")}
    _in_review(db, TEXT, "Access policy", {"content": V10}, {"content": V11})
    _in_review(db, WORD, "Retention policy", {"file_name": "a.docx", "file_path": str(tmp_path / "a.docx"), "file_type": "docx", "file_size": len(files["a"])},
               {"file_name": "b.docx", "file_path": str(tmp_path / "b.docx"), "file_type": "docx", "file_size": len(files["b"])})
    db.add(m.GovernanceDocument(id=DRAFT, tenant_id=1, title="Access standard", doc_type="standard", status="draft", current_version="1.0",
                                content="# Access standard\n", owner_id=OWNER, parent_document_id=TEXT, next_review_date=datetime(2026, 11, 1)))
    db.commit()
    app = FastAPI()
    app.include_router(router)
    me = {"id": STAFF}
    app.dependency_overrides[m.get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: db.get(m.GRCUser, me["id"])
    yield SimpleNamespace(db=db, me=me, http=TestClient(app), files=files, tmp=tmp_path)
    db.close()


def row_id(w, did, number):
    V = m.GovernanceDocumentVersion
    return w.db.query(V.id).filter(V.document_id == did, V.version_number == number).order_by(V.id).first()[0]


def rows(w, did):
    V = m.GovernanceDocumentVersion
    w.db.expire_all()
    return w.db.query(V).filter(V.document_id == did).order_by(V.id).all()


def doc(w, did):
    w.db.expire_all()
    return w.db.get(m.GovernanceDocument, did)


def get(w, who, path):
    w.me["id"] = who
    return w.http.get(path)


# ── which version each person reads ──────────────────────────────────────────
def test_staff_read_the_published_version_while_a_new_one_is_in_review(world):
    r10, r11 = row_id(world, TEXT, "1.0"), row_id(world, TEXT, "1.1")
    for who in (STAFF, ACKER, AUDITOR):                                     # the library right, asked to acknowledge, an auditor
        d = get(world, who, f"/governance/documents/{TEXT}").json()
        assert (d["content"], d["current_version"], d["version_in_force"], d["status"]) == (V10, "1.0", "1.0", "pending_review"), who
        assert d["newer_version"] == {"number": "1.1", "status": "pending_review"} and d["in_force_version_id"] == r10, who
        assert [v["id"] for v in d["versions"]] == [r10], who               # never the newer version's card
    for who in (OWNER, REVIEWER, PREPARER, EDITOR, ADMIN):                  # who may change, publish or sign it off: the newest
        d = get(world, who, f"/governance/documents/{TEXT}").json()
        assert (d["content"], d["current_version"], d["version_in_force"], d["in_force_version_id"]) == (V11, "1.1", "1.0", r10), who
        assert "newer_version" not in d and sorted(v["id"] for v in d["versions"]) == [r10, r11], who


def test_lists_and_the_tree_serve_staff_the_published_text_and_hide_a_draft_never_published(world):
    items = {d["id"]: d for d in get(world, STAFF, "/governance/documents").json()["items"]}
    assert sorted(items) == [TEXT, WORD]
    assert (items[TEXT]["content"], items[TEXT]["version_in_force"], items[WORD]["file_name"]) == (V10, "1.0", "a.docx")
    assert {d["id"]: d["content"] for d in world.http.get("/governance/documents/policies").json()["items"]} == {TEXT: V10, WORD: None}
    tree = {n["id"]: n for n in world.http.get("/governance/documents/hierarchy").json()}
    assert tree[TEXT]["content"] == V10 and [v["number"] for v in tree[TEXT]["versions"]] == ["1.0"] and tree[TEXT]["children"] == []
    assert world.http.get(f"/governance/documents/{TEXT}").json()["child_documents"] == []
    assert [r["document_id"] for r in world.http.get("/governance/documents/reviews/upcoming?days=100000").json()["reviews"]] == []

    items = {d["id"]: d for d in get(world, EDITOR, "/governance/documents").json()["items"]}
    assert sorted(items) == [TEXT, DRAFT, WORD] and items[TEXT]["content"] == V11 and items[WORD]["file_name"] == "b.docx"
    tree = {n["id"]: n for n in world.http.get("/governance/documents/hierarchy").json()}
    assert [v["number"] for v in tree[TEXT]["versions"]] == ["1.0", "1.1"] and [c["id"] for c in tree[TEXT]["children"]] == [DRAFT]
    assert [c["id"] for c in world.http.get(f"/governance/documents/{TEXT}").json()["child_documents"]] == [DRAFT]
    assert [r["document_id"] for r in world.http.get("/governance/documents/reviews/upcoming?days=100000").json()["reviews"]] == [DRAFT]

    items = {d["id"]: d for d in get(world, REVIEWER, "/governance/documents").json()["items"]}     # no library right: what they take part in
    assert sorted(items) == [TEXT, WORD] and items[TEXT]["content"] == V11
    r = world.http.get("/governance/documents/hierarchy")
    assert (r.status_code, r.json()["detail"]) == (403, P.LIBRARY_REFUSED)


def test_staff_get_the_published_file_of_a_word_document_and_never_a_newer_one(world):
    r11 = row_id(world, WORD, "1.1")
    base = f"/governance/documents/{WORD}"
    for who in (STAFF, ACKER):
        assert get(world, who, f"{base}/download-file").content == world.files["a"], who
        assert world.http.get(f"{base}/export").content == world.files["a"], who
        assert "Published wording." in world.http.get(f"{base}/view-html").json()["html"], who
        assert world.http.get(f"{base}/download-file?version_id={r11}").status_code == 404, who
    assert get(world, OWNER, f"{base}/download-file").content == world.files["b"]
    assert "New wording." in world.http.get(f"{base}/view-html").json()["html"]
    assert world.http.get(f"{base}/download-file?version_id={r11}").content == world.files["b"]


def test_comparing_two_documents_reads_each_as_the_person_reads_it(world):
    world.db.add(m.GovernanceDocument(id=8, tenant_id=1, title="Other policy", doc_type="policy", status="published", content="# Other\n"))
    world.db.commit()
    world.me["id"] = STAFF
    out = world.http.post("/governance/documents/8/compare-with-document", json={"target_document_id": TEXT}).json()
    assert out["template"]["content"] == V10.strip()
    r = world.http.post(f"/governance/documents/{TEXT}/compare-with-document", json={"target_document_id": DRAFT})
    assert (r.status_code, r.json()["detail"]) == (403, P.DRAFT_REFUSED)           # a draft never published is not theirs to read
    world.me["id"] = REVIEWER
    assert world.http.post(f"/governance/documents/{TEXT}/compare-with-document", json={"target_document_id": WORD}).json()["document"]["content"] == V11.strip()


def test_versions_and_the_change_chain_end_at_the_version_in_force_for_staff(world):
    r10, r11 = row_id(world, TEXT, "1.0"), row_id(world, TEXT, "1.1")
    listed = get(world, STAFF, f"/governance/versions/document/{TEXT}").json()
    assert ([v["id"] for v in listed["items"]], listed["total"], listed["current_version"]) == ([r10], 1, "1.0")
    history = world.http.get(f"/governance/versions/document/{TEXT}/history").json()
    assert [v["id"] for v in history["items"]] == [r10] and history["document"]["current_version"] == "1.0"
    assert world.http.get(f"/governance/versions/{r10}").json()["content"] == V10
    assert world.http.get(f"/governance/versions/{r11}").status_code == 404
    assert world.http.get(f"/governance/versions/compare/{r10}/{r11}").status_code == 404
    chain = world.http.get(f"/governance/documents/{TEXT}/change-chain").json()
    assert [v["number"] for v in chain["versions"]] == ["1.0"] and (chain["document"]["version"], chain["document"]["in_force"]) == ("1.0", "1.0")
    red = world.http.get(f"/governance/documents/{TEXT}/change-chain/redline?from={r10}").json()     # compared with 1.0, not the draft
    assert (red["to"]["id"], red["hunks"], [v["id"] for v in red["versions"]]) == (r10, [], [r10])
    for query in (f"from={r10}&to={r11}", f"from={r11}"):
        assert world.http.get(f"/governance/documents/{TEXT}/change-chain/redline?{query}").status_code == 404, query
    assert "every quarter" not in world.http.get(f"/governance/documents/{TEXT}/change-chain/export?format=csv").content.decode("utf-8-sig")

    assert len(get(world, REVIEWER, f"/governance/versions/document/{TEXT}").json()["items"]) == 2
    assert world.http.get(f"/governance/versions/{r11}").json()["content"] == V11
    assert world.http.get(f"/governance/versions/compare/{r10}/{r11}").json()["diff"]["content_changed"] is True
    assert [v["number"] for v in world.http.get(f"/governance/documents/{TEXT}/change-chain").json()["versions"]] == ["1.0", "1.1"]
    red = world.http.get(f"/governance/documents/{TEXT}/change-chain/redline?from={r10}").json()
    assert any("every quarter" in h["after_text"] for h in red["hunks"])
    assert "every quarter" in world.http.get(f"/governance/documents/{TEXT}/change-chain/export?format=csv").content.decode("utf-8-sig")


def test_the_text_before_and_after_an_edit_is_left_out_of_the_history_staff_read(world):
    for field, old, new in (("content", V10, V11), ("status", "published", "pending_review")):
        world.db.add(m.DocumentAuditLog(document_id=TEXT, tenant_id=1, action="edited", field_changed=field, old_value=old, new_value=new,
                                        performed_by=OWNER, performed_at=datetime.utcnow()))
    world.db.commit()
    logs = {log["field_changed"]: log for log in get(world, STAFF, f"/governance/documents/{TEXT}/audit-logs").json()}
    assert (logs["content"]["old_value"], logs["content"]["new_value"], logs["status"]["new_value"]) == (None, None, "pending_review")
    logs = {log["field_changed"]: log for log in get(world, REVIEWER, f"/governance/documents/{TEXT}/audit-logs").json()}
    assert (logs["content"]["old_value"], logs["content"]["new_value"]) == (V10, V11)


def test_staff_never_read_a_version_a_take_back_marked_reverted(world):
    db = world.db
    db.add(m.GovernanceDocument(id=9, tenant_id=1, title="Travel policy", doc_type="policy", status="pending_review", current_version="1.2",
                                published_version="1.1", content="v1.2", owner_id=OWNER))
    for number, status, content in (("1.0", "superseded", "v1.0"), ("1.1", "reverted", "the undone gap fix"),
                                    ("1.1", "superseded", "v1.1 as published"), ("1.2", "current", "v1.2")):
        db.add(m.GovernanceDocumentVersion(document_id=9, version_number=number, title="Travel policy", status=status, content=content))
    db.commit()
    assert get(world, STAFF, "/governance/documents/9").json()["content"] == "v1.1 as published"


# ── the version in force is recorded and kept, whichever route wrote the new one ──
def test_every_new_number_of_a_published_text_is_the_one_in_force(world):
    db = world.db
    for did in (10, 11, 12):
        _published(db, did, f"Policy {did}", content=SIGNED)
    db.commit()
    world.me["id"] = ADMIN
    assert world.http.post("/governance/documents/11/signoff", json={"signoffs": [{"role": "Approved by", "name": "Ada"}]}).status_code == 200
    assert world.http.post("/governance/versions/document/12", json={"change_type": "minor"}).status_code == 201     # a number, no text
    for did in (11, 12):
        d = doc(world, did)
        assert (d.status, d.published_version, d.last_published_version) == ("published", d.current_version, d.current_version), did
        assert d.current_version != "1.0", did
    # "published" sent with a new text is not a publication (it is published already): the new text is reviewed first, also when
    # the person may not publish (an editor) - it never becomes the text in force without a review
    _published(db, 16, "Policy 16", content=SIGNED)
    db.commit()
    for did, who in ((10, ADMIN), (16, EDITOR)):
        world.me["id"] = who
        assert world.http.put(f"/governance/documents/{did}", json={"status": "published", "content": SIGNED + "\nMore.\n"}).status_code == 200
        d = doc(world, did)
        assert d.status != "published" and d.published_version == "1.0" != d.current_version, who
        assert get(world, STAFF, f"/governance/documents/{did}").json()["content"] == SIGNED, who
    # a document with no versions yet keeps what it held as its first version
    assert [(v.version_number, v.change_type, v.status, v.content) for v in rows(world, 12)] == [
        ("1.0", "baseline", "superseded", SIGNED), ("1.1", "minor", "current", SIGNED)]


def test_a_new_text_written_over_a_published_document_without_versions_keeps_the_published_text_for_staff(world):
    _published(world.db, 13, "Policy 13", content=V10)
    world.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=13, role_type="reviewer", target_type="user", target_id=REVIEWER))
    world.db.commit()
    world.me["id"] = OWNER
    assert world.http.post("/governance/versions/document/13", json={"change_type": "minor", "content": V11}).status_code == 201
    d = doc(world, 13)
    assert d.status in P.ON_ITS_WAY and (d.current_version, d.published_version) == ("1.1", "1.0")
    assert get(world, STAFF, "/governance/documents/13").json()["content"] == V10
    assert get(world, REVIEWER, "/governance/documents/13").json()["content"] == V11


def test_uploading_a_file_over_a_published_document_keeps_the_published_file(world):
    _published(world.db, 14, "Policy 14", file_name="a.docx", file_path=str(world.tmp / "a.docx"), file_type="docx", file_size=len(world.files["a"]))
    world.db.commit()
    world.me["id"] = OWNER                                                 # asked to replace the file in place: the version in force stays
    r = world.http.post("/governance/documents/14/upload-file", data={"create_new_version": "false"},
                        files={"file": ("b.docx", world.files["b"], DOCX)})
    assert r.status_code == 200, r.text
    first, second = rows(world, 14)
    assert (first.version_number, first.change_type, first.status, first.file_path) == ("1.0", "baseline", "superseded", str(world.tmp / "a.docx"))
    assert second.status == "current" and second.file_path == doc(world, 14).file_path != first.file_path
    assert doc(world, 14).published_version == "1.0"
    assert get(world, STAFF, "/governance/documents/14/download-file").content == world.files["a"]
    assert get(world, OWNER, "/governance/documents/14/download-file").content == world.files["b"]


def test_a_file_put_in_place_is_kept_in_its_versions_row_so_it_stays_readable_once_replaced(world):
    world.me["id"] = OWNER
    made = world.http.post("/governance/documents", json={"title": "Leave policy", "doc_type": "policy"}).json()     # a 1.0 row, no file
    did = made["id"]
    assert world.http.post(f"/governance/documents/{did}/upload-file", files={"file": ("a.docx", world.files["a"], DOCX)}).status_code == 200
    (row,) = rows(world, did)                                              # the first file of version 1.0 goes into its row
    assert (row.version_number, row.status, row.file_name, row.file_path) == ("1.0", "current", "a.docx", doc(world, did).file_path)
    d = doc(world, did)
    d.status = "published"
    world.db.commit()
    world.http.post(f"/governance/documents/{did}/upload-file", files={"file": ("b.docx", world.files["b"], DOCX)})
    assert get(world, STAFF, f"/governance/documents/{did}/download-file").content == world.files["a"]

    # a document from before: its row 1.0 never held the file it published; replacing the file puts the old one there first
    _published(world.db, 15, "Policy 15", file_name="a.docx", file_path=str(world.tmp / "a.docx"), file_type="docx", file_size=len(world.files["a"]))
    world.db.add(m.GovernanceDocumentVersion(document_id=15, version_number="1.0", change_type="major", title="Policy 15", status="current"))
    world.db.commit()
    world.me["id"] = OWNER
    assert world.http.post("/governance/documents/15/upload-file", files={"file": ("b.docx", world.files["b"], DOCX)}).status_code == 200
    assert rows(world, 15)[0].file_path == str(world.tmp / "a.docx")
    assert get(world, STAFF, "/governance/documents/15/download-file").content == world.files["a"]


# ── who may read the document at all, on these routes ────────────────────────
def _reads(w):
    r10, r11 = row_id(w, TEXT, "1.0"), row_id(w, TEXT, "1.1")
    d, word = f"/governance/documents/{TEXT}", f"/governance/documents/{WORD}"
    return [("get", d, {}), ("get", f"{d}/review-history", {}), ("get", f"{d}/audit-logs", {}), ("get", f"{word}/download-file", {}),
            ("get", f"{word}/export", {}), ("get", f"{d}/view-html", {}), ("get", f"{d}/annotations", {}),
            ("post", f"{d}/compare-with-document", {"json": {"target_document_id": WORD}}),
            ("get", f"{d}/change-chain", {}), ("get", f"{d}/change-chain/redline?from={r10}", {}), ("get", f"{d}/change-chain/export?format=csv", {}),
            ("get", f"/governance/versions/document/{TEXT}", {}), ("get", f"/governance/versions/document/{TEXT}/history", {}),
            ("get", f"/governance/versions/{r10}", {}), ("get", f"/governance/versions/compare/{r10}/{r11}", {})]


def test_someone_neither_allowed_nor_asked_is_refused_every_read_of_the_document(world):
    reads = _reads(world)
    world.me["id"] = OUTSIDER
    for method, path, kw in reads:
        r = getattr(world.http, method)(path, **kw)
        assert (r.status_code, r.json()["detail"]) == (403, P.READ_REFUSED), path
    for path in ("/governance/documents/hierarchy", "/governance/documents/reviews/upcoming"):
        r = world.http.get(path)
        assert (r.status_code, r.json()["detail"]) == (403, P.LIBRARY_REFUSED), path
    assert world.http.get("/governance/documents").json()["items"] == [] and world.http.get("/governance/documents/policies").json()["items"] == []
    assert world.http.get("/governance/documents/404").status_code == 404                     # not there: not found first
    world.me["id"] = EDITOR                                                 # someone allowed reads every one of them
    for method, path, kw in reads:
        assert getattr(world.http, method)(path, **kw).status_code == 200, path


def test_auditors_open_a_single_document_and_remark_on_it_but_get_no_library_lists(world):
    world.me["id"] = AUDITOR
    assert world.http.get(f"/governance/documents/{TEXT}").json()["content"] == V10
    assert world.http.get(f"/governance/documents/{DRAFT}").status_code == 200
    assert world.http.get(f"/governance/documents/{WORD}/view-html").status_code == 200
    assert world.http.post(f"/governance/documents/{TEXT}/annotations", json={"comment": "Evidence checked."}).status_code == 201
    assert world.http.get(f"/governance/documents/{TEXT}/annotations").json()["total"] == 1
    assert world.http.get("/governance/documents/hierarchy").status_code == 403
    assert world.http.get("/governance/documents").json()["items"] == []


def test_a_draft_never_published_is_hidden_from_staff_but_not_from_its_owner_editors_and_administrators(world):
    r = get(world, STAFF, f"/governance/documents/{DRAFT}")
    assert (r.status_code, r.json()["detail"]) == (403, P.DRAFT_REFUSED)       # they read the library: it is a draft, not a right they lack
    assert world.http.get(f"/governance/documents/{DRAFT}/view-html").status_code == 403
    for who in (OWNER, EDITOR, ADMIN):
        assert get(world, who, f"/governance/documents/{DRAFT}").json()["content"] == "# Access standard\n", who


def test_remarks_are_written_by_editors_auditors_and_the_people_taking_part_only(world):
    url = f"/governance/documents/{TEXT}/annotations"
    for who in (STAFF, ACKER):                                             # reading it, or asked to acknowledge it, is not enough
        world.me["id"] = who
        r = world.http.post(url, json={"comment": "A remark."})
        assert (r.status_code, r.json()["detail"]) == (403, P.ANNOTATE_REFUSED), who
    ids = []
    for who in (REVIEWER, PREPARER, OWNER, EDITOR, AUDITOR):
        world.me["id"] = who
        r = world.http.post(url, json={"comment": f"From {who}."})
        assert r.status_code == 201, who
        ids.append(r.json()["id"])
    world.me["id"] = STAFF
    assert world.http.put(f"{url}/{ids[0]}", json={"status": "resolved"}).status_code == 403
    assert world.http.delete(f"{url}/{ids[0]}").status_code == 403
    world.me["id"] = PREPARER                                              # someone taking part may resolve another's remark
    assert world.http.put(f"{url}/{ids[0]}", json={"status": "resolved"}).json()["status"] == "resolved"
    assert world.http.get(url).json()["total"] == 5


# ── the review findings of 2026-10-17 (fix round 5) ──────────────────────────
def _signature():
    im = Image.new("RGB", (300, 120), "white")
    ImageDraw.Draw(im).line([(10, 100), (80, 20), (150, 90), (280, 30)], fill="#0f172a", width=3)
    out = BytesIO()
    im.save(out, "PNG")
    return "data:image/png;base64," + base64.b64encode(out.getvalue()).decode()


@pytest.mark.parametrize("status", ["expired", "exception_applied", "review_required", "draft", "approved"])   # archived: below
def test_staff_read_the_version_published_last_whatever_the_status(world, status):
    d = doc(world, TEXT)
    d.status = status                                                      # withdrawn, flagged for review, rejected, approved
    world.db.commit()
    for who in (STAFF, ACKER, AUDITOR):
        got = get(world, who, f"/governance/documents/{TEXT}").json()
        assert (got["content"], got["current_version"]) == (V10, "1.0"), (status, who)
        assert "every quarter" not in world.http.get(f"/governance/documents/{TEXT}/view-html").json()["html"], (status, who)
    assert {i["id"]: i["content"] for i in get(world, STAFF, "/governance/documents").json()["items"]}[TEXT] == V10
    assert get(world, OWNER, f"/governance/documents/{TEXT}").json()["content"] == V11
    if status == "review_required":                                       # still in force: it can be acknowledged
        world.me["id"] = ACKER
        r = world.http.post(f"/governance/documents/{TEXT}/lifecycle/acknowledge", json={"version": "1.0"})
        assert r.status_code == 200, r.text


def ids(w, who, path):
    return {i["id"] for i in get(w, who, path).json()["items"]}


def test_an_archived_document_is_not_in_force_staff_neither_list_nor_open_it_and_its_editors_still_do(world):
    world.me["id"] = EDITOR
    assert world.http.post("/governance/documents/bulk-archive", json={"document_ids": [TEXT]}).status_code == 200
    d = doc(world, TEXT)
    assert (d.status, d.published_version, d.last_published_version) == ("archived", None, "1.0")
    assert world.db.query(m.DocumentObligation).filter_by(document_id=TEXT, user_id=ACKER).one().status == "cancelled"
    for who, detail in ((STAFF, P.DRAFT_REFUSED), (ACKER, P.READ_REFUSED)):     # the library right | asked to acknowledge, cancelled
        for path in ("", "/view-html", "/change-chain", "/download-file"):
            r = get(world, who, f"/governance/documents/{TEXT}{path}")
            assert (r.status_code, r.json()["detail"]) == (403, detail), (who, path)
        assert TEXT not in ids(world, who, "/governance/documents") | ids(world, who, "/governance/documents?status=archived"), who
    assert TEXT not in {n["id"] for n in get(world, STAFF, "/governance/documents/hierarchy").json()}
    for who, text in ((EDITOR, V11), (OWNER, V11), (ADMIN, V11), (REVIEWER, V11), (AUDITOR, V10)):   # never staff's 1.1 to an auditor
        r = get(world, who, f"/governance/documents/{TEXT}")
        assert (r.status_code, r.json()["content"], r.json()["status"]) == (200, text, "archived"), who
    for who in (EDITOR, OWNER, ADMIN, REVIEWER):                           # the library's Archived list
        assert TEXT in ids(world, who, "/governance/documents?status=archived"), who


def test_pickers_never_offer_an_archived_document(world):
    doc(world, TEXT).status = "archived"
    world.db.commit()
    for who in (OUTSIDER, STAFF, EDITOR, OWNER, ADMIN):
        offered = ids(world, who, "/governance/documents?brief=1")
        assert TEXT not in offered and WORD in offered, who
    assert TEXT in ids(world, EDITOR, "/governance/documents")              # the library still lists it for its editors


def test_a_restored_document_stays_hidden_from_staff_until_it_is_published_again(world):
    world.me["id"] = EDITOR
    assert world.http.post("/governance/documents/bulk-archive", json={"document_ids": [TEXT]}).status_code == 200
    assert world.http.post(f"/governance/documents/{TEXT}/restore").status_code == 200
    d = doc(world, TEXT)
    assert (d.status, d.published_version, d.last_published_version) == ("draft", None, "1.0")      # nothing in force
    r = get(world, STAFF, f"/governance/documents/{TEXT}")
    assert (r.status_code, r.json()["detail"]) == (403, P.DRAFT_REFUSED)
    assert TEXT not in ids(world, STAFF, "/governance/documents") | ids(world, STAFF, "/governance/documents?brief=1")
    assert get(world, EDITOR, f"/governance/documents/{TEXT}").json()["content"] == V11
    world.me["id"] = ADMIN
    assert world.http.put(f"/governance/documents/{TEXT}/status", json={"status": "published"}).status_code == 200
    got = get(world, STAFF, f"/governance/documents/{TEXT}").json()
    assert (got["content"], got["current_version"], got["status"]) == (V11, "1.1", "published")
    assert TEXT in ids(world, STAFF, "/governance/documents") & ids(world, STAFF, "/governance/documents?brief=1")


def test_the_dashboard_summary_counts_only_the_documents_the_person_may_read(world):
    def summary(who):
        return get(world, who, "/governance/dashboard/summary").json()
    assert (summary(STAFF)["total_documents"], summary(EDITOR)["total_documents"]) == (2, 3)    # never the draft never published
    doc(world, TEXT).status = "archived"
    world.db.commit()
    assert summary(STAFF)["by_status"] == {"pending_review": 1}
    assert summary(EDITOR)["by_status"] == {"archived": 1, "draft": 1, "pending_review": 1}


def test_drafting_from_a_parent_needs_reading_it_and_drafts_from_the_text_the_person_reads(world, monkeypatch):
    import grc.db
    from grc.tasks import ai_drafting

    class Master:                                                          # never the real master database
        def query(self, *_a):
            return self

        def filter(self, *_a):
            return self

        def first(self):
            return ("cfsb",)

        def close(self):
            pass

    sent = []
    monkeypatch.setattr(grc.db, "MasterSession", Master)
    monkeypatch.setattr(ai_drafting, "create_job", lambda **_kw: "job-1")
    monkeypatch.setattr(ai_drafting, "dispatch_in_thread", lambda _slug, _job, payload: sent.append(payload))
    body = {"doc_type": "procedure", "title": "Access procedure"}
    for who, parent, detail in ((OUTSIDER, TEXT, P.READ_REFUSED), (STAFF, DRAFT, P.DRAFT_REFUSED)):
        world.me["id"] = who
        r = world.http.post("/governance/documents/ai-draft", json={**body, "parent_document_id": parent})
        assert (r.status_code, r.json()["detail"]) == (403, detail), who
    assert sent == []
    for who, text in ((STAFF, V10), (OWNER, V11)):
        world.me["id"] = who
        assert world.http.post("/governance/documents/ai-draft", json={**body, "parent_document_id": TEXT}).status_code == 200, who
        assert sent[-1]["parent_content"] == text, who


def test_the_brief_list_for_pickers_names_the_documents_in_force_to_anyone_and_nothing_more(world):
    keys = {"id", "title", "doc_type", "status", "current_version", "document_code", "parent_document_id", "tags"}   # no text
    for who, ids in ((OUTSIDER, [TEXT, WORD]), (STAFF, [TEXT, WORD]), (EDITOR, [TEXT, DRAFT, WORD]), (OWNER, [TEXT, DRAFT, WORD])):
        items = get(world, who, "/governance/documents?brief=1&sort_by=id&sort_order=asc").json()["items"]
        assert [i["id"] for i in items] == ids and all(set(i) == keys for i in items), who
        assert {i["id"]: i["current_version"] for i in items}[TEXT] == ("1.1" if who in (EDITOR, OWNER) else "1.0"), who
    assert get(world, OUTSIDER, "/governance/documents").json()["items"] == []        # the full list stays what they may read


def test_a_request_for_an_earlier_version_can_still_be_acknowledged(world):
    d = doc(world, TEXT)
    d.status = "published"                                                 # 1.1 published by a path that did not move requests up
    world.db.commit()
    world.me["id"] = ACKER
    url = f"/governance/documents/{TEXT}/lifecycle/acknowledge"
    assert world.http.post(url, json={"version": "0.9"}).status_code == 409   # neither what they read nor what they were asked
    assert world.http.post(url, json={"version": "1.0"}).status_code == 200   # the version their request names
    ob = world.db.query(m.DocumentObligation).filter_by(document_id=TEXT, user_id=ACKER).one()
    assert (ob.status, ob.document_version) == ("done", "1.1")


def test_a_new_file_keeps_the_file_it_replaces_in_the_row_of_that_version(world):
    a = str(world.tmp / "a.docx")
    for did in (17, 18):
        _published(world.db, did, f"Policy {did}", file_name="a.docx", file_path=a, file_type="docx", file_size=len(world.files["a"]))
    world.db.add(m.GovernanceDocumentVersion(document_id=17, version_number="1.0", change_type="major", title="Policy 17", status="current"))
    world.db.commit()                                                      # 17: a row 1.0 that never held its file
    service.write_file_version(world.db, doc(world, 17), world.db.get(m.GRCUser, OWNER), world.files["b"], change_type="minor",
                               summary="Edited", reason=None, done="Edited in the AI Studio")
    world.db.commit()
    assert rows(world, 17)[0].file_path == a
    assert get(world, STAFF, "/governance/documents/17/download-file").content == world.files["a"]
    world.me["id"] = OWNER                                                 # 18: a new number of a file document holds its file
    assert world.http.post("/governance/versions/document/18", json={"change_type": "minor"}).status_code == 201
    assert [(v.version_number, v.file_path) for v in rows(world, 18)] == [("1.0", a), ("1.1", a)]


def test_an_acknowledgement_signed_during_a_review_is_not_reported_as_edited_since(world):
    world.me["id"] = ACKER
    r = world.http.post(f"/governance/documents/{TEXT}/lifecycle/acknowledge", json={"version": "1.0", "signature_image": _signature()})
    assert r.status_code == 200, r.text
    esign_id = world.db.query(m.DocumentObligation).filter_by(document_id=TEXT, user_id=ACKER).one().esign_id
    seal = world.http.get(f"/governance/esignatures/{esign_id}").json()
    assert (seal["document_version"], seal["document_changed_since"], seal["valid"]) == ("1.0", False, True)
    d = doc(world, TEXT)
    d.status = "published"                                                 # 1.1 is in force now: the text acknowledged is not
    world.db.commit()
    assert world.http.get(f"/governance/esignatures/{esign_id}").json()["document_changed_since"] is True


def test_a_staff_session_exports_the_number_it_was_read_at(world):
    staff = service.create_session(world.db, world.db.get(m.GRCUser, STAFF), kind="library", document_id=TEXT)
    owner = service.create_session(world.db, world.db.get(m.GRCUser, OWNER), kind="library", document_id=TEXT)
    assert (staff.content, service.export_settings(world.db, staff)["version"]) == (V10, "1.0")
    assert (owner.content, service.export_settings(world.db, owner)["version"]) == (V11, "1.1")
