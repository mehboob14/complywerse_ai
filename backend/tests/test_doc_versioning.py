"""Version numbers of governance documents: a file starts at the version it states (its revision table, its document control,
a header or footer, else its file name, else 1.0), and each new version is a major or a minor change as Comply Agent decides
from what changed - by a rule when it cannot answer, and never holding a save up for longer than a few seconds."""
import io
import json
import time
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.governance import doc_version as V
from grc.modules.governance.routers import documents
from grc.routers.auth_router import require_auth

from test_studio_flow import reply, world          # noqa: F401  (world is a fixture)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
RULES = "# Access Policy\n\n## Purpose\n\nAccess is granted on need to know.\n\n## Rules\n\nAccounts are reviewed every year by the system owner.\n"


def word(path, *, table=None, paragraphs=(), footer=None):
    d = Document()
    d.add_heading("Access Policy", level=1)
    for text in paragraphs:
        d.add_paragraph(text)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, text in enumerate(row):
                t.cell(r, c).text = text
    d.add_paragraph("Access is granted on need to know.")
    if footer:
        d.sections[0].footer.paragraphs[0].text = footer
    d.save(path)
    return str(path)


# ── the version a file states ────────────────────────────────────────────────
def test_a_revision_table_gives_its_latest_version_whatever_the_cover_or_the_name_say(tmp_path):
    path = word(tmp_path / "a.docx", paragraphs=["Revision: 1.0"], table=[["Version", "Date", "Description"], ["1.0", "03.2012", "First issue"],
                                                                          ["2.1", "05.2014", "Annual review"], ["2.0", "04.2013", "New scope"]])
    assert V.detect(path, "Access Policy v9.9.docx") == ("2.1", "revision table")


def test_without_a_revision_table_the_document_states_its_version_on_the_cover_in_a_table_or_in_a_footer(tmp_path):
    assert V.detect(word(tmp_path / "a.docx", paragraphs=["MCB Bank Ltd.", "Revision:\t4.0 (Q1-2013)"]), "a.docx") == ("4.0", "document")
    assert V.detect(word(tmp_path / "b.docx", table=[["Document Code", "POL-7"], ["Version No.", "3.1"]]), "b.docx") == ("3.1", "document")
    assert V.detect(word(tmp_path / "c.docx", footer="Version 2"), "c.docx") == ("2.0", "document")
    assert V.detect(word(tmp_path / "d.docx", footer="POL-7 · v2.3"), "d.docx") == ("2.3", "document")
    # a year is not a version, and a sentence that mentions a revision is not where a document states its own
    assert V.detect(word(tmp_path / "e.docx", paragraphs=["Version 2013 of the standard applies.", "This policy follows the last revision 3 of the law."]), "e.docx") == ("1.0", "default")


def test_a_file_that_says_nothing_is_numbered_by_its_name_else_it_starts_at_one(tmp_path):
    path = word(tmp_path / "plain.docx")
    assert V.detect(path, "Information Technology Policy 2013 - v1 1 - Approved by BoD.docx") == ("1.1", "file name")
    assert V.detect(path, "IT Security Procedure Manual 4.0 - Q1-2013 - Approved.docx") == ("4.0", "file name")
    assert V.detect(path, "policy.docx") == ("1.0", "default")
    assert V.detect(None, "notes v2.md") == ("2.0", "file name")


def test_a_pdfs_revision_table_is_read_although_its_grid_shifts_a_cell_and_has_empty_rows():
    pdf_table = [["", "", "", "", ""], ["", "Version", "", "Date", "Document Status"], ["", "", "", "", ""], ["1.0", "", "", "3rd Dec 2013", "IT Policy"]]
    assert V._from_tables([pdf_table]) == ("1.0", None)


# ── major or minor ───────────────────────────────────────────────────────────
def test_without_comply_agent_a_rule_decides():
    added = V.decide(RULES, RULES + "\n## Exceptions\n\nThe CISO may grant one.\n")
    assert (added["change"], added["by"]) == ("major", "rule") and "section" in added["reason"]
    reworded = V.decide(RULES, RULES.replace("every year", "every twelve months"))
    assert (reworded["change"], reworded["by"]) == ("minor", "rule")
    rewritten = V.decide(RULES, RULES.replace("Accounts are reviewed every year by the system owner.", "Privileged accounts are reviewed monthly by security."))
    assert rewritten["change"] == "major"                                                      # more than a fifth of the words
    assert V.decide(RULES, RULES)["reason"] == "Only the file changed, not its text."
    assert [V.next_version("1.9", "minor"), V.next_version("1.9", "major"), V.next_version("2.0.1", "minor")] == ["1.10", "2.0", "2.1"]


def test_comply_agent_decides_from_what_changed_and_an_answer_it_cannot_give_in_time_or_in_shape_falls_back_to_the_rule(world, monkeypatch):
    world.respond(lambda body: reply(json.dumps({"change": "major", "reason": "Reviews become monthly."})))
    out = V.decide(RULES, RULES.replace("every year", "every month"), title="Access Policy")
    assert out == {"change": "major", "reason": "Reviews become monthly.", "by": "ai"}
    sent = world.wire.requests[-1].body
    assert "every month" in sent["messages"][0]["content"] and "need to know" not in sent["messages"][0]["content"]   # only what changed
    assert V.summary_of("Content edited", out) == "Content edited. Major change: Reviews become monthly."

    world.respond(lambda body: reply(json.dumps({"change": "huge", "reason": "?"})))
    assert V.decide(RULES, RULES.replace("every year", "every month"))["by"] == "rule"
    monkeypatch.setattr(V, "AI_SECONDS", 0.3)
    world.respond(lambda body: (time.sleep(2), reply(json.dumps({"change": "major", "reason": "late"})))[1])
    started = time.monotonic()
    late = V.decide(RULES, RULES.replace("every year", "every month"))
    assert late["by"] == "rule" and time.monotonic() - started < 1.5                           # a save never waits for a slow answer


# ── what a document's revision history says about a version ──────────────────
def test_the_revision_history_leaves_out_the_platforms_bookkeeping():
    d = V.described
    # the line that was in a document's Revision History: which module saved it and who picked minor is nothing a reader needs
    assert d("Edited in the Comply AI Studio. Minor change: Chosen by Admin when saving it to the library.") == ""
    assert d(V.summary_of("Edited in the Comply AI Studio", V.chosen("major", "J. Smith"))) == ""         # (a name may hold a full stop)
    assert d("Edited in the AI Studio: Updated the password rules. Minor change: Chosen by Admin when saving it to the library.") == "Updated the password rules."
    for plain in ("minor update", "signoff update", "Original document, before it was changed", "Baseline snapshot (before first edit)"):
        assert d(plain) == ""
    # a first version says it is the first, and whether it was uploaded; where its number came from is the Version column's business
    assert d(V.first_summary("1.0", "")) == d(V.first_summary("4.0", "revision table")) == d(V.first_summary("2.1", "file name")) == "Initial version (uploaded file)"
    assert d("Initial version") == "Initial version"
    # a file put in place: that it was, not how its version was read
    assert d("File updated: Policy v2.docx. The file states version 2.0 (from its file name).") == "File updated: Policy v2.docx."
    assert d("File updated: Policy.docx. Minor change: the new file's text could not be read to compare it with the last version.") == "File updated: Policy.docx."
    assert d("Updated the password rules · Proposed by Rana Reviewer, accepted by Olivia Owner") == "Updated the password rules"


def test_the_revision_history_keeps_what_the_change_was_without_its_label_or_its_note():
    d = V.described
    ai = {"change": "major", "reason": "A new section on retention was added.", "by": "ai"}
    assert d(V.summary_of("Edited in the Comply AI Studio", ai)) == "A new section on retention was added."          # the Change type column says "major"
    rule = {"change": "minor", "reason": "A limited change to the wording.", "by": "rule"}
    assert d(V.summary_of("Content edited", rule)) == "A limited change to the wording."                             # and not that a rule decided
    assert d(V.summary_of("Gap fix (clause replaced) for ISO/IEC 27001:2022 A.5.2", ai)) == "Gap fix (clause replaced) for ISO/IEC 27001:2022 A.5.2. A new section on retention was added."
    assert d("Took back 2 changes made since version 1.3: they were not agreed") == "Took back 2 changes made since version 1.3: they were not agreed"
    assert d("Temporary staff added to the scope") == "Temporary staff added to the scope"                            # a person's own summary is left as it is
    assert d("Vendor chosen by the board") == "Vendor chosen by the board" and d(None) == "" and d("") == ""


def test_a_person_who_gave_a_reason_has_it_as_the_summary_whatever_the_platform_wrote():
    said = V.summary_of("Edited in the Comply AI Studio: Added a retention schedule", V.chosen("minor", "Admin"))
    assert V.described(said, "Added a retention schedule") == "Added a retention schedule"
    assert V.described(V.summary_of("Edited in the Comply AI Studio", {"change": "major", "reason": "A section was added.", "by": "ai"}),
                       "  Retention\nperiods stated\n ") == "Retention periods stated"                          # one line, however it was typed


# ── the library ──────────────────────────────────────────────────────────────
@pytest.fixture
def library(tmp_path, monkeypatch):
    monkeypatch.setattr(documents, "GOVERNANCE_UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr(documents, "_resolve_tenant_slug_for_autoparse", lambda tenant_id: None)      # never the real master database
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    db.add(m.Tenant(id=1, name="CFSB", slug="cfsb"))
    db.add(m.GRCUser(id=1, username="olivia", email="o@cfsb.test", display_name="Olivia"))
    db.commit()
    app = FastAPI()
    app.include_router(documents.router)
    app.dependency_overrides[m.get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: db.get(m.GRCUser, 1)
    return SimpleNamespace(db=db, http=TestClient(app), tmp=tmp_path)


def _versions(db, doc_id):
    rows = db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc_id).order_by(m.GovernanceDocumentVersion.id).all()
    return [(v.version_number, v.change_type, v.status) for v in rows]


def test_a_file_uploaded_to_the_library_starts_at_its_own_version_and_a_new_file_moves_it_on(library):
    path = word(library.tmp / "manual.docx", paragraphs=["Revision: 4.0 (Q1-2013)"])
    with open(path, "rb") as f:
        made = library.http.post("/documents/upload-with-file", files={"file": ("manual.docx", f.read(), DOCX)},
                                 data={"title": "IT Security Procedure Manual", "doc_type": "procedure"})
    assert made.status_code == 201, made.text
    doc = library.db.get(m.GovernanceDocument, made.json()["document"]["id"])
    assert doc.current_version == "4.0" and _versions(library.db, doc.id) == [("4.0", "major", "current")]
    assert library.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id).one().change_summary == "Initial version: the uploaded file, version 4.0 as the document states it"

    # a new file that states a later version keeps the organisation's own numbering
    body = "Accounts are reviewed every year by the system owner, and access that is no longer needed is removed at once."
    later = word(library.tmp / "manual5.docx", paragraphs=["Revision: 5.0 (Q1-2014)", body])
    with open(later, "rb") as f:
        library.http.post(f"/documents/{doc.id}/upload-file", files={"file": ("manual5.docx", f.read(), DOCX)})
    library.db.refresh(doc)
    assert doc.current_version == "5.0" and _versions(library.db, doc.id)[-1] == ("5.0", "major", "current")

    # one that states no later version is numbered by what changed (here by the rule: no AI answer in this test)
    same = word(library.tmp / "manual5b.docx", paragraphs=["Revision: 5.0 (Q1-2014)", body.replace("every year", "every quarter")])
    with open(same, "rb") as f:
        library.http.post(f"/documents/{doc.id}/upload-file", files={"file": ("manual5b.docx", f.read(), DOCX)})
    library.db.refresh(doc)
    # its own version, the one before kept with its file (a draft too: every new file makes its own version)
    assert doc.current_version == "5.1" and _versions(library.db, doc.id) == [("4.0", "major", "superseded"), ("5.0", "major", "superseded"),
                                                                             ("5.1", "minor", "current")]
    five = library.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id, version_number="5.0").one()
    assert five.file_name == "manual5.docx" and "every year" in "\n".join(p.text for p in Document(five.file_path).paragraphs)


def test_a_new_file_for_a_published_document_sends_it_back_to_review_in_a_new_round(library):
    from datetime import datetime

    first, later = word(library.tmp / "a.docx", paragraphs=["Revision: 1.0"]), word(library.tmp / "b.docx", paragraphs=["Revision: 2.0"])
    library.db.add(m.GovernanceDocument(id=8, tenant_id=1, title="Access Policy", doc_type="policy", status="published", current_version="1.0",
                                        file_name="a.docx", file_path=first, file_type="docx", owner_id=1, content=None))
    library.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=8, role_type="reviewer", target_type="user", target_id=1, added_by=1))
    library.db.commit()
    doc = library.db.get(m.GovernanceDocument, 8)
    doc.signoff_round_at = datetime(2026, 1, 1)
    library.db.commit()
    with open(later, "rb") as f:
        out = library.http.post("/documents/8/upload-file", files={"file": ("b.docx", f.read(), DOCX)})
    assert out.status_code == 200, out.text
    library.db.refresh(doc)
    assert (doc.current_version, doc.status) == ("2.0", "pending_review") and doc.signoff_round_at > datetime(2026, 1, 1)
    assert library.db.query(m.DocumentObligation).filter_by(document_id=8, kind="review", user_id=1, status="pending").count() == 1
    assert library.db.query(m.DocumentAuditLog).filter_by(document_id=8, action="file_uploaded").one().action_details.endswith("; sent back to its reviewers")


def test_editing_a_documents_text_in_the_library_is_numbered_by_the_change(library):
    library.db.add(m.GovernanceDocument(id=7, tenant_id=1, title="Access Policy", doc_type="policy", status="draft", current_version="1.4", content=RULES, owner_id=1))
    library.db.commit()
    library.http.put("/documents/7", json={"content": RULES.replace("every year", "every twelve months"), "change_reason": "wording"})
    library.http.put("/documents/7", json={"content": RULES + "\n## Exceptions\n\nThe CISO may grant one.\n"})
    assert _versions(library.db, 7) == [("1.4", "baseline", "superseded"), ("1.5", "minor", "superseded"), ("2.0", "major", "current")]
    assert library.db.get(m.GovernanceDocument, 7).current_version == "2.0"


def test_every_change_makes_its_own_version_from_the_current_one_and_the_earlier_ones_stay_as_they_were(library):
    """Before and after a publication, each change is the next number after the document's current version (+0.1 minor, the
    next whole number major), and the version before keeps its number, text and why (the user, 2026-10-21: "1.2 gone??")."""
    library.db.add(m.GovernanceDocument(id=9, tenant_id=1, title="Access Policy", doc_type="policy", status="draft", current_version="1.0",
                                        content=RULES, owner_id=1))
    library.db.commit()

    def edit(text, reason):
        out = library.http.put("/documents/9", json={"content": text, "change_reason": reason})
        assert out.status_code == 200, out.text
        return library.db.get(m.GovernanceDocument, 9).current_version

    t1 = RULES.replace("every year", "every month")
    t2 = t1.replace("system owner", "data owner")
    t3 = t2 + "\n## Exceptions\n\nThe CISO may grant one.\n"
    assert [edit(t1, "monthly"), edit(t2, "owner"), edit(t3, "exceptions")] == ["1.1", "1.2", "2.0"]
    rows = library.db.query(m.GovernanceDocumentVersion).filter_by(document_id=9).order_by(m.GovernanceDocumentVersion.id).all()
    assert [(r.version_number, r.status, r.change_reason) for r in rows] == [("1.0", "superseded", None), ("1.1", "superseded", "monthly"),
                                                                             ("1.2", "superseded", "owner"), ("2.0", "current", "exceptions")]
    assert [r.content for r in rows] == [RULES, t1, t2, t3]                    # each keeps its own text

    doc = library.db.get(m.GovernanceDocument, 9)
    doc.status = "published"                                                   # 2.0 is in force
    library.db.commit()
    t4 = t3.replace("every month", "every week")
    assert [edit(t4, "weekly"), edit(t4 + "\n## Records\n\nReviews are kept for five years.\n", "records")] == ["2.1", "3.0"]
    assert [v for v, _, _ in _versions(library.db, 9)] == ["1.0", "1.1", "1.2", "2.0", "2.1", "3.0"]
    assert library.db.get(m.GovernanceDocument, 9).last_published_version == "2.0"
