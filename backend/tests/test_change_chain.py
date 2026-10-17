"""The change chain: who changed what, from what to what, where each change came from, taking a change back, and the catalog.

The core is tested on texts; the chain on a library document edited in different ways (a person, a sign-off stamp, a restored
version, the AI Studio); the studio's part (provenance recorded when a change is accepted, taking an accepted change back,
the tracked changes of a session) through the real studio routes with the AI as a mocked wire."""
import io
import json
import os
import zipfile
from datetime import datetime, timedelta

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

import grc.models as m
from grc.modules.governance import change_chain as C
from grc.modules.governance import change_chain_export as X
from grc.modules.governance.routers import change_chain as chain_routes
from grc.modules.governance.studio import docx_io, exporters
from grc.routers.auth_router import require_auth

from _claude_mock import reply
from test_studio_files import lib, make_library_doc                 # noqa: F401  (lib is a fixture: the library's own folder for files)
from test_studio_flow import ENTITY, answer, change, world          # noqa: F401  (world is a fixture)

POLICY = """# Access Control Policy

## 1. Purpose

This policy protects information assets.

## 2. Scope

All employees and contractors.

## 3. Access Control

Access to systems shall be reviewed every year.

Privileged access shall be approved by the CISO.

## 4. Approval Signoff

| Role | Name | Designation | Signature / Date |
| --- | --- | --- | --- |
| Reviewed by | | | |
"""


# ── cutting a difference into changes ────────────────────────────────────────
def test_a_difference_is_cut_into_changes_with_their_section_and_words():
    new = (POLICY.replace("every year", "every quarter")
           .replace("All employees and contractors.", "All employees and contractors.\n\nThird parties with access are in scope.")
           .replace("Privileged access shall be approved by the CISO.\n\n", ""))
    out = C.compare(POLICY, new)
    by = {h["status"]: h for h in out["hunks"]}
    assert [h["status"] for h in out["hunks"]] == ["inserted", "modified", "deleted"]
    assert by["inserted"]["section"] == "2. Scope" and by["inserted"]["after_text"] == "Third parties with access are in scope."
    assert by["modified"]["section"] == "3. Access Control" and by["modified"]["before_text"].endswith("every year.") and by["modified"]["after_text"].endswith("every quarter.")
    assert by["modified"]["words_added"] == 1 and by["modified"]["words_removed"] == 1
    assert by["deleted"]["before_text"] == "Privileged access shall be approved by the CISO." and by["deleted"]["words_removed"] == 8
    assert out["stats"] == {"inserted": 1, "deleted": 1, "modified": 1, "words_added": 8, "words_removed": 9}
    assert [r["hunk"] for r in out["rows"] if r["status"] != "equal"] == [0, 1, 2]
    assert all(r["hunk"] is None for r in out["rows"] if r["status"] == "equal")


def test_a_run_of_changed_blocks_is_split_at_each_heading_and_the_title_is_not_part_of_the_path():
    new = POLICY.replace("This policy protects information assets.", "Why.").replace("All employees and contractors.", "Everyone.")
    hunks = C.hunks_of(POLICY, new)
    assert [h["section"] for h in hunks] == ["1. Purpose", "2. Scope"]
    whole = C.hunks_of("", POLICY)                      # a first draft: one change per section, not one change for the document
    assert [h["section"] for h in whole] == ["", "1. Purpose", "2. Scope", "3. Access Control", "4. Approval Signoff"]
    assert all(h["status"] == "inserted" for h in whole) and whole[3]["blocks"] == 3
    paths = C.section_paths(__import__("grc.modules.governance.studio.blocks", fromlist=["x"]).parse("# T\n\n## A\n\n### A.1\n\ntext\n\n## B\n\nmore\n"))
    assert paths == ["", "A", "A › A.1", "A › A.1", "B", "B"]


def test_a_change_is_taken_back_alone_and_a_removed_part_goes_back_where_it_stood():
    new = (POLICY.replace("every year", "every quarter")
           .replace("All employees and contractors.", "All employees and contractors.\n\nThird parties with access are in scope.")
           .replace("Privileged access shall be approved by the CISO.\n\n", ""))
    hunks = C.hunks_of(POLICY, new)
    ids = {h["status"]: h["id"] for h in hunks}
    only_modified = C.revert(POLICY, new, [ids["modified"]])
    assert "every year" in only_modified and "Third parties with access" in only_modified and "approved by the CISO" not in only_modified
    back = C.revert(POLICY, new, [ids["deleted"]])
    assert back.index("Privileged access shall be approved by the CISO.") > back.index("every quarter")          # after the paragraph before it
    assert "Third parties" in back and "every quarter" in back
    gone = C.revert(POLICY, new, [ids["inserted"]])
    assert "Third parties" not in gone and "every quarter" in gone
    everything = C.revert(POLICY, new, ids.values())
    assert C.content_hash(everything) == C.content_hash(POLICY)
    with pytest.raises(ValueError):
        C.revert(POLICY, new, [99])
    with pytest.raises(ValueError):
        C.revert(POLICY, new, [])


def test_a_change_to_the_outline_says_how_it_changed_when_no_word_did():
    from grc.modules.governance.studio import blocks as B

    base = "# T\n\n## Scope\n\nText.\n\n## Rules\n\n- one\n- two\n\nSome words.\n\n1. first\n2. second\n3. third\n"
    cur = "# T\n\n### Scope\n\nText.\n\n## Rules\n\n- one\n  - two\n\nSome words.\n\n1. first\n2. second\n4. third\n"
    out = C.compare(base, cur)
    # a typed list number is not a change a reader sees (Markdown numbers a list from its first item; Word numbers lists itself)
    assert [h["structure"] for h in out["hunks"]] == [["Heading 2 → Heading 3"], ["List level 1 → 2"]]
    assert all(h["format_only"] and h["words_added"] == h["words_removed"] == 0 for h in out["hunks"])
    assert [r["structure"] for r in out["rows"] if r["status"] != "equal"] == [h["structure"] for h in out["hunks"]]
    assert B.locate(B.parse(base), B.parse(cur))[9] == 9                                 # an edit can still be placed on a renumbered item
    kinds = C.compare("Plain para.\n\n- a\n- b\n", "## Plain para.\n\n1. a\n2. b\n")
    assert [h["structure"] for h in kinds["hunks"]] == [["Paragraph → Heading 2"], ["Bullet list → numbered list"], ["Bullet list → numbered list"]]
    assert kinds["rows"][1]["was"] == {"kind": "list_item", "level": 0, "ordered": False}
    assert C.compare("Text.\n", "Text, more.\n")["hunks"][0]["structure"] == []
    stored = C.build_change_set(base, cur, [], default_kind="manual", by="Sara", by_id=1, at=None)
    assert stored["items"][0]["structure"] == ["Heading 2 → Heading 3"]


def test_the_fingerprint_of_a_document_ignores_white_space_and_notices_words():
    assert C.content_hash("# A\n\ntext\n") == C.content_hash("# A\r\n\r\n\r\ntext\r\n\r\n")
    assert C.content_hash("# A\n\ntext\n") != C.content_hash("# A\n\ntext!\n")


def test_a_change_is_attributed_to_the_latest_origin_that_wrote_it_and_says_when_it_was_edited_since():
    hunk = {"after_digests": ["a", "b"], "before_digests": ["x"]}
    first = C.make_origin("ai", by="Olivia", request="Write it", after=["a", "b"], sources=[{"title": "ISO 27001"}])
    later = C.make_origin("manual", by="Sara", after=["b"])
    out = C.attribute(hunk, [first, later])
    assert (out["origin"], out["by"], out["coverage"]) == ("manual", "Sara", 0.5)
    assert C.attribute(hunk, [first])["request"] == "Write it" and C.attribute(hunk, [first])["sources"] == [{"title": "ISO 27001"}]
    removal = {"after_digests": [], "before_digests": ["x"]}
    assert C.attribute(removal, [C.make_origin("ai", by="Olivia", before=["x"])])["by"] == "Olivia"
    assert C.attribute({"after_digests": ["zzz"], "before_digests": []}, [first]) is None


# ── a library document ───────────────────────────────────────────────────────
@pytest.fixture
def library(world, monkeypatch):
    """The studio world, plus the change chain routes over the same database."""
    monkeypatch.setattr(chain_routes, "get_user_tenants", lambda user, db: [1])
    app = FastAPI()
    app.include_router(chain_routes.router, prefix="/governance")
    app.dependency_overrides[m.get_db] = lambda: world.db
    app.dependency_overrides[require_auth] = lambda: world.db.get(m.GRCUser, world.me["id"])
    world.lib = TestClient(app)
    world.chain = lambda doc_id=10: world.lib.get(f"/governance/documents/{doc_id}/change-chain").json()
    world.db.add(m.GRCUser(id=3, username="khalid", email="k@cfsb.test", display_name="Khalid Approver"))
    world.db.commit()
    return world


def _edit(db, doc_id, new_content, *, user_id, reason=None, change_type="minor", at=None):
    """A content edit the way the document page makes one (a version row, the document updated)."""
    from grc.modules.governance.routers.documents import _snapshot_content_version

    doc = db.get(m.GovernanceDocument, doc_id)
    old_content, old_title, old_version = doc.content, doc.title, doc.current_version or "1.0"
    doc.content = new_content
    number = _snapshot_content_version(db, doc, old_content=old_content, old_title=old_title, old_version=old_version,
                                       change_type=change_type, change_reason=reason, user_id=user_id)
    db.flush()
    row = db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc_id, version_number=number).one()
    if at:
        row.created_at = at
        for baseline in db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc_id, change_type="baseline").all():
            if baseline.created_at > at:                    # the snapshot of the text before the first edit is older than the edit
                baseline.created_at = at - timedelta(seconds=1)
    db.commit()
    return row


def _audit(db, doc_id, user_id, action, at, *, status=None, old=None, details=""):
    db.add(m.DocumentAuditLog(document_id=doc_id, tenant_id=1, action=action, action_details=details, field_changed="status" if status else None,
                              old_value=old, new_value=status, performed_by=user_id, performed_at=at))
    db.commit()


def test_the_chain_lists_every_version_with_who_changed_what_the_stage_and_the_approvals_between(library):
    db = library.db
    doc = db.get(m.GovernanceDocument, 10)
    doc.content = POLICY
    doc.status = "draft"
    db.commit()
    t0 = datetime(2026, 10, 1, 9, 0)
    _edit(db, 10, POLICY.replace("every year", "every six months"), user_id=1, reason="Review frequency", at=t0)
    _audit(db, 10, 1, "sent_for_review", t0 + timedelta(hours=1), status="pending_review", old="draft")
    library.me["id"] = 2
    _edit(db, 10, db.get(m.GovernanceDocument, 10).content.replace("approved by the CISO", "approved by the CISO and logged"), user_id=2,
          reason="Evidence", at=t0 + timedelta(hours=2))
    _audit(db, 10, 2, "signed", t0 + timedelta(hours=3), details="Signed as Reviewed by → pending_approval")
    stamp = db.get(m.GovernanceDocument, 10).content.replace("| Reviewed by | | | |", "| Reviewed by | Rana | | 2026-10-01 · E-signed ab12cd34 |")
    _edit(db, 10, stamp, user_id=3, reason="Reviewed by: Rana", change_type="signoff", at=t0 + timedelta(hours=3, minutes=1))
    _audit(db, 10, 3, "signed", t0 + timedelta(hours=5), status="approved", old="pending_approval", details="Signed as Approved by")

    library.me["id"] = 1                                                # Sam may not read the document; Olivia administers
    c = library.chain()
    assert c["enabled"] is True
    # not published yet: the two edits and the sign-off stamp all went into one version (the user's rule, 2026-10-17)
    assert [v["number"] for v in c["versions"]] == ["1.0", "1.1"]
    baseline, pending = c["versions"]
    assert (baseline["origin"], baseline["changes"]) == ("initial", 0)
    assert pending["by"] == "Olivia Owner" and pending["changes"] == 3          # made by Olivia; three changes by three people
    who = {(i["by"], i["origin"]) for i in pending["items"]}
    assert who == {("Olivia Owner", "manual"), ("Sam Other", "manual"), ("Khalid Approver", "signoff")}     # each change keeps its author
    olivia = next(i for i in pending["items"] if i["by"] == "Olivia Owner")
    assert olivia["before_text"].endswith("every year.") and olivia["after_text"].endswith("every six months.") and olivia["section"] == "3. Access Control"
    assert pending["reason"].endswith("Review frequency")
    kinds = [e["kind"] for e in c["events"]]
    assert kinds == ["sent_for_review", "signed", "signed"] and c["events"][0]["actor"] == "Olivia Owner" and c["events"][2]["to"] == "approved"
    assert c["events"][1]["to"] == "pending_approval"                  # a signature that moved the document on says where to
    assert c["summary"]["versions"] == 2 and c["summary"]["changes"] == 3
    assert c["document"]["title"] == "Information Security Policy"


def test_a_stored_change_set_is_used_as_written_and_a_file_without_text_is_marked_not_comparable(library):
    db = library.db
    doc = db.get(m.GovernanceDocument, 10)
    doc.content = POLICY
    db.commit()
    row = _edit(db, 10, POLICY.replace("every year", "every quarter"), user_id=1, reason="x", change_type="studio")
    item = {"n": 1, "status": "modified", "section": "3. Access Control", "locked": False, "blocks": 1, "words_added": 1, "words_removed": 1, "format_only": False,
            "before_text": "old", "after_text": "new", "truncated": False, "before_digests": [], "after_digests": [], "origin": "ai", "by": "Olivia Owner",
            "by_id": 1, "at": "2026-10-02T10:00:00", "request": "Make it quarterly", "reason": "quarterly", "sources": [
                {"ref": "R1", "kind": "framework", "title": "ISO/IEC 27001", "locator": "A.5.1", "quote": "reviewed at planned intervals", "basis": "reference", "verified": True}],
            "ref": {"message_id": 4, "op_id": "c1"}, "coverage": 1.0}
    row.change_set = {"v": 1, "items": [item], "sources": item["sources"], "session_id": 7}
    db.commit()
    db.add(m.GovernanceDocumentVersion(document_id=10, version_number="1.2", change_type="minor", title="x", content=None, file_path="/nope/missing.docx",
                                       status="current", created_by=1, created_at=datetime.utcnow() + timedelta(hours=1)))
    db.commit()
    c = library.chain()
    stored = c["versions"][1]
    assert stored["origin"] == "ai" and stored["items"][0]["request"] == "Make it quarterly" and stored["session_id"] == 7
    assert c["sources"][0]["title"] == "ISO/IEC 27001" and c["sources"][0]["versions"] == ["1.1"]
    assert c["versions"][2]["comparable"] is False and c["versions"][2]["origin"] == "file"
    assert c["summary"]["ai_assisted"] == 1 and c["summary"]["sources"] == 1


def test_an_organisation_without_the_life_cycle_gets_nothing(library, monkeypatch):
    import grc.config as config

    monkeypatch.setattr(config, "DOCUMENT_LIFECYCLE_TENANTS", {"acme"})
    assert library.chain() == {"enabled": False}
    assert library.lib.get("/governance/documents/10/change-chain/export?format=csv").status_code == 404
    assert library.lib.post("/governance/documents/10/change-chain/revert", json={"from_version_id": 1, "hunks": [0]}).status_code == 404


# ── the document against an earlier version, and taking a change back ────────
def _three_versions(db):
    doc = db.get(m.GovernanceDocument, 10)
    doc.content, doc.status = POLICY, "published"
    db.commit()
    t0 = datetime(2026, 10, 1, 9, 0)
    _edit(db, 10, POLICY.replace("every year", "every quarter"), user_id=1, reason="Quarterly", at=t0)
    _edit(db, 10, db.get(m.GovernanceDocument, 10).content.replace("All employees and contractors.", "All staff, contractors and third parties."), user_id=2,
          reason="Scope", at=t0 + timedelta(hours=1))
    return db.query(m.GovernanceDocumentVersion).filter_by(document_id=10).order_by(m.GovernanceDocumentVersion.id).all()


def test_the_redline_against_an_earlier_version_says_who_made_each_change(library):
    versions = _three_versions(library.db)
    out = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[0].id}").json()
    by = {h["section"]: h for h in out["hunks"]}
    assert by["3. Access Control"]["by"]["by"] == "Olivia Owner" and by["3. Access Control"]["by"]["version"] == "1.1"
    assert by["2. Scope"]["by"]["by"] == "Sam Other" and by["2. Scope"]["by"]["reason"] == "Scope"
    assert out["can_revert"] is True and out["hash"] and [v["number"] for v in out["versions"]] == ["1.2", "1.1", "1.0"]
    since_last = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[1].id}").json()
    assert [h["section"] for h in since_last["hunks"]] == ["2. Scope"]
    assert library.lib.get("/governance/documents/10/change-chain/redline?from=9999").status_code == 404


def test_someone_who_reads_the_version_in_force_gets_the_chain_and_the_redline_up_to_it(library):
    db = library.db
    versions = _three_versions(db)                              # published at 1.0, then 1.1 and 1.2 written
    doc = db.get(m.GovernanceDocument, 10)
    doc.status = "pending_review"                               # the new version is on its way: 1.0 stays in force
    db.add(m.Role(id=12, tenant_id=1, name="Document readers"))
    db.add(m.Permission(id=12, name="governance:policies:view", resource="governance:policies", action="view"))
    db.add(m.RolePermission(role_id=12, permission_id=12))
    db.add(m.UserRole(user_id=2, role_id=12, tenant_id=1))       # Sam reads the library
    db.commit()
    assert doc.published_version == "1.0"
    library.me["id"] = 2
    c = library.chain()
    assert [v["number"] for v in c["versions"]] == ["1.0"] and (c["document"]["version"], c["document"]["in_force"]) == ("1.0", "1.0")
    red = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[0].id}").json()
    assert (red["to"]["id"], red["hunks"], [v["id"] for v in red["versions"]]) == (versions[0].id, [], [versions[0].id])
    assert library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[0].id}&to={versions[2].id}").status_code == 404
    library.me["id"] = 1
    assert [v["number"] for v in library.chain()["versions"]] == ["1.0", "1.1", "1.2"]


def test_the_changes_of_one_version_are_taken_back_from_the_document_as_it_is_now_unless_a_later_version_changed_them_again(library):
    db = library.db
    versions = _three_versions(db)                              # 1.0 baseline, 1.1 quarterly (Olivia), 1.2 scope (Sam)
    one = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[0].id}&to={versions[1].id}").json()
    assert one["to"]["number"] == "1.1" and [h["section"] for h in one["hunks"]] == ["3. Access Control"]
    hunk = one["hunks"][0]
    assert hunk["can_revert"] is True and hunk["current_id"] is not None and one["current_hash"] != one["hash"]       # the document is not version 1.1 any more
    back = library.lib.post("/governance/documents/10/change-chain/revert", json={"from_version_id": versions[0].id, "hunks": [hunk["current_id"]],
                                                                                   "hash": one["current_hash"]})
    assert back.status_code == 200
    content = db.get(m.GovernanceDocument, 10).content
    assert "every year" in content and "third parties" in content                   # Olivia's change is out, Sam's stays

    # the same paragraph is then changed again by someone else: Sam's own change can no longer be taken back alone
    _edit(db, 10, content.replace("All staff, contractors and third parties.", "Everyone with access."), user_id=1, reason="Simpler")
    latest = db.query(m.GovernanceDocumentVersion).filter_by(document_id=10).order_by(m.GovernanceDocumentVersion.id).all()
    sam = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[1].id}&to={versions[2].id}").json()
    assert sam["hunks"][0]["section"] == "2. Scope" and sam["hunks"][0]["can_revert"] is False and sam["hunks"][0]["current_id"] is None
    assert "changed again" in sam["why_not"] and len(latest) == 4             # the edit went into the version the take back made


def test_one_change_is_taken_back_as_a_new_version_and_a_published_document_goes_back_to_review(library):
    db = library.db
    versions = _three_versions(db)
    db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=10, role_type="reviewer", target_type="user", target_id=library.me["id"], added_by=library.me["id"]))
    db.commit()                                                                      # it has a reviewer, so it goes back to review
    red = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[0].id}").json()
    scope = next(h for h in red["hunks"] if h["section"] == "2. Scope")
    out = library.lib.post("/governance/documents/10/change-chain/revert", json={"from_version_id": versions[0].id, "hunks": [scope["id"]], "hash": red["hash"],
                                                                                 "reason": "Scope was not agreed"})
    assert out.status_code == 200 and out.json()["version"] == "1.3" and out.json()["status"] == "pending_review"
    doc = db.get(m.GovernanceDocument, 10)
    assert "All employees and contractors." in doc.content and "every quarter" in doc.content and "third parties" not in doc.content
    row = db.query(m.GovernanceDocumentVersion).filter_by(document_id=10, version_number="1.3").one()
    assert row.change_type == "minor" and row.change_reason == "Scope was not agreed" and row.created_by == library.me["id"]     # a kind, by the rule
    assert row.change_summary == "Took back 1 change made since version 1.0: Scope was not agreed"
    assert row.change_set["items"][0]["origin"] == "revert" and row.change_set["reverts"]["since_version"] == "1.0"
    chain = library.chain()
    assert chain["versions"][-1]["origin_label"] == "Taken back" and chain["events"][-1]["kind"] == "status_changed"
    assert chain["versions"][-1]["reverts"] == {"since_version": "1.0", "changes": 1} and chain["versions"][-1]["restored_from"] is None
    assert db.query(m.DocumentAuditLog).filter_by(document_id=10, action="status_changed", new_value="pending_review").count() == 1
    assert db.query(m.DocumentAuditLog).filter_by(document_id=10, action="content_edited", new_value="1.3").one().action_details.endswith(
        "; sent back to its reviewers")
    # the same request again: the document was edited meanwhile, so it is refused
    stale = library.lib.post("/governance/documents/10/change-chain/revert", json={"from_version_id": versions[0].id, "hunks": [scope["id"]], "hash": red["hash"]})
    assert stale.status_code == 409


def test_a_restore_and_a_take_back_are_major_or_minor_by_the_rule_and_keep_the_persons_reason(library, monkeypatch):
    from grc.modules.governance.routers import versions as version_routes

    db = library.db
    versions = _three_versions(db)                              # 1.0, 1.1 quarterly, 1.2 scope (both minor)
    monkeypatch.setattr(version_routes, "get_user_tenants", lambda user, db: [1])
    me = db.get(m.GRCUser, library.me["id"])

    def number(n):
        return db.query(m.GovernanceDocumentVersion).filter_by(document_id=10, version_number=n).one()

    # restoring 1.1 undoes a minor version and a few words: minor
    out = version_routes.rollback_to_version(document_id=10, version_id=versions[1].id, db=db, current_user=me,
                                             rollback_data=version_routes.VersionRollback(rollback_reason="Scope was not agreed"))
    assert (out["new_version"]["version_number"], out["new_version"]["change_type"], out["new_version"]["change_reason"]) == ("1.3", "minor", "Scope was not agreed")

    # a section is added (major): not published yet, so the pending 1.3 becomes 2.0 (one version up, now major)
    _edit(db, 10, db.get(m.GovernanceDocument, 10).content + "\n## 5. Exceptions\n\nNo exception is allowed.\n", user_id=1, reason="Exceptions", change_type="major")
    assert db.get(m.GovernanceDocument, 10).current_version == "2.0" and number("2.0").id == out["new_version"]["id"]
    red = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[2].id}").json()
    added = next(h for h in red["hunks"] if h["status"] == "inserted")
    back = library.lib.post("/governance/documents/10/change-chain/revert", json={"from_version_id": versions[2].id, "hunks": [added["id"]],
                                                                                  "hash": red["hash"], "reason": "Exceptions are in their own policy"})
    assert back.status_code == 200 and back.json()["version"] == "2.0"     # the take back goes into the same version
    row = number("2.0")
    assert row.change_type == "major" and row.change_reason.startswith("Exceptions are in their own policy")
    assert row.change_summary == "Took back 1 change made since version 1.2: Exceptions are in their own policy"
    items = library.chain()["versions"][-1]["items"]
    assert any(i["origin"] == "revert" or i["reason"] == "Exceptions are in their own policy" for i in items) or not any(
        i["status"] == "inserted" for i in items)                           # the section is gone; the version still says why


def test_a_sign_off_record_is_not_taken_back_and_a_file_that_cannot_be_read_cannot_be_taken_back_into(library):
    db = library.db
    doc = db.get(m.GovernanceDocument, 10)
    doc.content = POLICY
    db.commit()
    stamp = POLICY.replace("| Reviewed by | | | |", "| Reviewed by | Rana | | 2026-10-01 · E-signed ab12cd34 |")
    _edit(db, 10, stamp, user_id=2, reason="Reviewed by: Rana", change_type="signoff")
    versions = db.query(m.GovernanceDocumentVersion).filter_by(document_id=10).order_by(m.GovernanceDocumentVersion.id).all()
    red = library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[0].id}").json()
    assert red["hunks"][0]["locked"] is True and red["hunks"][0]["section"] == "4. Approval Signoff"
    assert "Reviewed by | Rana |  | 2026-10-01 · E-signed ab12cd34" in red["hunks"][0]["after_text"]         # a table reads as its rows
    assert red["hunks"][0]["after_text"].count("\n") == 1 and red["hunks"][0]["before_text"].startswith("Role | Name | Designation | Signature / Date")
    assert red["hunks"][0]["can_revert"] is False
    refused = library.lib.post("/governance/documents/10/change-chain/revert", json={"from_version_id": versions[0].id, "hunks": [red["hunks"][0]["id"]]})
    assert refused.status_code == 400 and "sign-off" in refused.json()["detail"]
    doc.content, doc.file_path = "", "/files/policy.docx"                # a document that is a file nobody can read any more
    db.commit()
    gone = library.lib.post("/governance/documents/10/change-chain/revert", json={"from_version_id": versions[0].id, "hunks": [0]})
    assert gone.status_code == 400 and "cannot be read" in gone.json()["detail"]
    assert library.lib.get(f"/governance/documents/10/change-chain/redline?from={versions[0].id}&to={versions[1].id}").json()["can_revert"] is False


# ── an uploaded original: the file is the document ───────────────────────────
def _edit_in_studio(world, doc_id, old, new, *, reason):
    """Open a library document in the studio, change one wording by hand and save it to the library: the next version."""
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": doc_id}).json()
    assert old in s["content"], s["content"][:300]
    put = world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"].replace(old, new), "base_revision": s["revision"]})
    assert put.status_code == 200
    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={"change_reason": reason})
    assert saved.status_code == 200, saved.text
    return saved.json()


def test_an_uploaded_word_original_keeps_its_history_and_a_change_is_taken_back_into_a_new_file(library, lib):
    world, db = library, library.db
    original = make_library_doc(world, lib)                                    # the uploaded file: version 1.0, published
    db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=20, role_type="reviewer", target_type="user", target_id=1, added_by=1))    # it has a reviewer, so a change sends it back to review
    db.commit()
    before = open(original, "rb").read()
    assert _edit_in_studio(world, 20, "Review it every year.", "Review it every six months.", reason="Review cycle")["version"] == "1.1"

    c = world.chain(20)
    assert c["document"]["is_file"] is True and [v["number"] for v in c["versions"]] == ["1.0", "1.1"]
    uploaded, edited = c["versions"]
    assert (uploaded["origin"], uploaded["has_file"], uploaded["file_name"]) == ("initial", True, "orig.docx")      # the original is the first file
    assert edited["changes"] == 1 and edited["has_file"] and edited["reason"] == "Review cycle"
    assert "every year" in edited["items"][0]["before_text"] and "every six months" in edited["items"][0]["after_text"]

    # the document page: tracked changes against the original, and the change can be taken back there
    red = world.lib.get(f"/governance/documents/20/change-chain/redline?from={uploaded['id']}").json()
    assert red["can_revert"] is True and red["why_not"] is None and len(red["hunks"]) == 1 and red["hunks"][0]["can_revert"] is True
    assert red["hunks"][0]["by"]["by"] == "Olivia Owner"
    doc = db.get(m.GovernanceDocument, 20)
    doc.status = "published"                                                   # approved again meanwhile
    db.commit()
    first_edit_file = doc.file_path
    out = world.lib.post("/governance/documents/20/change-chain/revert", json={"from_version_id": uploaded["id"], "hunks": [red["hunks"][0]["id"]],
                                                                               "hash": red["hash"], "reason": "The cycle was not agreed"})
    assert out.status_code == 200, out.text
    assert out.json()["version"] == "1.2" and out.json()["status"] == "pending_review" and out.json()["notes"] == []

    db.refresh(doc)
    assert doc.file_path not in (original, first_edit_file) and doc.file_type == "docx" and not (doc.content or "").strip()
    new_file = Document(doc.file_path)
    para = next(p for p in new_file.paragraphs if p.text.startswith("This policy applies"))
    assert "Review it every year." in para.text and "six months" not in para.text
    assert zipfile.ZipFile(doc.file_path).read("word/header1.xml") == zipfile.ZipFile(original).read("word/header1.xml")       # still the original's file
    assert open(original, "rb").read() == before                              # the uploaded file is untouched
    rows = {v.version_number: v for v in db.query(m.GovernanceDocumentVersion).filter_by(document_id=20)}
    assert rows["1.0"].file_path == original and rows["1.1"].file_path == first_edit_file and rows["1.2"].file_path == doc.file_path
    assert (rows["1.2"].status, rows["1.1"].status, rows["1.2"].change_type, rows["1.2"].change_reason) == ("current", "superseded", "minor", "The cycle was not agreed")
    assert rows["1.2"].change_set["items"][0]["origin"] == "revert" and rows["1.2"].change_set["reverts"]["since_version"] == "1.0"
    assert db.query(m.DocumentAuditLog).filter_by(document_id=20, action="status_changed", new_value="pending_review").count() == 2     # the studio save and this

    again = world.chain(20)
    assert [v["number"] for v in again["versions"]] == ["1.0", "1.1", "1.2"]
    assert again["versions"][-1]["origin_label"] == "Taken back" and again["versions"][-1]["has_file"] is True
    nothing = world.lib.get(f"/governance/documents/20/change-chain/redline?from={uploaded['id']}").json()
    assert nothing["hunks"] == []                                              # back to the original's wording
    stale = world.lib.post("/governance/documents/20/change-chain/revert", json={"from_version_id": uploaded["id"], "hunks": [0], "hash": red["hash"]})
    assert stale.status_code == 409                                            # the file changed since that redline was read


def test_restoring_the_original_file_of_an_uploaded_document_makes_it_the_documents_file_again_as_a_new_version(library, lib, monkeypatch):
    from grc.modules.governance.routers import versions as version_routes

    world, db = library, library.db
    original = make_library_doc(world, lib)                                    # the uploaded file, 1.0
    db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=20, role_type="reviewer", target_type="user", target_id=1, added_by=1))
    db.commit()
    _edit_in_studio(world, 20, "Review it every year.", "Review it every six months.", reason="Review cycle")       # 1.1, a new file
    doc = db.get(m.GovernanceDocument, 20)
    doc.status = "published"
    db.commit()
    edited_file = doc.file_path
    monkeypatch.setattr(version_routes, "get_user_tenants", lambda user, db: [1])
    listed = version_routes.list_versions(document_id=20, db=db, current_user=db.get(m.GRCUser, 1))
    rows = listed["items"] if isinstance(listed, dict) else listed
    first = next(v for v in rows if v["version_number"] == "1.0")
    assert (first["has_file"], first["file_name"], first["file_type"]) == (True, "orig.docx", "docx")

    out = version_routes.rollback_to_version(document_id=20, version_id=first["id"], rollback_data=version_routes.VersionRollback(rollback_reason="Back to what was approved"),
                                             db=db, current_user=db.get(m.GRCUser, 1))
    assert out["new_version"]["version_number"] == "1.2" and out["new_version"]["has_file"] is True and out["status"] == "pending_review"
    assert db.get(m.GovernanceDocumentVersion, out["new_version"]["id"]).change_reason == "Back to what was approved"            # the person's why
    db.refresh(doc)
    assert open(doc.file_path, "rb").read() == open(original, "rb").read() and doc.file_path not in (original, edited_file)    # a copy of the original
    assert doc.current_version == "1.2" and not (doc.content or "").strip() and os.path.exists(edited_file)                       # the edited file stays as 1.1
    assert db.query(m.DocumentAuditLog).filter_by(document_id=20, action="version_rollback").count() == 1
    chain = world.chain(20)
    assert chain["versions"][-1]["restored_from"] == "1.0" and chain["versions"][-1]["reverts"] is None
    assert chain["versions"][-1]["origin"] == "restore" and world.lib.get(
        f"/governance/documents/20/change-chain/redline?from={first['id']}").json()["hunks"] == []                           # the text is the original's again


def test_a_change_to_a_word_original_that_cannot_be_written_into_it_is_taken_back_into_a_clean_copy_and_the_person_is_told(library, lib, monkeypatch):
    world, db = library, library.db
    make_library_doc(world, lib, status="draft")
    _edit_in_studio(world, 20, "Review it every year.", "Review it every six months.", reason="Review cycle")
    chain = world.chain(20)
    red = world.lib.get(f"/governance/documents/20/change-chain/redline?from={chain['versions'][0]['id']}").json()

    def refuse(*args, **kwargs):
        raise docx_io.PatchProblem("block 9 reads differently")

    monkeypatch.setattr(docx_io, "export_in_place", refuse)
    out = world.lib.post("/governance/documents/20/change-chain/revert", json={"from_version_id": chain["versions"][0]["id"], "hunks": [red["hunks"][0]["id"]]})
    assert out.status_code == 200 and out.json()["version"] == "1.1" and out.json()["status"] == "draft"      # not published: still 1.1
    assert len(out.json()["notes"]) == 1 and "clean Word document" in out.json()["notes"][0]
    doc = db.get(m.GovernanceDocument, 20)
    assert "Review it every year." in "\n".join(p.text for p in Document(doc.file_path).paragraphs)


def test_a_change_is_written_into_a_word_file_with_tracked_changes_as_into_a_copy_with_them_accepted_and_no_copy_is_left_behind(library, lib):
    import _docx_fixtures as F
    from grc.modules.governance.studio import service as studio

    make_library_doc(library, lib, status="draft", file_fixture=F.tracked_policy)
    doc = library.db.get(m.GovernanceDocument, 20)
    text = C.document_text(doc)
    assert "seven years" in text and "five" not in text                           # the text is the accepted text
    data, notes = studio.file_bytes(library.db, doc, text.replace("seven years", "ten years"))
    assert notes == []                                                            # written into the file, not a rebuilt copy
    body = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "ten years" in body and "five" not in body and "Untouched paragraph." in body
    assert [f for f in os.listdir(lib) if "twin" in f] == []


def test_an_uploaded_pdf_original_is_kept_and_a_change_to_its_word_copy_is_taken_back_into_a_new_word_file(library, lib):
    world, db = library, library.db
    md = ("# Retention Policy\n\n## 1. Purpose\n\n" + "Records shall be kept as long as the law requires and no longer. " * 6 + "\n\n"
          "## 2. Rules\n\n1. Keep invoices for ten years.\n2. Keep emails for two years.\n")
    pdf = lib / "retention.pdf"
    pdf.write_bytes(exporters.to_pdf(md, title="Retention", entity_name="Capital Financial Services Bank", doc_code="POL-9", version="1.0"))
    size = os.path.getsize(pdf)
    db.add(m.GovernanceDocument(id=21, tenant_id=1, title="Retention Policy", doc_type="policy", status="draft", current_version="1.0", file_name="retention.pdf",
                                file_path=str(pdf), file_size=size, file_type="pdf", owner_id=1, content=None))
    db.add(m.GovernanceDocumentVersion(document_id=21, version_number="1.0", change_type="major", title="Retention Policy", file_name="retention.pdf", file_path=str(pdf),
                                       file_size=size, file_type="pdf", change_summary="Initial version", status="current", created_by=1))
    db.commit()
    assert _edit_in_studio(world, 21, "Keep emails for two years.", "Keep emails for three years.", reason="Longer")["version"] == "1.1"
    c = world.chain(21)
    uploaded, copy = c["versions"]
    assert uploaded["file_name"] == "retention.pdf" and copy["file_name"].endswith(".docx") and copy["changes"] == 1       # the PDF stays as the first file
    red = world.lib.get(f"/governance/documents/21/change-chain/redline?from={uploaded['id']}").json()
    assert red["can_revert"] is True and "three years" in red["hunks"][0]["after_text"]
    out = world.lib.post("/governance/documents/21/change-chain/revert", json={"from_version_id": uploaded["id"], "hunks": [red["hunks"][0]["id"]], "hash": red["hash"]})
    assert out.status_code == 200 and out.json()["version"] == "1.1"          # not published: the take back goes into 1.1
    doc = db.get(m.GovernanceDocument, 21)
    assert doc.file_type == "docx" and doc.file_path.endswith(".docx")
    text = "\n".join(p.text for p in Document(doc.file_path).paragraphs)
    assert "Keep emails for two years." in text and "three years" not in text
    assert os.path.exists(str(pdf)) and db.query(m.GovernanceDocumentVersion).filter_by(document_id=21, version_number="1.0").one().file_path == str(pdf)


# ── the catalog ──────────────────────────────────────────────────────────────
def test_the_catalog_is_exported_as_excel_pdf_and_csv_with_words_never_run_as_formulas(library):
    from openpyxl import load_workbook

    db = library.db
    doc = db.get(m.GovernanceDocument, 10)
    doc.content = POLICY
    db.commit()
    row = _edit(db, 10, POLICY.replace("every year", "=HYPERLINK(\"http://x\")"), user_id=1, reason="+ a reason that starts with a plus", change_type="studio")
    row.change_set = {"v": 1, "items": [{"n": 1, "status": "modified", "section": "3. Access Control", "locked": False, "blocks": 1, "words_added": 2, "words_removed": 2,
                                         "format_only": False, "before_text": "reviewed every year", "after_text": "=SUM(1,1) reviewed monthly", "truncated": False,
                                         "before_digests": [], "after_digests": [], "origin": "ai", "by": "Olivia Owner", "by_id": 1, "at": "2026-10-02T10:00:00",
                                         "request": "Make it monthly", "reason": "monthly", "sources": [{"ref": "R1", "kind": "framework", "title": "ISO/IEC 27001",
                                                                                                         "locator": "A.5.1", "quote": "planned intervals", "basis": "reference", "verified": True}],
                                         "ref": {}, "coverage": 1.0}], "sources": [], "session_id": 3}
    db.commit()
    xlsx = library.lib.get("/governance/documents/10/change-chain/export?format=xlsx")
    assert xlsx.status_code == 200 and xlsx.headers["content-type"] == X.XLSX and ".xlsx" in xlsx.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Summary", "Versions", "Changes", "Sources", "Approvals"]
    changes = list(wb["Changes"].iter_rows(values_only=True))
    assert changes[0][:7] == ("Ref", "Version", "Date", "Changed by", "Stage", "Section", "Change")
    assert changes[1][0] == "1.1.1" and changes[1][3] == "Olivia Owner" and changes[1][7] == "reviewed every year" and changes[1][8] == "=SUM(1,1) reviewed monthly"
    assert wb["Changes"].cell(row=2, column=9).data_type == "s" and changes[1][12] == "Make it monthly" and "ISO/IEC 27001 · A.5.1" in changes[1][14]
    assert wb["Sources"].max_row >= 1 and wb["Versions"].max_row == 3
    pdf = library.lib.get("/governance/documents/10/change-chain/export?format=pdf")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-" and len(pdf.content) > 2000
    csv_out = library.lib.get("/governance/documents/10/change-chain/export?format=csv")
    text = csv_out.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Document,Document version now,Ref,Version") and "'=SUM(1,1) reviewed monthly" in text
    assert library.lib.get("/governance/documents/10/change-chain/export?format=docx").status_code == 400


# ── the studio ───────────────────────────────────────────────────────────────
def _session_with_reference(world):
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 10}).json()
    fw = world.http.post(world.api(f"/sessions/{s['id']}/references"), json={"kind": "framework", "source_id": 1, "scope": {"mode": "all"}}).json()
    assert fw["key"] == "R1"
    return s


def _ask(world, sid, text, changes, used=()):
    world.respond(lambda body: reply(answer(changes, "Done.", used)))
    sent = world.http.post(world.api(f"/sessions/{sid}/messages"), json={"text": text}).json()
    return world.http.get(world.api(f"/sessions/{sid}/messages/{sent['id']}")).json()


USED = [{"ref": "R1", "locator": "A.8.2", "quote": "The allocation and use of privileged access rights shall be restricted and managed.", "how": "privileged access", "basis": "reference"}]


def test_an_accepted_change_keeps_who_accepted_it_what_it_replaced_and_what_it_rests_on(world):
    s = _session_with_reference(world)
    msg = _ask(world, s["id"], "Make access reviews quarterly.", [change("replace", "B7", markdown="Access to systems shall be reviewed every quarter.", reason="quarterly", refs=["R1"])], USED)
    out = world.http.post(world.api(f"/sessions/{s['id']}/messages/{msg['id']}/accept"), json={"op_ids": ["c1"]}).json()
    op = out["message"]["proposal"]["ops"][0]
    assert (op["status"], op["accepted_by"], op["revision"]) == ("accepted", 1, 2) and op["before_markdown"].strip() == "Access to systems shall be reviewed every year."

    from grc.modules.governance.studio import service
    session = world.db.get(m.StudioSession, s["id"])
    origins = service.session_origins(world.db, session)
    assert len(origins) == 1 and origins[0]["kind"] == "ai" and origins[0]["by"] == "Olivia Owner" and origins[0]["request"] == "Make access reviews quarterly."
    src = origins[0]["sources"][0]
    assert (src["ref"], src["kind"], src["locator"], src["verified"]) == ("R1", "framework", "A.8.2", True) and src["title"].startswith("ISO/IEC 27001")


def test_taking_an_accepted_change_back_restores_what_it_replaced_and_is_refused_once_the_text_was_edited(world):
    s = _session_with_reference(world)
    first = _ask(world, s["id"], "Quarterly.", [change("replace", "B7", markdown="Access to systems shall be reviewed every quarter.", reason="quarterly")])
    world.http.post(world.api(f"/sessions/{s['id']}/messages/{first['id']}/accept"), json={})
    reverted = world.http.post(world.api(f"/sessions/{s['id']}/messages/{first['id']}/ops/c1/revert")).json()
    assert "every year" in reverted["content"] and "every quarter" not in reverted["content"] and reverted["message"]["proposal"]["ops"][0]["status"] == "reverted"
    revisions = world.http.get(world.api(f"/sessions/{s['id']}/revisions")).json()["items"]
    assert revisions[0]["author_type"] == "revert" and revisions[0]["summary"].startswith("Took back a Comply Agent change")
    again = world.http.post(world.api(f"/sessions/{s['id']}/messages/{first['id']}/ops/c1/revert"))
    assert again.status_code == 400 and "accepted" in again.json()["detail"]

    # an inserted paragraph the person then edited cannot be found again
    second = _ask(world, s["id"], "Add contractors.", [change("insert", "B5", markdown="Contractors are in scope.", reason="scope")])
    world.http.post(world.api(f"/sessions/{s['id']}/messages/{second['id']}/accept"), json={})
    content = world.http.get(world.api(f"/sessions/{s['id']}")).json()
    rev = content["revision"]
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": content["content"].replace("Contractors are in scope.", "Contractors and visitors are in scope."), "base_revision": rev})
    clash = world.http.post(world.api(f"/sessions/{s['id']}/messages/{second['id']}/ops/c1/revert"))
    assert clash.status_code == 409 and "edited since" in clash.json()["detail"]


def test_a_removal_and_a_first_draft_are_taken_back_too(world):
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 10}).json()
    removal = _ask(world, s["id"], "Drop the scope paragraph.", [change("delete", "B5", reason="not needed")])
    world.http.post(world.api(f"/sessions/{s['id']}/messages/{removal['id']}/accept"), json={})
    assert "All employees and contractors." not in world.http.get(world.api(f"/sessions/{s['id']}")).json()["content"]
    back = world.http.post(world.api(f"/sessions/{s['id']}/messages/{removal['id']}/ops/c1/revert")).json()
    assert back["content"].index("All employees and contractors.") > back["content"].index("## 2. Scope")

    blank = world.http.post(world.api("/sessions"), json={"kind": "new", "title": "Memo", "template_id": "builtin:blank", "doc_type": "other"}).json()
    draft = _ask(world, blank["id"], "Write a memo.", [change("replace_all", markdown="# Memo\n\n## Purpose\n\nTell people things.\n")])
    world.http.post(world.api(f"/sessions/{blank['id']}/messages/{draft['id']}/accept"), json={})
    undone = world.http.post(world.api(f"/sessions/{blank['id']}/messages/{draft['id']}/ops/c1/revert")).json()
    assert undone["content"].strip() == "# Memo"


def test_the_tracked_changes_of_a_session_are_cut_into_changes_and_one_is_taken_back(world):
    s = _session_with_reference(world)
    msg = _ask(world, s["id"], "Quarterly and mention contractors.", [
        change("replace", "B7", markdown="Access to systems shall be reviewed every quarter.", reason="quarterly", refs=["R1"]),
        change("insert", "B5", markdown="Contractors are in scope.", reason="scope")], USED)
    world.http.post(world.api(f"/sessions/{s['id']}/messages/{msg['id']}/accept"), json={})
    diff = world.http.get(world.api(f"/sessions/{s['id']}/diff?from=0")).json()
    assert diff["can_revert"] is True and len(diff["hunks"]) == 2 and diff["stats"]["words_added"] >= 4
    by = {h["section"]: h for h in diff["hunks"]}
    assert by["3. Access Control"]["by"]["by"] == "Olivia Owner" and by["3. Access Control"]["by"]["sources"][0]["locator"] == "A.8.2"
    assert by["2. Scope"]["by"]["request"] == "Quarterly and mention contractors."
    old = world.http.get(world.api(f"/sessions/{s['id']}/diff?from=1&to=1")).json()
    assert old["can_revert"] is False
    bad = world.http.post(world.api(f"/sessions/{s['id']}/changes/revert"), json={"frm": 0, "hunks": [by["2. Scope"]["id"]], "hash": "0" * 16})
    assert bad.status_code == 409
    ok = world.http.post(world.api(f"/sessions/{s['id']}/changes/revert"), json={"frm": 0, "hunks": [by["2. Scope"]["id"]], "hash": diff["hash"]}).json()
    assert "Contractors are in scope" not in ok["content"] and "every quarter" in ok["content"]
    ops = {o["id"]: o["status"] for o in world.http.get(world.api(f"/sessions/{s['id']}/messages/{msg['id']}")).json()["proposal"]["ops"]}
    assert ops == {"c1": "accepted", "c2": "reverted"}                                    # the card of the change that was taken back says so
    revisions = world.http.get(world.api(f"/sessions/{s['id']}/revisions")).json()["items"]
    assert revisions[0]["author_type"] == "revert" and "Took back 1 change" in revisions[0]["summary"]


def test_saving_to_the_library_records_the_changes_with_their_sources_and_the_chain_shows_them(library):
    world = library
    doc = world.db.get(m.GovernanceDocument, 10)
    doc.content = POLICY
    doc.status = "published"
    world.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=10, role_type="reviewer", target_type="user", target_id=1, added_by=1))
    world.db.commit()
    _audit(world.db, 10, 1, "published", datetime(2026, 9, 1, 9, 0), status="published", old="approved")
    s = _session_with_reference(world)
    msg = _ask(world, s["id"], "Make access reviews quarterly and mention third parties.", [
        change("replace", "B7", markdown="Access to systems shall be reviewed every quarter.", reason="quarterly", refs=["R1"]),
        change("insert", "B5", markdown="Third parties with access are in scope.", reason="scope")], USED)
    world.http.post(world.api(f"/sessions/{s['id']}/messages/{msg['id']}/accept"), json={})
    cur = world.http.get(world.api(f"/sessions/{s['id']}")).json()
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": cur["content"].replace("approved by the CISO", "approved by the CISO and the CRO"), "base_revision": cur["revision"]})
    # Comply Agent says whether the new version is a major or a minor change
    world.respond(lambda body: reply(json.dumps({"change": "major", "reason": "Access reviews become quarterly and third parties come into scope."})))
    saved = world.http.post(world.api(f"/sessions/{s['id']}/save-to-document"), json={"change_reason": "Quarterly reviews"}).json()
    assert saved["version"] == "2.0" and saved["status"] == "pending_review"

    row = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=10, version_number="2.0").one()
    assert row.change_type == "major" and row.change_summary.endswith("Major change: Access reviews become quarterly and third parties come into scope.")
    assert sorted(i["section"] for i in row.change_set["items"]) == ["2. Scope", "3. Access Control", "3. Access Control"]
    quarterly = next(i for i in row.change_set["items"] if "quarter" in i["after_text"])
    assert (quarterly["origin"], quarterly["by"], quarterly["request"]) == ("ai", "Olivia Owner", "Make access reviews quarterly and mention third parties.")
    assert quarterly["sources"][0]["title"].startswith("ISO/IEC 27001") and quarterly["sources"][0]["verified"] is True
    by_hand = next(i for i in row.change_set["items"] if "the CRO" in i["after_text"])
    assert by_hand["origin"] == "manual" and by_hand["sources"] == []
    status_rows = world.db.query(m.DocumentAuditLog).filter_by(document_id=10, action="status_changed", new_value="pending_review").count()
    assert status_rows == 1

    c = world.chain()
    top = c["versions"][-1]
    assert (top["number"], top["origin"], top["stage"], top["changes"]) == ("2.0", "mixed", "published", len(row.change_set["items"]))
    assert c["sources"][0]["quote"].startswith("The allocation and use of privileged access rights")
    assert c["events"][-1]["kind"] == "status_changed" and c["events"][-1]["to"] == "pending_review"


def test_moving_a_session_to_the_library_records_the_first_version_with_its_sources(world, monkeypatch):
    monkeypatch.setattr("grc.modules.governance.routers.policy_parser.maybe_autoparse_document", lambda *a, **k: False)
    s = world.http.post(world.api("/sessions"), json={"kind": "new", "title": "Memo", "template_id": "builtin:blank", "doc_type": "other"}).json()
    world.http.post(world.api(f"/sessions/{s['id']}/references"), json={"kind": "framework", "source_id": 1, "scope": {"mode": "all"}})
    draft = _ask(world, s["id"], "Write a memo on privileged access.", [change("replace_all", markdown="# Memo\n\n## Purpose\n\nPrivileged access rights shall be restricted.\n", reason="draft", refs=["R1"])], USED)
    world.http.post(world.api(f"/sessions/{s['id']}/messages/{draft['id']}/accept"), json={})
    moved = world.http.post(world.api(f"/sessions/{s['id']}/move-to-library"), json={"doc_type": "policy"}).json()
    row = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=moved["document_id"]).one()
    assert row.version_number == "1.0" and row.content.startswith("# Memo")
    first = row.change_set["items"][-1]
    assert first["origin"] == "ai" and first["by"] == "Olivia Owner" and first["sources"][0]["locator"] == "A.8.2" and first["status"] == "inserted"
