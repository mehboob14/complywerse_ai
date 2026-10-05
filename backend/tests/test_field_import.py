"""Custom fields suggested from an uploaded template.

The person uploads the sheet, CSV or document the organisation already uses; AI reads it and proposes the fields;
they check them and add them in one click. What matters, and is tested here: a suggestion is only ever something the
file holds (anything the model adds is dropped), the model is not trusted on types or options its evidence does not
bear out, nothing is saved until the person confirms, and the AI being off still gives suggestions, said honestly.
"""
import csv
import importlib
import io
import json

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.models import get_db
from grc.routers.auth_router import require_auth
from grc.services import field_import as fi
from grc.services import module_settings as ms
from grc.services.assessment_evidence_ai import AIEmptyAnswer, AIUnavailable


def settings(existing=(), builtins=()):
    """What get_settings gives a module: its record, its live custom fields and its built-in fields."""
    return {
        "module": {"key": "assets", "label": "IT Asset Inventory", "record": "Asset"},
        "fields": [{"key": ms._slug(label), "label": label, "type": "text", "archived": False} for label in existing],
        "builtins": [{"key": ms._slug(label), "label": label} for label in builtins],
    }


def csv_bytes(rows, delimiter=","):
    out = io.StringIO()
    csv.writer(out, delimiter=delimiter).writerows(rows)
    return out.getvalue().encode("utf-8")


ASSETS = [
    ["No.", "Asset Name", "Owner", "Environment", "Go-live date", "Cost centre", "Internet facing?", "Notes"],
    [1, "Payments API", "A. Khan", "Production", "2024-03-01", 12345, "Yes", "Core payments gateway, PCI in scope"],
    [2, "Mail relay", "B. Ali", "Staging", "2023-11-15", 22345, "No", ""],
    [3, "HR portal", "C. Shah", "Production", "2022-06-30", 12345, "No", "Vendor hosted"],
    [4, "Data lake", "A. Khan", "Production", "2025-01-10", 33445, "No", ""],
    [5, "Test rig", "D. Raza", "Staging", "2025-02-02", 22345, "No", "Decommission after Q3"],
]


def by_label(result):
    return {f["label"]: f for f in result["fields"]}


# ── reading a sheet ──────────────────────────────────────────────────────────

def test_a_csv_gives_one_suggestion_per_heading_typed_from_its_values():
    got = fi.suggest("assets.csv", csv_bytes(ASSETS), settings(), use_ai=False)
    fields = by_label(got)
    assert got["by"] == "rules" and got["file"]["kind"] == "sheet"
    assert "No." not in fields and got["left_out"] == [{"label": "No.", "reason": "A row number"}]
    assert fields["Asset Name"]["type"] == "text"
    assert fields["Owner"]["type"] == "user"
    assert fields["Environment"]["type"] == "select" and fields["Environment"]["options"] == ["Production", "Staging"]
    assert fields["Go-live date"]["type"] == "date"
    assert fields["Cost centre"]["type"] == "number"
    assert fields["Internet facing?"]["type"] == "checkbox"
    assert fields["Notes"]["type"] in ("text", "textarea")
    assert fields["Go-live date"]["example"] == "2024-03-01" and fields["Go-live date"]["where"] == "column E"
    assert all(f["status"] == "new" and f["key"] for f in got["fields"])


def test_other_delimiters_and_encodings_are_read():
    semicolons = csv_bytes([["Name", "Risk level"], ["Alpha", "High"], ["Beta", "Low"], ["Gamma", "High"]], ";")
    assert [f["label"] for f in fi.suggest("a.csv", semicolons, settings(), use_ai=False)["fields"]] == ["Name", "Risk level"]
    latin = "Propriétaire,Écart\nAli,1\nBea,2\n".encode("latin-1")
    assert [f["label"] for f in fi.suggest("b.csv", latin, settings(), use_ai=False)["fields"]] == ["Propriétaire", "Écart"]


def test_a_title_above_the_headings_is_passed_over_and_excel_dropdowns_become_options():
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.append(["Asset register - prepared for audit"])
    ws.append([])
    ws.append(["Asset Name *", "Status", "Tier", "Criticality note"])
    inline = DataValidation(type="list", formula1='"Open,In progress,Closed"')
    ranged = DataValidation(type="list", formula1="=Lists!$A$2:$A$4")
    ws.add_data_validation(inline)
    ws.add_data_validation(ranged)
    inline.add("B4:B50")
    ranged.add("C4:C50")
    ws["D3"].comment = Comment("Why this asset is critical, in a sentence", "template author")
    lists = wb.create_sheet("Lists")
    for row in (["Tier"], ["Tier 1"], ["Tier 2"], ["Tier 3"]):
        lists.append(row)
    buf = io.BytesIO()
    wb.save(buf)

    got = fi.suggest("template.xlsx", buf.getvalue(), settings(), use_ai=False)
    fields = by_label(got)
    assert fields["Asset Name"]["required"] is True                      # the asterisk, not the label
    assert fields["Status"]["type"] == "select" and fields["Status"]["options"] == ["Open", "In progress", "Closed"]
    assert fields["Tier"]["options"] == ["Tier 1", "Tier 2", "Tier 3"]   # a range on another sheet
    assert fields["Criticality note"]["help"] == "Why this asset is critical, in a sentence"
    assert "Asset register" not in " ".join(fields)


def test_a_header_only_template_is_typed_from_its_headings():
    rows = [["Owner", "Review date", "Is it internet facing?", "Business justification", "Name"]]
    fields = by_label(fi.suggest("t.csv", csv_bytes(rows), settings(), use_ai=False))
    assert [fields[k]["type"] for k in ("Owner", "Review date", "Is it internet facing?", "Business justification", "Name")] == \
        ["user", "date", "checkbox", "textarea", "text"]
    assert fields["Name"]["example"] == ""


def test_legacy_spreadsheets_are_read_through_calamine():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Supplier", "Contract end"])
    wb.active.append(["Acme", "2026-01-31"])
    buf = io.BytesIO()
    wb.save(buf)
    got = fi.suggest("t.ods", buf.getvalue(), settings(), use_ai=False)    # the reader is picked by type; xlsx bytes open too
    assert [f["label"] for f in got["fields"]] == ["Supplier", "Contract end"]


# ── reading a document ───────────────────────────────────────────────────────

def make_docx():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Vendor onboarding form")
    doc.add_paragraph("Contract owner: ____________")
    doc.add_paragraph("Review date ........")
    doc.add_paragraph("Hosting location:")
    table = doc.add_table(rows=3, cols=2)
    for row, label in zip(table.rows, ("Data classification", "Sub-processors", "Exit plan")):
        row.cells[0].text = label
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_a_word_form_gives_its_label_lines_and_blank_form_rows():
    got = fi.suggest("form.docx", make_docx(), settings(), use_ai=False)
    labels = {f["label"] for f in got["fields"]}
    assert {"Contract owner", "Review date", "Hosting location", "Data classification", "Sub-processors", "Exit plan"} <= labels
    assert got["file"]["kind"] == "document"
    assert by_label(got)["Review date"]["type"] == "date" and by_label(got)["Contract owner"]["type"] == "user"


def make_pdf(lines):
    stream = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(f"({line}) Tj T*" for line in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{off:010d} 00000 n \n" for off in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_a_text_pdf_is_read_for_its_form_lines():
    got = fi.suggest("form.pdf", make_pdf(["Incident report", "Reported by: ________", "Date of incident: ________"]),
                     settings(), use_ai=False)
    assert {f["label"] for f in got["fields"]} == {"Reported by", "Date of incident"}


# ── only what the file holds ─────────────────────────────────────────────────

def ai(payload):
    """A model that answers with this JSON."""
    return lambda messages: json.dumps(payload)


def test_the_model_cannot_add_a_field_the_file_does_not_have_or_overclaim_a_type():
    answer = {"fields": [
        {"column": 1, "label": "Asset Name", "type": "text", "required": True, "help": "Invented guidance"},
        {"column": 3, "label": "Environment", "type": "select", "options": ["Production", "Staging", "Disaster recovery"]},
        {"column": 7, "label": "Notes", "type": "date"},                       # its values are sentences, not dates
        {"column": None, "label": "Business Continuity Plan", "type": "text"},   # not in the file at all
        {"column": 99, "label": "Data residency", "type": "select", "options": ["EU", "US"]},
    ], "left_out": [{"label": "No.", "reason": "Row counter"}]}
    got = fi.suggest("assets.csv", csv_bytes(ASSETS), settings(), complete=ai(answer))
    fields = by_label(got)
    assert got["by"] == "ai"
    assert set(fields) == {"Asset Name", "Environment", "Notes"}
    assert fields["Asset Name"]["required"] is False and fields["Asset Name"]["help"] == ""     # nothing in the file says so
    assert fields["Environment"]["options"] == ["Production", "Staging"]                          # the invented option is gone
    assert fields["Notes"]["type"] != "date"
    assert any("2 suggestions from AI were left out" in n for n in got["notes"])
    assert {"label": "No.", "reason": "Row counter"} in got["left_out"]


def test_a_dropdown_with_no_option_the_file_shows_becomes_text():
    answer = {"fields": [{"column": 1, "label": "Asset Name", "type": "select", "options": ["Alpha", "Beta"]}]}
    got = fi.suggest("assets.csv", csv_bytes(ASSETS), settings(), complete=ai(answer))
    assert by_label(got)["Asset Name"]["type"] == "text" and by_label(got)["Asset Name"]["options"] == []


def test_the_model_may_mark_required_only_where_the_file_does():
    rows = [["Owner *", "Region"], ["A", "East"]]
    answer = {"fields": [{"column": 0, "label": "Owner", "type": "user", "required": True},
                         {"column": 1, "label": "Region", "type": "text", "required": True}]}
    fields = by_label(fi.suggest("t.csv", csv_bytes(rows), settings(), complete=ai(answer)))
    assert fields["Owner"]["required"] is True and fields["Region"]["required"] is False


def test_in_a_document_a_label_has_to_be_in_its_text():
    answer = {"fields": [
        {"column": None, "label": "Contract owner", "type": "user", "required": False},
        {"column": None, "label": "Annual revenue", "type": "number"},
        {"column": None, "label": "Hosting location", "type": "select", "options": ["EU", "Pakistan"]},
    ]}
    got = fi.suggest("form.docx", make_docx(), settings(), complete=ai(answer))
    assert set(by_label(got)) == {"Contract owner", "Hosting location"}
    assert by_label(got)["Hosting location"]["type"] == "text"          # neither option appears in the document


def test_a_field_found_in_prose_shows_the_line_it_came_from():
    text = "Vendor form\nPlease state the data residency of the service (EU / UK / other).\nSign here"
    assert fi._where_in_text("data residency", text) == "“Please state the data residency of the service (EU / UK / other).”"
    assert fi._where_in_text("unrelated", text) == "the document text"


def test_a_very_wide_workbook_is_read_in_part_and_says_so():
    seen = {}

    def spy(messages):
        seen["columns"] = len(json.loads(messages[1]["content"])["columns"])
        return json.dumps({"fields": [{"column": 0, "label": "Heading 0", "type": "text"}]})

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "S0"
    for n in range(3):                                          # three sheets of 70 headings: 210 in all
        ws = wb.active if n == 0 else wb.create_sheet(f"S{n}")
        ws.append([f"Sheet {n} heading {i}" for i in range(70)])
        ws.append(["x"] * 70)
    buf = io.BytesIO()
    wb.save(buf)
    got = fi.suggest("wide.xlsx", buf.getvalue(), settings(), complete=spy)
    assert seen["columns"] == fi.MAX_AI_HEADINGS == 150
    assert any("first 150 of 210 headings" in n for n in got["notes"])
    normal = fi.suggest("a.csv", csv_bytes(ASSETS), settings(), complete=ai({"fields": [{"column": 1, "label": "Asset Name", "type": "text"}]}))
    assert normal["notes"] == []                                  # a normal file: no such note


def test_the_model_sees_headings_and_a_few_samples_never_the_rows():
    seen = {}

    def spy(messages):
        seen["user"] = messages[1]["content"]
        return json.dumps({"fields": [{"column": 1, "label": "Asset Name", "type": "text"}]})

    many = [["Name", "Secret"]] + [[f"row{i}", f"value-{i}"] for i in range(200)]
    fi.suggest("t.csv", csv_bytes(many), settings(), complete=spy)
    payload = json.loads(seen["user"])
    column = payload["columns"][1]
    assert column["heading"] == "Secret" and len(column["samples"]) <= fi.SAMPLE_VALUES
    assert "value-150" not in seen["user"] and "row199" not in seen["user"]


# ── no AI, or no usable answer ───────────────────────────────────────────────

@pytest.mark.parametrize("failure,words", [
    (AIUnavailable("no key"), "isn't set up"),
    (AIEmptyAnswer("cut off"), "could not read this file"),
    (RuntimeError("boom"), "could not be reached"),
])
def test_when_the_model_cannot_answer_the_rules_do_and_the_response_says_so(failure, words):
    def broken(messages):
        raise failure

    got = fi.suggest("assets.csv", csv_bytes(ASSETS), settings(), complete=broken)
    assert got["by"] == "rules" and len(got["fields"]) == 7
    assert any(words in n for n in got["notes"])


def test_the_licence_guard_refusing_a_prompt_falls_back_to_the_rules():
    from grc.services.licence_guard import LicenceRestrictedContent

    def refused(messages):
        raise LicenceRestrictedContent("SCF text")

    got = fi.suggest("assets.csv", csv_bytes(ASSETS), settings(), complete=refused)
    assert got["by"] == "rules" and any("cannot be sent to AI" in n for n in got["notes"])


@pytest.mark.parametrize("answer", ["not json at all", json.dumps({"fields": []}), json.dumps({"fields": [{"label": "Nothing here"}]})])
def test_an_unusable_answer_falls_back_to_the_rules(answer):
    got = fi.suggest("assets.csv", csv_bytes(ASSETS), settings(), complete=lambda messages: answer)
    assert got["by"] == "rules" and len(got["fields"]) == 7 and got["notes"]


# ── what is already on the form, and keys ────────────────────────────────────

def test_fields_already_on_the_form_are_shown_as_such_not_offered_again():
    got = fi.suggest("assets.csv", csv_bytes(ASSETS), settings(existing=["Date go live"], builtins=["Asset name", "Owner"]),
                     use_ai=False)
    fields = by_label(got)
    assert fields["Go-live date"]["status"] == "exists" and fields["Go-live date"]["exists_as"] == "Date go live"
    assert fields["Asset Name"]["status"] == "exists" and fields["Owner"]["status"] == "exists"
    assert fields["Environment"]["status"] == "new" and "key" in fields["Environment"] and "key" not in fields["Owner"]


def test_keys_are_valid_unique_and_clear_of_the_ones_taken():
    rows = [["2024 Budget", "Budget", "Budget code", "A very long heading " * 6, "Name"], ["1", "2", "3", "4", "5"]]
    got = fi.suggest("t.csv", csv_bytes(rows), settings(existing=["Budget"]), use_ai=False)
    keys = [f["key"] for f in got["fields"] if f["status"] == "new"]
    assert len(keys) == len(set(keys)) and all(ms.KEY.match(k) for k in keys)
    assert "budget" not in keys                    # the existing field's key


def test_the_same_heading_twice_is_one_field():
    got = fi.suggest("t.csv", csv_bytes([["Owner", "Owner", "Type"], ["a", "b", "c"]]), settings(), use_ai=False)
    assert [f["label"] for f in got["fields"]] == ["Owner", "Type"]


# ── files that cannot be used ────────────────────────────────────────────────

@pytest.mark.parametrize("name,data,words", [
    ("a.exe", b"MZ", "cannot be read"),
    ("a.doc", b"\xd0\xcf", "Save it as .docx"),
    ("a.csv", b"", "empty"),
    ("a.xlsx", b"not a zip", "not a valid Excel"),
    ("a.docx", b"not a zip", "not a valid Word"),
    ("a.pdf", b"hello", "not a valid PDF"),
    ("a.csv", b"\n\n\n", "empty"),
    ("a.csv", b"12,34\n56,78\n", "No column headings"),
])
def test_a_file_that_cannot_be_read_is_explained(name, data, words):
    with pytest.raises(fi.TemplateError, match=words):
        fi.suggest(name, data, settings(), use_ai=False)


def test_a_file_over_the_limit_is_refused():
    with pytest.raises(fi.TemplateError, match="larger than 5 MB"):
        fi.read_template("big.csv", b"a,b\n" + b"x" * (fi.MAX_BYTES + 1))


# ── adding what the person confirmed ─────────────────────────────────────────

@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    session = Session(engine)
    session.add(m.Tenant(id=1, name="Demo Bank", slug="demo"))
    session.commit()
    yield session
    session.close()


def test_confirmed_fields_are_saved_after_the_existing_ones_with_the_same_validation(db):
    ms.save_settings(db, 1, "assets", {"fields": [{"key": "cost_centre", "label": "Cost centre", "type": "number"}]})
    result = fi.add_fields(db, 1, "assets", [
        {"label": "Go-live date", "type": "date"},
        {"label": "Environment tier", "type": "select", "options": ["Tier 1", "Tier 2", "Tier 1"]},
        {"label": "Region", "type": "select", "options": []},                  # nothing to choose from: text
        {"label": "cost centre", "type": "number"},                            # already there
        {"label": "", "type": "text"},
        {"label": "Odd", "type": "wizard"},                                    # unknown type: text
    ], user_id=7)
    assert [a["label"] for a in result["added"]] == ["Go-live date", "Environment tier", "Region", "Odd"]
    assert {s["reason"] for s in result["skipped"]} == {"Already on the form", "No name"}
    saved = {f["label"]: f for f in ms.get_settings(db, 1, "assets")["fields"]}
    assert saved["Environment tier"]["options"] == ["Tier 1", "Tier 2"] and saved["Region"]["type"] == "text"
    assert saved["Odd"]["type"] == "text" and saved["Cost centre"]["type"] == "number"
    assert [f["label"] for f in sorted(saved.values(), key=lambda f: f["order"])][0] == "Cost centre"
    assert all(ms.KEY.match(f["key"]) for f in saved.values())


def test_nothing_new_is_not_a_write(db):
    ms.save_settings(db, 1, "assets", {"fields": [{"key": "cost_centre", "label": "Cost centre", "type": "number"}]})
    before = ms.get_settings(db, 1, "assets")["updated_at"]
    result = fi.add_fields(db, 1, "assets", [{"label": "Cost centre", "type": "number"}])
    assert result["added"] == [] and ms.get_settings(db, 1, "assets")["updated_at"] == before


def test_a_module_without_custom_fields_or_an_empty_request_is_refused(db):
    with pytest.raises(ValueError, match="Nothing to add"):
        fi.add_fields(db, 1, "assets", [])
    with pytest.raises(ValueError, match="Unknown module"):
        fi.add_fields(db, 1, "nonsense", [{"label": "x"}])


# ── the endpoints ────────────────────────────────────────────────────────────

@pytest.fixture
def api(db, monkeypatch):
    from grc.main import app

    router_module = importlib.import_module("grc.routers.module_settings_router")   # the module, not the APIRouter the package exports
    user = m.GRCUser(id=7, username="owner", email="owner@bank.example", display_name="Owner")
    db.add(user)
    db.commit()
    monkeypatch.setattr(router_module, "_check", lambda module_key, action, request, token, authorization, db: ms.spec(module_key))
    monkeypatch.setattr(router_module, "_tenant", lambda current_user, db: 1)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: user
    yield TestClient(app)
    app.dependency_overrides.clear()


def upload(http, name, data, module="assets"):
    return http.post(f"/module-settings/{module}/fields/suggest", files={"file": (name, data, "application/octet-stream")})


def test_uploading_a_template_returns_suggestions_and_saves_nothing(api, db, monkeypatch):
    monkeypatch.setattr(fi, "_complete", lambda messages: json.dumps({"fields": [
        {"column": 3, "label": "Environment", "type": "select", "options": ["Production", "Staging"]},
        {"column": 4, "label": "Go-live date", "type": "date"}]}))
    got = upload(api, "assets.csv", csv_bytes(ASSETS))
    assert got.status_code == 200, got.text
    body = got.json()
    fields = {f["label"]: f for f in body["fields"]}
    assert body["by"] == "ai" and set(fields) == {"Environment", "Go-live date"}
    assert fields["Go-live date"]["status"] == "new" and fields["Go-live date"]["key"] == "go_live_date"
    assert fields["Environment"]["status"] == "exists"                      # a built-in dropdown of the asset form
    assert ms.get_settings(db, 1, "assets")["fields"] == [] and db.query(m.ModuleSettings).count() == 0


def test_the_endpoint_explains_a_bad_file_and_a_big_one(api):
    assert upload(api, "a.exe", b"MZ").status_code == 400
    bad = upload(api, "a.xlsx", b"nope")
    assert bad.status_code == 400 and "not a valid Excel" in bad.json()["detail"]
    big = upload(api, "big.csv", b"a,b\n" + b"x" * (fi.MAX_BYTES + 10))
    assert big.status_code == 413


def test_importing_the_checked_fields_saves_them_and_reports_what_it_skipped(api, db):
    first = api.post("/module-settings/assets/fields/import", json={"fields": [
        {"label": "Go-live date", "type": "date", "required": True}, {"label": "Tier", "type": "select", "options": ["1", "2"]}]})
    assert first.status_code == 200, first.text
    body = first.json()
    assert [a["label"] for a in body["added"]] == ["Go-live date", "Tier"]
    assert {f["label"] for f in body["settings"]["fields"]} == {"Go-live date", "Tier"}
    again = api.post("/module-settings/assets/fields/import", json={"fields": [{"label": "tier", "type": "text"}]}).json()
    assert again["added"] == [] and again["skipped"] == [{"label": "tier", "reason": "Already on the form"}]
    assert api.post("/module-settings/assets/fields/import", json={"fields": []}).status_code == 400
    assert api.post("/module-settings/assets/fields/import", json={}).status_code == 400
