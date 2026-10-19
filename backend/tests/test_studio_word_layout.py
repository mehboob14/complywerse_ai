"""A Word file's own hierarchy in the studio: the numbers Word puts in front of paragraphs (which are not in their text), the
indents, clauses numbered through a heading style that look like body text, and lists that start deep under a heading -
which must still read back block for block, so the file stays edited in place. And its look: the typography Word gives each
paragraph (resolved through the styles as Word does) and the page, so the editor shows the document as Word does."""
import io

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor, Twips

from grc.modules.governance.studio import blocks as B
from grc.modules.governance.studio import docx_io as D
from grc.modules.governance.studio import word_layout as L

from test_studio_files import DOCX, upload
from test_studio_flow import answer, last_body, reply, world          # noqa: F401  (world is a fixture)


def _abstract(aid: int, levels: str) -> None:
    return parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="multilevel"/>{levels}</w:abstractNum>')


def _lvl(i: int, fmt: str, text: str, left: int, hanging: int) -> str:
    return (f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/>'
            f'<w:pPr><w:ind w:left="{left}" w:hanging="{hanging}"/></w:pPr></w:lvl>')


def _numbered(d, text, num, ilvl=0, style=None, *, plain=False, ind=None):
    p = d.add_paragraph(style=style)
    r = p.add_run(text)
    if plain:                                                   # a clause in a heading style that looks like body text
        r.bold, r.font.size = False, Pt(10)
    ppr = p._p.get_or_add_pPr()
    if num is not None:
        ppr.append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num}"/></w:numPr>'))
    if ind:
        ppr.append(parse_xml(f'<w:ind {nsdecls("w")} {ind}/>'))
    return p


def manual(path: str) -> str:
    """A procedure manual built the way the bank's is: Roman document-control items, headings numbered 1. / 1.1. / 1.1.1. by
    one multilevel list, clauses in a heading style typed as body text, a bullet list starting at Word's third level under a
    clause, a Roman list started again, and an indented note."""
    d = Document()
    numbering = d.part.numbering_part.element
    numbering.insert(0, _abstract(90, "".join(_lvl(i, "decimal", "".join(f"%{k + 1}." for k in range(i + 1)), 720 + 720 * i, 720) for i in range(3))))
    numbering.insert(1, _abstract(91, "".join(_lvl(i, "bullet", "", 720 * (i + 1), 360) for i in range(3))))
    numbering.insert(2, _abstract(92, _lvl(0, "upperRoman", "%1.", 1080, 720)))
    heads, bullets, roman = numbering.add_num(90).numId, numbering.add_num(91).numId, numbering.add_num(92).numId
    again = numbering.add_num(92)
    again.add_lvlOverride(ilvl=0).add_startOverride(1)
    _numbered(d, "Document Control", roman, style="List Paragraph")
    _numbered(d, "Revision History", roman, style="List Paragraph")
    _numbered(d, "ACCEPTABLE USE PROCEDURE", heads, 0, "Heading 1")
    _numbered(d, "Purpose", heads, 1, "Heading 2")
    _numbered(d, "All staff protect the information they use.", heads, 2, "Heading 3", plain=True)
    _numbered(d, "Passwords follow these rules:", heads, 2, "Heading 3", plain=True)
    _numbered(d, "Passwords are at least eight characters", bullets, 2)
    _numbered(d, "Passwords are changed every ninety days", bullets, 2)
    _numbered(d, "PASSWORD PROCEDURE", heads, 0, "Heading 1")
    _numbered(d, "Scope", heads, 1, "Heading 2")
    _numbered(d, "Annexes", again.numId, style="List Paragraph")
    _numbered(d, "An indented note.", None, ind='w:left="720" w:firstLine="360"')
    d.save(path)
    return path


def test_a_list_that_starts_deep_under_a_heading_reads_back_and_stays_edited_in_place(tmp_path):
    path = manual(str(tmp_path / "manual.docx"))
    imp = D.import_docx(path, twin=True)
    assert imp.native is not None                                                  # it used to read back as code: no editing in place
    assert "Passwords follow these rules:\n\n- Passwords are at least eight characters\n- Passwords are changed every ninety days\n" in imp.markdown
    md = imp.markdown.replace("- Passwords are changed every ninety days", "- Passwords are changed every ninety days\n- Passwords are never shared")
    data, _ = D.export_in_place(path, md)
    p = next(p for p in Document(io.BytesIO(data)).paragraphs if p.text == "Passwords are never shared")
    num_pr = p._p.pPr.find(qn("w:numPr"))
    assert num_pr.find(qn("w:ilvl")).get(qn("w:val")) == "2"                        # Word's level of the items beside it, not the first
    looks = L.for_blocks(data, md)
    look = looks[B.parse(md).index(next(b for b in B.parse(md) if b.text == "Passwords are never shared"))]
    assert (look["label"], look["left"], look["hang"]) == ("▪", 108.0, 18.0)


def test_each_block_gets_the_number_indent_and_look_word_gives_it(tmp_path):
    path = manual(str(tmp_path / "manual.docx"))
    md = D.import_docx(path, twin=True).markdown
    data, _ = D.export_in_place(path, md)
    looks = L.for_blocks(data, md + "\n---\n")                                    # a rule has no paragraph in Word: it gets nothing
    assert [x.get("label") for x in looks] == ["I.", "II.", "1.", "1.1.", "1.1.1.", "1.1.2.", "▪", "▪", "2.", "2.1.", "I.", None, None]
    assert [bool(x.get("body")) for x in looks][2:6] == [False, False, True, True]  # the clauses look like text, the headings do not
    assert looks[4]["left"] == 108.0 and looks[4]["hang"] == 36.0                    # the list level's indent (third level)
    assert (looks[11]["left"], looks[11]["hang"], "label" in looks[11]) == (36.0, -18.0, False) and looks[12] == {}   # a first-line indent; the rule


def test_the_studio_serves_the_layout_of_a_word_file_and_the_pdf_a_text_was_read_from(world, tmp_path):
    s = upload(world, manual(str(tmp_path / "manual.docx")), name="manual.docx", kind=DOCX)
    layout = world.http.get(world.api(f"/sessions/{s['id']}/layout")).json()
    assert layout["revision"] == s["revision"] and [x.get("label") for x in layout["blocks"]][:4] == ["I.", "II.", "1.", "1.1."]
    assert world.http.get(world.api(f"/sessions/{s['id']}/source-file")).status_code == 404       # a Word file is shown as Word draws it
    page = layout["page"]                                                                      # python-docx's Letter page, Cambria 11 pt
    assert (page["width"], page["left"], page["right"], page["top"], page["normal"]["font"], page["normal"]["size"]) == (612.0, 90.0, 90.0, 72.0, "Cambria", 11.0)
    assert layout["blocks"][2]["font"] == "Calibri" and layout["blocks"][2]["bold"] is True       # a heading: the theme's heading font
    from grc.modules.governance.studio import exporters

    pdf_path = tmp_path / "policy.pdf"
    pdf_path.write_bytes(exporters.to_pdf("# Policy\n\nThe bank protects information.\n", title="Policy"))
    p = upload(world, str(pdf_path), name="policy.pdf", kind="application/pdf")
    got = world.http.get(world.api(f"/sessions/{p['id']}/source-file"))
    assert got.status_code == 200 and got.headers["content-type"] == "application/pdf" and got.content == pdf_path.read_bytes()
    assert world.http.get(world.api(f"/sessions/{p['id']}/layout")).json() == {"revision": p["revision"], "blocks": [], "page": None}


def test_comply_agent_reads_the_numbers_word_shows_and_is_told_not_to_type_them(world, tmp_path):
    s = upload(world, manual(str(tmp_path / "manual.docx")), name="manual.docx", kind=DOCX)
    world.respond(lambda body: reply(answer([], "Clause 1.1.1 asks staff to protect the information they use.")))
    out = world.http.post(world.api(f"/sessions/{s['id']}/messages"), json={"text": "What does clause 1.1.1 say?"}).json()
    assert world.http.get(world.api(f"/sessions/{s['id']}/messages/{out['id']}")).json()["status"] == "done"
    prompt = last_body(world)["messages"][-1]["content"]
    assert "[B1 · I.] 1. Document Control" in prompt and "[B5 · 1.1.1.] ### All staff protect the information they use." in prompt
    assert "[B7] - Passwords are at least eight characters" in prompt               # a bullet: the Markdown shows it already
    assert "never write it into the text" in prompt


def _style(xml: str) -> None:
    return parse_xml(xml.replace("<w:style ", f"<w:style {nsdecls('w')} ", 1))


def typeset(path: str) -> str:
    """A file whose look comes from everywhere Word takes it: the document's defaults (python-docx's: the theme's minor font,
    11 pt, 10 pt after, 1.15 lines), a Normal style in Arial 10 pt, a style that names an East Asian font only (it must not
    change the Latin one), a style based on it, a theme font, runs that set their own look, a bold lead-in, spacing set on the
    paragraph, a theme colour, list items that leave out the space between them, a shaded table and an A4 page."""
    d = Document()
    d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Arial", Pt(10)
    styles = d.styles.element
    styles.append(_style('<w:style w:type="paragraph" w:styleId="Banner"><w:name w:val="Banner"/><w:basedOn w:val="Normal"/>'
                         '<w:pPr><w:spacing w:before="360" w:after="120" w:line="240" w:lineRule="auto"/><w:jc w:val="center"/></w:pPr>'
                         '<w:rPr><w:rFonts w:eastAsia="Times New Roman" w:cs="Arial"/><w:b/><w:sz w:val="32"/></w:rPr></w:style>'))
    styles.append(_style('<w:style w:type="paragraph" w:styleId="SubBanner"><w:name w:val="Sub Banner"/><w:basedOn w:val="Banner"/>'
                         '<w:rPr><w:sz w:val="24"/><w:u w:val="single"/></w:rPr></w:style>'))
    styles.append(_style('<w:style w:type="paragraph" w:styleId="Themed"><w:name w:val="Themed"/><w:basedOn w:val="Normal"/>'
                         '<w:rPr><w:rFonts w:asciiTheme="majorHAnsi" w:ascii="Courier New"/></w:rPr></w:style>'))
    d.add_paragraph("Big banner", style="Banner")
    d.add_paragraph("Sub banner", style="Sub Banner")
    r = d.add_paragraph(style="Banner").add_run("Plain banner")
    r.bold, r.font.size, r.font.color.rgb = False, Pt(10), RGBColor(0xFF, 0x00, 0x00)
    d.add_paragraph("Themed text", style="Themed")
    p = d.add_paragraph()
    p.add_run("Note:").bold = True
    p.add_run(" the rest of this paragraph is plain text.")
    p = d.add_paragraph("Justified with automatic space before.")
    ppr = p._p.get_or_add_pPr()
    ppr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:beforeAutospacing="1" w:line="280" w:lineRule="exact"/>'))
    ppr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="both"/>'))
    r = d.add_paragraph().add_run("Coloured by the theme")
    r.font.all_caps = True
    r._r.get_or_add_rPr().append(parse_xml(f'<w:color {nsdecls("w")} w:val="auto" w:themeColor="accent1"/>'))
    d.add_paragraph("First item", style="List Bullet")
    d.add_paragraph("Second item", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    for c, (head, body) in enumerate((("Role", "CISO"), ("Owner", "Board"))):
        t.cell(0, c).paragraphs[0].text = head
        t.cell(0, c)._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="1f497d"/>'))
        t.cell(1, c).paragraphs[0].add_run(body).font.size = Pt(9)
    sec = d.sections[0]
    sec.page_width, sec.left_margin, sec.right_margin, sec.top_margin = Twips(11906), Twips(1134), Twips(1134), Twips(1417)
    d.save(path)
    return path


def _looks(path: str):
    imp = D.import_docx(path, twin=True)
    assert imp.native is not None, imp.notes
    data, _ = D.export_in_place(path, imp.markdown)
    looks, page = L.read(data, imp.markdown)
    return {"table" if b.kind == "table" else b.text: look for b, look in zip(B.parse(imp.markdown), looks)}, page


def test_each_paragraph_gets_the_typography_word_resolves_through_its_styles(tmp_path):
    looks, _ = _looks(typeset(str(tmp_path / "typeset.docx")))
    # the style's own size, bold, centring and spacing; its East Asian font leaves Normal's Arial for Latin text
    assert looks["Big banner"] == {"font": "Arial", "size": 16.0, "bold": True, "align": "center", "before": 18.0, "after": 6.0, "line": 1.0}
    # a style based on it: what it sets itself wins, the rest comes from the style it is based on
    assert looks["Sub banner"] == {"font": "Arial", "size": 12.0, "bold": True, "underline": True, "align": "center", "before": 18.0, "after": 6.0, "line": 1.0}
    # the run's own size, colour and "not bold" beat the style
    assert looks["Plain banner"] == {"font": "Arial", "size": 10.0, "color": "#FF0000", "align": "center", "before": 18.0, "after": 6.0, "line": 1.0}
    # a theme font wins over a named font on the same element (python-docx's heading font is Calibri)
    assert looks["Themed text"]["font"] == "Calibri"
    # a bold lead-in is bold in the text itself (the Markdown shows it), not the whole paragraph
    assert "bold" not in looks["Note: the rest of this paragraph is plain text."]
    # justified, Word's automatic space before, an exact line height; the space after still from the document's defaults
    assert looks["Justified with automatic space before."] == {"font": "Arial", "size": 10.0, "align": "justify", "before": 14.0, "after": 10.0, "line": "14pt"}
    assert looks["Coloured by the theme"]["color"] == "#4F81BD" and looks["Coloured by the theme"]["caps"] is True
    # list items of one style leave out the space between them, as Word does (contextualSpacing)
    assert "after" not in looks["First item"] and looks["Second item"]["after"] == 10.0
    assert looks["table"] == {"head_fill": "#1F497D", "size": 9.0}                       # its header row's fill, its text's size


def test_the_page_is_the_files_last_section_with_its_normal_text(tmp_path):
    _, page = _looks(typeset(str(tmp_path / "typeset.docx")))
    assert page == {"width": 595.3, "left": 56.7, "right": 56.7, "top": 70.85, "normal": {"font": "Arial", "size": 10.0, "after": 10.0, "line": 1.15}}


def test_a_numbers_tab_to_a_list_stop_before_the_indent_is_where_its_text_starts(tmp_path):
    # Word's "1.3.1." at the margin with its indent at 0.85" and a list stop at 0.5": the text starts at 0.5" (as the View draws
    # it); a level with no such stop starts its text at the indent and gets no tab
    d = Document()
    numbering = d.part.numbering_part.element
    numbering.insert(0, _abstract(93, '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>'
                                      '<w:pPr><w:tabs><w:tab w:val="num" w:pos="720"/></w:tabs><w:ind w:left="1224" w:hanging="1224"/></w:pPr></w:lvl>'
                                      + _lvl(1, "decimal", "%1.%2.", 720, 720)))
    num = numbering.add_num(93).numId
    _numbered(d, "Usage Policy of Computing Resources", num, 0)
    _numbered(d, "All staff protect what they use.", num, 1)
    path = str(tmp_path / "tabs.docx")
    d.save(path)
    looks, _page = _looks(path)
    assert looks["Usage Policy of Computing Resources"]["tab"] == 36.0
    assert "tab" not in looks["All staff protect what they use."]
