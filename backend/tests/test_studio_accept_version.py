"""Accepting Comply Agent's changes stays in the studio: only Save to the library writes the library, as the minor or major version
the person chooses there (+0.1, or the next whole number, from the document's current version), and every version saved before
stays as it was - its number, text and file (the user, 2026-10-21: "it shouldn't auto saved to library once changes accepted
unless user clicked to save to library"; "made more changes and saved again it became version 2. and 1.2 gone??"). Without a
choice (an older screen, an API caller) Comply Agent decides, as before."""
import json

from docx import Document

import grc.models as m
from grc.modules.governance.routers import versions as version_routes

import _docx_fixtures as F
from _claude_mock import reply
from test_studio_files import lib, make_library_doc, upload          # noqa: F401  (lib is a fixture: the library's own folder for files)
from test_studio_flow import answer, change, world                  # noqa: F401  (world is a fixture)


def _ask(world, sid, changes, text="Change it."):
    world.respond(lambda body: reply(answer(changes, "Done.")))
    sent = world.http.post(world.api(f"/sessions/{sid}/messages"), json={"text": text}).json()
    return world.http.get(world.api(f"/sessions/{sid}/messages/{sent['id']}")).json()


def _accept(world, sid, msg):
    out = world.http.post(world.api(f"/sessions/{sid}/messages/{msg['id']}/accept"), json={})
    assert out.status_code == 200, out.text
    return out.json()


def _save(world, sid, change_type=None, reason=None):
    body = {k: v for k, v in (("change_type", change_type), ("change_reason", reason)) if v}
    out = world.http.post(world.api(f"/sessions/{sid}/save-to-document"), json=body)
    assert out.status_code == 200, out.text
    return out.json()


def _versions(db, doc_id):
    db.expire_all()
    rows = db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc_id).order_by(m.GovernanceDocumentVersion.id).all()
    return [(r.version_number, r.status) for r in rows]


def _row(db, doc_id, number):
    return db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc_id, version_number=number).one()


def _unsaved(world, sid):
    return world.http.get(world.api(f"/sessions/{sid}")).json()["unsaved_in_library"]


def test_an_accept_leaves_the_library_as_it_was_and_the_session_offers_save_to_the_library(world):
    draft = world.db.get(m.GovernanceDocument, 11).content
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 11}).json()
    out = _accept(world, s["id"], _ask(world, s["id"], [change("replace", "B3", markdown="Accounts shall be reviewed monthly.", reason="monthly")]))
    assert out["accepted"] == 1 and out["library_version"] is None and "reviewed monthly" in out["content"]     # the key stays for older screens
    world.db.expire_all()
    doc = world.db.get(m.GovernanceDocument, 11)
    assert (doc.current_version, doc.content, _versions(world.db, 11)) == ("1.0", draft, [])             # a draft too: nothing in the library
    assert _unsaved(world, s["id"]) is True                                                               # "not saved to the library yet"

    pub = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 10}).json()   # published: the same
    out = _accept(world, pub["id"], _ask(world, pub["id"], [change("replace", "B7", markdown="Access to systems shall be reviewed every quarter.", reason="quarterly")]))
    world.db.expire_all()
    published = world.db.get(m.GovernanceDocument, 10)
    assert out["library_version"] is None and "every quarter" in out["content"]
    assert (published.status, published.current_version, _versions(world.db, 10)) == ("published", "1.0", [])
    assert "every quarter" not in published.content and _unsaved(world, pub["id"]) is True


def test_each_save_is_the_version_the_person_chose_and_every_earlier_version_stays(world, monkeypatch):
    """The user's case: saved as 1.2, then saved again as a major change: 2.0 - and 1.2 is still there, as it was."""
    sid = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 11}).json()["id"]
    _accept(world, sid, _ask(world, sid, [change("replace", "B3", markdown="Accounts shall be reviewed monthly.", reason="monthly")],
                             text="Make the reviews monthly."))
    first = _save(world, sid, "minor", "Monthly reviews")
    assert (first["version"], first["change"], first["in_place"]) == ("1.1", "minor", False) and _unsaved(world, sid) is False
    _accept(world, sid, _ask(world, sid, [change("replace", "B3", markdown="Accounts shall be reviewed monthly by their owner.", reason="owner")]))
    assert _save(world, sid, "minor")["version"] == "1.2"
    one_one = _row(world.db, 11, "1.1")                                       # 1.1 kept: its number, its text, superseded
    assert (one_one.status, one_one.change_type, one_one.change_reason) == ("superseded", "minor", "Monthly reviews")
    assert "reviewed monthly." in one_one.content and "by their owner" not in one_one.content
    assert one_one.change_summary == "Edited in the Comply AI Studio: Monthly reviews. Minor change: Chosen by Olivia Owner when saving it to the library."
    item = one_one.change_set["items"][0]                                     # the accepted change keeps its author and the request
    assert (item["origin"], item["request"]) == ("ai", "Make the reviews monthly.")

    _accept(world, sid, _ask(world, sid, [change("insert", "B3", markdown="## 2. Exceptions\n\nExceptions need the CISO's approval.", reason="exceptions")]))
    third = _save(world, sid, "major")
    assert (third["version"], third["change"]) == ("2.0", "major")           # the next whole number from 1.2
    assert _versions(world.db, 11) == [("1.0", "superseded"), ("1.1", "superseded"), ("1.2", "superseded"), ("2.0", "current")]
    one_two = _row(world.db, 11, "1.2")
    assert "by their owner" in one_two.content and "Exceptions" not in one_two.content and one_two.status == "superseded"
    doc = world.db.get(m.GovernanceDocument, 11)
    assert doc.current_version == "2.0" and "Exceptions need the CISO's approval." in doc.content and _row(world.db, 11, "2.0").content == doc.content
    monkeypatch.setattr(version_routes, "get_user_tenants", lambda user, db: [1])
    listed = version_routes.list_versions(document_id=11, db=world.db, current_user=world.db.get(m.GRCUser, 1))
    assert sorted(v["version_number"] for v in listed["items"]) == ["1.0", "1.1", "1.2", "2.0"]      # every one of them is listed


def test_a_word_draft_saves_new_files_and_every_earlier_file_stays(world, lib):
    original = make_library_doc(world, lib, status="draft")
    before = open(original, "rb").read()
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 20}).json()
    block = next(i for i, part in enumerate(s["content"].split("\n\n"), start=1) if "Review it every year." in part)
    old = s["content"].split("\n\n")[block - 1]
    out = _accept(world, s["id"], _ask(world, s["id"], [change("replace", f"B{block}", markdown=old.replace("every year", "every six months"), reason="cycle")]))
    assert out["library_version"] is None and _unsaved(world, s["id"]) is True
    world.db.expire_all()
    assert world.db.get(m.GovernanceDocument, 20).file_path == original and _versions(world.db, 20) == [("1.0", "current")]

    assert _save(world, s["id"], "minor")["version"] == "1.1"
    doc = world.db.get(m.GovernanceDocument, 20)
    first_edit = doc.file_path
    assert first_edit != original and open(original, "rb").read() == before                       # the original file is untouched
    cur = world.http.get(world.api(f"/sessions/{s['id']}")).json()
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": cur["content"] + "\nStaff shall report incidents within one day.\n",
                                                                     "base_revision": cur["revision"]})
    assert _save(world, s["id"], "major")["version"] == "2.0"
    world.db.expire_all()
    rows = {r.version_number: r for r in world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=20)}
    assert {n: r.status for n, r in rows.items()} == {"1.0": "superseded", "1.1": "superseded", "2.0": "current"}
    assert (rows["1.0"].file_path, rows["1.1"].file_path, rows["2.0"].file_path) == (original, first_edit, world.db.get(m.GovernanceDocument, 20).file_path)
    texts = {n: "\n".join(p.text for p in Document(r.file_path).paragraphs) for n, r in rows.items()}
    assert "every six months" in texts["1.1"] and "report incidents" not in texts["1.1"]            # 1.1's file is as it was saved
    assert "every six months" in texts["2.0"] and "report incidents" in texts["2.0"] and "every six months" not in texts["1.0"]
    assert Document(rows["2.0"].file_path).sections[0].header.paragraphs[0].text == "ACME BANK - CONFIDENTIAL"   # still the same file


def test_without_a_choice_comply_agent_decides_as_before(world, monkeypatch):
    """An older screen or an API caller sends no change_type: Comply Agent decides, and the rule when it cannot answer."""
    sid = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 11}).json()["id"]
    _accept(world, sid, _ask(world, sid, [change("replace", "B3", markdown="Accounts shall be reviewed monthly.", reason="monthly")]))
    world.respond(lambda body: reply(json.dumps({"change": "major", "reason": "Reviews become monthly."})))
    out = _save(world, sid)
    assert (out["version"], out["change"]) == ("2.0", "major")
    assert _row(world.db, 11, "2.0").change_summary == "Edited in the Comply AI Studio. Major change: Reviews become monthly."

    monkeypatch.setattr("grc.services.ai_gateway.claude_client.configured", lambda: False)
    pub = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 10}).json()
    world.http.put(world.api(f"/sessions/{pub['id']}/content"), json={"content": pub["content"].replace("every year", "every twelve months"),
                                                                       "base_revision": pub["revision"]})
    assert _save(world, pub["id"])["version"] == "1.1"                         # a few words of a longer policy: minor, by the rule
    assert _row(world.db, 10, "1.1").change_summary.endswith("Minor change: A limited change to the wording. (decided by rule: Comply Agent was not available)")
    refused = world.http.post(world.api(f"/sessions/{sid}/save-to-document"), json={"change_type": "patch"})
    assert refused.status_code == 422                                          # minor or major only


def test_moving_an_upload_changed_in_the_studio_asks_for_the_second_version_and_a_new_document_stays_one_version(world, lib, tmp_path):
    path = F.corporate_policy(str(tmp_path / "policy.docx"))
    s = upload(world, path, name="Information Security Policy v2.0.docx")
    row = world.db.get(m.StudioSession, s["id"])                 # a session uploaded before an upload went into the library at once
    row.kind, row.document_id, row.settings = "upload", None, {}
    world.db.commit()
    assert world.http.get(world.api(f"/sessions/{s['id']}")).json()["original_version"] is None            # unchanged: one version
    paragraph = s["content"].split("\n\n")[1]
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": s["content"].replace(paragraph, paragraph.replace("every year", "every quarter")),
                                                                     "base_revision": s["revision"]})
    assert world.http.get(world.api(f"/sessions/{s['id']}")).json()["original_version"] == "1.0"         # its footer states v1.0
    moved = world.http.post(world.api(f"/sessions/{s['id']}/move-to-library"), json={"title": "Information Security Policy", "change_type": "major"}).json()
    first, second = world.db.query(m.GovernanceDocumentVersion).filter_by(document_id=moved["document_id"]).order_by(m.GovernanceDocumentVersion.id).all()
    assert (first.version_number, first.status, second.version_number, second.change_type, moved["version"]) == ("1.0", "superseded", "2.0", "major", "2.0")
    assert second.change_summary.endswith("Major change: Chosen by Olivia Owner when saving it to the library.")

    new = world.http.post(world.api("/sessions"), json={"kind": "new", "title": "Memo", "template_id": "builtin:blank", "doc_type": "other"}).json()
    world.http.put(world.api(f"/sessions/{new['id']}/content"), json={"content": "# Memo\n\nHello.\n", "base_revision": new["revision"]})
    assert world.http.get(world.api(f"/sessions/{new['id']}")).json()["original_version"] is None
    moved = world.http.post(world.api(f"/sessions/{new['id']}/move-to-library"), json={"change_type": "major"}).json()
    assert moved["version"] == "1.0" and _versions(world.db, moved["document_id"]) == [("1.0", "current")]     # nothing to choose


def test_a_document_not_in_the_library_has_no_versions(world):
    s = world.http.post(world.api("/sessions"), json={"kind": "new", "title": "Memo", "template_id": "builtin:blank", "doc_type": "other"}).json()
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": "# Memo\n\nHello.\n", "base_revision": s["revision"]})
    out = _accept(world, s["id"], _ask(world, s["id"], [change("replace", "B2", markdown="Hello, everyone.", reason="warmer")]))
    assert out["accepted"] == 1 and out["library_version"] is None


def test_a_request_to_format_offers_a_tidy_that_saves_as_the_next_version_and_can_be_taken_back(world, lib):
    from test_studio_word_fidelity import messy_policy

    make_library_doc(world, lib, status="draft", file_fixture=messy_policy)
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 20}).json()
    world.respond(lambda body: reply(answer([], "The headings and lists are right; only the spacing needs tidying.")))
    sent = world.http.post(world.api(f"/sessions/{s['id']}/messages"), json={"text": "Please format this document professionally."}).json()
    msg = world.http.get(world.api(f"/sessions/{s['id']}/messages/{sent['id']}")).json()
    (tidy,) = [o for o in msg["proposal"]["ops"] if o["kind"] == "tidy"]
    assert "15 extra spaces" in tidy["reason"] and "1 line break" in tidy["reason"] and "1 empty line at the end" in tidy["reason"]
    assert world.http.get(world.api(f"/sessions/{s['id']}/messages/{msg['id']}/redline")).json()["applied"] == []     # not a text change

    out = _accept(world, s["id"], msg)
    assert out["accepted"] == 1 and out["revision"] == s["revision"] + 1 and out["library_version"] is None
    assert _versions(world.db, 20) == [("1.0", "current")] and _unsaved(world, s["id"]) is True      # the tidied file waits for Save
    saved = _save(world, s["id"], "minor")
    assert (saved["version"], saved["change"]) == ("1.1", "minor")                                   # only the file changed
    doc = world.db.get(m.GovernanceDocument, 20)
    world.db.refresh(doc)
    texts = [p.text for p in Document(doc.file_path).paragraphs]
    assert texts[2] == "Purpose" and "This policy applies to all staff." in texts and texts[:2] == ["", ""]      # the layout lines stay

    session = world.db.get(m.StudioSession, s["id"])
    tidied = session.source_path
    back = world.http.post(world.api(f"/sessions/{s['id']}/messages/{msg['id']}/ops/tidy/revert"))
    assert back.status_code == 200, back.text
    world.db.expire_all()
    session = world.db.get(m.StudioSession, s["id"])
    assert session.source_path != tidied and "_tidied" not in session.source_path and (session.settings or {}).get("file_changed")

    plain = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 20}).json()
    world.respond(lambda body: reply(answer([], "Nothing to change.")))
    sent = world.http.post(world.api(f"/sessions/{plain['id']}/messages"), json={"text": "Is the scope clear?"}).json()
    assert not world.http.get(world.api(f"/sessions/{plain['id']}/messages/{sent['id']}")).json()["proposal"]["ops"]   # no tidy unasked


def test_tracked_changes_count_from_the_original_until_published_then_from_the_version_in_force(world):
    """The AI Studio's Tracked changes (GET .../tracked). Every save to the library moves where the session's own diff starts, so
    after a reload it showed nothing (the user, 2026-10-17: "it shows me empty even though the document is 1.2"). It counts from the
    library instead: the original until the document is published, then the version in force; with nothing since, what that
    version itself changed."""
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 11}).json()
    sid = s["id"]
    first = world.http.get(world.api(f"/sessions/{sid}/tracked")).json()
    assert first["hunks"] == [] and first["from_version_id"] is None              # no versions yet: since where the session started

    _accept(world, sid, _ask(world, sid, [change("replace", "B3", markdown="Accounts shall be reviewed monthly.", reason="monthly")]))
    _save(world, sid, "minor")
    _accept(world, sid, _ask(world, sid, [change("insert", "B3", markdown="## 2. Exceptions\n\nExceptions need the CISO's approval.",
                                                 reason="exceptions")]))
    _save(world, sid, "major")
    assert world.http.get(world.api(f"/sessions/{sid}/diff?from=0")).json()["hunks"] == []      # what it showed: nothing
    revisions = world.db.query(m.StudioRevision).filter_by(session_id=sid).count()
    t = world.http.get(world.api(f"/sessions/{sid}/tracked")).json()
    original = _row(world.db, 11, "1.0")
    assert t["label"] == "Changes since v1.0 (the original)" and t["from_version_id"] == original.id and t["can_revert"]
    assert len(t["hunks"]) == 2 and all(h["by"] and h["by"]["origin"] == "ai" for h in t["hunks"])    # who made each
    assert world.db.query(m.StudioRevision).filter_by(session_id=sid).count() == revisions           # a GET writes nothing

    doc = world.db.get(m.GovernanceDocument, 11)
    doc.status = "published"                                                       # 2.0 in force, nothing changed since
    world.db.commit()
    t = world.http.get(world.api(f"/sessions/{sid}/tracked")).json()
    assert t["label"] == "What v2.0 changed (since v1.1)" and len(t["hunks"]) == 1   # 1.1 is kept: 2.0 is compared with it

    # the change taken back, counted from that version: a new revision in the studio, not in the library yet
    exceptions = t["hunks"][0]
    out = world.http.post(world.api(f"/sessions/{sid}/changes/revert"), json={"hunks": [exceptions["id"]], "hash": t["hash"], "version_id": t["from_version_id"]})
    assert out.status_code == 200, out.text
    assert "reviewed monthly" in out.json()["content"] and "Exceptions need" not in out.json()["content"]
    t = world.http.get(world.api(f"/sessions/{sid}/tracked")).json()
    assert t["label"] == "Changes since v2.0, the version in force" and len(t["hunks"]) == 1
    wrong = world.http.post(world.api(f"/sessions/{sid}/changes/revert"), json={"hunks": [t["hunks"][0]["id"]], "hash": t["hash"], "version_id": 999999})
    assert wrong.status_code == 400                                                 # not a version of this document
