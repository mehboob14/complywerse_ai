"""The approval record (modules/governance/approval_record.py): every document carries - derived on demand, never stored - who reviewed
and approved this version, in which stages and order, when and with which drawn signature, and the workflow that ran from draft to
publication; it is exported on its own (PDF, Word) and added after the document's own pages (``export?with_record=true``).
Fixtures as in test_document_routes.py: Olivia owns and wrote the document, Rana and Rafay review, Mina and Mo approve, Vera only reads
the library. Workflow steps are driven through the sign-off endpoint functions; the record and its files over HTTP on in-memory SQLite."""
import hashlib
import io
import json
import re
from datetime import datetime, timedelta

import fitz
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.governance import approval_record
from grc.modules.governance import approval_record_export as X
from grc.modules.governance import doc_lifecycle as life
from grc.modules.governance import routes
from grc.modules.governance.router import router
from grc.modules.governance.routers import document_signoff as signoff
from grc.modules.governance.routers import documents
from grc.routers.auth_router import require_auth

from test_esignature import _png, _request

OLIVIA, RANA, RAFAY, MINA, MO, UMA, ADAM, VERA = 1, 2, 3, 4, 5, 6, 7, 8       # Vera only reads the library
APPROVERS = 10
TEXT = ("# Access policy\n\n## 1. Purpose\n\nThis policy sets how access to the bank's systems is requested, approved, granted, "
        "reviewed every quarter and removed when people move or leave, so that only the right people reach customer data.\n\n"
        "## 2. Scope\n\nAll staff, contractors, vendors and service accounts with access to production systems.\n")
SMALL = TEXT.replace("contractors, vendors", "contractors, temporary staff, vendors")
WHEN = re.compile(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2} UTC")
DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def _tenants(monkeypatch):
    monkeypatch.setenv("SESSION_SECRET", "unit-test-secret")
    monkeypatch.setattr(signoff, "get_user_tenants", lambda user, session: [1])


@pytest.fixture(scope="module")
def _schema():
    """The tables (about 450) made once for the module: each test starts from a copy of them, not from making them again."""
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    yield engine
    engine.dispose()


_TEST = {}                  # the database and the signed-in person of the test that is running
_app = FastAPI()            # the whole governance router is a few hundred routes: included once, not for every test
_app.include_router(router)
_app.dependency_overrides[m.get_db] = lambda: _TEST["db"]
_app.dependency_overrides[require_auth] = lambda: _TEST["db"].get(m.GRCUser, _TEST["me"]["id"])


@pytest.fixture
def world(tmp_path, _schema):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    source, target = _schema.raw_connection(), engine.raw_connection()
    try:
        source.driver_connection.backup(target.driver_connection)                 # sqlite copies the empty schema in a moment
    finally:
        source.close()
        target.close()
    db = sessionmaker(bind=engine)()
    db.add(m.Tenant(id=1, name="CFSB", slug="cfsb"))
    people = [(OLIVIA, "Olivia Owner", "Head of Policy"), (RANA, "Rana Reviewer", "Compliance Manager"), (RAFAY, "Rafay Reviewer", "Risk Officer"),
              (MINA, "Mina Manager", "CISO"), (MO, "Mo Manager", None), (UMA, "Uma User", "Analyst"), (ADAM, "Adam Admin", None),
              (VERA, "Vera Viewer", None)]
    for uid, name, title in people:
        db.add(m.GRCUser(id=uid, username=name.split()[0].lower(), email=f"{name.split()[0].lower()}@cfsb.test", display_name=name, designation=title))
    db.add(m.Role(id=APPROVERS, tenant_id=1, name="Approvers"))
    db.add(m.Role(id=11, tenant_id=1, name="Administrator"))
    for uid in (MINA, MO):
        db.add(m.UserRole(user_id=uid, role_id=APPROVERS, tenant_id=1))
    db.add(m.UserRole(user_id=ADAM, role_id=11, tenant_id=1))
    db.add(m.Role(id=900, tenant_id=1, name="Staff"))                       # everyone opens Task Management
    db.add(m.Permission(id=900, name=life.TASKS_VIEW, resource="critical_tasks:tasks", action="view"))
    db.add(m.RolePermission(role_id=900, permission_id=900))
    for uid in range(1, 9):
        db.add(m.UserRole(user_id=uid, role_id=900, tenant_id=1))
    db.add(m.Role(id=901, tenant_id=1, name="Library readers"))              # reads the library, sees no draft
    db.add(m.Permission(id=901, name="governance:policies:view", resource="governance:policies", action="view"))
    db.add(m.RolePermission(role_id=901, permission_id=901))
    db.add(m.UserRole(user_id=VERA, role_id=901, tenant_id=1))
    db.add(m.GovernanceDocument(id=5, tenant_id=1, title="Access policy", doc_type="policy", status="draft", current_version="1.0",
                                content=TEXT, owner_id=OLIVIA, author_id=OLIVIA, document_code="POL-005", classification="internal"))
    db.add(m.GovernanceDocumentVersion(document_id=5, version_number="1.0", change_type="major", title="Access policy", content=TEXT,
                                       change_summary="Initial version", status="current", created_by=OLIVIA, created_at=datetime(2026, 10, 1, 9, 0)))
    db.commit()
    me = {"id": OLIVIA}
    _TEST.update(db=db, me=me)                                                  # the app (made once) serves this test's database and person
    yield type("World", (), {"db": db, "engine": engine, "me": me, "http": TestClient(_app), "tmp": tmp_path})
    db.close()


# ── helpers ──────────────────────────────────────────────────────────────────
def user(w, uid):
    return w.db.get(m.GRCUser, uid)


def doc(w):
    return w.db.get(m.GovernanceDocument, 5)


def person(uid):
    return {"target_type": "user", "target_id": uid}


def role(rid=APPROVERS):
    return {"target_type": "role", "target_id": rid}


def stage(kind="review", people=(), **kw):
    return signoff.StageIn(kind=kind, people=[signoff.ParticipantIn(**p) for p in people], **kw)


def set_route(w, *stages, by=OLIVIA):
    return signoff.set_route(5, signoff.RouteUpdate(stages=list(stages)), w.db, user(w, by))


def usual(w, reviewers=(RANA,), approvers=(MINA,)):
    """The usual route: reviewers, then approvers."""
    return signoff.set_participants(5, signoff.ParticipantsUpdate(reviewers=[signoff.ParticipantIn(**person(u)) for u in reviewers],
                                                                  approvers=[signoff.ParticipantIn(**person(u)) for u in approvers]),
                                    w.db, user(w, OLIVIA))


def send(w, by=OLIVIA):
    return signoff.send_for_review(5, w.db, user(w, by))


def sign(w, uid, comment=None):
    return signoff.sign_document(5, signoff.SignRequest(signature_image=_png(), comment=comment), _request(), w.db, user(w, uid))


def reject(w, uid, comment="Not ready for approval."):
    return signoff.reject_document(5, signoff.RejectRequest(comment=comment), w.db, user(w, uid))


def ask_changes(w, uid, comment="Clarify who approves an exception in section 2."):
    return signoff.request_changes(5, signoff.ChangesRequest(comment=comment), w.db, user(w, uid))


def resubmit(w, major=False):
    return signoff.resubmit(5, signoff.ResubmitRequest(note="Section 2 now names the approver.", major=major), w.db, user(w, OLIVIA))


def publish(w, by=ADAM):
    documents.publish(w.db, doc(w), by, "Document published")
    w.db.commit()


def edit(w, text=SMALL, summary="Temporary staff added to the scope", by=OLIVIA, reason=None):
    """What a save of the author does: a new version (1.0 -> 1.1), and whatever the status says happens next. ``summary``: what the
    platform wrote about it; ``reason``: what the person wrote when they saved it (nothing, usually)."""
    d = doc(w)
    old, number, status = d.content, d.current_version, d.status
    d.content = text
    documents._snapshot_content_version(w.db, d, old_content=old, old_title=d.title, old_version=number, change_type="minor",
                                        change_reason=reason, user_id=by, summary=summary)
    life.changed_after_review(w.db, d, status, by, "Content edited")
    w.db.commit()


def approve_and_publish(w, reviewers=(RANA,), approvers=(MINA,)):
    usual(w, reviewers, approvers)
    send(w)
    for uid in reviewers:
        sign(w, uid)
    for uid in approvers:
        sign(w, uid)
    publish(w)


EVENTS = {"Prepared", "Sent for review", "Stage complete", "Changes requested", "Resubmitted", "Rejected", "Approved", "Published", "Archived"}
DECISIONS = {"signed", "rejected", "pending", "changes_requested", "delegated"}
STATES = {"published", "approved", "in_approval", "in_review", "draft", "archived"}


def check_contract(rec):
    """The JSON the screens code against: exactly these fields, and only these values where the values are a closed set."""
    assert list(rec) == ["reference", "generated_at", "document", "state", "round", "prepared_by", "stages", "milestones", "history", "notes",
                         "control", "revision_history", "approval_history"]
    assert set(rec["document"]) == {"id", "code", "title", "type", "classification", "version", "status", "department", "owner", "author",
                                    "effective_date", "next_review_date", "approved_at", "published_at", "published_by"}
    assert set(rec["control"]) == {"title", "code", "version", "type", "type_label", "classification", "owner", "department", "ownership",
                                   "approving_authority", "date_of_approval", "date_of_implementation", "review_frequency",
                                   "last_approval_date", "next_review_date", "status"}
    control = rec["control"]
    assert control["title"] == rec["document"]["title"] and control["version"] == rec["document"]["version"]
    assert control["date_of_approval"] == rec["document"]["approved_at"] or rec["state"] == "archived"       # (an archived document keeps the date of its approval)
    assert control["next_review_date"] == rec["document"]["next_review_date"]
    assert control["status"] in {"Published", "Approved, not published", "In approval", "In review", "Draft", "Archived"}      # words, never an id
    assert control["type_label"] is None or control["type_label"][:1].isupper()
    iso = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(\.\d+)?$")                                                        # naive UTC: no Z, no offset
    rows = rec["revision_history"]
    for r in rows:
        assert set(r) == {"version", "date", "change_type", "summary", "updated_by", "released"} and isinstance(r["released"], bool)
        assert r["change_type"] in {"Creation", "Major revision", "Minor revision", "Revision"} and (r["date"] is None or iso.match(r["date"]))
    assert all(r["released"] for r in rows[:-1]) and (not rows or rows[-1]["released"] or rows[-1]["version"] == rec["document"]["version"])    # not released: only the one worked on
    assert [r["change_type"] for r in rows][:1] in (["Creation"], [])
    assert {h["version"] for h in rec["history"]} <= {r["version"] for r in rows if r["released"]}
    assert [a["version"] for a in rec["approval_history"]] == [r["version"] for r in rows if r["released"]]              # one for each released version, in the same order
    for a in rec["approval_history"]:
        assert set(a) == {"version", "reviewed_by", "approved_by", "approval_date"} and (a["approval_date"] is None or iso.match(a["approval_date"]))
        assert all(set(x) == {"name", "designation"} and x["name"] for x in a["reviewed_by"])
        assert all(set(x) == {"name", "designation", "stage"} and x["name"] and x["stage"] not in ("Approval", "Approved by", "Review", "Reviewed by")
                   for x in a["approved_by"])                                                                                  # a stage is named by its route, never by its label
    for s in rec["stages"]:
        assert all(isinstance(p["name"], str) and p["name"] for p in s["people"])                                           # nobody is listed without a name
    assert set(rec["round"]) == {"version", "started_at", "completed_at", "complete"} and set(rec["prepared_by"]) == {"name", "designation", "at"}
    assert rec["state"] in STATES and isinstance(rec["round"]["complete"], bool) and all(isinstance(n, str) for n in rec["notes"])
    assert [s["position"] for s in rec["stages"]] == list(range(1, len(rec["stages"]) + 1)) and len(rec["stages"]) <= routes.MAX_STAGES
    for s in rec["stages"]:
        assert set(s) == {"position", "label", "name", "rule", "min_count", "status", "people"}
        assert s["rule"] in routes.RULES and s["status"] in {"complete", "pending", "changes_requested", "rejected"}
        for p in s["people"]:
            assert set(p) == {"name", "designation", "decision", "meaning", "signed_at", "comment", "on_behalf_of", "signature"} and p["decision"] in DECISIONS
            assert isinstance(p["name"], str) and isinstance(p["meaning"], str)
            assert p["signature"] is None or set(p["signature"]) == {"id", "drawn", "seal_short"}
    for x in rec["milestones"]:
        assert set(x) == {"at", "event", "actor", "detail"} and x["event"] in EVENTS and isinstance(x["detail"], str) and len(rec["milestones"]) <= 100
    assert [x["at"] for x in rec["milestones"]] == sorted(x["at"] for x in rec["milestones"])               # oldest first
    for h in rec["history"]:
        assert set(h) == {"version", "approved_at", "published_at", "summary", "prepared_by", "approvers"}
        assert all(set(a) == {"name", "designation", "meaning", "signed_at"} for a in h["approvers"])
    json.dumps(rec)                                                                                          # plain JSON all through
    return rec


def record(w, **kw):
    return check_contract(approval_record.build(w.db, doc(w), **kw))


def get(w, path, who=OLIVIA, **params):
    w.me["id"] = who
    return w.http.get(f"/governance/documents/5{path}", params=params)


def names(stage_):
    return [p["name"] for p in stage_["people"]]


def decisions(stage_):
    return [(p["name"], p["decision"]) for p in stage_["people"]]


def events(rec):
    return [x["event"] for x in rec["milestones"]]


def pdf_text(data):
    with fitz.open(stream=data, filetype="pdf") as pdf:
        return "\n".join(page.get_text() for page in pdf), pdf.page_count


def blob(rec):
    return json.dumps(rec, ensure_ascii=False)


# ── the root cause: signatures are tied to their version ────────────────────────
def test_a_signature_and_a_rejection_are_tied_to_the_version_row_they_were_given_on(world):
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    v10 = world.db.query(m.GovernanceDocumentVersion.id).filter_by(document_id=5, version_number="1.0").scalar()
    assert world.db.query(m.DocumentSignature).one().version_id == v10
    reject(world, MINA)                                                      # rejecting is tied to it too
    assert [s.version_id for s in world.db.query(m.DocumentSignature).order_by(m.DocumentSignature.id)] == [v10, v10]
    edit(world)                                                              # 1.1: the next signatures are given on that row
    send(world)
    sign(world, RANA)
    v11 = world.db.query(m.GovernanceDocumentVersion.id).filter_by(document_id=5, version_number="1.1").scalar()
    assert world.db.query(m.DocumentSignature).order_by(m.DocumentSignature.id.desc()).first().version_id == v11 != v10


def test_the_version_is_found_on_a_session_that_does_not_autoflush_even_when_the_stamp_has_just_made_it(world):
    """A document with no version row and an Approval Signoff table: the first sign-off stamp writes its baseline and its next version
    in the same transaction (documents._snapshot_content_version); tenant sessions do not flush by themselves (PITFALLS 74)."""
    table = "| Role | Name | Designation | Signature / Date |\n| --- | --- | --- | --- |\n| Reviewed by | ______ | ______ | ______ |"
    world.db.add(m.GovernanceDocument(id=6, tenant_id=1, title="Leave policy", doc_type="policy", status="draft", current_version="1.0",
                                      content=f"# Leave policy\n\n{table}\n", owner_id=OLIVIA, author_id=OLIVIA))
    world.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=6, role_type="reviewer", target_type="user", target_id=RANA))
    world.db.commit()
    db = Session(world.engine, autoflush=False)
    try:
        signoff.send_for_review(6, db, db.get(m.GRCUser, OLIVIA))
        signoff.sign_document(6, signoff.SignRequest(signature_image=_png()), _request(), db, db.get(m.GRCUser, RANA))
        sig = db.query(m.DocumentSignature).filter_by(document_id=6).one()
        row = db.get(m.GovernanceDocumentVersion, sig.version_id)
        assert row is not None and row.version_number == db.get(m.GovernanceDocument, 6).current_version      # the row the stamp made
    finally:
        db.close()


def test_a_document_without_version_rows_has_signatures_without_a_version_id(world):
    world.db.query(m.GovernanceDocumentVersion).delete()
    world.db.commit()
    usual(world)
    send(world)
    sign(world, RANA)
    assert world.db.query(m.DocumentSignature).one().version_id is None


# ── the usual route, complete ───────────────────────────────────────────────────
def test_the_usual_route_complete_names_both_stages_their_people_times_and_signatures(world):
    usual(world, (RANA, RAFAY), (MINA,))
    send(world)
    sign(world, RANA, "Looks right.")
    sign(world, RAFAY)
    sign(world, MINA)
    publish(world)
    rec = record(world)
    assert rec["state"] == "published"
    assert re.fullmatch(r"DAR-5-1\.0-[0-9A-F]{8}", rec["reference"])
    assert [(s["position"], s["label"], s["name"], s["rule"], s["status"]) for s in rec["stages"]] == [
        (1, "Reviewed by", None, "all", "complete"), (2, "Approved by", None, "all", "complete")]
    review, approval = rec["stages"]
    assert [(p["name"], p["designation"], p["decision"], p["meaning"]) for p in review["people"]] == [
        ("Rana Reviewer", "Compliance Manager", "signed", "Reviewed by"), ("Rafay Reviewer", "Risk Officer", "signed", "Reviewed by")]
    assert review["people"][0]["comment"] == "Looks right." and review["people"][1]["comment"] is None
    assert [(p["name"], p["designation"]) for p in approval["people"]] == [("Mina Manager", "CISO")]
    esigs = {e.record_id: e for e in world.db.query(m.ElectronicSignature).filter_by(record_type="document_signoff")}
    drawn = {user(world, s.signer_user_id).display_name: esigs[s.id] for s in world.db.query(m.DocumentSignature)}       # each signer's drawn signature
    for stage_ in rec["stages"]:
        for p in stage_["people"]:
            e = drawn[p["name"]]
            assert p["signature"] == {"id": e.id, "drawn": True, "seal_short": e.seal[:8].upper()} and p["signed_at"] and p["on_behalf_of"] is None
    assert rec["prepared_by"] == {"name": "Olivia Owner", "designation": "Head of Policy", "at": "2026-10-01T09:00:00"}
    d = rec["document"]
    assert (d["id"], d["code"], d["title"], d["type"], d["classification"], d["version"], d["status"], d["owner"], d["author"]) == (
        5, "POL-005", "Access policy", "policy", "internal", "1.0", "published", "Olivia Owner", "Olivia Owner")
    assert d["approved_at"] and d["published_at"] and d["published_by"] == "Adam Admin" and d["department"] is None
    assert rec["round"] == {"version": "1.0", "started_at": doc(world).signoff_round_at.isoformat(), "completed_at": d["approved_at"], "complete": True}
    assert events(rec) == ["Prepared", "Sent for review", "Stage complete", "Stage complete", "Approved", "Published"]
    assert [x["actor"] for x in rec["milestones"]][:2] == ["Olivia Owner", "Olivia Owner"] and rec["milestones"][-1]["actor"] == "Adam Admin"
    assert [(h["version"], h["summary"], h["prepared_by"], [a["name"] for a in h["approvers"]]) for h in rec["history"]] == [
        ("1.0", "Initial version", "Olivia Owner", ["Mina Manager"])]
    assert rec["history"][0]["approvers"][0]["meaning"] == "Approved by" and rec["history"][0]["published_at"] == d["published_at"]
    assert rec["notes"] == []
    assert approval_record.build(world.db, doc(world))["reference"] == rec["reference"]          # the same facts: the same reference


def test_the_reference_follows_the_content_and_the_times_are_naive_utc(world):
    usual(world)
    send(world)
    sign(world, RANA)
    before = record(world)
    sign(world, MINA)
    after = record(world)
    assert before["reference"] != after["reference"] and before["reference"].startswith("DAR-5-1.0-")
    body = {k: v for k, v in after.items() if k not in ("reference", "generated_at")}
    canon = json.dumps(body, sort_keys=True, separators=(",", ":"), ensure_ascii=False, default=str)
    assert after["reference"] == f"DAR-5-1.0-{hashlib.sha256(canon.encode('utf-8')).hexdigest()[:8].upper()}"
    for iso in (after["generated_at"], after["round"]["started_at"], after["milestones"][1]["at"], after["stages"][0]["people"][0]["signed_at"]):
        assert "Z" not in iso and "+" not in iso and datetime.fromisoformat(iso).tzinfo is None


def test_a_draft_nobody_signed_still_has_a_record(world):
    rec = record(world)
    assert rec["state"] == "draft" and rec["stages"] == [] and rec["history"] == [] and rec["notes"] == []
    assert rec["prepared_by"] == {"name": "Olivia Owner", "designation": "Head of Policy", "at": "2026-10-01T09:00:00"}
    assert rec["round"] == {"version": "1.0", "started_at": None, "completed_at": None, "complete": False}
    assert [(x["event"], x["actor"], x["detail"]) for x in rec["milestones"]] == [("Prepared", "Olivia Owner", "Initial version")]
    assert rec["document"]["approved_at"] is None and rec["document"]["published_at"] is None
    world.db.query(m.GovernanceDocumentVersion).delete()                      # not even a version row: the document's own author and date
    world.db.commit()
    assert approval_record.build(world.db, doc(world))["prepared_by"]["name"] == "Olivia Owner"


def test_a_document_sent_with_nobody_yet_signed_lists_who_is_asked(world):
    usual(world, (RANA,), (MINA,))
    send(world)
    rec = record(world)
    assert rec["state"] == "in_review" and rec["round"]["complete"] is False and rec["document"]["approved_at"] is None
    assert [(s["label"], s["status"], decisions(s)) for s in rec["stages"]] == [
        ("Reviewed by", "pending", [("Rana Reviewer", "pending")]), ("Approved by", "pending", [("Mina Manager", "pending")])]
    assert rec["stages"][0]["people"][0]["signed_at"] is None and rec["stages"][0]["people"][0]["signature"] is None
    assert rec["stages"][0]["people"][0]["designation"] == "Compliance Manager" and events(rec) == ["Prepared", "Sent for review"]


def test_designations_are_the_signatures_snapshot_and_the_current_profile_only_where_there_is_none(world):
    approve_and_publish(world, (RANA,), (MO,))
    world.db.get(m.GRCUser, RANA).designation = "Chief Risk Officer"           # promoted since: the signature keeps the title it was given with
    world.db.get(m.GRCUser, MO).designation = "Head of Operations"             # had none when signing: the current one, never an invented one
    world.db.get(m.GRCUser, UMA).designation = None
    world.db.commit()
    rec = record(world)
    assert rec["stages"][0]["people"][0]["designation"] == "Compliance Manager"
    assert rec["stages"][1]["people"][0]["designation"] == "Head of Operations"
    world.db.get(m.GRCUser, MO).designation = None
    world.db.commit()
    assert record(world)["stages"][1]["people"][0]["designation"] is None


# ── routes of their own ─────────────────────────────────────────────────────────
def test_a_route_of_three_stages_lists_every_stage_in_order_with_its_name_rule_and_people(world):
    set_route(world, stage("review", [person(RANA), person(RAFAY)], name="Compliance review", rule="at_least", min_count=1),
              stage("concurrence", [person(UMA)], name="Legal concurrence"),
              stage("approval", [role()], name="Board approval", rule="any"))
    send(world)
    sign(world, RAFAY)
    sign(world, UMA)
    sign(world, MO)
    rec = record(world)
    assert rec["state"] == "approved" and rec["round"]["complete"] is True
    assert [(s["position"], s["label"], s["name"], s["rule"], s["min_count"], s["status"]) for s in rec["stages"]] == [
        (1, "Reviewed by", "Compliance review", "at_least", 1, "complete"), (2, "Concurred by", "Legal concurrence", "all", None, "complete"),
        (3, "Approved by", "Board approval", "any", None, "complete")]
    assert [names(s) for s in rec["stages"]] == [["Rafay Reviewer"], ["Uma User"], ["Mo Manager"]]       # Rana was not needed, Mina not asked
    assert [p["meaning"] for s in rec["stages"] for p in s["people"]] == ["Reviewed by", "Concurred by", "Approved by"]
    assert events(rec) == ["Prepared", "Sent for review", "Stage complete", "Stage complete", "Stage complete", "Approved"]
    assert [x["detail"] for x in rec["milestones"] if x["event"] == "Stage complete"] == [
        "Compliance review: Rafay Reviewer", "Legal concurrence: Uma User", "Board approval: Mo Manager"]


def test_twelve_stages_all_render_in_order(world):
    assert routes.MAX_STAGES == 12
    crew = [RANA, RAFAY, UMA, MO]
    stages = [stage("review", [person(crew[i % 4])], name=f"Stage {i + 1}") for i in range(11)] + [stage("approval", [person(MINA)], name="Final approval")]
    set_route(world, *stages)
    send(world)
    for i in range(11):
        sign(world, crew[i % 4])
    mid = record(world)
    assert [s["status"] for s in mid["stages"]] == ["complete"] * 11 + ["pending"] and mid["state"] == "in_approval"
    sign(world, MINA)
    rec = record(world)
    assert [s["position"] for s in rec["stages"]] == list(range(1, 13)) and [s["name"] for s in rec["stages"]] == [f"Stage {i}" for i in range(1, 12)] + ["Final approval"]
    assert [names(s) for s in rec["stages"]] == [[user(world, crew[i % 4]).display_name] for i in range(11)] + [["Mina Manager"]]
    assert rec["round"]["complete"] is True and events(rec).count("Stage complete") == 12 and len(rec["history"][0]["approvers"]) == 1
    pdf, pages = pdf_text(X.to_pdf(rec, {}))
    assert "STAGE 12 — APPROVED BY · Final approval" in pdf and "STAGE 1 — REVIEWED BY · Stage 1" in pdf and pages >= 2


def test_people_who_have_not_signed_are_pending_with_a_group_as_one_entry(world):
    world.db.add(m.Team(id=1, tenant_id=1, name="Policy board", is_active=True))
    world.db.add(m.TeamMember(team_id=1, user_id=UMA))
    world.db.add(m.TeamMember(team_id=1, user_id=ADAM))
    world.db.commit()
    set_route(world, stage("review", [person(RANA), person(RAFAY), person(UMA)], name="Review", order="in_turn"),
              stage("approval", [role(), {"target_type": "team", "target_id": 1}], name="Approval"))
    send(world)
    sign(world, RANA)
    rec = record(world)
    assert rec["state"] == "in_review"
    review, approval = rec["stages"]
    assert decisions(review) == [("Rana Reviewer", "signed"), ("Rafay Reviewer", "pending"), ("Uma User", "pending")]      # in turn: Rafay is next
    assert decisions(approval) == [("Role: Approvers", "pending"), ("Department: Policy board", "pending")]                      # a group is one entry
    assert [s["status"] for s in rec["stages"]] == ["pending", "pending"]
    assert approval["people"][0]["designation"] is None and approval["people"][0]["signed_at"] is None and approval["people"][0]["meaning"] == "Approved by"
    sign(world, RAFAY)
    sign(world, UMA)
    rec = record(world)
    assert rec["state"] == "in_approval" and [s["status"] for s in rec["stages"]] == ["complete", "pending"]
    assert decisions(rec["stages"][1]) == [("Role: Approvers", "pending"), ("Department: Policy board", "pending")] and rec["document"]["approved_at"] is None
    sign(world, MINA)                                                          # a role is satisfied by one member: the entry becomes the signer
    rec = record(world)
    assert decisions(rec["stages"][1]) == [("Mina Manager", "signed"), ("Department: Policy board", "pending")] and rec["state"] == "in_approval"
    sign(world, ADAM)                                                          # ... and so is a team
    rec = record(world)
    assert decisions(rec["stages"][1]) == [("Mina Manager", "signed"), ("Adam Admin", "signed")] and rec["state"] == "approved"


def test_a_step_handed_over_shows_who_signed_for_whom(world):
    set_route(world, stage("review", [person(RANA)], name="Review"), stage("approval", [person(MINA)], name="Approval"))
    send(world)
    signoff.delegate(5, signoff.DelegateRequest(user_id=MO, note="On leave"), world.db, user(world, RANA))
    waiting = record(world)["stages"][0]["people"]
    assert [(p["name"], p["decision"], p["comment"]) for p in waiting] == [("Rana Reviewer", "delegated", "Handed over to Mo Manager")]
    sign(world, MO)
    done = record(world)["stages"][0]["people"]
    assert [(p["name"], p["decision"], p["on_behalf_of"], p["meaning"]) for p in done] == [("Mo Manager", "signed", "Rana Reviewer", "Reviewed by")]


def test_stage_names_are_read_from_the_route_only_while_that_stage_still_is_what_it_was(world):
    set_route(world, stage("review", [person(RANA)], name="Compliance review"), stage("approval", [person(MINA)], name="Board approval"))
    send(world)
    sign(world, RANA)
    sign(world, MINA)
    publish(world)
    assert [s["name"] for s in record(world)["stages"]] == ["Compliance review", "Board approval"]
    set_route(world, stage("review", [person(RAFAY)], name="Department review"), stage("approval", [person(MO)], name="Executive sign-off"))
    rec = record(world)                                                         # the route of the next version has the same shape: its names are not this round's
    assert [(s["label"], s["name"]) for s in rec["stages"]] == [("Reviewed by", None), ("Approved by", None)] and [names(s) for s in rec["stages"]] == [["Rana Reviewer"], ["Mina Manager"]]
    set_route(world, stage("approval", [person(MO)], name="CISO sign-off"))      # and one where stage 1 is another kind and stage 2 is gone
    rec = record(world)
    assert [(s["label"], s["name"], s["rule"]) for s in rec["stages"]] == [("Reviewed by", None, "all"), ("Approved by", None, "all")]
    assert [names(s) for s in rec["stages"]] == [["Rana Reviewer"], ["Mina Manager"]] and rec["round"]["complete"] is True      # the signatures speak for themselves
    set_route(world)                                                           # back to the usual route
    assert [(s["label"], s["name"]) for s in record(world)["stages"]] == [("Reviewed by", None), ("Approved by", None)]


# ── changes asked for, rejection, withdrawal ───────────────────────────────────
def test_changes_asked_for_hold_a_stage_and_a_resubmitted_stage_does_not_count_the_signature_it_voided(world):
    set_route(world, stage("review", [person(RANA), person(RAFAY), person(UMA)], name="Review", order="in_turn"),
              stage("approval", [role()], name="Approval", rule="any"))
    send(world)
    sign(world, RANA)
    ask_changes(world, RAFAY)
    held = record(world)
    assert held["state"] == "in_review" and [s["status"] for s in held["stages"]] == ["changes_requested", "pending"]
    assert [(p["name"], p["decision"]) for p in held["stages"][0]["people"]] == [
        ("Rana Reviewer", "signed"), ("Rafay Reviewer", "changes_requested"), ("Uma User", "pending")]
    asked = held["stages"][0]["people"][1]
    assert asked["comment"] == "Clarify who approves an exception in section 2." and asked["signed_at"] and asked["signature"] is None
    assert "Changes requested" in events(held) and "Resubmitted" not in events(held)

    edit(world)                                                               # a small change while it is held: it starts that stage again
    resubmit(world)
    again = record(world)
    assert [(p["name"], p["decision"]) for p in again["stages"][0]["people"]] == [
        ("Rana Reviewer", "pending"), ("Rafay Reviewer", "pending"), ("Uma User", "pending")]      # Rana's earlier signature no longer counts
    assert any("1 earlier signature was given before a stage was asked again" in n for n in again["notes"])
    assert events(again).count("Changes requested") == 1 and events(again).count("Resubmitted") == 1
    assert "Stage complete" not in events(again)

    for uid in (RANA, RAFAY, UMA):
        sign(world, uid)
    sign(world, MINA)
    rec = record(world)
    assert rec["state"] == "approved" and rec["round"]["complete"] is True
    assert [(p["name"], p["decision"]) for p in rec["stages"][0]["people"]] == [
        ("Rana Reviewer", "signed"), ("Rafay Reviewer", "signed"), ("Uma User", "signed")]               # each once: the voided signature is not listed
    assert world.db.query(m.DocumentSignature).filter_by(signer_user_id=RANA).count() == 2               # ... it is kept (append-only), in the Change log
    assert len([p for s in rec["stages"] for p in s["people"] if p["signature"]]) == 4
    assert events(rec).count("Approved") == 1 and events(rec).count("Changes requested") == 1 and events(rec).count("Resubmitted") == 1
    assert any("1 earlier signature was given" in n for n in rec["notes"])
    assert rec["history"][0]["version"] == "1.1" and [a["name"] for a in rec["history"][0]["approvers"]] == ["Mina Manager"]


def test_a_rejection_is_recorded_on_the_draft_and_belongs_to_the_version_it_was_given_on(world):
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    reject(world, MINA, "The scope is too wide.")
    rec = record(world)
    assert rec["state"] == "draft" and rec["round"]["complete"] is False and rec["document"]["approved_at"] is None
    assert [(s["label"], s["status"]) for s in rec["stages"]] == [("Reviewed by", "complete"), ("Approved by", "rejected")]
    rejected = rec["stages"][1]["people"][0]
    assert (rejected["name"], rejected["decision"], rejected["comment"], rejected["meaning"]) == ("Mina Manager", "rejected", "The scope is too wide.", "Approved by")
    assert rejected["signature"] is None and rejected["signed_at"]                       # a rejection is not drawn: no signature
    assert events(rec) == ["Prepared", "Sent for review", "Stage complete", "Rejected"] and rec["history"] == []
    edit(world)                                                                          # a new version made as a draft: nothing was signed on it
    fresh = record(world)
    assert fresh["document"]["version"] == "1.1" and fresh["stages"] == [] and fresh["state"] == "draft"
    old = record(world, version="1.0")
    assert old["stages"][1]["status"] == "rejected" and old["state"] == "draft"


def test_a_withdrawn_round_leaves_a_draft_and_a_new_round_starts_clean(world):
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    signoff.withdraw(5, signoff.WithdrawRequest(reason="Rewriting it"), world.db, user(world, OLIVIA))
    rec = record(world)
    assert rec["state"] == "draft" and [(s["label"], s["status"]) for s in rec["stages"]] == [("Reviewed by", "complete")]
    assert events(rec) == ["Prepared", "Sent for review", "Stage complete"]                  # a withdrawal is not one of the events the screens know
    send(world)                                                                          # sent again: a new round, nothing carried over
    again = record(world)
    assert again["state"] == "in_review" and decisions(again["stages"][0]) == [("Rana Reviewer", "pending")]
    assert events(again) == ["Prepared", "Sent for review"] and again["notes"] == []


def test_a_major_change_resubmitted_starts_a_new_round_and_the_record_starts_with_it(world):
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA)
    ask_changes(world, MINA)
    edit(world, TEXT + "\n## 3. Exceptions\n\nAn exception is approved by the CISO for at most ninety days.\n", summary="Exceptions added")
    assert resubmit(world)["change"] == "major"
    rec = record(world)
    assert rec["state"] == "in_review" and decisions(rec["stages"][0]) == [("Rana Reviewer", "pending")]       # signed again from the start
    assert "Stage complete" not in events(rec) and "Changes requested" not in events(rec) and "Resubmitted" in events(rec)
    assert rec["notes"] == []                                              # Rana's earlier signature belongs to the round before this one
    sign(world, RANA)
    sign(world, MINA)
    done = record(world)
    assert done["round"]["complete"] is True and [names(s) for s in done["stages"]] == [["Rana Reviewer"], ["Mina Manager"]]
    assert [h["version"] for h in done["history"]] == ["1.1"] and events(done).count("Approved") == 1


# ── versions, history, the life cycle again ──────────────────────────────────────
def test_two_approved_versions_history_has_both_newest_first_and_the_record_follows_the_new_round(world):
    approve_and_publish(world, (RANA,), (MINA,))
    first = record(world)
    edit(world)                                                                          # 1.1 goes through the life cycle again
    usual(world, (RAFAY,), (MO,))                                                        # with other people this time
    mid = record(world)
    assert mid["state"] == "in_review" and mid["document"]["version"] == "1.1" and mid["document"]["approved_at"] is None
    assert [names(s) for s in mid["stages"]] == [["Rafay Reviewer"], ["Mo Manager"]]
    assert [h["version"] for h in mid["history"]] == ["1.0"] and "Rana Reviewer" not in blob(mid["stages"]) and "Stage complete" not in events(mid)
    assert mid["prepared_by"]["at"] != first["prepared_by"]["at"] and mid["milestones"][0]["detail"] == "Temporary staff added to the scope"
    sign(world, RAFAY)
    sign(world, MO)
    publish(world)
    rec = record(world)
    assert rec["state"] == "published" and rec["round"]["version"] == "1.1" and rec["round"]["complete"] is True
    assert [names(s) for s in rec["stages"]] == [["Rafay Reviewer"], ["Mo Manager"]]
    assert [h["version"] for h in rec["history"]] == ["1.1", "1.0"]
    new, old = rec["history"]
    assert [a["name"] for a in new["approvers"]] == ["Mo Manager"] and [a["name"] for a in old["approvers"]] == ["Mina Manager"]
    assert (new["summary"], new["prepared_by"], old["summary"]) == ("Temporary staff added to the scope", "Olivia Owner", "Initial version")
    assert new["approved_at"] > old["approved_at"] and old["published_at"] and new["published_at"] and old["published_at"] < new["approved_at"]
    assert events(rec)[:2] == ["Prepared", "Sent for review"] and events(rec).count("Published") == 1             # this round only: not 1.0's publication
    # the record of the earlier version, asked for by number, is that version's round
    was = record(world, version="1.0")
    assert was["document"]["version"] == "1.0" and was["state"] == "published" and was["round"]["complete"] is True
    assert [names(s) for s in was["stages"]] == [["Rana Reviewer"], ["Mina Manager"]] and [h["version"] for h in was["history"]] == ["1.0"]
    assert was["document"]["published_at"] == first["document"]["published_at"] and was["document"]["published_by"] == "Adam Admin"
    assert was["prepared_by"] == first["prepared_by"] and events(was) == events(first)
    assert approval_record.build(world.db, doc(world), version="v1.1")["reference"] == rec["reference"]


def test_a_round_that_spanned_two_versions_reads_the_same_when_it_is_the_newest_and_when_it_is_an_earlier_version(world):
    """Stage 1 signed version 1.0 and stands (earlier stages' signatures do); stage 2 asked for changes, so version 1.1 was made and
    stage 2 started again (Uma's signature on 1.0 no longer counts); everyone signed 1.1. The record must read the same afterwards,
    when 1.2 has begun its own round and this one is read as an earlier version."""
    set_route(world, stage("review", [person(RANA)], name="First review"),
              stage("review", [person(RAFAY), person(UMA)], name="Second review"), stage("approval", [person(MINA)], name="Approval"))
    send(world)
    sign(world, RANA)
    sign(world, UMA)
    ask_changes(world, RAFAY)
    edit(world)
    resubmit(world)
    sign(world, RAFAY)
    sign(world, UMA)
    sign(world, MINA)
    publish(world)
    newest = record(world)
    assert newest["document"]["version"] == "1.1" and newest["round"]["complete"] is True
    assert [names(s) for s in newest["stages"]] == [["Rana Reviewer"], ["Rafay Reviewer", "Uma User"], ["Mina Manager"]]
    assert any("1 earlier signature was given" in n for n in newest["notes"])
    ids = {s.signer_user_id: s.version_id for s in world.db.query(m.DocumentSignature).order_by(m.DocumentSignature.id)}      # (the last one each)
    versions = dict(world.db.query(m.GovernanceDocumentVersion.id, m.GovernanceDocumentVersion.version_number))
    assert versions[ids[RANA]] == "1.0" and versions[ids[UMA]] == "1.1" and versions[ids[MINA]] == "1.1"      # the round holds both versions
    edit(world, SMALL + "\nAnother clause.\n")                                                              # 1.2 begins another round
    assert record(world)["document"]["version"] == "1.2" and record(world)["stages"][0]["status"] == "pending"
    was = record(world, version="1.1")
    for key in ("stages", "milestones", "history", "notes", "state", "prepared_by"):
        assert was[key] == newest[key], key
    assert {k: v for k, v in was["document"].items()} == newest["document"]
    assert {k: v for k, v in was["round"].items() if k != "started_at"} == {k: v for k, v in newest["round"].items() if k != "started_at"}
    assert abs(datetime.fromisoformat(was["round"]["started_at"]) - datetime.fromisoformat(newest["round"]["started_at"])) < timedelta(seconds=2)
    assert record(world, version="1.0")["stages"] == [] and [h["version"] for h in record(world, version="1.0")["history"]] == []      # 1.0 itself was never approved


def test_an_archived_document_keeps_its_record_and_says_when_it_was_archived(world):
    approve_and_publish(world, (RANA,), (MINA,))
    world.db.get(m.GovernanceDocument, 5).status = "archived"
    documents.create_audit_log(world.db, 5, 1, ADAM, "archived", "Archived by Adam Admin")
    world.db.commit()
    rec = record(world)
    assert rec["state"] == "archived" and events(rec)[-2:] == ["Published", "Archived"] and rec["milestones"][-1]["actor"] == "Adam Admin"
    assert [names(s) for s in rec["stages"]] == [["Rana Reviewer"], ["Mina Manager"]] and rec["document"]["status"] == "archived"


def test_a_version_that_does_not_exist_is_not_found_with_a_sentence(world):
    approve_and_publish(world)
    for path in ("/approval-record", "/approval-record/export"):
        r = get(world, path, version="9.9")
        assert r.status_code == 404 and r.json()["detail"] == "That version of the document was not found."
    r = world.http.get("/governance/documents/99/approval-record")
    assert r.status_code == 404 and r.json()["detail"] == "Document not found"
    with pytest.raises(Exception) as err:
        record(world, version="2.0")
    assert getattr(err.value, "status_code", None) == 404


def test_a_reader_who_may_not_see_drafts_gets_the_version_in_force_and_nothing_about_a_newer_round(world):
    approve_and_publish(world, (RANA,), (MINA,))
    edit(world)                                                                          # 1.1 is on its way: 1.0 stays in force
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    vera = get(world, "/approval-record", who=VERA)
    assert vera.status_code == 200
    rec = vera.json()
    assert rec["document"]["version"] == "1.0" and rec["document"]["status"] == "published" and rec["state"] == "published"
    assert rec["round"]["version"] == "1.0" and rec["round"]["complete"] is True
    assert [names(s) for s in rec["stages"]] == [["Rana Reviewer"], ["Mina Manager"]] and [h["version"] for h in rec["history"]] == ["1.0"]
    assert not any(n in vera.text for n in ("Rafay", "Mo Manager", "in_review", "pending_review", "Temporary staff"))
    assert not re.search(r"(?<![\d.])1\.1(?!\d)", vera.text)                  # no version 1.1 (not "41.1" in a time's fraction of a second)
    assert [r["version"] for r in rec["revision_history"]] == ["1.0"] and rec["control"]["version"] == "1.0" and rec["control"]["status"] == "Published"
    assert get(world, "/approval-record", who=VERA, version="1.1").status_code == 404           # nor a newer version by number
    assert get(world, "/approval-record/export", who=VERA, version="1.1").status_code == 404
    assert get(world, "/approval-record", who=VERA, version="1.0").status_code == 200
    olivia = get(world, "/approval-record", who=OLIVIA).json()                           # the owner reads the newest
    assert olivia["document"]["version"] == "1.1" and olivia["state"] == "in_approval" and names(olivia["stages"][0]) == ["Rafay Reviewer"]
    assert approval_record.build(world.db, doc(world), viewer=user(world, VERA)) == {**rec, "generated_at": approval_record.build(world.db, doc(world), viewer=user(world, VERA))["generated_at"]}
    uma = get(world, "/approval-record", who=UMA)                                        # asked to acknowledge it, no more: the version in force too
    assert uma.status_code == 200 and uma.json()["document"]["version"] == "1.0" and "Rafay" not in uma.text
    world.db.add(m.GRCUser(id=9, username="una", email="una@cfsb.test", display_name="Una Unrelated"))
    world.db.commit()
    nobody = get(world, "/approval-record", who=9)                                       # asked about nothing, no rights: not read at all
    assert nobody.status_code == 403 and get(world, "/approval-record/export", who=9).status_code == 403


def test_signatures_given_before_they_were_tied_to_a_version_are_read_through_their_e_signature_or_placed_as_unknown(world):
    """Rows from before version_id: the e-signature's version places them; with neither they are 'unknown'."""
    approve_and_publish(world, (RANA,), (MINA,))
    world.db.query(m.DocumentSignature).update({m.DocumentSignature.version_id: None})    # as they were before it was kept
    world.db.commit()
    edit(world)
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    sign(world, MO)
    old = record(world, version="1.0")                                                    # found through the e-signatures' document_version
    assert [names(s) for s in old["stages"]] == [["Rana Reviewer"], ["Mina Manager"]] and old["round"]["complete"] is True
    assert [h["version"] for h in record(world)["history"]] == ["1.1", "1.0"]
    # now the e-signatures of the first round go too (typed-name sign-offs from before drawn signatures): they cannot be placed
    ids = [s.id for s in world.db.query(m.DocumentSignature).filter(m.DocumentSignature.signer_user_id.in_((RANA, MINA)))]
    world.db.query(m.ElectronicSignature).filter(m.ElectronicSignature.record_id.in_(ids)).delete(synchronize_session=False)
    world.db.commit()
    rec = record(world)
    assert [h["version"] for h in rec["history"]] == ["1.1", "unknown"]
    assert [a["name"] for a in rec["history"][1]["approvers"]] == ["Mina Manager"] and rec["history"][1]["summary"] is None
    assert 'Earlier approvals whose version was not recorded are listed under version "unknown".' in rec["notes"]
    with pytest.raises(Exception):
        record(world, version="unknown")
    assert record(world, version="1.0")["stages"] == []                                   # nothing can be said of 1.0 any more


def test_signatures_of_the_current_round_with_no_version_are_listed_and_said_to_be(world):
    """A document approved before versions were kept with signatures (also the demo data): it is still approved by them."""
    d = doc(world)
    d.status, d.signoff_round_at, d.approved_by, d.approved_at = "published", datetime(2026, 9, 1), MINA, datetime(2026, 9, 2)
    d.published_at, d.published_by = datetime(2026, 9, 3), ADAM
    world.db.add(m.DocumentSignature(tenant_id=1, document_id=5, signer_user_id=MINA, role_type="approver", role_label="CISO", decision="signed",
                                     signature_text="mina", signed_at=datetime(2026, 9, 2)))
    world.db.commit()
    rec = record(world)
    assert rec["state"] == "published" and rec["round"] == {"version": "1.0", "started_at": "2026-09-01T00:00:00", "completed_at": "2026-09-02T00:00:00", "complete": True}
    assert [(s["label"], decisions(s)) for s in rec["stages"]] == [("Approved by", [("Mina Manager", "signed")])]
    person_ = rec["stages"][0]["people"][0]
    assert (person_["designation"], person_["meaning"], person_["signature"]) == ("CISO", "CISO", None)
    assert any("no version" in n for n in rec["notes"]) and rec["document"]["published_at"] == "2026-09-03T00:00:00" and rec["document"]["published_by"] == "Adam Admin"
    assert rec["history"][0]["version"] == "1.0" and rec["history"][0]["approved_at"] == "2026-09-02T00:00:00"


def test_a_document_approved_without_any_sign_off_says_so(world):
    d = doc(world)
    d.status = "approved"
    world.db.commit()
    rec = record(world)
    assert rec["state"] == "approved" and rec["stages"] == [] and rec["round"]["complete"] is False
    assert any(n.startswith("No sign-off is recorded") for n in rec["notes"])


def test_the_state_follows_the_status_of_the_document(world):
    seen = {}
    for status in ("draft", "pending_review", "pending_approval", "approved", "published", "expired", "archived", "exception_applied"):
        world.db.get(m.GovernanceDocument, 5).status = status
        world.db.commit()
        seen[status] = record(world)["state"]
    assert seen == {"draft": "draft", "pending_review": "in_review", "pending_approval": "in_approval", "approved": "approved", "published": "published",
                    "expired": "published", "archived": "archived", "exception_applied": "published"}
    assert approval_record.state_of(None) == "draft"


def test_milestones_are_at_most_one_hundred_and_the_first_is_kept(world):
    usual(world)
    send(world)
    for i in range(150):
        world.db.add(m.DocumentAuditLog(document_id=5, tenant_id=1, action="changes_requested", action_details=f"Round {i}", performed_by=RANA,
                                        performed_at=datetime.utcnow() + timedelta(seconds=i + 1)))
    world.db.commit()
    rec = record(world)
    assert len(rec["milestones"]) == 100 and rec["milestones"][0]["event"] == "Prepared" and rec["milestones"][-1]["detail"] == "Round 149"
    assert any("milestones" in n for n in rec["notes"])


# ── reading changes nothing ─────────────────────────────────────────────────────
def test_reading_the_record_and_its_files_writes_nothing(world, monkeypatch):
    set_route(world, stage("review", [person(RANA)], name="Review"), stage("approval", [person(MINA)], name="Approval"))
    send(world)
    sign(world, RANA)
    ask_changes(world, MINA)
    world.db.commit()
    flushed, commits = [], []
    event.listen(world.db, "after_flush", lambda session, ctx: flushed.append(1))
    monkeypatch.setattr(Session, "commit", lambda self: commits.append(1))
    tables = {t: world.db.execute(m.Base.metadata.tables[t].select()).fetchall() for t in
              ("grc_document_signatures", "grc_electronic_signatures", "grc_document_audit_logs", "grc_document_obligations",
               "grc_document_route_holds", "grc_governance_documents")}
    for path, params in (("/approval-record", {}), ("/approval-record", {"version": "1.0"}), ("/approval-record/export", {"format": "pdf"}),
                         ("/approval-record/export", {"format": "docx"}), ("/export", {"with_record": "true"}), ("/export", {})):
        assert get(world, path, **params).status_code == 200, path
    assert not world.db.new and not world.db.dirty and not world.db.deleted
    assert flushed == [] and commits == []
    assert tables == {t: world.db.execute(m.Base.metadata.tables[t].select()).fetchall() for t in tables}


# ── the files ───────────────────────────────────────────────────────────────────
def finished(world):
    """Published: Rana and Rafay review, Mina approves (three drawn signatures)."""
    approve_and_publish(world, (RANA, RAFAY), (MINA,))
    return record(world)


def test_the_control_pages_as_a_pdf_name_everyone_with_date_time_and_meaning_and_carry_the_signatures(world):
    rec = finished(world)
    r = get(world, "/approval-record/export")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"] == 'attachment; filename="POL-005-approval-record.pdf"'
    text, pages = pdf_text(r.content)
    for needle in ("Document Description", "Review and Approval", "Revision History", "Approval History", "Access policy", "POL-005",
                   "Rana Reviewer", "Compliance Manager", "Rafay Reviewer", "Mina Manager", "CISO", "STAGE 1 — REVIEWED BY", "STAGE 2 — APPROVED BY",
                   "Complete · 2 of 2 signed", "Signed", rec["reference"], "Initial version", "Olivia Owner", "PREPARED BY"):
        assert needle in text, needle
    assert re.search(r"Document control page 1 of %d" % pages, text) and WHEN.search(text)
    assert "generated by the platform" in text and "· Document control ·" in text                              # the foot says what the pages are
    assert rec["stages"][0]["people"][0]["signed_at"][:16].replace("T", " ") + " UTC" in text                  # date AND time of the signature
    with fitz.open(stream=r.content, filetype="pdf") as pdf:
        assert abs(pdf[0].rect.width - 595) < 2 and abs(pdf[0].rect.height - 842) < 2                          # A4
        assert sum(len(page.get_image_info()) for page in pdf) == 3                                            # the three drawn signatures


def test_the_control_pages_as_a_word_file_have_the_four_sections_the_grids_and_the_signatures_as_pictures(world):
    rec = finished(world)
    r = get(world, "/approval-record/export", format="docx")
    assert r.status_code == 200 and r.headers["content-type"] == DOCX_MEDIA
    assert r.headers["content-disposition"] == 'attachment; filename="POL-005-approval-record.docx"'
    word = Document(io.BytesIO(r.content))
    paragraphs = [p.text for p in word.paragraphs]
    assert [t for t in paragraphs if t in ("Document Description", "Review and Approval", "Revision History", "Approval History")] == [
        "Document Description", "Review and Approval", "Revision History", "Approval History"]
    assert all(p.style.name == "Normal" for p in word.paragraphs if p.text == "Review and Approval")           # never a Heading style: it could be numbered
    cells = "\n".join(c.text for t in word.tables for row in t.rows for c in row.cells)
    for needle in ("Rana Reviewer", "Compliance Manager", "Rafay Reviewer", "Mina Manager", "STAGE 1 — REVIEWED BY", "STAGE 2 — APPROVED BY", "Signed",
                   "Initial version", "Access policy", "POL-005", "PREPARED BY"):
        assert needle in cells, needle
    assert WHEN.search(cells) and len(word.inline_shapes) == 3 and "Drawn e-signature · seal" in cells
    footer = word.sections[0].footer.paragraphs[0]
    assert rec["reference"] in footer.text and "Document control page" in footer.text and "generated by the platform" in footer.text
    assert abs(word.sections[0].page_width.cm - 21.0) < 0.01 and abs(word.sections[0].page_height.cm - 29.7) < 0.01       # A4
    assert get(world, "/approval-record/export", format="odt").status_code == 400


def test_the_pdf_of_an_older_version_names_the_signers_of_that_version(world):
    approve_and_publish(world, (RANA,), (MINA,))
    edit(world)
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    sign(world, MO)
    publish(world)
    old, _ = pdf_text(get(world, "/approval-record/export", format="pdf", version="1.0").content)
    new, _ = pdf_text(get(world, "/approval-record/export", format="pdf").content)
    newest, _ = pdf_text(get(world, "/approval-record/export", format="pdf", version="1.1").content)
    assert "Rana Reviewer" in old and "Mina Manager" in old and "Rafay Reviewer" not in old and "Mo Manager" not in old and "DAR-5-1.0-" in old
    assert "Rafay Reviewer" in new and "Mo Manager" in new and "Rana Reviewer" not in new.split("Approval History")[0] and "DAR-5-1.1-" in new
    assert newest.replace(" ", "") != "" and "DAR-5-1.1-" in newest
    docx = Document(io.BytesIO(get(world, "/approval-record/export", format="docx", version="1.0").content))
    assert "Rana Reviewer" in "\n".join(c.text for t in docx.tables for row in t.rows for c in row.cells)


def test_a_content_only_document_exported_with_its_record_has_the_control_pages_where_its_own_sections_would_be(world, monkeypatch):
    from reportlab import rl_config

    monkeypatch.setattr(rl_config, "invariant", 1)                           # the same input, the same bytes
    rec = finished(world)
    plain = get(world, "/export")
    assert plain.status_code == 200 and plain.headers["content-disposition"] == 'attachment; filename="Access_policy.pdf"'
    from grc.routers._artifact_export import build_export

    today, _media, _ext = build_export("pdf", title="Access policy", content=TEXT)
    assert plain.content == today                                            # nothing changed for who does not ask
    assert get(world, "/export", with_record="false").content == today
    combined = get(world, "/export", with_record="true")
    assert combined.status_code == 200 and combined.headers["content-type"] == "application/pdf"
    assert combined.headers["content-disposition"] == 'attachment; filename="Access_policy-with-approval-record.pdf"'
    text, pages = pdf_text(combined.content)
    assert pages > pdf_text(today)[1]
    for needle in ("Document Description", "Review and Approval", "Revision History", "Rana Reviewer", "Rafay Reviewer", "Mina Manager", rec["reference"], "Purpose"):
        assert needle in text, needle
    # the document has none of the sections: the pages come in after its title, before its first section
    assert text.index("Access policy") < text.index("Document Description") < text.index("Review and Approval") < text.index("1. Purpose")
    with fitz.open(stream=combined.content, filetype="pdf") as pdf:
        counted = [page.get_text() for page in pdf if "Document control page" in page.get_text()]
        assert len(counted) >= 2 and all(f"of {len(counted)}" in t for t in counted)                                   # they count only themselves


def attach(world, name, data, kind):
    path = world.tmp / name
    path.write_bytes(data)
    d = doc(world)
    d.file_name, d.file_path, d.file_type, d.file_size = name, str(path), kind, len(data)
    world.db.commit()
    return path


def word_file(*lines):
    word = Document()
    word.add_heading("Access policy", 1)
    for line in lines:
        word.add_paragraph(line)
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(1, 1).text = "Control", "Owner"
    out = io.BytesIO()
    word.save(out)
    return out.getvalue()


def test_an_uploaded_word_file_exported_with_its_record_keeps_its_own_content_and_gains_the_record(world):
    rec = finished(world)
    original = word_file("Access is reviewed every quarter.", "Leavers lose access the same day.")
    path = attach(world, "Access policy.docx", original, "docx")
    plain = get(world, "/export")
    assert plain.status_code == 200 and plain.content == original and "x-approval-record" not in plain.headers       # as it always was
    assert get(world, "/export", with_record="false").content == original
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.headers["content-type"] == DOCX_MEDIA and "x-approval-record" not in r.headers
    assert r.headers["content-disposition"] == 'attachment; filename="Access policy-with-approval-record.docx"'
    assert path.read_bytes() == original                                                                              # the stored file is never modified
    word = Document(io.BytesIO(r.content))
    paragraphs = [p.text for p in word.paragraphs]
    assert paragraphs[:3] == ["Access policy", "Access is reviewed every quarter.", "Leavers lose access the same day."]
    assert "Document Description" in paragraphs and paragraphs.index("Document Description") > 3               # no cover (no page boundary): after its content
    cells = "\n".join(c.text for t in word.tables[1:] for row in t.rows for c in row.cells)
    assert word.tables[0].cell(0, 0).text == "Control" and "Rana Reviewer" in cells and "Mina Manager" in cells and WHEN.search(cells)
    assert len(word.inline_shapes) == 3 and rec["reference"] in "\n".join(paragraphs)
    assert any('w:br w:type="page"' in p._p.xml for p in word.paragraphs)                                              # the pages start on a new page


def test_a_real_corporate_word_file_and_one_with_tracked_changes_keep_every_part_of_their_own(world):
    """The record is added after the document; what the document has - its header and footer, lists, links, fields, pictures, merged
    cells, and a reviewer's tracked changes - stays exactly as it was."""
    from lxml import etree

    from _docx_fixtures import corporate_policy, parts, tracked_policy

    rec = finished(world)
    for build in (corporate_policy, tracked_policy):
        path = build(str(world.tmp / f"{build.__name__}.docx"))
        original = open(path, "rb").read()
        attach(world, f"{build.__name__}.docx", original, "docx")
        r = get(world, "/export", with_record="true")
        assert r.status_code == 200 and "x-approval-record" not in r.headers, build.__name__

        def body(data):
            root = etree.fromstring(parts(data).read("word/document.xml"))
            return [c for c in root.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")]

        before, after = body(original), body(r.content)
        w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        kept_before = [etree.tostring(c) for c in before if c.tag != w + "sectPr"]
        assert [etree.tostring(c) for c in after[:len(kept_before)]] == kept_before, build.__name__          # its own content, untouched
        assert after[-1].tag == w + "sectPr" and len(after) > len(before)                                    # no cover: the pages came in before the section settings
        names_before, names_after = set(parts(original).namelist()), set(parts(r.content).namelist())
        assert names_before <= names_after, build.__name__                                                  # no part of the file was dropped
        for part in sorted(n for n in names_before if n.startswith("word/header") or n.startswith("word/footer")):
            assert parts(r.content).read(part) == parts(original).read(part), part
        word = Document(io.BytesIO(r.content))
        assert any(p.text == "Document Description" for p in word.paragraphs) and rec["reference"] in "\n".join(p.text for p in word.paragraphs)
        assert len(word.inline_shapes) >= 3
    xml = parts(r.content).read("word/document.xml").decode()
    assert xml.count("<w:ins ") >= 2 and xml.count("<w:del ") >= 2 and "rPrChange" in xml                       # the reviewer's changes are still tracked


def test_a_word_file_without_the_usual_styles_still_gets_its_record(world):
    """A file made by another tool: no Heading or Table Grid styles. The record brings its own formatting."""
    finished(world)
    bare = Document()
    bare.add_paragraph("Plain text.")
    for name in ("Heading 1", "Heading 2", "Heading 3", "Table Grid"):
        style = bare.styles[name].element
        style.getparent().remove(style)
    out = io.BytesIO()
    bare.save(out)
    attach(world, "bare.docx", out.getvalue(), "docx")
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and "x-approval-record" not in r.headers
    word = Document(io.BytesIO(r.content))
    assert any(p.text == "Document Description" for p in word.paragraphs) and "Rana Reviewer" in "\n".join(c.text for t in word.tables for row in t.rows for c in row.cells)


def pdf_file(*pages):
    pdf = fitz.open()
    for text in pages:
        pdf.new_page(width=595, height=842).insert_text((72, 100), text)
    data = pdf.tobytes()
    pdf.close()
    return data


def test_an_uploaded_pdf_exported_with_its_record_keeps_its_pages_and_gains_the_control_pages_after_its_cover(world):
    rec = finished(world)
    original = pdf_file("Original page one", "Original page two")                      # page 1 has hardly any text: a cover
    path = attach(world, "Access policy.pdf", original, "pdf")
    assert get(world, "/export").content == original
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and "x-approval-record" not in r.headers
    assert r.headers["content-disposition"] == 'attachment; filename="Access policy-with-approval-record.pdf"'
    assert path.read_bytes() == original
    text, pages = pdf_text(r.content)
    assert pages > 2 and text.startswith("Original page one") and text.rstrip().endswith("Original page two")
    with fitz.open(stream=r.content, filetype="pdf") as pdf:
        assert "Original page one" in pdf[0].get_text() and "Document Description" in pdf[1].get_text() and "Original page two" in pdf[pages - 1].get_text()
        assert [p for p in range(pages) if "Document control page" in pdf[p].get_text()] == list(range(1, pages - 1))
    assert re.search(r"Document control page 1 of %d\b" % (pages - 2), text)
    for needle in ("Rana Reviewer", "Rafay Reviewer", "Mina Manager", rec["reference"]):
        assert needle in text, needle


def test_a_file_of_another_kind_is_served_unchanged_and_says_its_record_is_apart(world):
    finished(world)
    original = b"PK\x03\x04 pretend spreadsheet"
    attach(world, "Register.xlsx", original, "xlsx")
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.content == original and r.headers["x-approval-record"] == "separate"
    assert r.headers["content-disposition"].startswith('attachment; filename="Register.xlsx"')
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert get(world, "/export").content == original and "x-approval-record" not in get(world, "/export").headers


def test_a_word_or_pdf_file_that_cannot_be_read_is_served_unchanged_with_its_record_apart(world):
    finished(world)
    attach(world, "broken.docx", b"this is not a zip file", "docx")
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.content == b"this is not a zip file" and r.headers["x-approval-record"] == "separate"
    attach(world, "broken.pdf", b"%PDF-1.4 truncated", "pdf")
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.content == b"%PDF-1.4 truncated" and r.headers["x-approval-record"] == "separate"
    locked = fitz.open()                                                      # opens without a password but is protected: rewriting it would drop that
    locked.new_page().insert_text((72, 72), "Protected")
    protected = locked.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="", permissions=fitz.PDF_PERM_ACCESSIBILITY)
    attach(world, "locked.pdf", protected, "pdf")
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.content == protected and r.headers["x-approval-record"] == "separate"


def test_a_content_only_pdf_whose_record_cannot_be_made_is_served_as_it_was_with_its_record_apart(world, monkeypatch):
    finished(world)

    def boom(*args, **kwargs):
        raise RuntimeError("the record could not be built")

    monkeypatch.setattr(approval_record, "build", boom)
    r = get(world, "/export", with_record="true")
    assert r.status_code == 200 and r.headers["x-approval-record"] == "separate"
    assert r.headers["content-disposition"] == 'attachment; filename="Access_policy.pdf"' and "Approval record" not in pdf_text(r.content)[0]


def test_who_reads_the_version_in_force_exports_that_versions_record(world):
    approve_and_publish(world, (RANA,), (MINA,))
    edit(world)
    usual(world, (RAFAY,), (MO,))
    sign(world, RAFAY)
    text, _ = pdf_text(get(world, "/export", who=VERA, with_record="true").content)            # the content-only PDF of 1.0, and its record
    assert "Rana Reviewer" in text and "Mina Manager" in text and "Rafay Reviewer" not in text and "Mo Manager" not in text
    text, _ = pdf_text(get(world, "/export", who=OLIVIA, with_record="true").content)
    assert "Rafay Reviewer" in text and "Mo Manager" in text


def test_the_files_are_named_after_the_code_else_the_title_and_survive_any_name(world):
    rec = finished(world)
    assert X.stem(rec) == "POL-005"
    d = doc(world)
    d.document_code = None
    d.title = "Política de acceso — يوم"
    world.db.commit()
    r = get(world, "/approval-record/export")
    disposition = r.headers["content-disposition"]
    assert r.status_code == 200 and disposition.startswith("attachment; filename=\"") and "filename*=UTF-8''" in disposition
    assert disposition.encode("latin-1") and "approval-record.pdf" in disposition
    assert X.content_disposition("a.pdf") == 'attachment; filename="a.pdf"'


def test_whatever_a_person_wrote_or_a_name_is_written_in_the_files_still_open(world):
    """A comment may carry any character a person pasted, a name any script: Word's XML refuses some, a PDF cannot print others."""
    world.db.get(m.GRCUser, RANA).display_name = "رنا Reviewer <b>&</b>"
    world.db.commit()
    usual(world, (RANA,), (MINA,))
    send(world)
    sign(world, RANA, "Fine\x0b\x1f, <i>not italic</i> & ok “quoted” " + "long" * 400)
    sign(world, MINA)
    for fmt in ("pdf", "docx"):
        r = get(world, "/approval-record/export", format=fmt)
        assert r.status_code == 200, fmt
    word = Document(io.BytesIO(get(world, "/approval-record/export", format="docx").content))
    cells = "\n".join(c.text for t in word.tables for row in t.rows for c in row.cells)
    assert "Fine, <i>not italic</i> & ok" in cells and "<b>&</b>" in cells and "\x0b" not in cells
    text, _ = pdf_text(get(world, "/approval-record/export").content)
    assert "<i>not italic</i>" in text and "Mina Manager" in text and " …" in text.replace("\n", " ")             # a long comment is cut, not overflowed
    assert X.scrub({"a": ["x\x00y", ("p\x1fq",)], "n": 3}) == {"a": ["xy", ["pq"]], "n": 3}


def test_an_unreadable_document_has_no_record_and_the_routes_are_guarded_like_its_neighbours():
    from fastapi.routing import APIRoute

    from grc.modules.governance import permissions

    found = {(r.path, tuple(sorted(r.methods))): r for r in router.routes if isinstance(r, APIRoute) and "approval-record" in r.path}
    assert set(found) == {("/governance/documents/{document_id}/approval-record", ("GET",)),
                          ("/governance/documents/{document_id}/approval-record/export", ("GET",))}

    def calls(dependant):
        for d in dependant.dependencies:
            yield d.call
            yield from calls(d)

    assert all(permissions.require_doc_reader in set(calls(r.dependant)) for r in found.values())
    export = next(r for r in router.routes if isinstance(r, APIRoute) and r.path == "/governance/documents/{document_id}/export")
    assert [q.name for q in export.dependant.query_params] == ["with_record"]
    assert next(q for q in export.dependant.query_params).default is False


# ── what the files print ────────────────────────────────────────────────────────
def test_every_time_in_the_files_is_a_date_with_a_time_and_utc_and_the_formatters_are_plain(world):
    assert X.when("2026-10-21T10:15:30.123456") == "2026-10-21 10:15 UTC" and X.when(None) == "-" and X.when("not a date") == "not a date"
    assert X.day("2026-10-21T00:00:00") == "2026-10-21" and X.day(None) == "-"
    assert X.rule_text({"rule": "at_least", "min_count": 2}) == "at least 2 sign" and X.rule_text({"rule": "any", "min_count": None}) == "any one signs"
    assert X.stage_title({"position": 2, "label": "Approved by", "name": "Board approval"}) == "STAGE 2 — APPROVED BY · Board approval"
    assert X.stage_title({"position": 1, "label": "Reviewed by", "name": None}) == "STAGE 1 — REVIEWED BY"
    assert X.fit(b"not a picture") is None
    rec = finished(world)
    text, _ = pdf_text(X.to_pdf(rec, {}))
    assert not re.search(r"\d{4}-\d{2}-\d{2}T\d{2}", text) and len(WHEN.findall(text)) >= 5
    assert "Drawn e-signature" in text                                         # the seal is printed even when a picture is not available
