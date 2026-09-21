"""The monthly pack: every section of the client's SUMMARY sheet.

Their pack reports, per source, the roll-forward, where open issues stand
(Not Started … Extension), how far past due they are, what closed this month,
and the recommendations they do not track. Ours is computed from the register;
Past Due follows their own definition (not implemented by the agreed date);
their figures are read from the workbook's SUMMARY so differences show.
"""
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import grc.models as m
from grc.modules.issue_management.audit_register.export import build_workbook
from grc.modules.issue_management.audit_register.parser import parse_workbook
from grc.modules.issue_management.audit_register.service import apply_workbook
from grc.modules.issue_management.audit_register.summary import effective_status, pack, reports

SOURCE_ROW = ["Source ", "Regulator", "CFSB Audit-EY", "Total"]
REG_HEADERS = ["#", "Regulator", "Report #", "Report Name", "Report  Date", "Issue #",
               "Risk Rating", "Issue Name", "Issue", "Owner", "LOB", "Target Date",
               "Revised Target Date", "Date of Validation by CFSB IA",
               "Status (NS/IP/DE/ PD/ EXT/CD )"]
REC_HEADERS = ["#", "Source (Mercadien/ EY)", "Report Number", "Audit Report Name",
               "Audit Report Date", "Issue #", "Risk Rating", "Issue Name", "Recommendation",
               "Owner", "Target Date", "Status (NS/IP/DE/ PD/ EXT/CD )"]
ARCHIVE_HEADERS = ["Project Name-Year (20XX)", "Issue ID", "Issue / Recommendation",
                   "Recommendation State", "Source", "Issue Subject", "Relative Risk",
                   "Original Target Date"]
OPENED = date(2025, 1, 10)


def _reg(n, status, target=None, revised=None, validated=None):
    return [n, "OCC", "R-1", "Exam", OPENED, f"MRA-{n}", "High", f"Finding {n}", "Detail",
            "Rivera, P", "ERM", target, revised, validated, status]


def _workbook(tmp_path):
    wb = Workbook()
    s = wb.active
    s.title = "SUMMARY"
    s.append(["CFSB AUDIT SERVICES"])
    s.append(["Summary: Issue & Status     Recommendation & Status"])
    s.append([])
    s.append(["April 2026 Issue Summary"])
    s.append(SOURCE_ROW)
    s.append(["Beginning Number of Issues", 7, 0, 7])
    s.append(["Add: Issues opened this month", 0, 0, 0])
    s.append(["Less: Issues closed this month", 1, 0, 1])
    s.append(["Ending Number of Issues ", 6, 0, 6])
    s.append([])
    s.append(["Open Issue status as of April 30 2026"])
    s.append(["STATUS OF ISSUE RESOLUTION PROGRESS:"])
    s.append(["Not Started", 99])            # their unfilled section: not the status table
    s.append(["In Progress"])
    s.append([])
    s.append(SOURCE_ROW)
    s.append(["Not Started", 1, 0, 1])
    s.append(["In Progress", 2, 0, 2])
    s.append(["Delayed", 1, 0, 1])
    s.append(["Past Due", 0, 0, 0])
    s.append(["Extension", 1, 0, 1])
    s.append(["Total Open Issues", 5, 0, 5])
    s.append([])
    s.append(SOURCE_ROW)
    s.append(["Recommendations (not tracked)", 0, 1, 1])

    reg = wb.create_sheet("Regulatory")
    reg.append(["AS OF", date(2026, 4, 30), "", "",
                "Past Due: Action Plan has not been implemented by the agreed upon target date."])
    reg.append(["", "", "", "",
                "Not Started CFSB Audit Services has not started the validation yet."])
    reg.append(REG_HEADERS)
    reg.append(_reg(1, "IP", target=date(2026, 6, 30)))                 # in progress
    reg.append(_reg(2, "IP", target=date(2026, 3, 31)))                 # past its date → PD
    reg.append(_reg(3, "DE", target=date(2026, 1, 15)))                 # delayed stays delayed
    reg.append(_reg(4, "EXT", target=date(2026, 2, 28), revised=date(2026, 7, 31)))
    reg.append(_reg(5, "RM"))                                           # not one of their codes
    reg.append(_reg(6, "CD", target=date(2026, 3, 31), validated=date(2026, 4, 15)))
    reg.append(_reg(7, "NS", target=date(2026, 12, 31)))

    rec = wb.create_sheet("IA-EY Recommendation")
    rec.append(REC_HEADERS)
    rec.append([1, "EY", "EY-1", "Governance review", OPENED, "R1", "Low", "Document the charter",
                "Write it down", "Hughes, M", date(2026, 9, 30), "IP"])

    closed = wb.create_sheet("IssueClosed042026&later")
    closed.append(ARCHIVE_HEADERS)
    closed.append(["REF-2023", "REF3", "Issue", "Closed-Validated", "Internal Audit Report",
                   "Monitoring gap", "High", date(2024, 1, 31)])
    closed.append([])
    closed.append(["Count of Status (NS/IP/DE/ PD/ EXT/ CD)"])     # a pivot under the table
    closed.append(["10", "10"])
    closed.append(REC_HEADERS)                                      # a second, empty table
    path = tmp_path / "april.xlsx"
    wb.save(path)
    return str(path)


@pytest.fixture
def db(tmp_path):
    engine = create_engine("sqlite://")
    m.Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(m.Tenant(id=1, name="CFSB", slug="cfsb"))
        session.commit()
        apply_workbook(session, 1, parse_workbook(_workbook(tmp_path)))
        session.commit()
        yield session


def test_the_workbook_s_own_pack_is_read_section_by_section(tmp_path):
    book = parse_workbook(_workbook(tmp_path))

    assert book.summary_month == "2026-04"
    assert book.summary["regulator"]["ending"] == 6
    assert book.summary["regulator"]["NS"] == 1          # the table, not the "99" above it
    assert book.summary["regulator"]["open"] == 5
    assert book.summary["ey"]["recommendations"] == 1
    assert book.status_definitions["Regulatory"]["PD"].startswith("Past Due: Action Plan")
    assert book.status_definitions["Regulatory"]["NS"].startswith("Not Started")
    # One real row on the archive sheet: the pivot and the empty second table are not findings.
    assert book.sheet_counts["IssueClosed042026&later"] == 1
    assert not [w for w in book.warnings if "IssueClosed" in w]


def test_status_follows_their_definitions(db):
    by_ref = {p.issue_ref: effective_status(p, p.issue, date(2026, 4, 30))
              for p in db.query(m.AuditIssueProfile).filter(m.AuditIssueProfile.source == "regulator")}

    assert by_ref == {"MRA-1": "IP", "MRA-2": "PD", "MRA-3": "DE", "MRA-4": "EXT",
                      "MRA-5": "unstated", "MRA-6": "CD", "MRA-7": "NS"}


def test_a_plan_reported_done_is_in_progress_but_the_pen_test_sheet_still_ages_it(db):
    as_of = date(2026, 4, 30)
    mra = db.query(m.AuditIssueProfile).filter(m.AuditIssueProfile.issue_ref == "MRA-2").one()
    mra.issue.workflow_state = "closure_review"          # management: "Completed", past its date
    assert effective_status(mra, mra.issue, as_of) == "IP"

    archive = m.AuditIssueProfile(source="it_pen", source_sheet="IT Pen", record_type="issue",
                                  recommendation_state="Closed-Pending Validation",
                                  target_date=date(2024, 4, 5))
    pending = m.Issue(workflow_state="closure_review")
    assert effective_status(archive, pending, as_of) == "PD"   # their pack: 91 of these, Past Due


def test_a_finding_never_closes_before_it_opens(db):
    from grc.modules.issue_management.audit_register.summary import timeline
    profile = m.AuditIssueProfile()
    issue = m.Issue(workflow_state="closed", detected_at=datetime(2026, 5, 10),
                    closed_at=datetime(2026, 4, 15))

    assert timeline(profile, issue, date(2026, 4, 1)) == (date(2026, 4, 15), date(2026, 4, 15))
    assert all(r["ending"] >= 0 for r in pack(db, 1, "2026-04")["roll_forward"]["rows"])


def test_the_pack_has_every_section_and_reconciles(db):
    april = pack(db, 1, "2026-04")

    assert april["title"] == "April 2026 Issue Summary" and april["as_of"] == "2026-04-30"
    rows = {r["status"]: r["by_source"]["regulator"] for r in april["status_rows"]}
    assert rows == {"NS": 1, "IP": 1, "DE": 1, "PD": 1, "EXT": 1, "unstated": 1}
    assert april["total_open"]["by_source"]["regulator"] == 6
    assert april["total_open"]["check"] == 0                     # matches the roll-forward ending
    assert april["resolution"] == {
        "not_started": 1, "in_progress": 1, "delayed_past_due": 2, "extended": 1, "unstated": 1,
        "closed_this_month": 2, "subtotal": 8, "less_closed": 2, "total": 6}
    aging = {a["bucket"]: a["total"] for a in april["aging"]}
    assert aging == {"0-30": 1, "30-60": 0, "60-90": 0, "90-180": 1, "180+": 0}
    # MRA-6 validated in April, and the archive sheet's issue closed as of the pack.
    assert [c["reference"] for c in april["closed"]] == ["REF3", "MRA-6"]
    assert april["recommendations"]["by_source"]["ey"] == 1       # counted apart, not tracked
    # Their own figures sit beside ours: they had 2 In Progress where we count 1.
    in_progress = next(r for r in april["status_rows"] if r["status"] == "IP")
    assert in_progress["reported"]["regulator"] == 2


def test_the_export_is_laid_out_like_their_summary_sheet(db):
    content = build_workbook(pack(db, 1, "2026-04"), "CFSB")
    sheet = load_workbook(BytesIO(content))["SUMMARY"]
    labels = [row[0] for row in sheet.iter_rows(values_only=True)]

    assert sheet["A1"].value == "CFSB AUDIT SERVICES"
    assert sheet["A4"].value == "April 2026 Issue Summary"
    for label in ("Beginning Number of Issues", "Ending Number of Issues ",
                  "STATUS OF ISSUE RESOLUTION PROGRESS:", "ISSUE AGING ANALYSIS",
                  "List of issues closed this month: ", "Total Open Issues",
                  "Recommendations (not tracked)"):
        assert label in labels
    total_open = next(r for r in sheet.iter_rows(values_only=True) if r[0] == "Total Open Issues")
    assert total_open[1] == 6                                     # the Regulator column


def test_the_reports_view_groups_findings_by_report(db):
    lines = {(r["source"], r["report"]): r for r in reports(db, 1)}

    exam = lines[("regulator", "Exam")]
    assert exam["findings"] == 7 and exam["closed"] == 1 and exam["open"] == 6
    assert exam["regulator"] == "OCC"


def test_the_blank_template_is_their_workbook_and_imports_back(db, tmp_path):
    from grc.modules.issue_management.audit_register.export import build_template

    book = parse_workbook(_workbook(tmp_path))
    blank = tmp_path / "blank.xlsx"
    blank.write_bytes(build_template(book.sheet_headers, book.status_definitions, "CFSB"))

    back = parse_workbook(str(blank))
    assert list(back.sheet_counts) == list(book.sheet_counts)          # their sheets, in order
    assert not [w for w in back.warnings if "no register columns" in w]
    assert back.sheet_headers == book.sheet_headers                    # their header rows exactly
    assert back.status_definitions["Regulatory"]["PD"].startswith("Past Due")
    assert not back.rows                                               # nothing filled in yet


def test_without_an_upload_the_standard_sheets_import_once_filled(tmp_path):
    from openpyxl import load_workbook as load
    from grc.modules.issue_management.audit_register.export import build_template
    from grc.modules.issue_management.audit_register.template import (
        ENTRY_TEMPLATES, default_headers, layout_for)

    headers = {sheet: default_headers(layout_for(src, sheet)) for _, _, sheet, src, _ in ENTRY_TEMPLATES}
    path = tmp_path / "standard.xlsx"
    path.write_bytes(build_template(headers, {}, "Bank"))
    wb = load(path)
    for _, _, sheet, _, _ in ENTRY_TEMPLATES:                          # one finding per sheet
        ws = wb[sheet]
        cols = [c.value for c in ws[ws.max_row]]
        row = ["" for _ in cols]
        for i, h in enumerate(cols):
            if h in ("Issue Name", "Issue Subject"):
                row[i] = f"A {sheet} finding"
            elif h in ("Issue #", "Issue ID"):
                row[i] = "X-1"
            elif h in ("Risk Rating", "Relative Risk", "Risk Level"):
                row[i] = "High"
            elif h in ("Issue", "Description", "Condition"):
                row[i] = "What was found"
        ws.append(row)
    wb.save(path)

    back = parse_workbook(str(path))
    assert {r.sheet: r.source for r in back.rows} == {s: src for _, _, s, src, _ in ENTRY_TEMPLATES}
