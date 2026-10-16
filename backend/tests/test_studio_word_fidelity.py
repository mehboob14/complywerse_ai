"""A Word file reads as Word shows it: a table of contents, field results, symbols, text boxes, SmartArt, nested tables, rows in
content controls and page breaks are read; hidden text is not; line breaks at the edges and italics next to punctuation never
leave stray marks. What Word makes or draws itself (contents, boxes, diagrams) is marked locked, and the file stays editable in
place."""
import io

import pytest
from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from grc.modules.governance.studio import blocks as B
from grc.modules.governance.studio import docx_io as D
from grc.modules.governance.studio import word_layout as L

MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
WPS = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
V = 'xmlns:v="urn:schemas-microsoft-com:vml"'
DIAGRAM = ('<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>'
           '<dgm:pt modelId="3"><dgm:t><a:p><a:r><a:t>Low incident</a:t></a:r></a:p></dgm:t></dgm:pt>'
           '<dgm:pt modelId="0" type="doc"/>'
           '<dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>Critical incident</a:t></a:r></a:p></dgm:t></dgm:pt>'
           '<dgm:pt modelId="2"><dgm:t><a:p><a:r><a:t>Response </a:t></a:r><a:r><a:t>team</a:t></a:r></a:p></dgm:t></dgm:pt>'
           '<dgm:pt modelId="9" type="pres"><dgm:t><a:p><a:r><a:t>layout only</a:t></a:r></a:p></dgm:t></dgm:pt>'
           '</dgm:ptLst><dgm:cxnLst>'
           '<dgm:cxn modelId="10" srcId="0" destId="3" srcOrd="1"/><dgm:cxn modelId="11" srcId="0" destId="1" srcOrd="0"/>'
           '<dgm:cxn modelId="12" srcId="1" destId="2" srcOrd="0"/><dgm:cxn modelId="13" type="presOf" srcId="1" destId="9"/>'
           '</dgm:cxnLst></dgm:dataModel>')


def _box(text: str) -> str:
    inner = f"<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
    return (f'<w:r><mc:AlternateContent {MC}><mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData '
            f'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp {WPS}><wps:txbx>{inner}</wps:txbx>'
            f'</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape {V}><v:textbox>'
            f'{inner}</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>')


@pytest.fixture
def fancy(tmp_path):
    d = Document()
    d.add_paragraph("Escalation Manual", style="Title")
    part = Part(PackURI("/word/diagrams/data1.xml"), "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
                DIAGRAM.encode(), d.part.package)
    rid = d.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData")
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    n = nsdecls("w", "wp", "a", "r", "dgm")
    for xml in (
        # a table of contents: Word's entries (field results), its own page numbers
        '<w:p><w:pPr><w:pStyle w:val="TOC1"/></w:pPr><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:hyperlink w:anchor="_Toc1"><w:r><w:t>Purpose</w:t></w:r><w:r><w:tab/></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGEREF _Toc1 \\h </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>2</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:hyperlink></w:p>',
        '<w:p><w:pPr><w:pStyle w:val="TOC1"/></w:pPr><w:hyperlink w:anchor="_Toc2"><w:r><w:t>Scope</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>3</w:t></w:r></w:hyperlink>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>',
        '<w:p><w:r><w:t xml:space="preserve">Owner: </w:t></w:r><w:fldSimple w:instr=" DOCPROPERTY Owner "><w:r><w:t>Head of IT</w:t></w:r></w:fldSimple>'
        '<w:r><w:rPr><w:vanish/></w:rPr><w:t xml:space="preserve"> secret note</w:t></w:r><w:r><w:t xml:space="preserve"> </w:t></w:r>'
        '<w:r><w:sym w:font="Wingdings" w:char="F0FC"/></w:r><w:r><w:t xml:space="preserve"> approved.</w:t></w:r></w:p>',
        '<w:p><w:r><w:t>Ends with breaks</w:t><w:br/><w:br/></w:r></w:p>',
        '<w:p><w:r><w:t>the clarification</w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t xml:space="preserve">, and O</w:t></w:r>'
        '<w:r><w:t>NET-based data.</w:t></w:r></w:p>',
        '<w:p><w:r><w:t>Before the page</w:t></w:r><w:r><w:br w:type="page"/></w:r><w:r><w:t>after the page.</w:t></w:r></w:p>',
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Appendix A</w:t></w:r><w:r><w:br w:type="page"/></w:r><w:r><w:t>Appendix B</w:t></w:r></w:p>',
        f'<w:p><w:r><w:t>The anchor paragraph.</w:t></w:r>{_box("Text in a box.")}</w:p>',
        f'<w:p><w:r><w:drawing><wp:inline><wp:extent cx="100" cy="100"/><wp:docPr id="7" name="Diagram 1"/><a:graphic><a:graphicData '
        f'uri="http://schemas.openxmlformats.org/drawingml/2006/diagram"><dgm:relIds r:dm="{rid}" r:lo="" r:qs="" r:cs=""/></a:graphicData>'
        '</a:graphic></wp:inline></w:drawing></w:r></w:p>',
        '<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>'
        '<w:tr><w:tc><w:p><w:r><w:t>Req</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Text</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:p><w:r><w:t>FR3</w:t></w:r></w:p></w:tc><w:tc><w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc>'
        '<w:p><w:r><w:t>Nested words</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p/></w:tc></w:tr>'
        '<w:sdt><w:sdtPr/><w:sdtContent><w:tr><w:tc><w:p><w:r><w:t>FR4</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>In a repeating section</w:t></w:r></w:p></w:tc></w:tr>'
        '</w:sdtContent></w:sdt></w:tbl>',
        '<w:p><w:r><w:t>The end.</w:t></w:r></w:p>',
    ):
        el = parse_xml(xml.replace("<w:p>", f"<w:p {n}>", 1) if xml.startswith("<w:p>") else
                       xml.replace("<w:p><w:pPr>", f"<w:p {n}><w:pPr>", 1) if xml.startswith("<w:p><w:pPr>") else
                       xml.replace("<w:tbl>", f"<w:tbl {n}>", 1))
        sect.addprevious(el)
    path = str(tmp_path / "fancy.docx")
    d.save(path)
    return path, D.import_docx(path, twin=True)


def test_everything_word_shows_is_read_and_nothing_else(fancy):
    _, imp = fancy
    md = imp.markdown
    assert imp.native is not None, imp.notes                                     # still edited in place
    assert "Purpose\t2\n\nScope\t3" in md                                         # the contents, with Word's page numbers
    assert "Owner: Head of IT ✓ approved." in md and "secret" not in md           # a property field, a symbol; hidden text is not shown
    assert "Ends with breaks\n" in md and "Ends with breaks\\" not in md          # breaks at the end leave no backslash
    assert "the clarification, *and O*NET-based data." in md                     # the comma sits outside the italics
    assert "Before the page\\\nafter the page." in md                              # a page break starts a new line
    assert "## Appendix A Appendix B" in md                                       # ... and a heading stays one line
    assert "The anchor paragraph.\n\nText in a box." in md                        # a text box after the paragraph it is drawn in, once
    assert "Critical incident\\\nResponse team\\\nLow incident" in md and "layout only" not in md   # SmartArt in reading order
    assert "| FR3 | Nested words |" in md and "| FR4 | In a repeating section |" in md
    texts = [b.text for b in B.parse(md)]
    assert "the clarification, and ONET-based data." in texts                     # no literal stars


def test_what_word_makes_or_draws_is_locked(fancy):
    path, imp = fancy
    looks = L.for_blocks(open(path, "rb").read(), imp.markdown)
    blocks = B.parse(imp.markdown)
    locked = [b.text.split("\n")[0] for b, look in zip(blocks, looks) if look.get("locked")]
    assert locked == ["Purpose\t2", "Scope\t3", "Text in a box.", "Critical incident"]


def test_a_text_box_is_edited_inside_it_and_new_text_goes_after_its_paragraph(fancy):
    path, imp = fancy
    new = imp.markdown.replace("Text in a box.", "Text in a box, edited.\n\nA new paragraph.")
    data, _ = D.export_in_place(path, new)
    doc = Document(io.BytesIO(data))
    boxes = doc.element.body.xpath(".//w:txbxContent")
    assert [[t.text for t in box.iter(qn("w:t"))] for box in boxes] == [["Text in a box, edited."]] * 2    # the copy for old Words too
    after = [p.text for p in doc.paragraphs]
    assert after.index("A new paragraph.") == after.index("The anchor paragraph.") + 1
    assert D.check_blocks(D.walk(Document(io.BytesIO(data)))[0], B.parse(new), lenient=True) is None


def test_a_diagram_is_kept_and_cannot_be_rewritten_here(fancy):
    path, imp = fancy
    data, _ = D.export_in_place(path, imp.markdown.replace("The end.", "The very end."))
    assert "Critical incident" in D.render_nodes(D.walk(Document(io.BytesIO(data)))[0])
    with pytest.raises(D.PatchProblem):
        D.export_in_place(path, imp.markdown.replace("Low incident", "Minor incident"))


def test_a_cell_with_a_nested_table_keeps_its_lines_when_edited(fancy):
    path, imp = fancy
    data, _ = D.export_in_place(path, imp.markdown.replace("Nested words", "Nested words, edited"))
    doc = Document(io.BytesIO(data))
    nested = doc.element.body.xpath(".//w:tbl//w:tbl")
    assert len(nested) == 1 and [t.text for t in nested[0].iter(qn("w:t"))] == ["Nested words, edited"]   # edited where it is
    data, _ = D.export_in_place(path, imp.markdown.replace("| Nested words |", "| Nested words<br>and a second line |"))
    doc = Document(io.BytesIO(data))
    assert not doc.element.body.xpath(".//w:tbl//w:tbl")                         # new lines: the cell is rewritten, once
    assert "| FR3 | Nested words<br>and a second line |" in D.render_nodes(D.walk(doc)[0])


# ── tidying the spacing ──────────────────────────────────────────────────────
def messy_policy(path: str) -> str:
    """A Word file typed the way many are: empty lines for space, spaces and tabs to indent, double spaces, a line break at the end."""
    d = Document()
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    n = nsdecls("w")
    for xml in (
        "<w:p/>", "<w:p/>",                                                                   # empty lines that lay out the page stay
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t xml:space="preserve">  Purpose  </w:t></w:r></w:p>',
        '<w:p><w:r><w:tab/><w:t xml:space="preserve">This  policy   applies to </w:t></w:r><w:r><w:rPr><w:b/></w:rPr>'
        '<w:t xml:space="preserve"> all staff</w:t></w:r><w:r><w:t xml:space="preserve">.  </w:t><w:br/></w:r></w:p>',
        "<w:p/>", "<w:p/>", "<w:p/>",                                                         # so do these
        '<w:p><w:r><w:br w:type="page"/></w:r></w:p>',                                        # a page break is kept
        '<w:p><w:bookmarkStart w:id="1" w:name="keep"/><w:bookmarkEnd w:id="1"/></w:p>',     # so is a bookmark
        '<w:p><w:r><w:t>Scope  text.</w:t></w:r></w:p>',
        '<w:p><w:r><w:t xml:space="preserve">Signed:      Date:</w:t></w:r></w:p>',                # spaces that line up a form stay
        '<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t xml:space="preserve"> A  cell </w:t></w:r></w:p></w:tc></w:tr></w:tbl>',
        "<w:p/>", "<w:p/>",                                                                   # at the end, after the table: one stays
    ):
        sect.addprevious(parse_xml(xml.replace("<w:p", f"<w:p {n}", 1) if xml.startswith("<w:p") else xml.replace("<w:tbl>", f"<w:tbl {n}>", 1)))
    d.save(path)
    return path


@pytest.fixture
def messy(tmp_path):
    return messy_policy(str(tmp_path / "messy.docx"))


def test_the_tidy_takes_out_stray_spacing_and_keeps_what_matters(messy, tmp_path):
    before = D.import_docx(messy, twin=True)
    data, n = D.tidy(messy)
    assert n == {"spaces": 15, "line_breaks": 1, "empty_lines": 1}, n                         # 4 + 7 + 1 + 3 spaces, the tab counted
    out = tmp_path / "tidied.docx"
    out.write_bytes(data)
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert texts[:4] == ["", "", "Purpose", "This policy applies to all staff."] and texts[4:7] == ["", "", ""]
    assert len(texts) == 12 and texts[9:11] == ["Scope text.", "Signed:      Date:"] and texts[11] == ""   # the form keeps its spaces
    assert doc.tables[0].cell(0, 0).text == "A cell" and len(doc.element.body.xpath("./w:p")) == len(texts)
    assert len(doc.element.body.xpath('.//w:br[@w:type="page"]')) == 1 and doc.element.body.xpath(".//w:bookmarkStart")
    assert doc.element.body[-2].tag == qn("w:p")                                             # a paragraph still follows the table
    assert [r.bold for r in doc.paragraphs[3].runs if r.text.strip()] == [None, True, None]  # formatting stays on its words
    after = D.import_docx(str(out), twin=True)
    assert after.native is not None and [B.normalize(b.text) for b in B.parse(after.markdown)] == [B.normalize(b.text) for b in B.parse(before.markdown)]
    assert D.tidy(str(out)) == (None, {"spaces": 0, "line_breaks": 0, "empty_lines": 0})  # nothing more to do
    assert D.tidy(messy, write=False) == (None, n)                                           # the counts alone
