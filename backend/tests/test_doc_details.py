"""Creating or uploading a document asks for its details and reads them from the file: the document-control table, the revision
table, the first heading, the Purpose section, the words of its title and its file name. Comply Agent is asked only for the title,
type and description the file leaves open, from the document's opening; nothing is saved until the person confirms. The save
routes take those details, and a new document's next review is its effective date (or today) plus its review cycle."""
import json
import tempfile
from datetime import datetime

from dateutil.relativedelta import relativedelta
from docx import Document

import grc.models as m
from grc.modules.governance import doc_details as D

from test_doc_versioning import DOCX, library          # noqa: F401  (library is a fixture)
from test_studio_flow import reply, world               # noqa: F401  (world is a fixture)

NAME = "Acceptable Use Policy (Final) 24Sept2026.docx"


def control_docx(path, *, title="Acceptable Use Policy", owner="Chief Information Security Officer (CISO)", purpose=True, tail=""):
    d = Document()
    if title:
        d.add_heading(title, level=0)
    rows = [("Owner", owner), ("Approval Date", "September 24, 2026"), ("Review Cycle", "Every two years or when required"),
            ("Next Review Date", "September 2028")]
    t = d.add_table(rows=len(rows), cols=2)
    for r, (label, value) in enumerate(rows):
        t.cell(r, 0).text, t.cell(r, 1).text = label, value
    if purpose:
        d.add_heading("1. Purpose", level=1)
        d.add_paragraph("This policy sets the rules for using the bank's information systems. It applies to all staff. It is reviewed every year.")
    d.add_heading("2. Scope", level=1)
    for _ in range(40):
        d.add_paragraph("All employees, contractors and third parties who use the bank's systems and data.")
    if tail:
        d.add_paragraph(tail)
    d.save(path)
    return str(path)


def _people(db):
    db.add(m.GRCUser(id=2, username="cyrus", email="c@cfsb.test", display_name="Cyrus", designation="CISO", department="Information Security"))
    db.add(m.GRCDepartment(id=5, tenant_id=1, name="Information Security", code="IS"))
    db.commit()


def values(out):
    return {k: f["value"] for k, f in out["fields"].items()}


def test_the_file_gives_its_details_without_comply_agent(library, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    _people(library.db)
    out = D.read(control_docx(library.tmp / "aup.docx"), NAME, library.db)
    assert out["ai"] == "not_needed"
    assert values(out) == {
        "title": "Acceptable Use Policy", "doc_type": "policy",
        "description": "This policy sets the rules for using the bank's information systems. It applies to all staff.",
        "classification": "internal", "owner_id": 2, "department_id": 5, "review_cycle_months": 24,
        "effective_date": "2026-09-24", "next_review_date": "2028-09-01"}
    f = out["fields"]
    assert f["owner_id"] == {"value": 2, "from": "file", "stated": "Chief Information Security Officer (CISO)"}      # the role, by designation
    assert f["effective_date"]["stated"] == "Approval date: September 24, 2026" and f["classification"]["from"] == "default"

    # an owner role nobody holds: no one is chosen, and the screen can say whom the file names
    nobody = D.read(control_docx(library.tmp / "b.docx", owner="Head of Compliance"), NAME, library.db)["fields"]
    assert nobody["owner_id"] == {"value": None, "from": "default", "stated": "Head of Compliance"} and nobody["department_id"]["value"] is None

    # no title heading and no Purpose, no AI: the title is the cleaned file name, the type its words; the description stays open
    bare = D.read(control_docx(library.tmp / "c.docx", title=None, purpose=False), "IT Security Procedure Manaual 4.0 - Q1-2013 - Approved.docx", library.db)
    assert bare["ai"] == "unavailable"
    assert (values(bare)["title"], values(bare)["doc_type"], bare["fields"]["description"]) == (
        "IT Security Procedure Manaual", "procedure", {"value": None, "from": "default"})


def test_the_words_a_file_uses_for_its_name_cycle_and_dates():
    assert D.clean_name("Information Technology Policy 2013 - v1 1 - Approved by BoD.pdf") == "Information Technology Policy"
    assert D.clean_name("CFSB Data Governance Policy (Final) September 24, 2026.pdf") == "CFSB Data Governance Policy"
    assert [D.months_of(t) for t in ("Semi-annually", "Annually or when otherwise required", "Biennial", "Every three years", "Quarterly")] == [6, 12, 24, 36, None]
    assert D.type_of("Information Security Policy and Procedure Manual") == "procedure" and D.type_of("Board Charter") == "charter"
    assert D.parse_day("24/09/2026") == datetime(2026, 9, 24) and D.parse_day("When required") is None


def test_the_route_reads_a_file_and_saves_nothing(library, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    scratch = library.tmp / "tmp"
    scratch.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(scratch))
    with open(control_docx(library.tmp / "aup.docx"), "rb") as f:
        r = library.http.post("/documents/read-details", files={"file": (NAME, f.read(), DOCX)})
    assert r.status_code == 200, r.text
    assert r.json()["fields"]["title"] == {"value": "Acceptable Use Policy", "from": "file"} and r.json()["ai"] == "not_needed"
    assert list(scratch.iterdir()) == [] and library.db.query(m.GovernanceDocument).count() == 0          # the copy is gone, nothing saved


def test_comply_agent_fills_only_what_the_file_leaves_open_from_its_opening(world, monkeypatch, tmp_path):
    monkeypatch.setattr(D, "OPENING_CHARS", 1500)
    monkeypatch.setattr(D, "AI_SECONDS", 180)              # the guard's first call in a test process can be slow to start
    path = control_docx(tmp_path / "scan.docx", title=None, purpose=False, tail="TAIL-MARKER far into the document")
    world.respond(lambda body: reply(json.dumps({"title": "Acceptable Use Policy", "doc_type": "policy", "description": "Rules for using the bank's systems."})))
    out = D.read(path, "scan_0042.docx", world.db)
    assert out["ai"] == "used"
    f = out["fields"]
    assert (f["title"], f["doc_type"], f["description"]) == ({"value": "Acceptable Use Policy", "from": "ai"}, {"value": "policy", "from": "ai"},
                                                            {"value": "Rules for using the bank's systems.", "from": "ai"})
    assert f["effective_date"]["from"] == "file" and f["review_cycle_months"] == {"value": 24, "from": "file", "stated": "Every two years or when required"}
    sent = world.wire.requests[-1].body["messages"][0]["content"]
    assert "September 24, 2026" in sent and "TAIL-MARKER" not in sent                       # the opening only

    # Comply Agent failing or not set up: the file's own details, and the screen is told
    world.respond(lambda body: reply("not json"))
    failed = D.read(path, "scan_0042.docx", world.db)
    assert failed["ai"] == "unavailable" and failed["fields"]["title"] == {"value": "scan 0042", "from": "file"}


# ── saving ───────────────────────────────────────────────────────────────────
def _today():
    return datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0)


def test_the_library_saves_the_details_and_works_out_the_next_review(library):
    _people(library.db)
    with open(control_docx(library.tmp / "aup.docx"), "rb") as f:
        data = f.read()
    made = library.http.post("/documents/upload-with-file", files={"file": (NAME, data, DOCX)},
                             data={"title": "Acceptable Use Policy", "doc_type": "policy", "owner_id": "2", "department_id": "5",
                                   "review_cycle_months": "24", "effective_date": "2026-09-24"})
    assert made.status_code == 201, made.text
    doc = library.db.get(m.GovernanceDocument, made.json()["document"]["id"])
    assert (doc.owner_id, doc.author_id, doc.department_id, doc.review_cycle_months) == (2, 1, 5, 24)
    assert (doc.effective_date, doc.next_review_date) == (datetime(2026, 9, 24), datetime(2028, 9, 24))

    plain = library.http.post("/documents/upload-with-file", files={"file": ("b.docx", data, DOCX)}, data={"title": "B", "doc_type": "policy"})
    doc = library.db.get(m.GovernanceDocument, plain.json()["document"]["id"])
    assert (doc.owner_id, doc.review_cycle_months, doc.next_review_date) == (1, 12, _today() + relativedelta(months=12))
    assert library.http.post("/documents/upload-with-file", files={"file": ("c.docx", data, DOCX)},
                             data={"title": "C", "doc_type": "policy", "parent_document_id": "999"}).status_code == 404

    # a document created without a file has its first version from the start, and its next review too
    made = library.http.post("/documents", json={"title": "Access Policy", "doc_type": "policy", "content": "# Access Policy\n", "review_cycle_months": 6})
    assert made.status_code == 201, made.text
    doc = library.db.get(m.GovernanceDocument, made.json()["id"])
    rows = library.db.query(m.GovernanceDocumentVersion).filter_by(document_id=doc.id).all()
    assert [(v.version_number, v.status, v.content) for v in rows] == [("1.0", "current", "# Access Policy\n")]
    assert doc.next_review_date == _today() + relativedelta(months=6)


def test_the_studio_upload_and_move_to_the_library_save_the_details(world, monkeypatch):
    monkeypatch.setattr("grc.modules.governance.routers.policy_parser.maybe_autoparse_document", lambda *a, **k: False)
    up = world.http.post(world.api("/sessions/upload"), files={"file": ("retention.md", b"# Retention Standard\n\nKeep records.\n", "text/markdown")},
                         data={"doc_type": "standard", "title": "Retention Standard", "classification": "confidential", "description": "How long records are kept.",
                               "owner_id": "2", "review_cycle_months": "36", "effective_date": "2026-01-31"})
    assert up.status_code == 200, up.text
    doc = world.db.get(m.GovernanceDocument, up.json()["document_id"])
    assert (doc.title, doc.doc_type, doc.classification, doc.description, doc.owner_id, doc.author_id) == (
        "Retention Standard", "standard", "confidential", "How long records are kept.", 2, 1)
    assert (doc.review_cycle_months, doc.next_review_date) == (36, datetime(2029, 1, 31))

    s = world.http.post(world.api("/sessions"), json={"kind": "new", "title": "Access Control Policy", "template_id": "builtin:blank"}).json()
    world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": "# Access Control Policy\n\nControl access.\n", "base_revision": 1})
    moved = world.http.post(world.api(f"/sessions/{s['id']}/move-to-library"), json={"doc_type": "policy", "review_cycle_months": 6, "owner_id": 2}).json()
    doc = world.db.get(m.GovernanceDocument, moved["document_id"])
    assert (doc.owner_id, doc.review_cycle_months, doc.next_review_date) == (2, 6, _today() + relativedelta(months=6))
