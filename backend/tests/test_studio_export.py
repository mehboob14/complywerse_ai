"""AI Studio files: the logo upload, and the document exported as Word, PDF, Markdown and HTML with the entity name, logo,
header and footer, numbered lists that restart, and every placeholder still to be filled in highlighted."""
import io
import re
import zipfile

import pytest

import grc.models as m
from grc.modules.governance.studio import exporters as X

from test_studio_flow import answer, change, world, reply          # noqa: F401  (world is a fixture)

MD = """# Information Security Policy

Owner: [[Policy owner]]. Approved: ____________.

## 1. Purpose

The bank **shall** protect *all* information, including <u>customer data</u>, see [ECC](https://example.com/ecc) and [bad](javascript:alert(1)).

1. First rule
2. Second rule
   - nested bullet
3. Third rule

Some other text.

1. Restarted one
2. Restarted two

- plain bullet

## 2. Roles

| Role | Responsibility |
|------|----------------|
| CISO | [[Name]] |
| Board | Approves |

> A quotation.

<script>alert(1)</script>
"""


def png(color=(200, 30, 30, 255), size=(120, 60)):
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGBA", size, color).save(out, "PNG")
    return out.getvalue()


@pytest.fixture
def logo(tmp_path):
    path = tmp_path / "logo.png"
    path.write_bytes(png())
    return str(path)


def docx_parts(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    return z, z.read("word/document.xml").decode()


# ── the files ────────────────────────────────────────────────────────────────
def test_word_has_the_headings_lists_tables_header_footer_and_highlighted_placeholders(logo):
    from docx import Document

    data = X.to_docx(MD, title="Information Security Policy", entity_name="Capital Financial Services Bank", logo_path=logo,
                     doc_code="POL-001", version="1.0", classification="internal")
    doc = Document(io.BytesIO(data))
    styles = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert styles[0] == ("Title", "Information Security Policy")
    assert ("Heading 1", "1. Purpose") in styles and ("Heading 1", "2. Roles") in styles
    assert [t for s, t in styles if s == "List Bullet 2"] == ["nested bullet"]
    assert doc.tables[0].style.name == "Table Grid"
    z, xml = docx_parts(data)
    assert "w:tblHeader" in xml and xml.count('<w:highlight w:val="yellow"/>') == 2          # [[Policy owner]] and [[Name]], not the signature line
    assert "Approved: ____________" in xml and "<script" not in xml.lower().replace("&lt;script", "")
    numbering = [int(n) for n in re.findall(r'<w:numId w:val="(\d+)"/>', xml)]
    assert len(set(numbering)) == 2 and numbering[:3] == [numbering[0]] * 3 and numbering[3:] == [numbering[3]] * 2    # the second list is its own list
    num_xml = z.read("word/numbering.xml").decode()
    for n in set(numbering):
        assert re.search(rf'<w:num w:numId="{n}">.*?<w:startOverride w:val="1"/>', num_xml, re.S)
    header, footer = z.read("word/header1.xml").decode(), z.read("word/footer1.xml").decode()
    assert "Capital Financial Services Bank" in header and "<w:drawing>" in header
    assert "POL-001 · v1.0 · Internal" in footer and "PAGE" in footer and "NUMPAGES" in footer
    assert any(n.startswith("word/media/") for n in z.namelist())
    assert doc.core_properties.title == "Information Security Policy" and doc.core_properties.author == "Capital Financial Services Bank"
    link = [r for r in doc.part.rels.values() if r.is_external]
    assert [r.target_ref for r in link] == ["https://example.com/ecc"]                       # the javascript: link is text, not a link


def test_the_logo_can_sit_at_the_foot_of_the_page_and_a_watermark_can_go_across_every_page(logo):
    import fitz

    word = zipfile.ZipFile(io.BytesIO(X.to_docx(MD, entity_name="Capital Financial Services Bank", logo_path=logo, logo_position="footer-right", watermark="DRAFT")))
    header, footer = word.read("word/header1.xml").decode(), word.read("word/footer1.xml").decode()
    assert "Capital Financial Services Bank" in header and "<w:drawing>" not in header                 # the name stays at the top, the logo moved down
    assert "<w:drawing>" in footer and footer.index("NUMPAGES") < footer.index("<w:drawing>") and '<w:jc w:val="right"/>' in footer
    assert 'string="DRAFT"' in header and "rotation:315" in header                                       # Word's own watermark shape, in the header
    plain = zipfile.ZipFile(io.BytesIO(X.to_docx(MD, logo_path=logo))).read("word/header1.xml").decode()
    assert "<w:drawing>" in plain and "textpath" not in plain                                            # by default: the logo at the top, no watermark

    pdf = X.to_pdf(MD, entity_name="Capital Financial Services Bank", logo_path=logo, logo_position="footer-center", watermark="Draft – not approved")
    page = fitz.open(stream=pdf, filetype="pdf")[0]
    spot = page.get_image_rects(page.get_images()[0][0])[0]
    assert spot.y0 > page.rect.height * 0.9 and abs((spot.x0 + spot.x1) / 2 - page.rect.width / 2) < 2        # at the foot, centred
    text = page.get_text()
    assert "Draft – not approved" in text and "Capital Financial Services Bank" in text and "Page 1 of 1" in text

    html = X.export(MD, "html", entity_name="Capital Bank", logo_path=logo, logo_position="footer-left", watermark='Draft "1"').decode()
    assert '<footer style="text-align:left"><img src="data:image/png;base64,' in html and 'data-watermark="Draft &quot;1&quot;"' in html
    assert '<header style="justify-content:flex-start"><span>Capital Bank</span></header>' in html
    assert "data-watermark" not in X.export(MD, "html").decode()


_COLOURS = re.compile(r'(?:<w:color w:val|w:color|w:fill|fillcolor)="#?([0-9A-Fa-f]{6})"')


def coloured(hexes):
    """The colours among ``hexes`` (``RRGGBB`` or ``RGB``) that are not black, white or a grey."""
    full = {(h if len(h) == 6 else "".join(c * 2 for c in h)).upper() for h in hexes}
    return {h for h in full if not h[0:2] == h[2:4] == h[4:6]}


def test_a_document_made_here_is_black_and_white_with_light_grey_table_headers(logo):
    """Professional documents are not colourful: title, headings and entity name black (none of the theme colours of Word's own
    template, no blue line under the title), table headers and rules light grey; only links and placeholders to fill have colour."""
    import fitz

    md = MD + "\n---\n\n```\nretain = 7\n```\n"
    word = zipfile.ZipFile(io.BytesIO(X.to_docx(md, entity_name="Capital Financial Services Bank", logo_path=logo, doc_code="POL-001", watermark="DRAFT")))
    styles = word.read("word/styles.xml").decode()
    for sid in ("Title", "Heading1", "Heading2", "Heading3", "Heading4", "TitleChar", "Heading1Char", "Heading2Char", "Heading3Char", "Heading4Char"):
        style = re.search(rf'<w:style [^>]*w:styleId="{sid}".*?</w:style>', styles, re.S).group(0)
        assert '<w:color w:val="000000"/>' in style and "w:themeColor" not in style and "w:themeShade" not in style and "<w:pBdr>" not in style, sid
    assert re.search(r'w:styleId="Heading1".*?<w:sz w:val="32"/>', styles, re.S) and re.search(r'w:styleId="Title".*?<w:sz w:val="48"/>', styles, re.S)   # sizes kept
    shown = "".join(word.read(n).decode() for n in ("word/document.xml", "word/header1.xml", "word/footer1.xml"))
    assert coloured(_COLOURS.findall(shown)) == {"0563C1"}                                            # the link, nothing else
    assert 'w:fill="F2F2F2"' in shown and '<w:highlight w:val="yellow"/>' in shown and 'fillcolor="#c0c0c0"' in shown

    page = fitz.open(stream=X.to_pdf(md, entity_name="Capital Financial Services Bank", logo_path=logo, doc_code="POL-001", watermark="DRAFT"), filetype="pdf")[0]
    spans = [s for b in page.get_text("dict")["blocks"] for line in b.get("lines", []) for s in line["spans"]]
    assert coloured(f"{s['color']:06X}" for s in spans) == {"0563C1"}
    fills = {"%02X%02X%02X" % tuple(round(c * 255) for c in d[k]) for d in page.get_drawings() for k in ("fill", "color") if d.get(k)}
    assert coloured(fills) == {"FFF2A8"} and "F2F2F2" in fills and "000000" in fills               # the placeholders' highlight; grey header, black grid

    html = X.export(md, "html", entity_name="Capital Bank", logo_path=logo, watermark="DRAFT").decode()
    css = re.search(r"<style>(.*?)</style>", html, re.S).group(1)
    assert coloured(re.findall(r"#([0-9a-fA-F]{6}|[0-9a-fA-F]{3})\b", css)) == {"FFF2A8"}           # the placeholders' <mark>
    assert "h1,h2,h3,h4{color:#000}" in css and "th{background:#f2f2f2" in css and "rgba(0,0,0,.06)" in css


def test_a_numbered_list_keeps_the_number_it_starts_at():
    from docx import Document

    data = X.to_docx("# T\n\n5. Fifth\n6. Sixth\n")
    z, xml = docx_parts(data)
    n = re.search(r'<w:numId w:val="(\d+)"/>', xml).group(1)
    assert re.search(rf'<w:num w:numId="{n}">.*?<w:startOverride w:val="5"/>', z.read("word/numbering.xml").decode(), re.S)
    assert [p.text for p in Document(io.BytesIO(data)).paragraphs if p.style.name == "List Number"] == ["Fifth", "Sixth"]


def test_pdf_has_the_entity_name_page_numbers_nested_lists_and_refuses_right_to_left_text(logo):
    import fitz

    data = X.to_pdf(MD, title="Policy", entity_name="Capital Financial Services Bank", logo_path=logo, doc_code="POL-001", version="1.0", classification="internal")
    assert data.startswith(b"%PDF")
    text = "".join(page.get_text() for page in fitz.open(stream=data, filetype="pdf"))
    assert "Capital Financial Services Bank" in text and "Page 1 of 1" in text and "POL-001" in text
    assert re.search(r"1\.\s+First rule\s+2\.\s+Second rule\s+.\s*nested bullet\s+3\.\s+Third rule", text)      # nesting keeps the count going
    assert re.search(r"1\.\s+Restarted one\s+2\.\s+Restarted two", text)
    with pytest.raises(X.ExportProblem, match="right-to-left"):
        X.to_pdf("# مرحبا\n\nنص عربي\n")


def test_a_long_document_gets_page_x_of_y_on_every_page():
    import fitz

    data = X.to_pdf("# Long\n\n" + "\n\n".join(f"Paragraph {i} " + "word " * 60 for i in range(120)))
    pages = fitz.open(stream=data, filetype="pdf")
    assert len(pages) > 2
    assert all(f"Page {i + 1} of {len(pages)}" in p.get_text() for i, p in enumerate(pages))


def test_html_is_escaped_marks_placeholders_and_embeds_the_logo(logo):
    page = X.export(MD, "html", title="Policy", entity_name="Capital <Bank>", logo_path=logo).decode()
    assert "<script" not in page.lower() and "<p>alert(1)</p>" in page                              # a tag in the text is dropped, never run
    assert "<mark>[[Policy owner]]</mark>" in page and "<mark>____" not in page
    assert 'href="https://example.com/ecc"' in page and "javascript:" not in page.replace("[bad](javascript:alert(1))", "")
    assert "data:image/png;base64," in page and "Capital &lt;Bank&gt;" in page
    assert "<table>" in page and "<ol>" in page and "<ul>" in page and "<u>customer data</u>" in page


def test_markdown_is_the_text_as_it_is():
    assert X.export(MD, "md") == MD.encode("utf-8")
    assert X.export("# T", "md") == b"# T\n"
    with pytest.raises(X.ExportProblem):
        X.export(MD, "exe")


# ── logos ────────────────────────────────────────────────────────────────────
def upload(world, data, name="logo.png", **form):
    return world.http.post(world.api("/assets"), files={"file": (name, data, "image/png")}, data=form)


def test_a_logo_is_checked_cleaned_listed_and_served(world):
    r = upload(world, png(size=(3000, 1500)), name="Bank Logo.png")
    assert r.status_code == 200
    a = r.json()
    assert a["name"] == "Bank Logo" and a["mime"] == "image/png" and a["is_default"] is True and (a["width"], a["height"]) == (1600, 800)   # shrunk to the box
    second = upload(world, png(), name="Dark.png", make_default="true").json()
    items = world.http.get(world.api("/assets")).json()["items"]
    assert [(i["id"], i["is_default"]) for i in items] == [(second["id"], True), (a["id"], False)]
    served = world.http.get(world.api(f"/assets/{second['id']}/file"))
    assert served.status_code == 200 and served.headers["content-type"] == "image/png" and served.content[:4] == b"\x89PNG"
    new = world.http.post(world.api("/sessions"), json={"kind": "new", "title": "Memo", "template_id": "builtin:blank", "doc_type": "other"}).json()
    assert new["instructions"]["logo_asset_id"] == second["id"]                      # the default logo goes on new documents


def test_something_that_is_not_a_usable_picture_is_refused(world):
    assert upload(world, b"<svg xmlns='http://www.w3.org/2000/svg'><script>alert(1)</script></svg>", name="logo.svg").status_code == 400
    assert upload(world, b"MZ\x90\x00 not a picture", name="logo.png").status_code == 400
    assert upload(world, b"").status_code == 400
    assert upload(world, b"\x89PNG\r\n\x1a\n" + b"0" * (4 * 1024 * 1024)).status_code == 400
    assert world.http.get(world.api("/assets")).json()["items"] == []


def test_a_logo_uploaded_with_trailing_junk_is_saved_again_clean(world):
    a = upload(world, png() + b"<?php evil(); ?>").json()
    body = world.http.get(world.api(f"/assets/{a['id']}/file")).content
    assert b"<?php" not in body and body.startswith(b"\x89PNG")


def test_only_the_uploader_or_an_administrator_changes_or_removes_a_logo(world):
    a = upload(world, png()).json()
    world.me["id"] = 2
    assert world.http.put(world.api(f"/assets/{a['id']}"), json={"name": "Mine now"}).status_code == 403
    assert world.http.delete(world.api(f"/assets/{a['id']}")).status_code == 403
    assert world.http.get(world.api(f"/assets/{a['id']}/file")).status_code == 200                    # everyone may use it
    world.me["id"] = 1
    assert world.http.put(world.api(f"/assets/{a['id']}"), json={"name": "Primary"}).json()["name"] == "Primary"
    assert world.http.delete(world.api(f"/assets/{a['id']}")).json() == {"ok": True}
    assert world.http.get(world.api(f"/assets/{a['id']}/file")).status_code == 404 and world.http.get(world.api("/assets")).json()["items"] == []


# ── the export endpoints ─────────────────────────────────────────────────────
def start(world, content=MD, title="Information Security Policy"):
    s = world.http.post(world.api("/sessions"), json={"kind": "new", "title": title, "template_id": "builtin:blank", "doc_type": "other"}).json()
    assert world.http.put(world.api(f"/sessions/{s['id']}/content"), json={"content": content, "base_revision": s["revision"]}).status_code == 200
    return s["id"]


def test_the_document_downloads_as_each_format_with_the_logo_and_entity_name(world):
    logo_id = upload(world, png()).json()["id"]
    sid = start(world)
    for fmt, mime, magic in (("docx", "officedocument.wordprocessingml", b"PK"), ("pdf", "application/pdf", b"%PDF"), ("md", "text/markdown", b"# Information"), ("html", "text/html", b"<!doctype")):
        r = world.http.post(world.api(f"/sessions/{sid}/export"), json={"format": fmt})
        assert r.status_code == 200 and mime in r.headers["content-type"] and r.content.startswith(magic), fmt
        assert r.headers["content-disposition"].startswith('attachment; filename="Information-Security-Policy.' + fmt) and r.headers["cache-control"] == "no-store"
    z, _ = docx_parts(world.http.post(world.api(f"/sessions/{sid}/export"), json={"format": "docx"}).content)
    assert "Capital Financial Services Bank" in z.read("word/header1.xml").decode() and any(n.startswith("word/media/") for n in z.namelist()) and logo_id
    plain = world.http.post(world.api(f"/sessions/{sid}/export"), json={"format": "docx", "no_logo": True, "entity_name": "Other Bank"}).content
    z, _ = docx_parts(plain)
    assert "Other Bank" in z.read("word/header1.xml").decode() and not any(n.startswith("word/media/") for n in z.namelist())


def test_a_document_in_the_library_is_exported_with_its_code_version_and_classification(world):
    world.db.get(m.GovernanceDocument, 10).document_code = "POL-001"
    world.db.commit()
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 10}).json()
    r = world.http.post(world.api(f"/sessions/{s['id']}/export"), json={"format": "docx"})
    assert r.headers["content-disposition"].startswith('attachment; filename="Information-Security-Policy-v1.0.docx"')
    z, _ = docx_parts(r.content)
    assert "POL-001 · v1.0" in z.read("word/footer1.xml").decode()


def test_the_preflight_names_what_is_left_to_do_and_what_cannot_be_exported(world):
    sid = start(world)
    check = world.http.post(world.api(f"/sessions/{sid}/export/check"), json={"format": "docx"}).json()
    assert check["can_export"] is True
    rules = {n["rule"]: n for n in check["notes"]}
    assert rules["placeholders"]["count"] == 2 and "highlighted" in rules["placeholders"]["text"] and rules["logo"]["level"] == "info"
    assert "entity" not in rules                                                                      # the tenant's legal entity is already set
    world.http.patch(world.api(f"/sessions/{sid}"), json={"instructions": {"entity_name": "", "logo_required": True}})
    rules = {n["rule"]: n for n in world.http.post(world.api(f"/sessions/{sid}/export/check"), json={"format": "pdf"}).json()["notes"]}
    assert rules["entity"]["level"] == "warning" and rules["logo"]["level"] == "warning"
    arabic = start(world, "# مرحبا\n\nنص عربي\n", "Arabic")
    blocked = world.http.post(world.api(f"/sessions/{arabic}/export/check"), json={"format": "pdf"}).json()
    assert blocked["can_export"] is False and blocked["notes"][-1]["level"] == "error"
    assert world.http.post(world.api(f"/sessions/{arabic}/export"), json={"format": "pdf"}).status_code == 400
    assert world.http.post(world.api(f"/sessions/{arabic}/export"), json={"format": "docx"}).status_code == 200      # Word can
    assert world.http.post(world.api(f"/sessions/{sid}/export"), json={"format": "exe"}).status_code == 400


def test_changes_the_ai_proposed_and_nobody_accepted_are_flagged_and_left_out(world):
    s = world.http.post(world.api("/sessions"), json={"kind": "library", "document_id": 10}).json()
    world.respond(lambda body: reply(answer([change("replace", "B7", markdown="Access to systems shall be reviewed every quarter.")], "Made it quarterly.")))
    world.http.post(world.api(f"/sessions/{s['id']}/messages"), json={"text": "Quarterly please."})
    check = world.http.post(world.api(f"/sessions/{s['id']}/export/check"), json={"format": "md"}).json()
    assert [(n["rule"], n["count"]) for n in check["notes"] if n["rule"] == "pending"] == [("pending", 1)]
    assert b"every quarter" not in world.http.post(world.api(f"/sessions/{s['id']}/export"), json={"format": "md"}).content


def test_another_persons_session_cannot_be_exported(world):
    sid = start(world)
    world.me["id"] = 2
    assert world.http.post(world.api(f"/sessions/{sid}/export"), json={"format": "md"}).status_code == 404
    assert world.http.post(world.api(f"/sessions/{sid}/export/check"), json={"format": "md"}).status_code == 404
