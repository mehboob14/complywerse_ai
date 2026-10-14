"""Policy statements <-> the original document (cfsb): statements are held to the file's own wording, and an edit to a
statement is written back into the file as the next document version."""
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.governance import doc_sync
from grc.modules.governance.routers import documents, policy_parser
from grc.routers.auth_router import require_auth

S1 = "All administrators shall use multi-factor authentication."
S2 = "Passwords must be changed every 90 days."
S3 = "Backups shall be tested “quarterly”."          # curly quotes in the file
LONG = "Access rights must be reviewed at least once every year by the system owner."


def make_docx(path, extra=()):
    doc = Document()
    doc.add_paragraph("Information Security Policy")
    p = doc.add_paragraph()
    p.add_run("1. ").bold = True
    p.add_run(S1)
    doc.add_paragraph(S2 + " " + "Accounts that are idle are disabled.")
    doc.add_paragraph(S3)
    wrapped = doc.add_paragraph("Access rights must be reviewed at least")
    wrapped.add_run(" once every year by the system owner.")
    for line in extra:
        doc.add_paragraph(line)
    doc.save(path)


@pytest.fixture
def world(tmp_path, monkeypatch):
    monkeypatch.setattr(documents, "GOVERNANCE_UPLOAD_DIR", str(tmp_path))
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    m.Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    path = str(tmp_path / "policy.docx")
    make_docx(path)
    with factory() as db:
        db.add(m.Tenant(id=1, name="CFSB", slug="cfsb"))
        db.add(m.GRCUser(id=1, username="auditor", email="a@cfsb.test", display_name="A"))
        db.add(m.GovernanceDocument(id=5, tenant_id=1, title="InfoSec policy", doc_type="policy", status="approved",
                                    current_version="1.0", file_name="policy.docx", file_path=path, file_type="docx",
                                    file_size=1, content=None))
        for i, text in enumerate([S1, S2, S3.replace("“", '"').replace("”", '"'), LONG], start=1):
            db.add(m.PolicyStatement(tenant_id=1, document_id=5, statement_code=f"PS-0005-{i:03d}", statement_text=text, status="active"))
        db.commit()
    return SimpleNamespace(factory=factory, path=path, tmp=tmp_path, user=SimpleNamespace(id=1))


def statement(db, code):
    return db.query(m.PolicyStatement).filter_by(statement_code=code).one()


def texts(path):
    return [p.text for p in Document(path).paragraphs]


def test_statements_that_are_in_the_file_are_matched_whatever_the_quotes_or_line_breaks(world):
    with world.factory() as db:
        counts = doc_sync.anchor(db, db.get(m.GovernanceDocument, 5))
        assert counts == {"exact": 4, "corrected": 0, "not_found": 0, "total": 4}
        s3 = statement(db, "PS-0005-003")
        assert s3.source_text == S3 and s3.source_match == "exact"         # the file's curly quotes, kept as the file has them
        assert statement(db, "PS-0005-004").source_text == LONG              # written over two runs in the file


def test_a_paraphrase_is_corrected_to_the_files_wording_and_an_invented_statement_is_flagged(world):
    with world.factory() as db:
        a = statement(db, "PS-0005-001")
        a.statement_text = "All administrators must use multi-factor authentication."          # the AI changed one word
        db.add(m.PolicyStatement(tenant_id=1, document_id=5, statement_code="PS-0005-009", status="active",
                                 statement_text="Staff shall wear badges at all times in every building."))
        db.commit()
        counts = doc_sync.anchor(db, db.get(m.GovernanceDocument, 5))
        assert (counts["corrected"], counts["not_found"]) == (1, 1)
        assert a.statement_text == S1 and a.source_match == "corrected"
        assert [v.change_type for v in a.versions] == ["verbatim_correction"]
        ghost = statement(db, "PS-0005-009")
        assert ghost.source_match == "not_found" and ghost.source_text is None and ghost.statement_text.startswith("Staff")


def test_the_same_sentence_twice_in_a_file_is_two_places(world):
    make_docx(world.path, extra=[S1])
    with world.factory() as db:
        db.add(m.PolicyStatement(tenant_id=1, document_id=5, statement_code="PS-0005-010", statement_text=S1, status="active"))
        db.commit()
        doc_sync.anchor(db, db.get(m.GovernanceDocument, 5))
        starts = {statement(db, c).source_start for c in ("PS-0005-001", "PS-0005-010")}
        assert len(starts) == 2


def test_an_edited_statement_is_written_into_the_word_file_and_makes_the_next_version(world):
    with world.factory() as db:
        doc = db.get(m.GovernanceDocument, 5)
        doc_sync.anchor(db, doc)
        assert doc_sync.state(db, doc)["pending"] == []
        statement(db, "PS-0005-001").statement_text = "All administrators and contractors shall use multi-factor authentication."
        statement(db, "PS-0005-004").statement_text = LONG.replace("every year", "every six months")
        db.commit()
        info = doc_sync.state(db, doc)
        assert [p["code"] for p in info["pending"]] == ["PS-0005-001", "PS-0005-004"] and info["can_apply"]

        out = doc_sync.apply(db, doc, world.user, "Annual review became six-monthly")
        assert (out["version"], out["updated"], out["sent_to_review"], out["word_copy"]) == ("1.1", 2, True, False)
        assert doc.status == "pending_review" and doc.current_version == "1.1" and doc.file_path != world.path

        new = texts(doc.file_path)
        assert "1. All administrators and contractors shall use multi-factor authentication." in new
        assert "Access rights must be reviewed at least once every six months by the system owner." in new
        assert S2 + " Accounts that are idle are disabled." in new and "Information Security Policy" in new
        runs = Document(doc.file_path).paragraphs[1].runs
        assert runs[0].text == "1. " and runs[0].bold                         # the numbering run kept its formatting
        assert texts(world.path)[1] == "1. " + S1                              # the earlier file is untouched

        versions = db.query(m.GovernanceDocumentVersion).filter_by(document_id=5).order_by(m.GovernanceDocumentVersion.id).all()
        assert [(v.version_number, v.status) for v in versions] == [("1.0", "superseded"), ("1.1", "current")]
        assert versions[0].file_path == world.path and versions[1].file_path == doc.file_path
        assert "PS-0005-001" in versions[1].change_summary and versions[1].change_reason.startswith("Annual")
        assert db.query(m.DocumentAuditLog).filter_by(document_id=5, action="statements_written_to_document").count() == 1

        assert doc_sync.state(db, doc)["pending"] == []                           # in step again, and still found in the new file
        assert statement(db, "PS-0005-001").source_match == "exact"
        with pytest.raises(LookupError):
            doc_sync.apply(db, doc, world.user)


def test_a_draft_document_stays_a_draft_and_an_unmatched_statement_cannot_be_written(world):
    with world.factory() as db:
        doc = db.get(m.GovernanceDocument, 5)
        doc.status = "draft"
        db.add(m.PolicyStatement(tenant_id=1, document_id=5, statement_code="PS-0005-009", status="active",
                                 statement_text="Staff shall wear badges at all times in every building."))
        db.commit()
        doc_sync.anchor(db, doc)
        statement(db, "PS-0005-009").statement_text = "Staff shall wear badges."
        statement(db, "PS-0005-002").statement_text = "Passwords must be changed every 60 days."
        db.commit()
        info = doc_sync.state(db, doc)
        assert [p["code"] for p in info["pending"]] == ["PS-0005-002"] and [b["code"] for b in info["not_found"]] == ["PS-0005-009"]
        assert doc_sync.apply(db, doc, world.user)["status"] == "draft"


def test_a_pdf_gets_a_word_copy_and_stays_as_the_earlier_version(world, monkeypatch):
    pdf = str(world.tmp / "policy.pdf")
    open(pdf, "wb").write(b"%PDF-1.4")
    text = "\n".join(["Information Security Policy", S1, S2 + " Accounts that are idle", "are disabled.", S3, LONG])
    monkeypatch.setattr(policy_parser, "extract_text_from_file", lambda path, ft: text)
    with world.factory() as db:
        doc = db.get(m.GovernanceDocument, 5)
        doc.file_name, doc.file_path, doc.file_type = "policy.pdf", pdf, "pdf"
        db.commit()
        assert doc_sync.anchor(db, doc)["exact"] == 4
        statement(db, "PS-0005-002").statement_text = "Passwords must be changed every 60 days."
        db.commit()
        assert doc_sync.state(db, doc)["creates_word_copy"]
        out = doc_sync.apply(db, doc, world.user)
        assert out["word_copy"] and doc.file_type == "docx" and doc.file_name == "policy.docx"
        assert "Passwords must be changed every 60 days. Accounts that are idle" in " ".join(texts(doc.file_path))
        base = db.query(m.GovernanceDocumentVersion).filter_by(document_id=5, version_number="1.0").one()
        assert base.file_path == pdf and base.file_type == "pdf"


def test_only_the_configured_tenants_have_it(world):
    with world.factory() as db:
        db.query(m.Tenant).one().slug = "acme"
        db.commit()
        assert not doc_sync.enabled(db)
        doc_sync.anchor_after_parse(db, 5)
        assert statement(db, "PS-0005-001").source_match is None


def test_the_routes_report_check_and_apply(world):
    from grc.modules.governance.routers.policy_parser import router
    from grc.models import get_db

    app = FastAPI()
    app.include_router(router)
    db = world.factory()
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: SimpleNamespace(id=1)
    client = TestClient(app)

    assert client.post("/documents/5/document-sync/apply", json={}).status_code == 409           # nothing edited yet
    info = client.post("/documents/5/document-sync/check").json()
    assert info["counts"]["exact"] == 4 and info["pending"] == []
    listing = client.get("/documents/5/policy-statements").json()["statements"]
    assert {s["source_match"] for s in listing} == {"exact"} and not any(s["pending_in_document"] for s in listing)

    client.put("/documents/5/statements/%d" % statement(db, "PS-0005-002").id,
               json={"statement_text": "Passwords must be changed every 45 days.", "change_reason": "tighter"})
    assert client.get("/documents/5/document-sync").json()["pending"][0]["code"] == "PS-0005-002"
    done = client.post("/documents/5/document-sync/apply", json={"change_reason": "tighter"}).json()
    assert done["version"] == "1.1" and done["sent_to_review"]
    assert any("every 45 days" in t for t in texts(db.get(m.GovernanceDocument, 5).file_path))
    db.query(m.Tenant).one().slug = "acme"
    db.commit()
    assert client.get("/documents/5/document-sync").json() == {"enabled": False}
    assert client.post("/documents/5/document-sync/check").status_code == 404
