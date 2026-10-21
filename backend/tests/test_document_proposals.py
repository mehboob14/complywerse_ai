"""A change to a governance document - its text, its file, or one of its statements - by anyone who is not its owner or author waits
for them to accept or reject it (DECISIONS D35, the user's asks of 2026-10-21; modules/governance/proposals.py). Every writer keeps
it as a proposal and changes nothing; the owner accepts it (applied as the next version, credited to the proposer) or rejects it
with a note (the proposer revises it or leaves it). In review a proposed change holds the route: accepted, the stage starts again
(minor) or the route (major); not accepted, its proposer signs it as it is (the others' signatures stand) or proposes again. The
owner's and author's own changes are direct, as they always were. Through the HTTP routes, on a session that does not autoflush
(as tenant sessions do not) with foreign keys on."""
import io
import os
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.governance import change_chain as C
from grc.modules.governance import doc_lifecycle as life
from grc.modules.governance import doc_sync, proposals
from grc.modules.governance.routers import document_signoff
from grc.modules.governance.studio import service as studio_service
from grc.modules.workflow_engine.services.action_handlers import WorkflowActionHandlers
from grc.routers.auth_router import require_auth

import _docx_fixtures as F
from test_esignature import _png

OLIVIA, RANA, RAFAY, MINA, ADAM, UMA, SAM = 1, 2, 3, 4, 5, 6, 7     # Olivia owns and wrote it; Rana and Uma may edit; Adam administers
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT = ("# Access policy\n\n## 1. Purpose\n\nThis policy sets how access to the bank's systems is requested, approved, granted, "
        "reviewed every quarter and removed when people move or leave, so that only the right people reach customer data.\n\n"
        "## 2. Scope\n\nAll staff, contractors, vendors and service accounts with access to production systems.\n")
SMALL = TEXT.replace("contractors, vendors", "contractors, temporary staff, vendors")          # a minor change
BIG = TEXT + "\n## 3. Exceptions\n\nExceptions need the CISO's written approval.\n"             # a major one: a section added
WHY = "Temporary staff were missing from the scope."
S1 = "All administrators shall use multi-factor authentication."
S1_NOW = "All administrators and service owners shall use multi-factor authentication."


@pytest.fixture
def w(monkeypatch, tmp_path):
    monkeypatch.setenv("SESSION_SECRET", "unit-test-secret")
    folder = tmp_path / "gov"
    folder.mkdir()
    monkeypatch.setattr("grc.modules.governance.routers.documents.GOVERNANCE_UPLOAD_DIR", str(folder))
    monkeypatch.setattr("grc.modules.governance.routers.documents._resolve_tenant_slug_for_autoparse", lambda tenant_id: None)
    monkeypatch.setattr(studio_service, "UPLOAD_DIR", str(tmp_path / "studio"))
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    event.listen(engine, "connect", lambda conn, _: conn.execute("PRAGMA foreign_keys=ON"))
    m.Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine, autoflush=False)()          # as tenant sessions are (PITFALLS 74)
    db.add(m.Tenant(id=1, name="CFSB", slug="cfsb"))
    for uid, name in [(OLIVIA, "Olivia Owner"), (RANA, "Rana Reviewer"), (RAFAY, "Rafay Reviewer"), (MINA, "Mina Manager"),
                      (ADAM, "Adam Admin"), (UMA, "Uma Editor"), (SAM, "Sam Staff")]:
        db.add(m.GRCUser(id=uid, username=name.split()[0].lower(), email=f"{name.split()[0].lower()}@cfsb.test", display_name=name))
    db.flush()
    db.add(m.Role(id=10, tenant_id=1, name="Editors"))
    db.add(m.Permission(id=10, name="governance:policies:edit", resource="governance:policies", action="edit"))
    db.add(m.Role(id=11, tenant_id=1, name="Administrator"))
    db.add(m.Role(id=900, tenant_id=1, name="Staff"))                       # everyone opens Task Management
    db.add(m.Permission(id=900, name=life.TASKS_VIEW, resource="critical_tasks:tasks", action="view"))
    db.flush()
    db.add(m.RolePermission(role_id=10, permission_id=10))
    db.add(m.RolePermission(role_id=900, permission_id=900))
    for uid in (RANA, UMA):
        db.add(m.UserRole(user_id=uid, role_id=10, tenant_id=1))
    db.add(m.UserRole(user_id=ADAM, role_id=11, tenant_id=1))
    for uid in range(1, 8):
        db.add(m.UserRole(user_id=uid, role_id=900, tenant_id=1))
    db.add(m.GovernanceDocument(id=5, tenant_id=1, title="Access policy", doc_type="policy", status="draft", current_version="1.0",
                                content=TEXT, owner_id=OLIVIA, author_id=OLIVIA))
    db.commit()

    from grc.modules.governance.router import router

    app = FastAPI()
    app.include_router(router)
    me = {"id": OLIVIA}
    app.dependency_overrides[m.get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: db.get(m.GRCUser, me["id"])
    http = TestClient(app)

    def call(method, path, uid=OLIVIA, **kw):
        me["id"] = uid
        return http.request(method, f"/governance{path}", **kw)

    yield SimpleNamespace(db=db, call=call, folder=folder)
    db.close()


# ── helpers ──────────────────────────────────────────────────────────────────
def doc(w, document_id=5):
    w.db.expire_all()
    return w.db.get(m.GovernanceDocument, document_id)


def put(w, uid, document_id=5, **body):
    return w.call("PUT", f"/documents/{document_id}", uid, json=body)


def propose_edit(w, uid=RANA, content=SMALL, why=WHY):
    out = put(w, uid, content=content, change_reason=why)
    assert out.status_code == 200 and out.json()["proposed"] is True, out.text
    return out.json()["proposal"]["id"]


def accept(w, pid, uid=OLIVIA, document_id=5, note=None):
    return w.call("POST", f"/documents/{document_id}/proposals/{pid}/accept", uid, json={"note": note})


def reject(w, pid, note, uid=OLIVIA, document_id=5):
    return w.call("POST", f"/documents/{document_id}/proposals/{pid}/reject", uid, json={"note": note})


def withdraw(w, pid, uid=RANA, document_id=5):
    return w.call("POST", f"/documents/{document_id}/proposals/{pid}/withdraw", uid)


def listing(w, uid=OLIVIA, document_id=5):
    return w.call("GET", f"/documents/{document_id}/proposals", uid).json()


def versions(w, document_id=5):
    w.db.expire_all()
    rows = w.db.query(m.GovernanceDocumentVersion).filter_by(document_id=document_id).order_by(m.GovernanceDocumentVersion.id).all()
    return [(v.version_number, v.created_by) for v in rows]


def version_row(w, number, document_id=5):
    w.db.expire_all()
    return (w.db.query(m.GovernanceDocumentVersion).filter_by(document_id=document_id, version_number=number)
            .order_by(m.GovernanceDocumentVersion.id.desc()).first())


def alerts(w, uid):
    return [n.subject for n in w.db.query(m.WorkflowNotification).filter_by(user_id=uid).order_by(m.WorkflowNotification.id)]


def last_alert(w, uid):
    return w.db.query(m.WorkflowNotification).filter_by(user_id=uid).order_by(m.WorkflowNotification.id.desc()).first()


def pending(w, kind=None, document_id=5):
    w.db.expire_all()
    q = w.db.query(m.DocumentObligation).filter_by(document_id=document_id, status="pending")
    return sorted((o.kind, o.user_id) for o in (q.filter_by(kind=kind) if kind else q))


def person(uid):
    return {"target_type": "user", "target_id": uid}


def send(w, *stages):
    out = w.call("PUT", "/documents/5/signoff/route", OLIVIA, json={"stages": list(stages)})
    assert out.status_code == 200, out.text
    out = w.call("POST", "/documents/5/signoff/send-for-review", OLIVIA)
    assert out.status_code == 200, out.text


def review_then_approval(w, reviewers=(RANA, RAFAY, SAM)):
    send(w, {"kind": "review", "name": "Review", "people": [person(u) for u in reviewers]},
         {"kind": "approval", "name": "Approval", "people": [person(MINA)]})


def sign(w, uid):
    return w.call("POST", "/documents/5/signoff/sign", uid, json={"signature_image": _png()})


def signoff(w, uid=OLIVIA):
    return w.call("GET", "/documents/5/signoff", uid).json()


def hold(w):
    w.db.expire_all()
    return w.db.query(m.DocumentRouteHold).filter_by(document_id=5).order_by(m.DocumentRouteHold.id.desc()).first()


def on(dt):
    return f"{dt.day} {dt:%b %Y}"


def word_doc(w, document_id=6, status="draft"):
    """A document that is its Word file (no text of its own), with its version row."""
    path = F.corporate_policy(str(w.folder / f"orig_{document_id}.docx"))
    size = os.path.getsize(path)
    w.db.add(m.GovernanceDocument(id=document_id, tenant_id=1, title="Uploaded policy", doc_type="policy", status=status, current_version="1.0",
                                  file_name="orig.docx", file_path=path, file_size=size, file_type="docx", content=None,
                                  owner_id=OLIVIA, author_id=OLIVIA))
    w.db.add(m.GovernanceDocumentVersion(document_id=document_id, version_number="1.0", change_type="major", title="Uploaded policy",
                                         file_name="orig.docx", file_path=path, file_size=size, file_type="docx",
                                         change_summary="Initial version", status="current", created_by=OLIVIA))
    w.db.commit()
    return path


def longer_passwords(original, path):
    """The same Word file with one sentence changed."""
    d = Document(original)
    p = next(p for p in d.paragraphs if "twelve characters" in p.text)
    p.runs[0].text = p.runs[0].text.replace("twelve", "fourteen")
    d.save(path)
    return str(path)


def file_text(path_or_bytes):
    d = Document(io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes)
    return "\n".join(p.text for p in d.paragraphs)


def chain_events(w, document_id=5):
    return [e["kind"] for e in C.chain(w.db, doc(w, document_id))["events"]
            if e["kind"] in ("change_proposed", "proposal_accepted", "proposal_rejected", "proposal_withdrawn")]


# ── the rule ─────────────────────────────────────────────────────────────────
def test_an_editors_change_waits_for_the_owner_and_changes_nothing_while_the_owners_own_is_direct(w):
    out = put(w, RANA, content=SMALL)                                   # no why: refused, nothing kept
    assert out.status_code == 400 and out.json()["detail"] == proposals.MESSAGE_REQUIRED
    assert w.db.query(m.DocumentChangeProposal).count() == 0

    out = put(w, RANA, content=SMALL, change_reason=WHY, description="Who may reach what")
    body = out.json()
    assert out.status_code == 200 and body["proposed"] is True
    assert body["message"] == "Sent to Olivia Owner to accept. The document does not change until then."
    p = body["proposal"]
    assert (p["kind"], p["source"], p["status"], p["message"], p["stale"], p["holds_review"]) == ("document", "edit", "pending", WHY, False, False)
    assert p["proposed_by"] == {"id": RANA, "name": "Rana Reviewer"} and p["base"] == {"version_id": None, "number": "1.0"}
    assert (p["may_accept"], p["may_reject"], p["may_withdraw"], p["statement"], p["has_file"]) == (False, False, True, None, False)
    d = doc(w)
    assert (d.content, d.current_version, d.file_path) == (TEXT, "1.0", None) and versions(w) == []
    assert d.description == "Who may reach what" and body["description"] == "Who may reach what"     # details are saved as always

    mine = put(w, OLIVIA, content=BIG, change_reason="Exceptions").json()        # the owner's own edit: as it always was
    assert "proposed" not in mine and mine["current_version"] == "2.0" and doc(w).content == BIG


def test_the_owner_accepts_it_as_the_next_version_credited_to_whoever_proposed_it(w):
    pid = propose_edit(w)
    listed = listing(w)
    assert (listed["can_decide"], listed["pending"], [i["id"] for i in listed["items"]]) == (True, 1, [pid])
    assert listed["items"][0]["may_accept"] is True and listing(w, RANA)["can_decide"] is False
    cmp = w.call("GET", f"/documents/5/proposals/{pid}/compare", OLIVIA).json()
    assert cmp["base"] == {"label": "v1.0 · in the library", "version_id": None}
    assert cmp["proposed"]["label"].startswith("Proposed by Rana Reviewer · ")
    assert cmp["stats"]["modified"] == 1 and len(cmp["changes"]) == 1 and "temporary staff" in cmp["changes"][0]["after_text"]
    assert any(r.get("hunk") == cmp["changes"][0]["id"] for r in cmp["rows"])
    assert w.call("GET", f"/documents/5/proposals/{pid}/file", OLIVIA).status_code == 404       # a text proposal has no file
    assert accept(w, pid, uid=RANA).status_code == 403                                          # only the owner (or author) decides

    out = accept(w, pid, note="Thanks")
    assert out.status_code == 200, out.text
    assert out.json()["version"]["number"] == "1.1" and out.json()["proposal"]["status"] == "accepted"
    assert out.json()["proposal"]["applied_version"]["number"] == "1.1" and out.json()["proposal"]["decision_note"] == "Thanks"
    d = doc(w)
    assert (d.content, d.current_version) == (SMALL, "1.1")
    row = version_row(w, "1.1")
    assert row.created_by == RANA and row.change_reason == WHY
    assert row.change_summary.endswith("· Proposed by Rana Reviewer, accepted by Olivia Owner")
    assert alerts(w, OLIVIA) == ["Change proposed: Access policy"]
    assert last_alert(w, OLIVIA).message == f'Rana Reviewer proposed a change: "{WHY}". Accept or reject it. It is in your tasks.'
    told = last_alert(w, RANA)
    assert (told.subject, told.message) == ("Change accepted: Access policy", 'Your change to "Access policy" was accepted (v1.1).')
    asked = w.db.query(m.DocumentObligation).filter_by(kind="decide").one()
    assert (asked.user_id, asked.status, asked.proposal_id) == (OLIVIA, "done", pid) and pending(w) == []
    task = w.db.get(m.CriticalTask, asked.critical_task_id)
    assert (task.title, task.status) == ("Decide on: Access policy", "Completed")
    shown = life.task_request(w.db, task)
    assert (shown["kind"], shown["proposal_id"], shown["action_label"]) == ("decide", pid, "Decide on a proposed change")
    assert chain_events(w) == ["change_proposed", "proposal_accepted"]
    assert [C.EVENT_LABEL[k] for k in chain_events(w)] == ["Change proposed", "Proposed change accepted"]
    assert accept(w, pid).status_code == 409                                                    # decided once


def test_not_accepted_needs_a_note_keeps_everything_and_asks_its_proposer_to_revise_it_or_leave_it(w):
    pid = propose_edit(w)
    assert reject(w, pid, "  ").status_code == 400
    assert reject(w, pid, "No", uid=RANA).status_code == 403
    out = reject(w, pid, "Contractors cover them.")
    assert out.status_code == 200 and out.json()["proposal"]["status"] == "rejected"
    assert doc(w).content == TEXT and versions(w) == []
    told = last_alert(w, RANA)
    assert told.subject == "Change not accepted: Access policy"
    assert told.message == 'Your change to "Access policy" was not accepted: "Contractors cover them.". Revise it and propose it again, or leave it. It is in your tasks.'
    assert told.link == f"/governance/documents/5?tab=proposals&proposal={pid}"
    assert pending(w) == [("revise", RANA)]
    revise = w.db.query(m.DocumentObligation).filter_by(kind="revise").one()
    task = w.db.get(m.CriticalTask, revise.critical_task_id)
    shown = life.task_request(w.db, task)
    assert task.title == "Revise your change to: Access policy"
    assert (shown["proposal_id"], shown["action_href"], shown["action_label"]) == (pid, f"/governance/documents/5?tab=proposals&proposal={pid}",
                                                                                   "Your change was not accepted")
    mine = w.call("GET", "/documents/lifecycle/mine", RANA).json()["items"]
    assert [(i["kind"], i["proposal_id"]) for i in mine] == [("revise", pid)]

    assert withdraw(w, pid, uid=OLIVIA).status_code == 403                  # only its proposer leaves it
    out = withdraw(w, pid)
    assert out.json()["proposal"]["status"] == "withdrawn" and pending(w) == []
    assert w.db.get(m.CriticalTask, task.id).status == "Completed"
    assert chain_events(w) == ["change_proposed", "proposal_rejected", "proposal_withdrawn"]


def test_withdrawing_one_that_waits_ends_the_owners_request_and_tells_them(w):
    pid = propose_edit(w)
    assert pending(w) == [("decide", OLIVIA)]
    assert withdraw(w, pid).status_code == 200
    assert pending(w) == [] and alerts(w, OLIVIA)[-1] == "Proposed change withdrawn: Access policy"
    assert withdraw(w, pid).status_code == 409


def test_a_change_made_on_an_older_version_is_stale_it_cannot_be_accepted_but_can_be_rejected(w):
    pid = propose_edit(w)
    put(w, OLIVIA, content=BIG, change_reason="Exceptions")             # the owner's own edit lands first: 2.0
    out = accept(w, pid)
    assert out.status_code == 409
    assert out.json() == {"detail": "This change was made on v1.0; the document is now v2.0. Ask Rana Reviewer to propose it again.",
                          "code": "proposal_stale"}
    item = listing(w)["items"][0]
    assert (item["stale"], item["may_accept"], item["may_reject"], item["status"]) == (True, False, True, "pending")
    assert doc(w).content == BIG
    assert reject(w, pid, "Overtaken by 2.0").json()["proposal"]["status"] == "rejected"


def test_when_neither_the_owner_nor_the_author_is_an_active_user_an_administrator_decides_and_nothing_is_held(w):
    pid = propose_edit(w)
    w.db.get(m.GRCUser, OLIVIA).is_active = False
    w.db.commit()
    assert listing(w, ADAM)["can_decide"] is True and listing(w, ADAM)["items"][0]["may_accept"] is True
    assert accept(w, pid, uid=RANA).status_code == 403
    assert accept(w, pid, uid=ADAM).status_code == 200 and doc(w).content == SMALL
    out = put(w, UMA, content=BIG, change_reason="Exceptions")           # with nobody to decide, a change is direct
    assert "proposed" not in out.json() and doc(w).content == BIG


# ── every writer ─────────────────────────────────────────────────────────────
def test_a_new_file_from_an_editor_waits_and_becomes_the_documents_file_once_accepted(w):
    original = word_doc(w)
    before = open(original, "rb").read()
    new = longer_passwords(original, w.folder / "upload.docx")
    data = open(new, "rb").read()
    files_before = set(os.listdir(w.folder))
    out = w.call("POST", "/documents/6/upload-file", RANA, files={"file": ("policy v2.docx", data, DOCX)})
    assert out.status_code == 400 and set(os.listdir(w.folder)) == files_before          # no why: refused before the file is kept

    out = w.call("POST", "/documents/6/upload-file", RANA, files={"file": ("policy v2.docx", data, DOCX)},
                 data={"change_reason": "Longer passwords"})
    body = out.json()
    assert out.status_code == 200 and body["proposed"] is True
    p = body["proposal"]
    assert (p["source"], p["has_file"], p["message"]) == ("file", True, "Longer passwords")
    d = doc(w, 6)
    assert (d.file_path, d.current_version) == (original, "1.0") and len(versions(w, 6)) == 1 and body["document"]["current_version"] == "1.0"
    got = w.call("GET", f"/documents/6/proposals/{p['id']}/file", OLIVIA)
    assert got.status_code == 200 and got.content == data
    assert w.call("GET", f"/documents/5/proposals/{p['id']}/file", OLIVIA).status_code == 404       # not this document's
    cmp = w.call("GET", f"/documents/6/proposals/{p['id']}/compare", OLIVIA).json()
    assert any("fourteen" in c["after_text"] for c in cmp["changes"]) and cmp["base"]["label"] == "v1.0 · in the library"

    out = accept(w, p["id"], document_id=6).json()
    d = doc(w, 6)
    assert out["version"]["number"] == "1.1" and d.current_version == "1.1"
    assert d.file_path != original and open(d.file_path, "rb").read() == data and open(original, "rb").read() == before
    assert version_row(w, "1.1", 6).created_by == RANA


def test_a_version_and_a_restore_from_an_editor_wait_and_a_version_without_text_is_only_a_number(w):
    out = w.call("POST", "/versions/document/5", RANA, json={"change_type": "minor", "content": SMALL, "change_summary": "Scope",
                                                             "change_reason": WHY})
    assert out.status_code == 200 and out.json()["proposal"]["source"] == "version" and out.json()["proposal"]["change_type"] == "minor"
    assert doc(w).current_version == "1.0" and versions(w) == []
    accept(w, out.json()["proposal"]["id"])
    assert (doc(w).content, doc(w).current_version, version_row(w, "1.1").created_by) == (SMALL, "1.1", RANA)

    plain = w.call("POST", "/versions/document/5", RANA, json={"change_type": "minor", "change_summary": "Renumbered"})
    assert plain.status_code == 201 and doc(w).current_version == "1.2"                  # the same text: direct, as before

    first = version_row(w, "1.0").id
    out = w.call("POST", f"/versions/document/5/rollback/{first}", RANA, json={"rollback_reason": "Back to the reviewed text"})
    assert out.status_code == 200 and out.json()["proposal"]["source"] == "restore" and doc(w).content == SMALL
    accept(w, out.json()["proposal"]["id"])
    assert (doc(w).content, doc(w).current_version, version_row(w, "1.3").created_by) == (TEXT, "1.3", RANA)


def test_taking_changes_back_waits_for_the_owner(w):
    put(w, OLIVIA, content=SMALL, change_reason="Scope")                               # 1.1, the owner's
    first = version_row(w, "1.0").id
    red = w.call("GET", f"/documents/5/change-chain/redline?from={first}", RANA).json()
    out = w.call("POST", "/documents/5/change-chain/revert", RANA, json={"from_version_id": first, "hunks": [red["hunks"][0]["id"]],
                                                                          "hash": red["current_hash"], "reason": "Not agreed"})
    body = out.json()
    assert body["proposed"] is True and body["changes_taken_back"] == 0 and body["proposal"]["source"] == "take_back"
    assert doc(w).content == SMALL and doc(w).current_version == "1.1"
    accept(w, body["proposal"]["id"])
    row = version_row(w, "1.2")
    assert C.content_hash(doc(w).content) == C.content_hash(TEXT) and row.created_by == RANA
    assert row.change_set["items"][0]["origin"] == "revert" and row.change_set["items"][0]["reason"] == "Not agreed"


def test_saving_from_the_ai_studio_waits_and_the_draft_stays_unsaved_until_the_owner_accepts(w):
    s = w.call("POST", "/studio/sessions", RANA, json={"kind": "library", "document_id": 5}).json()
    w.call("PUT", f"/studio/sessions/{s['id']}/content", RANA, json={"content": SMALL, "base_revision": s["revision"]})
    out = w.call("POST", f"/studio/sessions/{s['id']}/save-to-document", RANA, json={"change_type": "minor", "change_reason": WHY}).json()
    assert out["proposed"] is True and out["version"] == "1.0" and out["change"] is None
    p = out["proposal"]
    assert (p["source"], p["change_type"], p["message"], p["studio_session_id"]) == ("studio", "minor", WHY, s["id"]) and doc(w).content == TEXT
    view = w.call("GET", f"/studio/sessions/{s['id']}", RANA).json()
    assert view["unsaved_in_library"] is True and view["proposal"]["id"] == p["id"] and view["proposal"]["status"] == "pending"

    accept(w, p["id"])
    row = version_row(w, "1.1")
    assert doc(w).content == SMALL and row.created_by == RANA
    assert "Chosen by Rana Reviewer when saving it to the library" in row.change_summary and "accepted by Olivia Owner" in row.change_summary
    assert row.change_set["session_id"] == s["id"] and row.change_set["items"][0]["by"] == "Rana Reviewer"   # who made what, as proposed
    view = w.call("GET", f"/studio/sessions/{s['id']}", RANA).json()
    assert view["unsaved_in_library"] is False and view["proposal"] is None and view["base_label"] == "Access policy v1.1"


def test_a_word_file_saved_from_the_studio_by_an_editor_is_kept_beside_the_library_until_accepted(w):
    original = word_doc(w)
    before = open(original, "rb").read()
    s = w.call("POST", "/studio/sessions", RANA, json={"kind": "library", "document_id": 6}).json()
    assert s["source"]["in_place"] is True
    new = s["content"].replace("Review it every year.", "Review it every six months.")
    w.call("PUT", f"/studio/sessions/{s['id']}/content", RANA, json={"content": new, "base_revision": s["revision"]})
    out = w.call("POST", f"/studio/sessions/{s['id']}/save-to-document", RANA, json={"change_type": "minor", "change_reason": "Twice a year"}).json()
    assert out["proposed"] is True and out["proposal"]["has_file"] is True
    assert doc(w, 6).file_path == original and open(original, "rb").read() == before
    got = w.call("GET", f"/documents/6/proposals/{out['proposal']['id']}/file", OLIVIA)
    assert "every six months" in file_text(got.content)

    assert accept(w, out["proposal"]["id"], document_id=6).status_code == 200
    d = doc(w, 6)
    assert d.file_path != original and "every six months" in file_text(d.file_path) and open(original, "rb").read() == before
    assert version_row(w, "1.1", 6).created_by == RANA
    w.db.expire_all()
    assert w.db.get(m.StudioSession, s["id"]).settings["library_file"] == d.file_path        # the draft is based on the new file


def _finding(w):
    w.db.add(m.PolicyGapAnalysisRun(id=1, tenant_id=1, document_id=5))
    w.db.add(m.PolicyGapFinding(id=1, tenant_id=1, analysis_run_id=1, document_id=5, framework_name="ISO/IEC 27001", clause_reference="A.5.18",
                                clause_title="Access rights", gap_description="Access rights are not reviewed.",
                                missing_requirement="A periodic review of access rights.", compliance_status="not_addressed",
                                remediation_status="open"))
    w.db.commit()


QUARTERLY = "Access rights are reviewed every quarter by their owners."


def test_a_gap_fix_from_an_editor_waits_with_the_gap_analysis_why_and_its_undo_waits_too(w):
    _finding(w)
    out = w.call("POST", "/gap-analysis/findings/1/apply-fix", RANA, json={"mode": "append", "proposed_text": QUARTERLY}).json()
    assert out["proposed"] is True and out["applied_at"] is None
    p = out["proposal"]
    assert p["message"] == "A.5.18 Access rights: Access rights are not reviewed. A periodic review of access rights."
    assert p["finding"] == {"id": 1, "clause_reference": "A.5.18", "clause_title": "Access rights"} and p["source"] == "gap_fix"
    w.db.expire_all()
    f = w.db.get(m.PolicyGapFinding, 1)
    assert (f.applied_at, f.compliance_status, f.remediation_status) == (None, "not_addressed", "open")
    assert doc(w).content == TEXT and w.db.query(m.PolicyStatement).count() == 0

    assert accept(w, p["id"]).status_code == 200
    w.db.expire_all()
    f = w.db.get(m.PolicyGapFinding, 1)
    assert (f.applied_by, f.compliance_status, f.remediation_status) == (RANA, "fully_compliant", "in_progress")
    assert QUARTERLY in doc(w).content and f.applied_version_id == version_row(w, doc(w).current_version).id
    assert w.db.query(m.PolicyStatement).one().statement_text == QUARTERLY            # the statement the fix added

    out = w.call("PUT", "/gap-analysis/findings/1/override", RANA, json={"override_status": "not_addressed",
                                                                          "override_justification": "Wrong clause"}).json()
    assert out["proposed"] is True and out["proposal"]["source"] == "gap_undo" and out["applied_at"] is not None
    assert QUARTERLY in doc(w).content and out["is_overridden"] is False
    assert accept(w, out["proposal"]["id"]).status_code == 200
    w.db.expire_all()
    f = w.db.get(m.PolicyGapFinding, 1)
    assert QUARTERLY not in doc(w).content and f.applied_at is None and (f.is_overridden, f.override_justification) == (True, "Wrong clause")
    assert w.db.query(m.PolicyStatement).one().status == "superseded"


def test_writing_statements_into_the_document_from_an_editor_waits_for_the_owner(w):
    path = str(w.folder / "sync.docx")
    d = Document()
    d.add_paragraph("Information Security Policy")
    d.add_paragraph(S1)
    d.save(path)
    w.db.add(m.GovernanceDocument(id=7, tenant_id=1, title="InfoSec policy", doc_type="policy", status="draft", current_version="1.0",
                                  file_name="sync.docx", file_path=path, file_type="docx", file_size=os.path.getsize(path), content=None,
                                  owner_id=OLIVIA, author_id=OLIVIA))
    w.db.add(m.PolicyStatement(id=70, tenant_id=1, document_id=7, statement_code="PS-0007-001", statement_text=S1, status="active"))
    w.db.commit()
    doc_sync.anchor(w.db, doc(w, 7))
    assert w.call("PUT", "/documents/7/statements/70", OLIVIA, json={"statement_text": S1_NOW}).status_code == 200     # the owner's: direct
    assert w.call("POST", "/documents/7/document-sync/apply", RANA, json={}).status_code == 400          # a why is required
    out = w.call("POST", "/documents/7/document-sync/apply", RANA, json={"change_reason": "Service owners too"}).json()
    assert out["proposed"] is True and out["proposal"]["source"] == "write_back" and out["proposal"]["has_file"] is True
    assert (doc(w, 7).file_path, doc(w, 7).current_version) == (path, "1.0") and S1_NOW not in file_text(path)

    assert accept(w, out["proposal"]["id"], document_id=7).status_code == 200
    d = doc(w, 7)
    assert d.file_path != path and S1_NOW in file_text(d.file_path) and version_row(w, d.current_version, 7).created_by == RANA
    assert w.db.get(m.PolicyStatement, 70).source_text == S1_NOW                     # written into the file: not pending any more


def test_statement_changes_from_an_editor_wait_with_their_justification_and_the_owners_stay_direct(w):
    new = {"statement_text": "Access is reviewed every quarter.", "category": "Access", "priority": "medium"}
    assert w.call("POST", "/documents/5/statements", RANA, json=new).status_code == 400        # a justification is required
    out = w.call("POST", "/documents/5/statements", RANA, json={**new, "change_reason": "Needed for the audit"}).json()
    assert out["proposed"] is True and out["message"] == "Sent to Olivia Owner to accept. The statement does not change until then."
    st = out["proposal"]["statement"]
    assert (st["id"], st["code"], st["before"], st["after"]["statement_text"]) == (None, None, None, new["statement_text"])
    assert out["proposal"]["kind"] == "statement" and w.db.query(m.PolicyStatement).count() == 0
    assert accept(w, out["proposal"]["id"]).json()["version"] is None
    w.db.expire_all()
    added = w.db.query(m.PolicyStatement).one()
    assert (added.statement_code, added.created_by, added.versions[0].change_reason) == ("PS-0005-001", RANA, "Needed for the audit")
    assert listing(w)["items"][0]["statement"]["code"] == "PS-0005-001"
    own = w.call("POST", "/documents/5/statements", OLIVIA, json={"statement_text": "Owners review access yearly."}).json()
    assert own["statement_code"] == "PS-0005-002" and "proposed" not in own              # the owner's own: no justification needed

    out = w.call("PUT", f"/documents/5/statements/{added.id}", RANA, json={"statement_text": "Access is reviewed monthly.", "priority": "high",
                                                                            "change_reason": "Monthly now"}).json()
    st = out["proposal"]["statement"]
    assert (st["before"]["statement_text"], st["after"]["statement_text"], st["after"]["priority"]) == (new["statement_text"], "Access is reviewed monthly.", "high")
    w.db.expire_all()
    assert w.db.get(m.PolicyStatement, added.id).statement_text == new["statement_text"]
    accept(w, out["proposal"]["id"])
    w.db.expire_all()
    edited = w.db.get(m.PolicyStatement, added.id)
    assert (edited.statement_text, edited.priority) == ("Access is reviewed monthly.", "high")
    assert max(edited.versions, key=lambda v: v.version_number).change_reason == "Monthly now"

    first = min(edited.versions, key=lambda v: v.version_number)
    out = w.call("POST", f"/documents/5/statements/{added.id}/rollback", RANA, json={"version_id": first.id, "change_reason": "Quarterly after all"}).json()
    assert out["proposed"] is True and out["proposal"]["statement"]["after"]["statement_text"] == new["statement_text"]
    accept(w, out["proposal"]["id"])
    w.db.expire_all()
    assert w.db.get(m.PolicyStatement, added.id).statement_text == new["statement_text"]

    out = w.call("DELETE", f"/documents/5/statements/{added.id}", RANA, json={"change_reason": "A duplicate"}).json()
    assert out["proposed"] is True and out["proposal"]["statement"]["after"] is None
    w.db.expire_all()
    assert w.db.get(m.PolicyStatement, added.id) is not None
    assert accept(w, out["proposal"]["id"]).status_code == 200
    w.db.expire_all()
    assert w.db.get(m.PolicyStatement, added.id) is None                                  # removed; the proposals stay, unlinked
    assert w.db.query(m.DocumentChangeProposal).filter_by(statement_id=added.id).count() == 0
    kinds = [e["kind"] for e in C.chain(w.db, doc(w))["events"] if e["kind"].startswith("statement")]
    assert kinds == ["statement_added", "statement_added", "statement_edited", "statement_rolled_back", "statement_deleted"]


def test_the_why_comes_by_each_writers_own_field_or_as_message(w):
    out = put(w, RANA, content=SMALL, message="By its other name").json()
    assert out["proposal"]["message"] == "By its other name"
    own = w.call("POST", "/documents/5/statements", OLIVIA, json={"statement_text": "Owners review access yearly."}).json()
    out = w.call("DELETE", f"/documents/5/statements/{own['id']}", RANA, json={"message": "A duplicate"}).json()
    assert out["proposal"]["message"] == "A duplicate"
    out = w.call("DELETE", f"/documents/5/statements/{own['id']}?change_reason=Again", RANA).json()
    assert out["proposal"]["message"] == "Again"
    first = w.db.query(m.PolicyStatementVersion).filter_by(statement_id=own["id"]).first()
    out = w.call("POST", f"/documents/5/statements/{own['id']}/rollback", RANA, json={"version_id": first.id, "message": "Back"}).json()
    assert out["proposal"]["message"] == "Back"
    assert w.call("DELETE", f"/documents/5/statements/{own['id']}", RANA).status_code == 400        # none at all: refused


def test_a_statement_proposal_is_stale_once_the_statement_changed(w):
    own = w.call("POST", "/documents/5/statements", OLIVIA, json={"statement_text": "Owners review access yearly."}).json()
    out = w.call("PUT", f"/documents/5/statements/{own['id']}", RANA, json={"statement_text": "Owners review access monthly.",
                                                                             "change_reason": "Monthly"}).json()
    w.call("PUT", f"/documents/5/statements/{own['id']}", OLIVIA, json={"statement_text": "Owners review access twice a year."})
    stale = accept(w, out["proposal"]["id"])
    assert stale.status_code == 409 and stale.json()["code"] == "proposal_stale"
    assert stale.json()["detail"] == "This change was made to PS-0005-001 as it was then; it has changed since. Ask Rana Reviewer to propose it again."


def test_statements_are_read_again_after_an_accepted_change_never_for_a_proposal(w, monkeypatch):
    started = []
    monkeypatch.setattr("grc.modules.governance.statement_reparse.start", lambda slug, ids, session: started.append((slug, list(ids))))
    w.db.info["tenant_slug"] = "cfsb"                                    # as a real tenant session carries it
    pid = propose_edit(w)
    assert started == []
    accept(w, pid)
    assert started == [("cfsb", [5])]


# ── in review ────────────────────────────────────────────────────────────────
def test_in_review_a_proposed_change_holds_the_route_and_accepting_it_starts_the_stage_again(w):
    review_then_approval(w)
    assert sign(w, RAFAY).status_code == 200
    rafay_signed = w.db.query(m.DocumentSignature).filter_by(signer_user_id=RAFAY).one().signed_at
    pid = propose_edit(w)                                                # Rana reviews it and may edit it
    v = signoff(w)
    h = v["route"]["hold"]
    assert (h["kind"], h["proposal_id"], h["proposal_status"], h["requested_by"], h["comment"]) == ("proposal", pid, "pending", "Rana Reviewer", WHY)
    assert v["can_resubmit"] is False and v["decide"] == {"count": 1, "first_id": pid} and v["my_step"] is None
    assert listing(w)["items"][0]["holds_review"] is True
    refused = sign(w, SAM)
    assert refused.status_code == 403 and refused.json()["detail"] == document_signoff.ON_HOLD_PROPOSAL
    assert w.call("POST", "/documents/5/signoff/resubmit", OLIVIA, json={}).status_code == 409
    assert pending(w) == [("decide", OLIVIA)]                            # every review waits; the owner decides first
    assert {o.user_id: o.status for o in w.db.query(m.DocumentObligation).filter_by(kind="review")} == {RANA: "held", RAFAY: "done", SAM: "held"}
    told = last_alert(w, SAM)
    assert (told.subject, told.message) == ("On hold: Access policy", "Rana Reviewer proposed a change; Olivia Owner decides first. "
                                                                      "You are asked again when it is decided.")
    assert w.call("POST", "/documents/5/signoff/send-for-review", OLIVIA).status_code == 409

    assert accept(w, pid).status_code == 200
    h = hold(w)
    assert (h.kind, h.outcome, h.change, h.resolved_by) == ("proposal", "resumed", "minor", OLIVIA)
    assert (doc(w).status, doc(w).current_version) == ("pending_review", "1.1")
    assert pending(w) == [("review", RANA), ("review", RAFAY), ("review", SAM)]      # the stage starts again for everyone
    stage = signoff(w)["route"]["stages"][0]
    assert stage["progress"] == {"required": 3, "signed": 0} and stage["restarted"] is True
    again = last_alert(w, RAFAY)
    assert again.subject == "Review again: Access policy"
    assert again.message.startswith(f"You reviewed this on {on(rafay_signed)} (v1.0). Olivia Owner made changes since (v1.0 → v1.1). Please review it again.")
    rana = [n for n in w.db.query(m.WorkflowNotification).filter_by(user_id=RANA).order_by(m.WorkflowNotification.id)][-2:]
    assert [n.subject for n in rana] == ["Review again: Access policy", "Change accepted: Access policy"]
    assert rana[0].message.startswith(f'You proposed a change on {on(datetime.utcnow())}: "{WHY}". Olivia Owner accepted it (v1.1) — see what changed.')
    for uid in (RANA, RAFAY, SAM):
        assert sign(w, uid).status_code == 200
    assert doc(w).status == "pending_approval"


def test_a_major_change_accepted_in_review_starts_the_whole_route_again(w):
    send(w, {"kind": "review", "name": "Department head", "people": [person(RAFAY)]},
         {"kind": "review", "name": "Compliance", "people": [person(RANA)]}, {"kind": "approval", "name": "Approval", "people": [person(MINA)]})
    sign(w, RAFAY)
    assert doc(w).route_stage == 2
    pid = propose_edit(w, content=BIG, why="Exceptions are missing")
    assert accept(w, pid).json()["version"]["number"] == "2.0"
    h = hold(w)
    assert (h.outcome, h.change) == ("restarted", "major")
    assert (doc(w).status, doc(w).route_stage) == ("pending_review", 1) and pending(w) == [("review", RAFAY)]


def test_not_accepted_in_review_its_proposer_signs_it_as_it_is_and_the_others_signatures_stand(w):
    review_then_approval(w, reviewers=(RANA, RAFAY))
    sign(w, RAFAY)
    pid = propose_edit(w)
    reject(w, pid, "Keep it as it is")
    v = signoff(w, RANA)
    assert v["route"]["hold"]["proposal_status"] == "rejected" and v["my_actionable_role"] is None
    step = v["my_step"]
    assert (step["may_sign_as_is"], step["role_type"], step["can_reject"], step["can_delegate"]) == (True, "reviewer", False, False)
    assert signoff(w, RAFAY)["my_step"] is None and signoff(w, MINA)["my_step"] is None
    asked = w.call("POST", "/documents/5/signoff/request-changes", RANA, json={"comment": "More"})
    assert asked.status_code == 403 and asked.json()["detail"] == document_signoff.AS_IS_ONLY        # signing it as it is, or propose again
    told = last_alert(w, RANA)
    assert told.message == 'Your change to "Access policy" was not accepted: "Keep it as it is". Sign it as it is, or propose another change. It is in your tasks.'
    assert pending(w) == [("revise", RANA)]

    out = sign(w, RANA).json()
    assert out["status"] == "pending_approval"                           # Rafay's signature stands: Rana's completes the stage
    h = hold(w)
    assert (h.outcome, h.change, h.resolved_by) == ("as_is", "none", RANA)
    assert pending(w) == [("approval", MINA)]
    reviews = {o.user_id: o.status for o in w.db.query(m.DocumentObligation).filter_by(kind="review")}
    assert reviews == {RANA: "done", RAFAY: "done"} and w.db.query(m.DocumentObligation).filter_by(kind="revise").one().status == "done"
    assert (doc(w).content, doc(w).current_version) == (TEXT, "1.0")


def test_not_accepted_in_review_its_proposer_may_propose_again_on_the_same_hold_and_the_owner_decides_again(w):
    review_then_approval(w)
    sign(w, RAFAY)
    first = propose_edit(w)
    reject(w, first, "Say interns too")
    again = TEXT.replace("contractors, vendors", "contractors, temporary staff, interns, vendors")
    second = propose_edit(w, content=again, why="Interns too")
    h = hold(w)
    assert (h.proposal_id, h.resolved_at, h.comment) == (second, None, "Interns too") and w.db.query(m.DocumentRouteHold).count() == 1
    assert pending(w) == [("decide", OLIVIA)]                            # the first one's revision is done
    assert signoff(w, RANA)["my_step"] is None                           # it waits for the owner again
    accept(w, second)
    assert hold(w).outcome == "resumed" and doc(w).content == again


def test_withdrawing_the_change_that_holds_the_review_lets_it_go_on_where_it_was(w):
    review_then_approval(w, reviewers=(RANA, RAFAY))
    sign(w, RAFAY)
    pid = propose_edit(w)
    assert pending(w, "review") == []
    assert withdraw(w, pid).status_code == 200
    h = hold(w)
    assert (h.outcome, h.change) == ("as_is", "none")
    assert pending(w) == [("review", RANA)] and signoff(w, RANA)["my_step"]["may_sign_as_is"] is False
    assert sign(w, RANA).json()["status"] == "pending_approval"           # Rafay's signature still counts


def test_outside_review_a_proposal_holds_nothing_and_a_published_document_is_reviewed_again_once_accepted(w):
    d = doc(w)
    d.status = "published"
    w.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=5, role_type="reviewer", target_type="user", target_id=RAFAY, added_by=OLIVIA))
    w.db.commit()
    pid = propose_edit(w)
    assert doc(w).status == "published" and w.db.query(m.DocumentRouteHold).count() == 0 and pending(w, "review") == []
    accept(w, pid)
    assert doc(w).status == "pending_review" and pending(w, "review") == [("review", RAFAY)]


def test_a_document_is_not_sent_for_review_while_a_proposed_change_waits(w):
    pid = propose_edit(w)
    w.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=5, role_type="reviewer", target_type="user", target_id=RAFAY, added_by=OLIVIA))
    w.db.commit()
    out = w.call("POST", "/documents/5/signoff/send-for-review", OLIVIA)
    assert out.status_code == 409
    assert out.json()["detail"] == "A change Rana Reviewer proposed waits for Olivia Owner: accept or reject it before it is sent for review."
    reject(w, pid, "Not now")
    assert w.call("POST", "/documents/5/signoff/send-for-review", OLIVIA).status_code == 200


# ── two fixes on the way ─────────────────────────────────────────────────────
def test_resubmitting_a_word_document_is_sized_by_its_files_text_not_its_empty_text(w):
    original = word_doc(w, document_id=5 + 3)                            # 8: a document that is its Word file
    w.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=8, role_type="reviewer", target_type="user", target_id=RAFAY, added_by=OLIVIA))
    w.db.add(m.DocumentSignoffAssignment(tenant_id=1, document_id=8, role_type="approver", target_type="user", target_id=MINA, added_by=OLIVIA))
    w.db.commit()
    assert w.call("POST", "/documents/8/signoff/send-for-review", OLIVIA).status_code == 200
    assert w.call("POST", "/documents/8/signoff/request-changes", RAFAY, json={"comment": "Longer passwords"}).status_code == 200
    new = longer_passwords(original, w.folder / "fourteen.docx")
    assert w.call("POST", "/documents/8/upload-file", OLIVIA, files={"file": ("policy.docx", open(new, "rb").read(), DOCX)}).status_code == 200
    out = w.call("POST", "/documents/8/signoff/resubmit", OLIVIA, json={"note": "Done"}).json()
    assert (out["change"], out["outcome"]) == ("minor", "resumed")                 # was "none": sized by the empty text
    assert document_signoff.change_since(w.db, doc(w, 8), hold_of(w, 8)) == "minor"


def hold_of(w, document_id):
    w.db.expire_all()
    return w.db.query(m.DocumentRouteHold).filter_by(document_id=document_id).one()


def test_the_workflow_draft_action_only_fills_an_empty_drafts_text(w):
    word_doc(w, document_id=6)
    w.db.add(m.GovernanceDocument(id=9, tenant_id=1, title="Empty draft", doc_type="policy", status="draft", content=None, owner_id=OLIVIA))
    w.db.add(m.GovernanceDocument(id=10, tenant_id=1, title="In review", doc_type="policy", status="pending_review", content=None, owner_id=OLIVIA))
    w.db.commit()

    def run(document_id):
        instance = SimpleNamespace(tenant_id=1, context={}, trigger_payload={})
        return WorkflowActionHandlers._execute_governance_document_action(w.db, instance, SimpleNamespace(name="Draft it"), "execute",
                                                                          "generate_policy_ai_draft", {"document_id": document_id})
    assert run(6)["result"] == "left_as_it_is" and doc(w, 6).content is None              # its file is its text
    assert run(10)["result"] == "left_as_it_is" and doc(w, 10).content is None            # being reviewed
    assert run(9)["result"] == "draft_generated" and doc(w, 9).content.startswith("# Empty draft")


# ── around it ────────────────────────────────────────────────────────────────
def test_the_daily_sweep_reminds_and_tells_on_proposed_changes_too(w):
    pid = propose_edit(w)
    asked = w.db.query(m.DocumentObligation).filter_by(kind="decide").one()
    assert round((asked.due_date - asked.requested_at).total_seconds() / 86400) == 3          # Governance > Settings: 3 days by default
    asked.due_date = datetime.utcnow() - timedelta(days=1)
    w.db.commit()
    out = life.sweep(w.db)
    assert out["owner_told"] == 1
    late = next(n for n in w.db.query(m.WorkflowNotification).filter_by(user_id=OLIVIA) if n.subject.startswith("Overdue: "))
    assert late.subject == "Overdue: Decide on Access policy" and late.link == f"/governance/documents/5?tab=proposals&proposal={pid}"


def test_reading_proposed_changes_writes_nothing(w):
    pid = propose_edit(w)
    commits = []
    event.listen(w.db, "after_commit", lambda session: commits.append(1))
    for path in ("/documents/5/proposals", f"/documents/5/proposals/{pid}", f"/documents/5/proposals/{pid}/compare", "/documents/5/signoff",
                 "/documents/5/lifecycle", "/documents/lifecycle/mine"):
        assert w.call("GET", path, OLIVIA).status_code == 200, path
    assert commits == [] and not w.db.new and not w.db.dirty


def test_a_document_with_proposed_changes_can_still_be_deleted(w):
    propose_edit(w)
    own = w.call("POST", "/documents/5/statements", OLIVIA, json={"statement_text": "Owners review access yearly."}).json()
    w.call("PUT", f"/documents/5/statements/{own['id']}", RANA, json={"statement_text": "Monthly.", "change_reason": "Monthly"})
    assert w.db.query(m.DocumentChangeProposal).count() == 2
    assert w.call("DELETE", "/documents/5", OLIVIA).status_code == 204
    assert w.db.query(m.DocumentChangeProposal).count() == 0 and doc(w) is None


def test_a_proposal_never_sets_the_documents_text_or_file(w):
    """statement_reparse (models _13) follows every set of a document's content or file: a proposal must set neither."""
    original = word_doc(w)
    upload = open(longer_passwords(original, w.folder / "y.docx"), "rb").read()
    seen = []
    listener = lambda target, value, old, initiator: seen.append(target.id)          # noqa: E731
    event.listen(m.GovernanceDocument.content, "set", listener)
    event.listen(m.GovernanceDocument.file_path, "set", listener)
    try:
        propose_edit(w)
        out = w.call("POST", "/documents/6/upload-file", RANA, files={"file": ("x.docx", upload, DOCX)}, data={"change_reason": "Longer"})
        assert out.json()["proposed"] is True and seen == []
    finally:
        event.remove(m.GovernanceDocument.content, "set", listener)
        event.remove(m.GovernanceDocument.file_path, "set", listener)
