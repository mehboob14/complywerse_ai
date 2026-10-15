"""Word files that look like the ones organisations really have, built with python-docx and raw XML: direct run formatting,
a title, headings, numbered and bulleted lists, a link, a field, a soft line break, tables (one with merged cells), an image,
a header and a footer; and a file with tracked changes."""
import io
import zipfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor


def png(size=(40, 20), color=(30, 90, 160, 255)) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGBA", size, color).save(out, "PNG")
    return out.getvalue()


def run(paragraph, text, *, bold=None, italic=None, size=11, font="Calibri", underline=None):
    r = paragraph.add_run(text)
    r.bold, r.italic, r.underline = bold, italic, underline
    r.font.size, r.font.name = Pt(size), font
    return r


def link(paragraph, url, text, size=11):
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, attrs in (("w:color", {"w:val": "0563C1"}), ("w:sz", {"w:val": str(size * 2)}), ("w:u", {"w:val": "single"})):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def field(paragraph, instruction, shown):
    for kind, text in (("begin", None), (None, instruction), ("separate", None), ("t", shown), ("end", None)):
        r = OxmlElement("w:r")
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r.append(fc)
        elif kind == "t":
            t = OxmlElement("w:t")
            t.text = text
            r.append(t)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {text} "
            r.append(it)
        paragraph._p.append(r)


def corporate_policy(path: str) -> str:
    d = Document()
    d.sections[0].header.paragraphs[0].text = "ACME BANK - CONFIDENTIAL"
    d.sections[0].footer.paragraphs[0].text = "POL-001 v1.0"
    d.add_paragraph("Information Security Policy", style="Title")
    p = d.add_paragraph()
    run(p, "This policy applies to ")
    run(p, "all employees", bold=True)
    run(p, " and contractors of the bank, see ")
    link(p, "https://example.com/ecc", "the ECC")
    run(p, ". Review it every year.")
    d.add_heading("1. Purpose", level=1)
    p = d.add_paragraph()
    run(p, "The purpose is to protect ")
    run(p, "information", italic=True)
    run(p, " assets.", size=11)
    d.add_heading("2. Rules", level=1)
    d.add_paragraph("Use strong passwords", style="List Number")
    d.add_paragraph("Lock your screen", style="List Number")
    d.add_paragraph("Report incidents", style="List Number")
    d.add_paragraph("Hardware is covered", style="List Bullet")
    d.add_paragraph("Software is covered", style="List Bullet")
    d.add_heading("2.1 Passwords", level=2)
    p = d.add_paragraph()
    run(p, "Passwords shall be at least twelve characters long.", size=10.5, font="Arial")
    p = d.add_paragraph()
    run(p, "First line")
    p.runs[-1].add_break()
    run(p, "second line")
    d.add_paragraph().add_run().add_picture(io.BytesIO(png()))
    p = d.add_paragraph()
    run(p, "Page ")
    field(p, "PAGE", "3")
    run(p, " of this policy.")
    d.add_heading("3. Roles", level=1)
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for c, h in enumerate(("Role", "Responsibility", "Owner")):
        cell = t.cell(0, c)
        cell.paragraphs[0].text = ""
        run(cell.paragraphs[0], h, bold=True, size=10)
    for r, row in enumerate((("CISO", "Owns the policy", "Board"), ("IT", "Implements controls", "CISO")), start=1):
        for c, v in enumerate(row):
            run(t.cell(r, c).paragraphs[0], v, size=10)
    d.add_paragraph("After the table.")
    m = d.add_table(rows=3, cols=3)
    m.style = "Table Grid"
    m.cell(0, 0).merge(m.cell(0, 1))
    m.cell(0, 0).paragraphs[0].text = "Merged header"
    m.cell(0, 2).paragraphs[0].text = "Status"
    for r, row in enumerate((("A", "B", "Open"), ("C", "D", "Closed")), start=1):
        for c, v in enumerate(row):
            m.cell(r, c).paragraphs[0].text = v
    d.add_paragraph("The end.")
    d.save(path)
    return path


def _xml(fragment: str):
    return parse_xml(fragment.replace("<w:p>", f"<w:p {nsdecls('w')}>", 1) if fragment.startswith("<w:p>") else fragment)


AUTH = 'w:author="Reviewer" w:date="2026-01-01T00:00:00Z"'


def tracked_policy(path: str) -> str:
    d = Document()
    d.add_paragraph("Retention Policy", style="Title")
    body = d.element.body
    sect = body.find(qn("w:sectPr"))

    def add(xml: str) -> None:
        el = parse_xml(xml)
        (sect.addprevious(el) if sect is not None else body.append(el))

    n = nsdecls("w")
    add(f'<w:p {n}><w:r><w:t xml:space="preserve">Records shall be kept for </w:t></w:r><w:del w:id="1" {AUTH}><w:r><w:delText>five</w:delText></w:r></w:del>'
        f'<w:ins w:id="2" {AUTH}><w:r><w:t>seven</w:t></w:r></w:ins><w:r><w:t xml:space="preserve"> years.</w:t></w:r></w:p>')
    add(f'<w:p {n}><w:pPr><w:rPr><w:ins w:id="3" {AUTH}/></w:rPr></w:pPr><w:ins w:id="4" {AUTH}><w:r><w:t>A paragraph the reviewer added.</w:t></w:r></w:ins></w:p>')
    add(f'<w:p {n}><w:pPr><w:rPr><w:del w:id="5" {AUTH}/></w:rPr></w:pPr><w:del w:id="6" {AUTH}><w:r><w:delText>A paragraph the reviewer removed.</w:delText></w:r></w:del></w:p>')
    add(f'<w:p {n}><w:r><w:t xml:space="preserve">Make this </w:t></w:r><w:r><w:rPr><w:b/><w:rPrChange w:id="7" {AUTH}><w:rPr/></w:rPrChange></w:rPr><w:t>important</w:t></w:r></w:p>')
    add(f'<w:p {n}><w:r><w:t>Untouched paragraph.</w:t></w:r></w:p>')
    d.save(path)
    return path


def parts(path_or_bytes) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, (bytes, bytearray)) else path_or_bytes)
